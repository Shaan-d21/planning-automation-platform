"""Tests for Planning form-based Excel reports."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from app.clients.epm_client import EPMClient
from app.models.data_validation import FormGrid, FormGridRow, FormLayout
from app.models.report import FormReportRequest
from app.services.data_slice_report_service import DataSliceReportService
from app.services.excel_report_renderer import ExcelFormReportRenderer
from app.services.form_service import PlanningFormService
from app.services.report_catalog_service import ReportCatalogService
from app.services.report_service import FormReportService
from app.utils.exceptions import APIRequestError, ReportGenerationError


def _client() -> Mock:
    client = Mock(spec=EPMClient)
    client.application_name = "Vision"
    return client


def _layout() -> FormLayout:
    return FormLayout(
        page_dimensions=("Scenario", "Version"),
        row_dimensions=("Account",),
        column_dimensions=("Period",),
        allowed_page_members=(
            ("Scenario", ("Forecast", "Plan")),
            ("Version", ("Working", "Final")),
        ),
        pov=(
            ("Scenario", "Forecast"),
            ("Version", "Working"),
        ),
    )


def _grid() -> FormGrid:
    return FormGrid(
        row_dimensions=("Account",),
        column_dimensions=("Period",),
        columns=(("Jan",), ("Feb",)),
        rows=(
            FormGridRow(
                headers=("Revenue",),
                data=(1000, "1,250.50"),
            ),
            FormGridRow(
                headers=("Units",),
                data=("#Missing", 20),
            ),
        ),
        pov=(
            ("Scenario", "Plan"),
            ("Version", "Working"),
        ),
    )


def test_report_service_resolves_pov_and_creates_workbook(
    tmp_path: Path,
) -> None:
    form_service = Mock(spec=PlanningFormService)
    form_service.get_form_layout.return_value = _layout()
    form_service.export_form.return_value = _grid()
    output = tmp_path / "Revenue_Report.xlsx"
    service = FormReportService(
        _client(),
        form_service=form_service,
    )

    result = service.generate(
        FormReportRequest(
            form_name="Revenue Form",
            output_path=output,
            title="Revenue Forecast",
            page_member_overrides=(("Scenario", "Plan"),),
        )
    )

    assert result.output_path == output.resolve()
    assert result.row_count == 2
    assert result.data_cell_count == 4
    form_service.export_form.assert_called_once_with(
        "Revenue Form",
        page_members=("Plan", "Working"),
        filter_members=(),
    )
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["Report", "Report Metadata"]
    assert workbook["Report"]["A1"].value == "Revenue Forecast"
    assert workbook["Report"]["B11"].value == 1000
    assert workbook["Report"]["C11"].value == 1250.5
    assert workbook["Report"]["B12"].value is None
    assert workbook["Report Metadata"]["B5"].value == "Revenue Form"


def test_report_service_rejects_unknown_page_dimension(
    tmp_path: Path,
) -> None:
    form_service = Mock(spec=PlanningFormService)
    form_service.get_form_layout.return_value = _layout()
    service = FormReportService(
        _client(),
        form_service=form_service,
    )

    with pytest.raises(ReportGenerationError, match="not on the form"):
        service.generate(
            FormReportRequest(
                form_name="Revenue Form",
                output_path=tmp_path / "report.xlsx",
                title="Report",
                page_member_overrides=(("Year", "FY26"),),
            )
        )

    form_service.export_form.assert_not_called()


def test_excel_renderer_does_not_overwrite_without_permission(
    tmp_path: Path,
) -> None:
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing")

    with pytest.raises(ReportGenerationError, match="already exists"):
        ExcelFormReportRenderer().render(
            application_name="Vision",
            form_name="Revenue Form",
            title="Revenue Forecast",
            grid=_grid(),
            output_path=output,
            generated_at=datetime.now().astimezone(),
            overwrite=False,
        )

    assert output.read_bytes() == b"existing"


def test_excel_renderer_can_return_downloadable_bytes() -> None:
    content = ExcelFormReportRenderer().render_bytes(
        application_name="Vision",
        form_name="Plan1 data slice",
        title="Plan1 Data Review",
        grid=_grid(),
        generated_at=datetime.now().astimezone(),
    )

    workbook = load_workbook(BytesIO(content), data_only=True)

    assert workbook["Report"]["A1"].value == "Plan1 Data Review"
    assert workbook["Report Metadata"]["B5"].value == "Plan1 data slice"


def test_report_service_uses_catalog_data_slice(
    tmp_path: Path,
) -> None:
    catalog_file = tmp_path / "reports.json"
    catalog_file.write_text(
        """
        {
          "reports": [{
            "name": "Revenue Form",
            "title": "Revenue",
            "cube": "VisASO",
            "pov": {"Scenario": "Actual", "Year": "FY22"},
            "columns": [{
              "dimensions": ["Period"],
              "members": [["Jan", "Feb"]]
            }],
            "rows": [{
              "dimensions": ["Account"],
              "members": [["Revenue", "Units"]]
            }]
          }]
        }
        """,
        encoding="utf-8",
    )
    form_service = Mock(spec=PlanningFormService)
    data_slice_service = Mock(spec=DataSliceReportService)
    data_slice_service.export.return_value = _grid()
    service = FormReportService(
        _client(),
        form_service=form_service,
        data_slice_service=data_slice_service,
        catalog_service=ReportCatalogService(),
        catalog_file=catalog_file,
    )

    layout = service.get_form_layout("Revenue Form")
    result = service.generate(
        FormReportRequest(
            form_name="Revenue Form",
            output_path=tmp_path / "catalog-report.xlsx",
            title="Revenue",
            page_member_overrides=(("Scenario", "Forecast"),),
        ),
        layout=layout,
    )

    assert result.data_cell_count == 4
    form_service.get_form_layout.assert_not_called()
    form_service.export_form.assert_not_called()
    definition = data_slice_service.export.call_args.args[0]
    assert definition.cube == "VisASO"
    assert data_slice_service.export.call_args.kwargs["pov"] == (
        ("Scenario", "Forecast"),
        ("Year", "FY22"),
    )


def test_cloud_registered_report_uses_stable_data_slice_definition(
    tmp_path: Path,
) -> None:
    catalog_file = tmp_path / "reports.json"
    catalog_file.write_text(
        """
        {
          "reports": [{
            "name": "Revenue Form",
            "title": "Revenue",
            "cube": "Plan1",
            "pov": {"Scenario": "Forecast", "Version": "Working"},
            "columns": [{
              "dimensions": ["Period"],
              "members": [["Jan", "Feb"]]
            }],
            "rows": [{
              "dimensions": ["Account"],
              "members": [["Revenue", "Units"]]
            }]
          }]
        }
        """,
        encoding="utf-8",
    )
    client = _client()
    client.is_cloud_environment = True
    form_service = Mock(spec=PlanningFormService)
    data_slice_service = Mock(spec=DataSliceReportService)
    data_slice_service.export.return_value = _grid()
    service = FormReportService(
        client,
        form_service=form_service,
        data_slice_service=data_slice_service,
        catalog_service=ReportCatalogService(),
        catalog_file=catalog_file,
    )

    result = service.generate(
        FormReportRequest(
            form_name="Revenue Form",
            output_path=tmp_path / "cloud-form.xlsx",
            title="Revenue",
        )
    )

    assert result.data_cell_count == 4
    form_service.get_form_layout.assert_not_called()
    form_service.export_form.assert_not_called()
    data_slice_service.export.assert_called_once()


def test_unregistered_form_404_explains_cloud_api_and_report_difference(
    tmp_path: Path,
) -> None:
    catalog_file = tmp_path / "reports.json"
    catalog_file.write_text(
        '{"reports": []}',
        encoding="utf-8",
    )
    client = _client()
    client.is_cloud_environment = True
    form_service = Mock(spec=PlanningFormService)
    not_found = APIRequestError("Not Found", status_code=404)
    form_service.get_form_layout.side_effect = not_found
    service = FormReportService(
        client,
        form_service=form_service,
        catalog_service=ReportCatalogService(),
        catalog_file=catalog_file,
    )

    with pytest.raises(
        ReportGenerationError,
        match="Oracle Reports artifacts are not Planning forms",
    ):
        service.get_form_layout("Revenue Form")
