"""Tests for downloadable validation evidence workbooks."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.models.data_validation import (
    DataComparisonCell,
    DataQualityIssue,
    DataQualityResult,
    DataQualityRules,
    DataValidationResult,
)
from app.services.data_validation_excel_renderer import (
    DataValidationExcelRenderer,
)


def test_quality_validation_workbook_contains_summary_and_exception() -> None:
    result = DataQualityResult(
        status="FAIL",
        checked_cells=2,
        passed_cells=1,
        issue_count=1,
        missing_count=1,
        zero_count=0,
        below_minimum_count=0,
        above_maximum_count=0,
        non_numeric_count=0,
        issues=(
            DataQualityIssue(
                code="MISSING",
                severity="ERROR",
                message="No data.",
                pov=(("Scenario", "Forecast"),),
                row_headers=("Revenue",),
                column_headers=("Jan",),
                raw_value="#Missing",
                numeric_value=None,
            ),
        ),
        truncated=False,
        rules=DataQualityRules(),
    )

    content = DataValidationExcelRenderer().render_quality(
        application_name="Vision",
        cube="Plan1",
        pov={"Scenario": "Forecast"},
        result=result,
    )

    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == ["Validation Summary", "Exceptions"]
    assert workbook["Validation Summary"]["B6"].value == "FAIL"
    assert workbook["Exceptions"]["B4"].value == "Missing"
    assert workbook["Exceptions"]["D4"].value == "Revenue"


def test_comparison_workbook_contains_all_cell_evidence() -> None:
    cell = DataComparisonCell(
        row_headers=("Revenue",),
        column_headers=("Jan",),
        source_value=Decimal("100"),
        target_value=Decimal("99"),
        difference=Decimal("1"),
        matches=False,
    )
    result = DataValidationResult(
        source_form="Plan1 source slice",
        target_form="Rpt target slice",
        compared_cells=1,
        matched_cells=0,
        mismatches=(),
        tolerance=Decimal("0"),
        cells=(cell,),
    )

    content = DataValidationExcelRenderer().render_comparison(
        application_name="Vision",
        source_cube="Plan1",
        target_cube="Rpt",
        source_pov={"Scenario": "Forecast"},
        target_pov={"Scenario": "Forecast"},
        result=result,
    )

    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == ["Comparison Summary", "Cell Comparison"]
    assert workbook["Cell Comparison"]["A4"].value == "DIFFERENCE"
    assert workbook["Cell Comparison"]["F4"].value == "1"
