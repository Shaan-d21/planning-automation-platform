"""Excel evidence workbooks for live Planning data validation."""

from __future__ import annotations

from collections.abc import Mapping
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.data_validation import (
    DataComparisonCell,
    DataQualityResult,
    DataValidationResult,
)
from app.utils.exceptions import DataValidationError


class DataValidationExcelRenderer:
    """Render quality checks and reconciliations as portable XLSX evidence."""

    _TITLE_FILL = PatternFill("solid", fgColor="17365D")
    _HEADER_FILL = PatternFill("solid", fgColor="2F5EE7")
    _PASS_FILL = PatternFill("solid", fgColor="DDF4EA")
    _WARNING_FILL = PatternFill("solid", fgColor="FFF0CE")
    _ERROR_FILL = PatternFill("solid", fgColor="FCE0E3")
    _WHITE_BOLD = Font(color="FFFFFF", bold=True)

    def render_quality(
        self,
        *,
        application_name: str,
        cube: str,
        pov: Mapping[str, str],
        result: DataQualityResult,
    ) -> bytes:
        """Return a quality-check workbook as bytes."""
        workbook = Workbook()
        summary = workbook.active
        assert summary is not None
        summary.title = "Validation Summary"
        self._title(summary, "Oracle EPM Data Quality Validation", 7)
        details = (
            ("Application", application_name),
            ("Cube", cube),
            ("Generated (UTC)", self._generated_at()),
            ("Status", result.status),
            ("Checked cells", result.checked_cells),
            ("Passed cells", result.passed_cells),
            ("Issues", result.issue_count),
            ("Missing", result.missing_count),
            ("Zeros", result.zero_count),
            ("Below minimum", result.below_minimum_count),
            ("Above maximum", result.above_maximum_count),
            ("Non-numeric", result.non_numeric_count),
            ("Minimum", self._value(result.rules.minimum)),
            ("Maximum", self._value(result.rules.maximum)),
            ("Issue list truncated", "Yes" if result.truncated else "No"),
        )
        self._details(summary, details, start_row=3)
        pov_row = 3 + len(details) + 2
        summary.cell(pov_row, 1, "Point of View").font = Font(bold=True)
        self._table_header(summary, pov_row + 1, ("Dimension", "Member"))
        for offset, (dimension, member) in enumerate(pov.items(), 1):
            summary.cell(pov_row + 1 + offset, 1, self._value(dimension))
            summary.cell(pov_row + 1 + offset, 2, self._value(member))
        self._fit(summary)

        exceptions = workbook.create_sheet("Exceptions")
        self._title(exceptions, "Validation Exceptions", 8)
        headers = (
            "Severity",
            "Check",
            "POV",
            "Row intersection",
            "Column intersection",
            "Value",
            "Numeric value",
            "Explanation",
        )
        self._table_header(exceptions, 3, headers)
        for row_number, issue in enumerate(result.issues, 4):
            values = (
                issue.severity,
                issue.code.replace("_", " ").title(),
                self._pairs(issue.pov),
                " | ".join(issue.row_headers),
                " | ".join(issue.column_headers),
                self._value(issue.raw_value),
                self._value(issue.numeric_value),
                issue.message,
            )
            for column, value in enumerate(values, 1):
                exceptions.cell(row_number, column, self._value(value))
            exceptions.cell(row_number, 1).fill = (
                self._ERROR_FILL
                if issue.severity == "ERROR"
                else self._WARNING_FILL
            )
        if not result.issues:
            exceptions.cell(4, 1, "No exceptions were found.")
            exceptions.cell(4, 1).fill = self._PASS_FILL
        exceptions.freeze_panes = "A4"
        exceptions.auto_filter.ref = f"A3:H{max(3, len(result.issues) + 3)}"
        self._fit(exceptions)
        return self._bytes(workbook)

    def render_comparison(
        self,
        *,
        application_name: str,
        source_cube: str,
        target_cube: str,
        source_pov: Mapping[str, str],
        target_pov: Mapping[str, str],
        result: DataValidationResult,
    ) -> bytes:
        """Return a source-target reconciliation workbook as bytes."""
        workbook = Workbook()
        summary = workbook.active
        assert summary is not None
        summary.title = "Comparison Summary"
        mismatch_count = result.compared_cells - result.matched_cells
        self._title(summary, "Oracle EPM Source-to-Target Validation", 6)
        self._details(
            summary,
            (
                ("Application", application_name),
                ("Source cube", source_cube),
                ("Target cube", target_cube),
                ("Generated (UTC)", self._generated_at()),
                ("Status", "PASS" if mismatch_count == 0 else "FAIL"),
                ("Compared cells", result.compared_cells),
                ("Matched cells", result.matched_cells),
                ("Differences", mismatch_count),
                ("Tolerance", self._value(result.tolerance)),
                ("Source POV", self._mapping(source_pov)),
                ("Target POV", self._mapping(target_pov)),
            ),
            start_row=3,
        )
        self._fit(summary)

        differences = workbook.create_sheet("Cell Comparison")
        self._title(differences, "Source-to-Target Cell Comparison", 7)
        headers = (
            "Status",
            "Row intersection",
            "Column intersection",
            "Source",
            "Target",
            "Difference",
            "Tolerance",
        )
        self._table_header(differences, 3, headers)
        cells = result.cells or tuple(
            _mismatch_cell(item) for item in result.mismatches
        )
        for row_number, cell in enumerate(cells, 4):
            values = (
                "MATCH" if cell.matches else "DIFFERENCE",
                " | ".join(cell.row_headers),
                " | ".join(cell.column_headers),
                self._value(cell.source_value),
                self._value(cell.target_value),
                self._value(cell.difference),
                self._value(result.tolerance),
            )
            for column, value in enumerate(values, 1):
                differences.cell(row_number, column, self._value(value))
            differences.cell(row_number, 1).fill = (
                self._PASS_FILL if cell.matches else self._ERROR_FILL
            )
        differences.freeze_panes = "A4"
        differences.auto_filter.ref = f"A3:G{max(3, len(cells) + 3)}"
        self._fit(differences)
        return self._bytes(workbook)

    def _title(self, sheet: Any, title: str, columns: int) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        cell = sheet.cell(1, 1, title)
        cell.fill = self._TITLE_FILL
        cell.font = Font(color="FFFFFF", bold=True, size=15)
        cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 27
        sheet.sheet_view.showGridLines = False

    @staticmethod
    def _details(sheet: Any, values: tuple[tuple[str, Any], ...], *, start_row: int) -> None:
        for offset, (label, value) in enumerate(values):
            row = start_row + offset
            sheet.cell(row, 1, label).font = Font(bold=True)
            sheet.cell(row, 2, DataValidationExcelRenderer._value(value))

    def _table_header(self, sheet: Any, row: int, headers: tuple[str, ...]) -> None:
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(row, column, header)
            cell.fill = self._HEADER_FILL
            cell.font = self._WHITE_BOLD
            cell.alignment = Alignment(wrap_text=True)

    @staticmethod
    def _fit(sheet: Any) -> None:
        for column in range(1, sheet.max_column + 1):
            width = max(
                (len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 45)

    @staticmethod
    def _bytes(workbook: Workbook) -> bytes:
        stream = BytesIO()
        try:
            workbook.save(stream)
        except (OSError, ValueError) as exc:
            raise DataValidationError(
                "The validation Excel workbook could not be created."
            ) from exc
        return stream.getvalue()

    @staticmethod
    def _generated_at() -> str:
        return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")

    @staticmethod
    def _value(value: Any) -> Any:
        if value is None:
            return "Missing"
        if isinstance(value, (int, float, bool)):
            return value
        text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
        if text.startswith(("=", "+", "-", "@")):
            try:
                float(text.replace(",", ""))
            except ValueError:
                return f"'{text}"
        return text

    @staticmethod
    def _pairs(values: tuple[tuple[str, str], ...]) -> str:
        return " | ".join(f"{dimension}: {member}" for dimension, member in values)

    @staticmethod
    def _mapping(values: Mapping[str, str]) -> str:
        return " | ".join(f"{dimension}: {member}" for dimension, member in values.items())


def _mismatch_cell(item) -> DataComparisonCell:
    """Adapt one bounded mismatch to the full-cell renderer contract."""
    return DataComparisonCell(
        row_headers=item.row_headers,
        column_headers=item.column_headers,
        source_value=item.source_value,
        target_value=item.target_value,
        difference=item.difference,
        matches=False,
    )
