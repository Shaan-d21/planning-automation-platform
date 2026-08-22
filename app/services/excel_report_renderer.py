"""Professional Excel rendering for Planning form reports."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

from app.models.data_validation import FormGrid
from app.utils.exceptions import ReportGenerationError


class ExcelFormReportRenderer:
    """Render one normalized Planning form grid as a readable workbook."""

    _TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
    _SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
    _HEADER_FILL = PatternFill("solid", fgColor="4472C4")
    _SUBHEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
    _WHITE_FONT = Font(color="FFFFFF", bold=True)
    _LABEL_FONT = Font(color="1F1F1F", bold=True)
    _THIN_GRAY = Side(style="thin", color="D9E2F3")
    _NUMBER_FORMAT = '#,##0.00;[Red](#,##0.00);-'
    _MISSING_VALUES = {"", "#missing", "missing", "none", "null"}

    def render(
        self,
        *,
        application_name: str,
        form_name: str,
        title: str,
        grid: FormGrid,
        output_path: Path,
        generated_at: datetime,
        overwrite: bool,
    ) -> Path:
        """Create and save an XLSX workbook without silently overwriting."""
        path = Path(output_path).expanduser()
        if path.suffix.lower() != ".xlsx":
            raise ReportGenerationError(
                "Report output filename must use the .xlsx extension."
            )
        if path.exists() and not overwrite:
            raise ReportGenerationError(
                f"Report output already exists: '{path}'. Use overwrite "
                "only when replacement is intentional."
            )

        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            workbook = self._workbook(
                application_name=application_name,
                form_name=form_name,
                title=title,
                grid=grid,
                generated_at=generated_at,
            )
            workbook.save(path)
        except ReportGenerationError:
            raise
        except (OSError, ValueError) as exc:
            raise ReportGenerationError(
                f"Unable to create Excel report '{path}': {exc}"
            ) from exc
        return path.resolve()

    def _workbook(
        self,
        *,
        application_name: str,
        form_name: str,
        title: str,
        grid: FormGrid,
        generated_at: datetime,
    ) -> Workbook:
        """Build the shared workbook used by file and HTTP exports."""
        workbook = Workbook()
        report_sheet = workbook.active
        if report_sheet is None:
            raise ReportGenerationError(
                "Unable to create the Excel report worksheet."
            )
        report_sheet.title = "Report"
        self._render_report_sheet(
            report_sheet,
            application_name=application_name,
            form_name=form_name,
            title=title,
            grid=grid,
            generated_at=generated_at,
        )
        self._render_metadata_sheet(
            workbook.create_sheet("Report Metadata"),
            application_name=application_name,
            form_name=form_name,
            title=title,
            grid=grid,
            generated_at=generated_at,
        )
        return workbook

    def render_bytes(
        self,
        *,
        application_name: str,
        form_name: str,
        title: str,
        grid: FormGrid,
        generated_at: datetime,
    ) -> bytes:
        """Render a workbook in memory for an authenticated HTTP download."""
        try:
            workbook = self._workbook(
                application_name=application_name,
                form_name=form_name,
                title=title,
                grid=grid,
                generated_at=generated_at,
            )
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        except ReportGenerationError:
            raise
        except (OSError, ValueError) as exc:
            raise ReportGenerationError(
                f"Unable to create the Excel report: {exc}"
            ) from exc

    def _render_report_sheet(
        self,
        sheet: Any,
        *,
        application_name: str,
        form_name: str,
        title: str,
        grid: FormGrid,
        generated_at: datetime,
    ) -> None:
        row_dimension_count = max(len(grid.row_dimensions), 1)
        column_count = max(len(grid.columns), 1)
        total_columns = row_dimension_count + column_count

        sheet.sheet_view.showGridLines = False
        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=total_columns,
        )
        title_cell = sheet.cell(1, 1, self._safe_text(title))
        title_cell.fill = self._TITLE_FILL
        title_cell.font = Font(color="FFFFFF", bold=True, size=16)
        title_cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 28

        sheet.cell(2, 1, "Application").font = self._LABEL_FONT
        sheet.cell(2, 2, self._safe_text(application_name))
        sheet.cell(3, 1, "Planning Form").font = self._LABEL_FONT
        sheet.cell(3, 2, self._safe_text(form_name))
        sheet.cell(4, 1, "Generated").font = self._LABEL_FONT
        if total_columns >= 3:
            sheet.merge_cells(
                start_row=4,
                start_column=2,
                end_row=4,
                end_column=3,
            )
        generated_cell = sheet.cell(4, 2, generated_at.replace(tzinfo=None))
        generated_cell.number_format = "yyyy-mm-dd hh:mm:ss"

        current_row = 6
        if grid.pov:
            sheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=total_columns,
            )
            pov_title = sheet.cell(current_row, 1, "Point of View")
            pov_title.fill = self._SECTION_FILL
            pov_title.font = self._LABEL_FONT
            current_row += 1
            for dimension, member in grid.pov:
                sheet.cell(current_row, 1, self._safe_text(dimension)).font = (
                    self._LABEL_FONT
                )
                sheet.cell(current_row, 2, self._safe_text(member))
                current_row += 1
            current_row += 1

        column_header_depth = max(
            len(grid.column_dimensions),
            max((len(column) for column in grid.columns), default=1),
            1,
        )
        header_start = current_row
        data_start = header_start + column_header_depth

        for row_index in range(
            header_start,
            header_start + column_header_depth,
        ):
            for column_index in range(1, row_dimension_count + 1):
                cell = sheet.cell(row_index, column_index)
                cell.fill = self._HEADER_FILL
                cell.font = self._WHITE_FONT
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        for index in range(row_dimension_count):
            label = (
                grid.row_dimensions[index]
                if index < len(grid.row_dimensions)
                else "Member"
            )
            sheet.cell(header_start, index + 1, self._safe_text(label))

        for column_index, headers in enumerate(
            grid.columns,
            start=row_dimension_count + 1,
        ):
            for level in range(column_header_depth):
                value = headers[level] if level < len(headers) else ""
                cell = sheet.cell(
                    header_start + level,
                    column_index,
                    self._safe_text(value),
                )
                cell.fill = self._HEADER_FILL
                cell.font = self._WHITE_FONT
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        for row_offset, grid_row in enumerate(grid.rows):
            row_index = data_start + row_offset
            for header_index in range(row_dimension_count):
                value = (
                    grid_row.headers[header_index]
                    if header_index < len(grid_row.headers)
                    else ""
                )
                cell = sheet.cell(
                    row_index,
                    header_index + 1,
                    self._safe_text(value),
                )
                cell.fill = self._SUBHEADER_FILL
                cell.alignment = Alignment(vertical="top")
            for data_index, raw_value in enumerate(grid_row.data):
                cell = sheet.cell(
                    row_index,
                    row_dimension_count + data_index + 1,
                    self._excel_value(raw_value),
                )
                if isinstance(cell.value, (int, float)):
                    cell.number_format = self._NUMBER_FORMAT
                    cell.alignment = Alignment(horizontal="right")

        final_row = max(data_start + len(grid.rows) - 1, header_start)
        data_range = sheet.iter_rows(
            min_row=header_start,
            max_row=final_row,
            min_col=1,
            max_col=total_columns,
        )
        for row in data_range:
            for cell in row:
                cell.border = Border(bottom=self._THIN_GRAY)

        sheet.freeze_panes = sheet.cell(data_start, row_dimension_count + 1)
        sheet.auto_filter.ref = (
            f"A{header_start}:"
            f"{get_column_letter(total_columns)}{final_row}"
        )
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = (
            f"{header_start}:{data_start - 1}"
        )
        sheet.sheet_view.zoomScale = 90
        self._size_columns(sheet, total_columns, row_dimension_count)

    def _render_metadata_sheet(
        self,
        sheet: Any,
        *,
        application_name: str,
        form_name: str,
        title: str,
        grid: FormGrid,
        generated_at: datetime,
    ) -> None:
        sheet.sheet_view.showGridLines = False
        sheet["A1"] = "Report Metadata"
        sheet["A1"].fill = self._TITLE_FILL
        sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
        sheet.merge_cells("A1:B1")

        metadata = [
            ("Report title", title),
            ("Application", application_name),
            ("Planning form", form_name),
            ("Generated", generated_at.replace(tzinfo=None)),
            ("Row dimensions", ", ".join(grid.row_dimensions)),
            ("Column dimensions", ", ".join(grid.column_dimensions)),
            ("Data rows", len(grid.rows)),
            ("Data columns", len(grid.columns)),
            ("Data cells", len(grid.rows) * len(grid.columns)),
        ]
        for row_index, (label, value) in enumerate(metadata, start=3):
            sheet.cell(row_index, 1, self._safe_text(label)).font = (
                self._LABEL_FONT
            )
            cell = sheet.cell(row_index, 2, value)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

        pov_start = 3 + len(metadata) + 2
        sheet.cell(pov_start, 1, "Point of View").fill = self._SECTION_FILL
        sheet.cell(pov_start, 1).font = self._LABEL_FONT
        sheet.cell(pov_start, 2).fill = self._SECTION_FILL
        for offset, (dimension, member) in enumerate(
            grid.pov,
            start=1,
        ):
            sheet.cell(
                pov_start + offset,
                1,
                self._safe_text(dimension),
            ).font = self._LABEL_FONT
            sheet.cell(
                pov_start + offset,
                2,
                self._safe_text(member),
            )
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 42

    @staticmethod
    def _size_columns(
        sheet: Any,
        total_columns: int,
        row_dimension_count: int,
    ) -> None:
        for column_index in range(1, total_columns + 1):
            width = 22 if column_index <= row_dimension_count else 15
            dimension: ColumnDimension = sheet.column_dimensions[
                get_column_letter(column_index)
            ]
            dimension.width = width

    @classmethod
    def _excel_value(cls, value: Any) -> Any:
        if value is None:
            return None
        if isinstance(value, bool):
            return value
        if isinstance(value, Decimal):
            return float(value)
        if isinstance(value, (int, float)):
            return value
        normalized = str(value).strip()
        if normalized.casefold() in cls._MISSING_VALUES:
            return None
        try:
            return float(Decimal(normalized.replace(",", "")))
        except InvalidOperation:
            return cls._safe_text(normalized)

    @staticmethod
    def _safe_text(value: Any) -> str:
        text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
        if text.startswith("="):
            return f"'{text}"
        return text
