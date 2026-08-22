"""Build Document 01: Product Introduction and Beginner's Guide."""

from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-01"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / (
    "BISP_EPM_Automation_Document_01_Product_Introduction.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

FONT_REGULAR = Path(r"C:\Windows\Fonts\calibri.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\calibrib.ttf")

NAVY = "0B1F3A"
BLUE = "2F5DE0"
BLUE_DARK = "2047B8"
BLUE_LIGHT = "EDF3FF"
ORANGE = "F05A36"
ORANGE_LIGHT = "FFF2EC"
TEAL = "00866A"
TEAL_LIGHT = "E8F7F2"
GOLD = "B76E00"
GOLD_LIGHT = "FFF4D8"
RED = "B42318"
RED_LIGHT = "FDEDEC"
INK = "19243A"
MUTED = "66748C"
LINE = "D8E1EF"
SURFACE = "F6F8FC"
WHITE = "FFFFFF"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, *, top=100, start=130, bottom=100, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, *, color: str = LINE, size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_DXA:
        raise ValueError(f"Table widths must total {CONTENT_DXA} DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_picture_alt_text(run, description: str) -> None:
    drawing = run._element.find(qn("w:drawing"))
    if drawing is None:
        return
    for node in drawing.iter():
        if node.tag.endswith("docPr"):
            node.set("descr", description)
            node.set("title", description)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    number_run = OxmlElement("w:r")
    number_text = OxmlElement("w:t")
    number_text.text = "1"
    number_run.append(number_text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, number_run, end))


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 18, NAVY, 18, 10),
        ("Heading 2", 14, BLUE_DARK, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title.paragraph_format.space_after = Pt(10)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = rgb(BLUE_DARK)
    subtitle.paragraph_format.space_after = Pt(12)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    for name in ("Kicker", "Lead", "Small", "Table Text", "Table Header", "Quote Example"):
        if name not in styles:
            styles.add_style(name, 1)

    kicker = styles["Kicker"]
    kicker.font.name = "Calibri"
    kicker._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    kicker._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    kicker.font.size = Pt(9)
    kicker.font.bold = True
    kicker.font.color.rgb = rgb(BLUE)
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(3)
    kicker.paragraph_format.keep_with_next = True

    lead = styles["Lead"]
    lead.font.name = "Calibri"
    lead._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    lead._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    lead.font.size = Pt(12.5)
    lead.font.color.rgb = rgb(MUTED)
    lead.paragraph_format.space_after = Pt(12)
    lead.paragraph_format.line_spacing = 1.2

    small = styles["Small"]
    small.font.name = "Calibri"
    small._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    small._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    small.font.size = Pt(9)
    small.font.color.rgb = rgb(MUTED)
    small.paragraph_format.space_after = Pt(4)

    table_text = styles["Table Text"]
    table_text.font.name = "Calibri"
    table_text._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_text._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_text.font.size = Pt(9.5)
    table_text.font.color.rgb = rgb(INK)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.1

    table_header = styles["Table Header"]
    table_header.font.name = "Calibri"
    table_header._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_header._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_header.font.size = Pt(9.5)
    table_header.font.bold = True
    table_header.font.color.rgb = rgb(WHITE)
    table_header.paragraph_format.space_after = Pt(0)

    quote = styles["Quote Example"]
    quote.font.name = "Calibri"
    quote._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    quote._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    quote.font.size = Pt(12)
    quote.font.italic = True
    quote.font.color.rgb = rgb(NAVY)
    quote.paragraph_format.left_indent = Inches(0.3)
    quote.paragraph_format.right_indent = Inches(0.3)
    quote.paragraph_format.space_before = Pt(6)
    quote.paragraph_format.space_after = Pt(6)


def configure_page(section, *, first_page: bool = False) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = first_page


def configure_header_footer(section) -> None:
    for header in (section.header, section.even_page_header):
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(
            "BISP SOLUTIONS  /  ORACLE EPM AUTOMATION PLATFORM"
        )
        set_run_font(run, size=8.5, color=MUTED, bold=True)

    for footer in (section.footer, section.even_page_footer):
        table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
        set_table_geometry(table, [7000, 2360], indent_dxa=0)
        table.cell(0, 0).paragraphs[0].text = (
            "Document 01  |  Product Introduction and Beginner's Guide"
        )
        set_run_font(
            table.cell(0, 0).paragraphs[0].runs[0],
            size=8.5,
            color=MUTED,
        )
        add_page_number(table.cell(0, 1).paragraphs[0])
        for cell in table.rows[0].cells:
            set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            top = OxmlElement("w:top")
            top.set(qn("w:val"), "single")
            top.set(qn("w:sz"), "6")
            top.set(qn("w:color"), LINE)
            borders.append(top)
            tc_pr.append(borders)


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    if number == 4:
        add_page_break(doc)
    kicker = doc.add_paragraph(f"CHAPTER {number}", style="Kicker")
    kicker.paragraph_format.page_break_before = number != 4
    kicker.paragraph_format.space_before = Pt(6)
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_para(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead_run = paragraph.add_run(bold_lead)
        set_run_font(lead_run, bold=True, color=NAVY)
        body = paragraph.add_run(text[len(bold_lead):])
        set_run_font(body)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run)


def restarted_numbering_id(doc: Document) -> int:
    """Create a real Word numbering instance that restarts at one."""
    numbering = doc.part.numbering_part.element
    style_num_id = int(
        doc.styles["List Number"]._element.pPr.numPr.numId.val
    )
    abstract_id = None
    existing_ids: list[int] = []
    for child in numbering:
        if child.tag == qn("w:num"):
            current_id = int(child.get(qn("w:numId")))
            existing_ids.append(current_id)
            if current_id == style_num_id:
                abstract = child.find(qn("w:abstractNumId"))
                abstract_id = int(abstract.get(qn("w:val")))
    if abstract_id is None:
        raise RuntimeError("List Number style is missing its numbering definition")
    num_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def add_numbered(doc: Document, items: list[str]) -> None:
    num_id = restarted_numbering_id(doc)
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num_id
        run = paragraph.add_run(item)
        set_run_font(run)


def add_callout(
    doc: Document,
    label: str,
    text: str,
    *,
    fill: str = BLUE_LIGHT,
    accent: str = BLUE,
) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_border(cell, color=accent, size=10)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(label.upper())
    set_run_font(run, size=9, color=accent, bold=True)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, path: Path, caption: str, alt_text: str, *, width=6.35) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_picture_alt_text(run, alt_text)
    doc.add_paragraph(caption, style="Caption")


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    mark_header_row(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_fill(cell, NAVY)
        set_cell_border(cell, color=NAVY, size=8)
        paragraph = cell.paragraphs[0]
        paragraph.style = "Table Header"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run(header)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if row_index % 2 == 0:
                set_cell_fill(cell, SURFACE)
            set_cell_border(cell, color=LINE, size=6)
            paragraph = cell.paragraphs[0]
            paragraph.style = "Table Text"
            paragraph.add_run(value)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def font(size: int, *, bold: bool = False):
    path = FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size)


def wrap(draw: ImageDraw.ImageDraw, text: str, font_obj, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font_obj)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw, box, title, subtitle, *, fill, title_color="#FFFFFF"):
    x1, y1, x2, y2 = box
    title_font = font(38, bold=True)
    body_font = font(23)
    title_lines = wrap(draw, title, title_font, x2 - x1 - 70)
    body_lines = wrap(draw, subtitle, body_font, x2 - x1 - 70)
    total = len(title_lines) * 48 + 12 + len(body_lines) * 31
    y = y1 + ((y2 - y1) - total) / 2
    for line in title_lines:
        bbox = draw.textbbox((0, 0), line, font=title_font)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=title_font, fill=title_color)
        y += 48
    y += 12
    for line in body_lines:
        bbox = draw.textbbox((0, 0), line, font=body_font)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=body_font, fill=title_color)
        y += 31


def arrow(draw, start, end, *, color=BLUE, width=8):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    for offset in (math.pi * 0.83, -math.pi * 0.83):
        point = (
            end[0] + length * math.cos(angle + offset),
            end[1] + length * math.sin(angle + offset),
        )
        draw.line([end, point], fill=color, width=width)


def make_capability_map(path: Path) -> None:
    image = Image.new("RGB", (2000, 1180), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((50, 45, 1950, 1135), radius=48, fill="#F6F8FC", outline="#D8E1EF", width=4)
    draw.text((105, 85), "ONE PLATFORM, FIVE BUSINESS OUTCOMES", font=font(34, bold=True), fill="#2F5DE0")
    draw.text((105, 142), "The platform brings daily Planning work and Oracle automation into one governed experience.", font=font(27), fill="#66748C")
    colors = ["#2F5DE0", "#00866A", "#B76E00", "#7C3AED", "#F05A36"]
    items = [
        ("PLAN", "Cycles, assignments, approvals and notifications"),
        ("AUTOMATE", "Rules, Pipelines, loads, Data Maps and administration jobs"),
        ("REVIEW", "Live cube data, POV grids and source-to-target comparison"),
        ("UNDERSTAND", "Reports, job evidence, logs and current activity"),
        ("ASSIST", "Guided AI conversations with human approval"),
    ]
    top = 245
    height = 148
    gap = 28
    for index, ((title, subtitle), color) in enumerate(zip(items, colors, strict=True)):
        y1 = top + index * (height + gap)
        y2 = y1 + height
        draw.rounded_rectangle((105, y1, 1895, y2), radius=26, fill="#FFFFFF", outline="#D8E1EF", width=3)
        draw.rounded_rectangle((105, y1, 355, y2), radius=26, fill=color)
        draw.rectangle((330, y1, 355, y2), fill=color)
        draw_centered_text(draw, (115, y1, 345, y2), title, "", fill=color)
        draw.text((410, y1 + 39), subtitle, font=font(29), fill="#19243A")
        draw.text((410, y1 + 87), "Available through role-aware screens and monitored workflows", font=font(22), fill="#66748C")
    image.save(path, quality=95)


def make_execution_flow(path: Path) -> None:
    image = Image.new("RGB", (2100, 1080), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((85, 70), "FROM REQUEST TO TRUSTED RESULT", font=font(38, bold=True), fill="#0B1F3A")
    draw.text((85, 126), "Every state-changing action follows the same controlled path.", font=font(27), fill="#66748C")
    steps = [
        ("1", "Choose", "Select a task, operation or assistant request", "#2F5DE0"),
        ("2", "Verify", "Check role, Oracle connection and live artifacts", "#3977D8"),
        ("3", "Prepare", "Collect only the inputs required for this run", "#00866A"),
        ("4", "Review", "Show the exact target, files, scope and impact", "#B76E00"),
        ("5", "Approve", "Require an authorized human decision", "#7C3AED"),
        ("6", "Execute", "Queue, run, monitor and retain evidence", "#F05A36"),
    ]
    x = 90
    y = 290
    box_w = 290
    box_h = 470
    gap = 46
    for index, (number, title, body, color) in enumerate(steps):
        x1 = x + index * (box_w + gap)
        x2 = x1 + box_w
        draw.rounded_rectangle((x1, y, x2, y + box_h), radius=34, fill="#F6F8FC", outline="#D8E1EF", width=3)
        draw.ellipse((x1 + 92, y + 34, x1 + 198, y + 140), fill=color)
        number_font = font(42, bold=True)
        bbox = draw.textbbox((0, 0), number, font=number_font)
        draw.text((x1 + 145 - (bbox[2]-bbox[0])/2, y + 56), number, font=number_font, fill="#FFFFFF")
        title_font = font(34, bold=True)
        bbox = draw.textbbox((0, 0), title, font=title_font)
        draw.text((x1 + box_w/2 - (bbox[2]-bbox[0])/2, y + 180), title, font=title_font, fill="#19243A")
        lines = wrap(draw, body, font(24), box_w - 58)
        line_y = y + 245
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font(24))
            draw.text((x1 + box_w/2 - (bbox[2]-bbox[0])/2, line_y), line, font=font(24), fill="#66748C")
            line_y += 34
        if index < len(steps) - 1:
            arrow(draw, (x2 + 8, y + box_h/2), (x2 + gap - 8, y + box_h/2), color="#9AACCB", width=7)
    draw.rounded_rectangle((210, 855, 1890, 1000), radius=34, fill="#EDF3FF", outline="#2F5DE0", width=3)
    draw.text((290, 887), "RESULT", font=font(26, bold=True), fill="#2F5DE0")
    draw.text((290, 927), "A monitored Oracle job, durable status, diagnostic evidence and a clear next action", font=font(30, bold=True), fill="#0B1F3A")
    image.save(path, quality=95)


def make_lifecycle(path: Path) -> None:
    image = Image.new("RGB", (2100, 1250), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((90, 70), "A TYPICAL PLANNING CYCLE", font=font(38, bold=True), fill="#0B1F3A")
    draw.text((90, 126), "The platform supports the work around Oracle Planning without replacing Oracle's configured business design.", font=font(26), fill="#66748C")
    stages = [
        ("D-2", "Prepare", "Confirm source files, metadata, mappings and owners", "#2F5DE0"),
        ("D-1", "Open", "Create the cycle, publish tasks and verify readiness", "#3977D8"),
        ("Day 1", "Load", "Import metadata, actuals and planning assumptions", "#00866A"),
        ("Day 2", "Calculate", "Run Business Rules and Oracle Pipelines", "#7C3AED"),
        ("Day 3", "Publish", "Push approved data through Data Maps", "#B76E00"),
        ("Day 4", "Review", "Compare cubes, investigate differences and approve", "#F05A36"),
        ("Close", "Report", "Generate outputs, retain evidence and complete tasks", "#0B1F3A"),
    ]
    center_y = 680
    left = 120
    right = 1980
    draw.line((left, center_y, right, center_y), fill="#CBD5E6", width=12)
    span = (right - left) / (len(stages) - 1)
    for index, (when, title, body, color) in enumerate(stages):
        cx = left + index * span
        draw.ellipse((cx - 33, center_y - 33, cx + 33, center_y + 33), fill=color, outline="#FFFFFF", width=5)
        above = index % 2 == 0
        box_y1 = 260 if above else 770
        box_y2 = 580 if above else 1090
        box_x1 = max(35, cx - 145)
        box_x2 = min(2065, cx + 145)
        draw.rounded_rectangle((box_x1, box_y1, box_x2, box_y2), radius=28, fill="#F6F8FC", outline=color, width=3)
        draw.text((box_x1 + 24, box_y1 + 24), when, font=font(25, bold=True), fill=color)
        draw.text((box_x1 + 24, box_y1 + 66), title, font=font(31, bold=True), fill="#19243A")
        lines = wrap(draw, body, font(21), int(box_x2 - box_x1 - 48))
        ty = box_y1 + 120
        for line in lines:
            draw.text((box_x1 + 24, ty), line, font=font(21), fill="#66748C")
            ty += 30
        connector_end_y = box_y2 if above else box_y1
        draw.line((cx, center_y + (-38 if above else 38), cx, connector_end_y), fill=color, width=4)
    image.save(path, quality=95)


def make_navigation(path: Path) -> None:
    image = Image.new("RGB", (2000, 1260), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((90, 70), "WHERE SHOULD I START?", font=font(38, bold=True), fill="#0B1F3A")
    draw.text((90, 128), "Choose the statement that best describes what you need to do.", font=font(27), fill="#66748C")
    root = (690, 220, 1310, 360)
    draw.rounded_rectangle(root, radius=34, fill="#2F5DE0")
    draw_centered_text(draw, root, "What needs your attention?", "Start with Home for a personalized summary", fill="#2F5DE0")
    options = [
        ((85, 525, 560, 775), "Assigned work", "MY WORK", "Tasks, due dates and readiness", "#00866A"),
        ((585, 525, 1060, 775), "One Oracle action", "OPERATIONS", "Rules, loads, Pipelines, maps and admin jobs", "#F05A36"),
        ((1085, 525, 1560, 775), "Check the numbers", "DATA REVIEW", "Live grids and source-to-target comparison", "#7C3AED"),
        ((1585, 525, 1915, 775), "Need guidance", "EPM ASSISTANT", "Ask, clarify and prepare a governed action", "#B76E00"),
    ]
    for box, title, destination, detail, color in options:
        x1, y1, x2, y2 = box
        arrow(draw, ((root[0]+root[2])//2, root[3]), ((x1+x2)//2, y1-22), color="#9AACCB", width=5)
        draw.rounded_rectangle(box, radius=28, fill="#F6F8FC", outline=color, width=4)
        draw.text((x1+28, y1+28), title, font=font(27, bold=True), fill="#19243A")
        draw.text((x1+28, y1+82), destination, font=font(25, bold=True), fill=color)
        ty = y1 + 132
        for line in wrap(draw, detail, font(21), x2-x1-56):
            draw.text((x1+28, ty), line, font=font(21), fill="#66748C")
            ty += 30
    footer = (250, 935, 1750, 1135)
    draw.rounded_rectangle(footer, radius=34, fill="#EDF3FF", outline="#2F5DE0", width=3)
    draw.text((320, 975), "AFTER EXECUTION", font=font(25, bold=True), fill="#2F5DE0")
    draw.text((320, 1025), "Open Jobs & Activity for status and evidence, or Reports for a downloadable business output.", font=font(29, bold=True), fill="#0B1F3A")
    image.save(path, quality=95)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "capabilities": ASSET_DIR / "capability-map.png",
        "execution": ASSET_DIR / "governed-execution-flow.png",
        "lifecycle": ASSET_DIR / "planning-lifecycle.png",
        "navigation": ASSET_DIR / "navigation-guide.png",
    }
    make_capability_map(assets["capabilities"])
    make_execution_flow(assets["execution"])
    make_lifecycle(assets["lifecycle"])
    make_navigation(assets["navigation"])
    return assets


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    configure_page(section, first_page=True)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(38)
    run = paragraph.add_run()
    run.add_picture(str(LOGO), width=Inches(2.35))
    set_picture_alt_text(run, "BISP Solutions company logo")

    paragraph = doc.add_paragraph("ORACLE EPM AUTOMATION PLATFORM", style="Kicker")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(34)
    paragraph.paragraph_format.space_after = Pt(12)

    paragraph = doc.add_paragraph("Product Introduction", style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph = doc.add_paragraph("and Beginner's Guide", style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragraph = doc.add_paragraph(
        "A plain-language guide to what the platform does, who it serves, "
        "and how it supports a governed Oracle Planning lifecycle.",
        style="Subtitle",
    )
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Inches(0.45)
    paragraph.paragraph_format.right_indent = Inches(0.45)

    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120], indent_dxa=0)
    values = [
        ("DOCUMENT", "01 of the project handbook"),
        ("VERSION", "1.0  |  14 August 2026"),
        ("AUDIENCE", "Business users, consultants and administrators"),
    ]
    for cell, (label, value) in zip(table.rows[0].cells, values, strict=True):
        set_cell_fill(cell, BLUE_LIGHT)
        set_cell_border(cell, color=LINE, size=6)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        set_run_font(run, size=8.5, color=BLUE, bold=True)
        paragraph = cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, size=9.5, color=NAVY, bold=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    add_callout(
        doc,
        "Scope of this volume",
        "This guide documents the current modern application and the shared governed services that support it.",
        fill=ORANGE_LIGHT,
        accent=ORANGE,
    )

    paragraph = doc.add_paragraph("Prepared by BISP Solutions")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.runs[0]
    set_run_font(run, size=10, color=MUTED, bold=True)


def add_front_matter(doc: Document) -> None:
    add_page_break(doc)
    doc.add_paragraph("ABOUT THIS GUIDE", style="Kicker")
    doc.add_paragraph("How to use Document 01", style="Heading 1")
    doc.add_paragraph(
        "This volume is the starting point for the BISP Solutions Oracle EPM Automation Platform documentation. "
        "It explains the product in business language before later documents explore installation, individual operations, data review, administration, the AI assistant, and software architecture.",
        style="Lead",
    )

    add_heading(doc, "Who should read it")
    add_bullets(doc, [
        "A business user who wants to know where to begin.",
        "A finance or Planning team member who wants a simpler way to complete assigned work.",
        "An Oracle EPM consultant evaluating which activities can be automated.",
        "An administrator responsible for access, schedules, jobs, and environment health.",
        "A project sponsor who needs a clear view of scope, control, and business value.",
    ])

    add_heading(doc, "Document control")
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Document", "01 - Product Introduction and Beginner's Guide"],
            ["Product", "BISP Solutions Oracle EPM Automation Platform"],
            ["Version", "1.0"],
            ["Publication date", "14 August 2026"],
            ["Status", "Current implemented baseline"],
            ["Scope", "Modern React application and its shared production backend"],
        ],
        [2700, 6660],
    )

    guide_heading = doc.add_paragraph("Guide map", style="Heading 2")
    guide_heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["Chapter", "What you will learn"],
        [
            ["1. Welcome", "What the platform is and the value it provides."],
            ["2. Why it exists", "Which manual problems and risks it reduces."],
            ["3. People and roles", "Who uses the platform and what each role sees."],
            ["4. Capabilities", "What has been implemented in the modern application."],
            ["5. How it works", "The technology and governed execution journey."],
            ["6. Planning lifecycle", "How the features support a real Planning cycle."],
            ["7. Navigation", "Where a first-time user should go for each need."],
            ["8. Trust boundaries", "Controls, limitations, and Oracle responsibilities."],
            ["9. First visit", "A simple step-by-step introduction."],
            ["10. Glossary", "Plain-language meanings of common terms."],
        ],
        [2200, 7160],
    )
    add_heading(doc, "How the guide is written")
    add_bullets(doc, [
        "Business language appears before technical terms.",
        "Every capability is described by its purpose, not only by its screen name.",
        "Controls and limitations are stated openly so users know when Oracle Planning configuration is still required.",
    ])


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    add_front_matter(doc)

    add_chapter(
        doc,
        1,
        "Welcome to the Oracle EPM Automation Platform",
        "One workspace that helps people complete Planning work, run approved Oracle automation, review results, and understand what happened.",
    )
    add_heading(doc, "The platform in one sentence")
    add_callout(
        doc,
        "Plain-language definition",
        "The platform is a secure web workspace that connects business users to approved Oracle EPM Planning processes and operations without requiring them to navigate many technical screens.",
    )
    add_para(
        doc,
        "Oracle EPM Planning remains the system where the Planning application, cubes, dimensions, forms, rules, integrations, Pipelines, and Data Maps are configured. The BISP platform sits in front of those capabilities and provides a simpler, role-aware way to use them.",
    )
    add_heading(doc, "What changes for the user")
    add_table(
        doc,
        ["Instead of...", "The user can..."],
        [
            ["Remembering where every Oracle job is located", "Open one searchable Operations catalog."],
            ["Typing artifact names from memory", "Choose synchronized Oracle artifacts from live selectors."],
            ["Checking several screens for status", "Follow progress and evidence in Jobs & Activity."],
            ["Repeating file and period setup", "Use guided file choices, periods, and governed defaults."],
            ["Guessing what to do next", "Use Home, My Work, and the EPM Assistant for context."],
        ],
        [4180, 5180],
    )
    add_heading(doc, "What the platform is not")
    add_bullets(doc, [
        "It is not a replacement for the Oracle Planning application design.",
        "It does not invent business rules, mappings, forms, or Pipelines without an approved Oracle design.",
        "It does not bypass Oracle security or the platform's own role permissions.",
        "It is not an uncontrolled bot; state-changing work is validated, reviewed, approved, queued, and audited.",
    ])

    add_chapter(
        doc,
        2,
        "Why the platform was created",
        "Planning teams need speed and simplicity, but finance automation also needs evidence, repeatability, and control.",
    )
    add_heading(doc, "The everyday problem")
    add_para(doc, "A typical Planning cycle involves several different kinds of work. Files must be prepared, metadata and data must be loaded, calculations must run, data may be published to reporting cubes, results must be checked, and reports must be produced. Each step can involve a different Oracle screen, job name, file, period, and responsible person.")
    add_para(doc, "That complexity creates avoidable manual effort. A business user may know the required business outcome but not the exact technical path. An administrator may spend time answering status questions or identifying which job failed. A consultant may repeat the same preparation steps for every cycle.")
    add_heading(doc, "The product response")
    add_bullets(doc, [
        "Bring related Planning activities into one consistent experience.",
        "Show users only the screens and actions allowed for their role.",
        "Discover current Oracle artifacts so selectors stay aligned with the connected application.",
        "Guide the user through required inputs instead of presenting a technical command line.",
        "Require review and approval before an operation changes Oracle state.",
        "Run long work in the background and retain status, diagnostics, and outputs.",
        "Make common tasks available through mouse-driven choices wherever Oracle exposes enough information.",
        "Provide an AI assistant that explains and prepares governed work without bypassing controls.",
    ])
    add_figure(
        doc,
        assets["capabilities"],
        "Figure 1. The five business outcomes provided by the modern platform.",
        "Capability map showing Plan, Automate, Review, Understand, and Assist as five outcomes of the platform.",
    )
    add_callout(
        doc,
        "Design principle",
        "The product should reduce manual effort without hiding important business impact. Simplicity comes from better guidance and automation, not from removing control.",
        fill=TEAL_LIGHT,
        accent=TEAL,
    )

    add_chapter(
        doc,
        3,
        "Who uses the platform",
        "The interface adapts to the user's responsibility. People see the work they need, not every capability in the system.",
    )
    add_heading(doc, "The four business-facing roles")
    add_table(
        doc,
        ["Role", "Primary responsibility", "Typical experience"],
        [
            ["Service Administrator", "Operate and govern the platform.", "Environment health, users, cycles, schedules, Oracle operations, catalog synchronization, jobs, and evidence."],
            ["Power User", "Coordinate and review Planning work.", "Approvals, operational execution, Data Review, reports, activity monitoring, and the assistant."],
            ["User / Planner", "Complete assigned Planning work.", "My Work, approved process actions, Data Review, reports, notifications, and guided assistance."],
            ["Viewer / Executive", "Consume approved business outputs.", "A read-focused Home experience, reports, approved context, notifications, and assistant guidance."],
        ],
        [1900, 2700, 4760],
    )
    add_heading(doc, "Role-aware does not mean separate products")
    add_para(doc, "Every role uses the same platform and the same governed backend. Navigation is filtered by server-enforced permissions. A hidden button is not the security control; the API checks the user's permission again before allowing the action.")
    add_heading(doc, "How the homepage changes")
    add_bullets(doc, [
        "Administrators see environment health, failures, active cycles, and operational priorities.",
        "Power Users see reviews, exceptions, due work, and progress in their scope.",
        "Users and Planners see assigned tasks ordered by urgency and readiness.",
        "Viewers and Executives see approved business context and reporting activity without operational controls.",
    ])
    add_callout(
        doc,
        "Important",
        "Platform roles control the BISP application experience. Oracle EPM permissions still control what the configured Oracle account can read or execute.",
        fill=GOLD_LIGHT,
        accent=GOLD,
    )

    add_chapter(
        doc,
        4,
        "What has been implemented",
        "The modern application combines business workspaces, standalone Oracle operations, data review, reporting, scheduling, activity evidence, and governed AI assistance.",
    )
    add_heading(doc, "Planning workspaces")
    add_table(
        doc,
        ["Workspace", "Purpose"],
        [
            ["Home", "Personalized priorities, cycle progress, quick actions, and recent activity."],
            ["My Work", "Assigned tasks, readiness, due dates, execution links, completion, and submission for approval."],
            ["Planning Cycles", "Open business cycles, define context, publish stages and assignments, and monitor cycle progress."],
            ["Approvals", "Review submitted work, approve it, or return it with guidance."],
            ["Notifications", "Receive task, approval, execution, and operational updates."],
        ],
        [2500, 6860],
    )
    add_heading(doc, "Standalone Oracle operations")
    add_table(
        doc,
        ["Operation", "What it does"],
        [
            ["Metadata Import", "Uploads or reuses a metadata file and runs an existing Planning metadata job, with optional error output and Cube Refresh."],
            ["Planning Data Import", "Uploads or reuses a data file and runs an existing native Planning data-import job."],
            ["Data Integration", "Runs a configured Data Integration for selected periods and modes, using a configured, existing, or newly uploaded file."],
            ["Business Rules", "Runs a deployed Calculation Manager rule with defaults or supplied runtime prompt values."],
            ["Pipelines", "Runs a configured Oracle Pipeline using live variables, stages, periods, and stage-specific file requirements."],
            ["Data Maps", "Pushes Planning data to a target cube with governed clear and member-override controls."],
            ["Substitution Variables", "Reads application- or cube-scoped variables and permits authorized updates or explicit creation."],
            ["Cube Refresh", "Runs an existing saved Planning database refresh job and monitors its terminal result."],
            ["Report Generation", "Reads an approved Planning data slice or supported form export and creates a downloadable Excel workbook."],
        ],
        [2500, 6860],
    )
    add_heading(doc, "Cross-cutting capabilities")
    add_bullets(doc, [
        "Live Oracle catalog synchronization and stale-artifact hiding.",
        "Oracle Inbox discovery, local upload, safe replacement, and configured-file reuse.",
        "Year and period selectors where the operation exposes those inputs.",
        "Read-only preflight and explicit confirmation before execution.",
        "Durable background execution with status polling and duplicate-run protection.",
        "Job evidence, errors, logs, generated artifacts, and searchable activity history.",
        "Schedules for repeatable approved operations.",
        "Optional success and failure email notifications.",
        "Excel-triggered Pipeline execution through the versioned API.",
    ])

    add_chapter(
        doc,
        5,
        "How the platform works",
        "The screen is only the starting point. A controlled backend validates identity, reads live Oracle context, queues the work, monitors it, and records the outcome.",
    )
    add_heading(doc, "Technology in simple language")
    add_table(
        doc,
        ["Layer", "Technology", "Why it is used"],
        [
            ["User experience", "React, TypeScript and Vite", "Provides the responsive, role-aware modern web application."],
            ["Application API", "Python and FastAPI", "Receives secure requests and coordinates business services."],
            ["Oracle connection", "Oracle Planning REST APIs and Requests", "Authenticates, discovers artifacts, starts jobs, and reads results."],
            ["Business records", "PostgreSQL, SQLAlchemy and Alembic", "Stores users, cycles, tasks, schedules, executions, conversations, and auditable state using versioned schemas."],
            ["Background work", "Durable execution worker", "Runs long Oracle jobs outside the browser request and safely claims queued work."],
            ["AI assistance", "LangGraph with Gemini or Groq", "Provides stateful conversations and tool-governed preparation while keeping providers replaceable."],
        ],
        [1900, 2500, 4960],
    )
    add_heading(doc, "One backend, several safe entry points")
    add_para(
        doc,
        "The modern web application, approved schedules, Excel integration, and the EPM Assistant all use the same application services and execution queue. This prevents each interface from inventing a different way to run Oracle work.",
    )
    add_callout(
        doc,
        "Why this matters",
        "A Pipeline started from the web, Excel, a schedule, or an approved AI request receives the same permission checks, monitoring, evidence, and recovery behavior.",
        fill=TEAL_LIGHT,
        accent=TEAL,
    )
    add_heading(doc, "The governed execution journey")
    add_figure(
        doc,
        assets["execution"],
        "Figure 2. A common controlled path is used for state-changing Oracle work.",
        "Six-step governed execution flow: Choose, Verify, Prepare, Review, Approve, and Execute.",
    )
    add_heading(doc, "Why a background worker is important")
    add_para(doc, "Oracle jobs can take seconds or many minutes. The browser should not need to remain responsible for that work. The platform first records an approved request in PostgreSQL. A worker claims it, communicates with Oracle, renews its lease, monitors progress, and records terminal evidence.")
    add_callout(
        doc,
        "Recovery safety",
        "If a worker disappears during an Oracle action, the platform does not blindly repeat the request. The execution is marked for recovery review because Oracle may already have accepted the job.",
        fill=RED_LIGHT,
        accent=RED,
    )

    add_chapter(
        doc,
        6,
        "How it supports a Planning lifecycle",
        "The platform connects business responsibilities with approved Oracle jobs from readiness through reporting and close.",
    )
    add_figure(
        doc,
        assets["lifecycle"],
        "Figure 3. Example sequence from pre-cycle readiness through reporting and completion.",
        "Timeline showing D-2 Prepare, D-1 Open, Day 1 Load, Day 2 Calculate, Day 3 Publish, Day 4 Review, and Close Report.",
    )
    add_heading(doc, "The lifecycle in business language")
    add_numbered(doc, [
        "Prepare the source files, Oracle configuration, owners, dates, and expected business context.",
        "Open the Planning cycle so assigned work becomes visible in My Work.",
        "Load approved metadata, actual data, and planning assumptions through the appropriate Oracle operation.",
        "Run calculations and multi-stage Pipelines that apply the organization's Planning logic.",
        "Publish approved data to reporting or downstream cubes through Data Maps.",
        "Review live data, compare source and target values, investigate differences, and complete approvals.",
        "Generate Excel outputs, retain job evidence, resolve remaining exceptions, and complete the cycle.",
    ])
    add_heading(doc, "Oracle Pipeline versus the platform")
    add_para(doc, "An Oracle Pipeline remains the best place to configure a repeatable multi-stage technical sequence such as integrations, rules, Data Maps, file stages, and Oracle notifications. The BISP platform should not duplicate work that Oracle already owns inside that Pipeline.")
    add_para(doc, "The platform adds the business experience around it: role-aware access, task context, live input discovery, file choice, approval, scheduling, monitoring, evidence, notifications, reporting, and AI-assisted preparation.")
    add_table(
        doc,
        ["Oracle Pipeline owns", "The BISP platform adds"],
        [
            ["The approved technical stage order", "A business-friendly place to find and start it"],
            ["Configured integrations, rules, Data Maps and files", "Live discovery of runtime choices and required inputs"],
            ["Oracle-side execution of each stage", "Role checks, approval, queueing, monitoring and evidence"],
            ["Oracle-native technical notifications when configured", "Platform notifications, task context, reports and audit history"],
        ],
        [4680, 4680],
    )
    add_callout(
        doc,
        "Practical effect",
        "A business user can run an approved Pipeline with the required context while the Oracle consultant keeps the technical design in the system where it belongs.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )

    add_chapter(
        doc,
        7,
        "Finding your way around",
        "A first-time user should not need to understand the entire system. Start with the outcome you need.",
    )
    add_figure(
        doc,
        assets["navigation"],
        "Figure 4. A simple decision guide for choosing the right modern workspace.",
        "Navigation decision diagram routing assigned work to My Work, one Oracle action to Operations, checking numbers to Data Review, and guidance to EPM Assistant.",
    )
    add_heading(doc, "Navigation reference")
    add_table(
        doc,
        ["Area", "Use it when...", "Typical next step"],
        [
            ["Home", "You have just signed in or need priorities.", "Open the first ready item or a role-specific quick action."],
            ["My Work", "A cycle task is assigned to you.", "Review readiness, start the linked action, and complete or submit the task."],
            ["Operations", "You need one independent Oracle service.", "Choose the service, artifact, inputs, review, and execute."],
            ["Data Review", "You need to inspect or compare numbers.", "Choose a live cube, dimensions, members, and load the grid."],
            ["Reports", "You need a downloadable business output.", "Choose the approved definition or supported form export and generate Excel."],
            ["Jobs & Activity", "A run is queued, running, completed, or failed.", "Inspect steps, Oracle IDs, errors, logs, and artifacts."],
            ["EPM Assistant", "You need an explanation or help preparing work.", "Ask in business language, confirm the artifact, provide guided inputs, and approve."],
        ],
        [1800, 3500, 4060],
    )
    add_heading(doc, "A useful rule of thumb")
    add_callout(
        doc,
        "Start small",
        "Use My Work when a task was assigned. Use Operations for one independent service. Use Data Review to check numbers. Use Jobs & Activity after execution. Use the Assistant when you do not know which path fits.",
    )

    add_chapter(
        doc,
        8,
        "Trust, control, and honest limitations",
        "Production automation is useful only when users understand which controls are guaranteed and which work still belongs in Oracle Planning.",
    )
    add_heading(doc, "Controls implemented in the platform")
    add_bullets(doc, [
        "Server-side sign-in, secure password hashing, signed sessions, and CSRF protection.",
        "Four platform roles with API-enforced permissions.",
        "Server-managed Oracle credentials that are not returned to the browser.",
        "Live Oracle catalog synchronization for jobs, rules, maps, cubes, Pipelines, and Data Integrations where supported.",
        "Read-only discovery and preflight before state-changing execution.",
        "Explicit human approval and a visible summary of the target and supplied inputs.",
        "A durable PostgreSQL queue, worker leases, monitored Oracle status, and retained evidence.",
        "Audit attribution for manual, scheduled, API, Excel, and AI-agent requests.",
        "AI tool permissions that cannot exceed the signed-in user's platform permissions.",
    ])
    add_heading(doc, "What still requires Oracle configuration")
    add_table(
        doc,
        ["Configure in Oracle Planning", "Use from the BISP platform"],
        [
            ["Application, cubes, dimensions, members, forms, security and valid intersections", "Discover and use authorized cubes, dimensions, members, and supported form/report data."],
            ["Metadata and data job definitions", "Select, supply files, run, monitor, and retain evidence."],
            ["Import formats, mappings, categories, periods and integration options", "Run the configured Data Integration with business-friendly period, file, and mode choices."],
            ["Calculation Manager rules and their RTP definitions", "Discover rule names, choose defaults or enter known exact RTP names and values, then execute."],
            ["Pipeline stages, variables, file references and technical order", "Inspect the live definition, resolve runtime inputs, approve, run, and monitor."],
            ["Data Map dimensional scope and mapping", "Choose the map and optional supported clear or override controls, then validate results."],
        ],
        [4700, 4660],
    )
    add_heading(doc, "Important limitations")
    add_bullets(doc, [
        "Oracle APIs differ across Cloud and on-premises versions. The platform uses capability detection and supported fallbacks, but an endpoint unavailable in Oracle cannot be created by the UI.",
        "Some Oracle artifacts must already exist before they can be run. The platform focuses on governed execution, not complete Oracle application design.",
        "The Data Review grid is designed for review and reconciliation. It is not a full replacement for every Smart View or Planning-form feature.",
        "The EPM Assistant can misunderstand a request. It must use permitted tools, live choices, guided inputs, and human approval; users should still verify the exact artifact and scope.",
        "The platform cannot guarantee that source business data is correct. It can load, compare, monitor, and expose differences, but business ownership remains essential.",
    ])
    add_heading(doc, "When a capability is not available")
    add_numbered(doc, [
        "Check the connected environment and synchronize the Oracle catalog again.",
        "Confirm that the artifact exists in the current Oracle application and that the configured Oracle account can access it.",
        "Use Oracle Planning for design or configuration work that its public APIs do not expose.",
        "Record the gap as a product improvement only after the Oracle and permission checks are complete.",
    ])

    add_chapter(
        doc,
        9,
        "Your first visit: a seven-step walkthrough",
        "This short tour helps a new user understand the product without running a production operation.",
    )
    add_numbered(doc, [
        "Sign in with the platform username provided by the Service Administrator. This is the BISP platform identity, not a request to expose Oracle credentials in the browser.",
        "Read Home. Confirm your role label, connected Planning application, priorities, active cycle, and recent activity.",
        "Open My Work. Look for a task marked ready. Read its description, dates, stage, entity, and linked action before doing anything.",
        "Open Operations. Browse the service catalog and read the risk badge and description. Do not start a run during this introductory tour.",
        "Open Data Review. Notice that cubes, dimensions, and members are obtained from the connected Oracle application before a grid is loaded.",
        "Open Jobs & Activity. Review how completed or failed work exposes status and evidence. This is where you return after a background operation starts.",
        "Open EPM Assistant. Ask: 'What can I do in this platform?' Use the answer for orientation, but remember that any proposed state-changing action still requires governed inputs and approval.",
    ])
    add_heading(doc, "A safe first assistant question")
    paragraph = doc.add_paragraph(style="Quote Example")
    paragraph.add_run("“Explain the Planning lifecycle and tell me which workspace I should use to review data.”")
    add_para(doc, "This question asks for guidance only. It helps the user understand the environment without preparing or executing an Oracle operation.")
    add_callout(
        doc,
        "Success for a first-time user",
        "After this tour, the user should know where assigned work appears, where independent operations are run, where data is reviewed, where job evidence is stored, and when the Assistant can help.",
        fill=TEAL_LIGHT,
        accent=TEAL,
    )

    add_chapter(
        doc,
        10,
        "Beginner's glossary",
        "These terms appear throughout the application and the remaining project documentation.",
    )
    glossary = [
        ["Application", "The Oracle Planning solution that contains cubes, dimensions, forms, rules, jobs, and security."],
        ["Approval", "An authorized human decision that permits or returns submitted work."],
        ["Artifact", "A named Oracle object such as a rule, job, Pipeline, Data Map, cube, or integration."],
        ["Cube / plan type", "A multidimensional store of Planning data, such as revenue, workforce, or reporting data."],
        ["Cube Refresh", "A Planning administration job that synchronizes application metadata with the underlying cube."],
        ["Data Integration", "A configured process that imports, maps, validates, and exports source data to a Planning target."],
        ["Data Map", "An Oracle definition that moves a selected slice of Planning data from one cube to another."],
        ["Data Push", "The act of publishing Planning data through a Data Map. It is different from importing an external file."],
        ["Dimension", "A business perspective used to organize data, such as Account, Entity, Period, Product, Scenario, Version, or Year."],
        ["EPM", "Enterprise Performance Management: planning, forecasting, consolidation, reporting, and related performance processes."],
        ["Execution", "One recorded attempt to run an operation or scheduled action."],
        ["Inbox", "Oracle's server-side file location used by jobs and integrations."],
        ["Job", "A unit of work accepted and tracked by Oracle Planning."],
        ["Metadata", "The business structures that describe the model, including dimensions, members, aliases, properties, and hierarchies."],
        ["Pipeline", "An Oracle-defined multi-stage sequence that can combine integrations, rules, mappings, files, and other steps."],
        ["Planning cycle", "A governed business period containing stages, assignments, due dates, approvals, and automation work."],
        ["POV", "Point of view: the fixed members that define the context of a data grid, such as Scenario, Version, Entity, and Year."],
        ["RTP", "Runtime prompt: a value supplied when a Business Rule starts, such as Entity, Scenario, or Year."],
        ["Substitution variable", "A named value stored at application or cube scope and reused by forms, rules, calculations, or other Oracle objects."],
        ["Task", "One assigned unit of Planning work with an owner, status, readiness, date, and optional linked action."],
    ]
    add_table(doc, ["Term", "Meaning"], glossary[:10], [2300, 7060])
    second_glossary = doc.add_paragraph(
        "Platform and workflow terms",
        style="Heading 2",
    )
    second_glossary.paragraph_format.page_break_before = True
    add_table(doc, ["Term", "Meaning"], glossary[10:], [2300, 7060])

    continue_heading = doc.add_paragraph("CONTINUE LEARNING", style="Kicker")
    continue_heading.paragraph_format.page_break_before = True
    doc.add_paragraph("What comes next", style="Heading 1")
    doc.add_paragraph(
        "Document 01 provides the common language needed for the rest of the handbook. Later volumes explain each area in practical detail.",
        style="Lead",
    )
    add_table(
        doc,
        ["Next document", "Focus"],
        [
            ["02 - Solution Architecture and Technology Stack", "How the frontend, FastAPI backend, PostgreSQL database, worker, Oracle APIs, and AI providers work together."],
            ["03 - Installation, Configuration, and Startup", "How to prepare an environment and start every required service safely."],
            ["04 - Modern Web Application User Guide", "Screen-by-screen instructions for the role-aware business experience."],
            ["05 - Oracle EPM Automation Operations", "Detailed instructions and prerequisites for every implemented Oracle operation."],
            ["06 - Planning Workspace and Lifecycle", "A complete business cycle from readiness to close."],
            ["07 - Data Review and Reporting", "Live grids, comparison, validation, exports, and reports."],
            ["08 - EPM Assistant and AI Agent", "Conversation, tool use, artifact selection, approvals, execution, safety, and limitations."],
        ],
        [3400, 5960],
    )
    add_callout(
        doc,
        "Key takeaway",
        "The BISP Solutions Oracle EPM Automation Platform does not replace Planning expertise. It turns approved expertise and Oracle configuration into a simpler, repeatable, monitored, and role-aware working experience.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    doc.settings.odd_and_even_pages_header_footer = True
    for index, section in enumerate(doc.sections):
        configure_page(section, first_page=(index == 0))
        configure_header_footer(section)

    properties = doc.core_properties
    properties.title = "BISP Oracle EPM Automation Platform - Product Introduction"
    properties.subject = "Document 01 - Product Introduction and Beginner's Guide"
    properties.author = "BISP Solutions"
    properties.keywords = "Oracle EPM, Planning, automation, beginner guide, BISP Solutions"
    properties.comments = "Documents the current modern application only."

    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
