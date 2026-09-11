"""Build Document 10: Engineering, Testing, Deployment, and Roadmap."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    TEAL,
    TEAL_LIGHT,
    add_bullets,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    font,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    wrap,
)
from docs.builders.build_document_03 import (  # noqa: E402
    add_code_block,
    compact_trailing_empty_paragraph,
    configure_styles,
    patch_list_numbering,
)
from docs.builders.build_document_09 import (  # noqa: E402
    _arrow,
    _card,
    _title,
    add_chapter,
    add_heading,
    add_note,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-10"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = (
    OUTPUT_DIR
    / "BISP_EPM_Automation_Document_10_Engineering_Testing_Deployment_and_Roadmap.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
MUTED_HEX = f"#{MUTED}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"
RED_HEX = f"#{RED}"


def create_extension_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 820), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "One business capability crosses five controlled layers",
        "New Oracle operations are added through stable contracts rather than placing Oracle calls directly in a page or agent prompt.",
    )
    steps = [
        ((45, 170, 275, 480), "Business contract", "Typed request, response, risk, permission, and validation rules.", BLUE_HEX),
        ((325, 170, 555, 480), "Oracle service", "REST or EPM Automate adapter owns transport and Oracle-specific payloads.", ORANGE_HEX),
        ((605, 170, 835, 480), "Application use case", "Preflight, governance, execution payload, evidence, and notification policy.", TEAL_HEX),
        ((885, 170, 1115, 480), "API and React", "Versioned endpoint plus a mouse-friendly workflow for review and submission.", BLUE_HEX),
        ((1165, 170, 1360, 480), "Assistant tool", "Least-privilege discovery, input collection, approval, and durable execution.", NAVY_HEX),
    ]
    for index, (bounds, heading, body, color) in enumerate(steps, 1):
        _card(draw, bounds, heading, body, color=color, label=str(index))
        if index < len(steps):
            _arrow(draw, (bounds[2] + 5, 325), (steps[index][0][0] - 5, 325), color="#9AA9BF", width=4)
    _card(
        draw,
        (220, 585, 1180, 735),
        "Shared execution and evidence",
        "Every write-capable entry point submits normalized work to the same PostgreSQL queue and records the same workflow, actor, Oracle job, and recovery evidence.",
        color=TEAL_HEX,
        fill=f"#{TEAL_LIGHT}",
    )
    image.save(path)


def create_test_pyramid(path: Path) -> None:
    image = Image.new("RGB", (1400, 850), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Testing combines fast deterministic checks with controlled live UAT",
        "Most defects should be found without Oracle or model-provider access; a smaller live layer proves the external environment.",
    )
    levels = [
        ((160, 595, 1240, 735), "Unit and service tests", "Python services, repositories, validation, payloads, security, matching, and React components using mocks and isolated test state.", BLUE_HEX, f"#{BLUE_LIGHT}"),
        ((260, 435, 1140, 565), "Integration and contract tests", "FastAPI routes, PostgreSQL migration smoke test, queue behavior, Excel contract, frontend API behavior, and production build.", TEAL_HEX, f"#{TEAL_LIGHT}"),
        ((390, 275, 1010, 405), "Deterministic agent release evaluation", "Versioned intent-routing and artifact-ranking cases with 100 percent current thresholds and no Oracle or LLM calls.", ORANGE_HEX, f"#{ORANGE_LIGHT}"),
        ((535, 145, 865, 245), "Live environment UAT", "Selected cloud and on-premises Oracle cases plus Gemini or Groq availability, ambiguity, approval, failure, and recovery scenarios.", NAVY_HEX, "#FFFFFF"),
    ]
    for bounds, heading, body, color, fill in levels:
        _card(draw, bounds, heading, body, color=color, fill=fill)
    image.save(path)


def create_release_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 800), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "A release moves through evidence gates",
        "Source, schemas, frontend assets, API processes, workers, and Oracle-facing behavior are promoted as one versioned product.",
    )
    steps = [
        (45, "Review", "Short-lived branch, pull request, and current documentation.", BLUE_HEX),
        (315, "CI gates", "Python 3.11 and 3.13 tests, PostgreSQL migrations, agent evaluation, React tests and build.", ORANGE_HEX),
        (585, "Protect state", "Back up PostgreSQL, stop writers, and reconcile running Oracle work.", ORANGE_HEX),
        (855, "Deploy", "Install tagged code, apply Alembic, initialize LangGraph, build React, start API and workers.", BLUE_HEX),
        (1125, "Verify", "Readiness, environment health, catalog access, queue processing, and a low-risk smoke test.", TEAL_HEX),
    ]
    for index, (x, heading, body, color) in enumerate(steps, 1):
        _card(draw, (x, 175, x + 220, 535), heading, body, color=color, label=f"STEP {index}")
        if index < len(steps):
            _arrow(draw, (x + 225, 355), (x + 255, 355), color="#9AA9BF", width=4)
    draw.text((170, 655), "Release evidence", font=font(18, bold=True), fill=RED_HEX)
    draw.text((340, 651), "Tag, commit, CI result, backup, schema revisions, operator, health checks, smoke test, and rollback decision.", font=font(19, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_topology(path: Path) -> None:
    image = Image.new("RGB", (1400, 880), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Production separates request handling from Oracle execution",
        "The browser remains stateless with respect to long-running jobs; PostgreSQL and controlled runtime storage connect independently supervised processes.",
    )
    _card(draw, (55, 160, 350, 345), "Users and clients", "React browser, approved Excel macro, and controlled external clients over HTTPS.", color=BLUE_HEX, fill=f"#{BLUE_LIGHT}")
    _card(draw, (510, 140, 890, 365), "Reverse proxy and FastAPI", "Same-origin delivery, authentication, authorization, CSRF, validation, API contracts, health endpoints, and queue submission.", color=BLUE_HEX)
    _card(draw, (1050, 160, 1345, 345), "Oracle EPM", "Planning REST endpoints and optional EPM Automate for operations explicitly configured to use it.", color=NAVY_HEX)
    _card(draw, (140, 525, 500, 745), "PostgreSQL", "Application state, durable execution queue, workflow evidence, schedules, identities, catalogs, and LangGraph checkpoints.", color=TEAL_HEX, fill=f"#{TEAL_LIGHT}")
    _card(draw, (570, 525, 930, 745), "One or more workers", "Atomic claims, leases, heartbeats, Oracle submission and monitoring, terminal status, cleanup, and recovery quarantine.", color=ORANGE_HEX, fill=f"#{ORANGE_LIGHT}")
    _card(draw, (1000, 525, 1340, 745), "Runtime storage", "Temporary uploads, generated reports, job evidence, and optional shared storage required by scaled workers.", color=TEAL_HEX)
    _arrow(draw, (355, 250), (505, 250), color="#9AA9BF", width=5)
    _arrow(draw, (895, 250), (1045, 250), color="#9AA9BF", width=5)
    _arrow(draw, (690, 375), (690, 510), color="#9AA9BF", width=5)
    _arrow(draw, (505, 625), (555, 625), color="#9AA9BF", width=5)
    _arrow(draw, (935, 625), (985, 625), color="#9AA9BF", width=5)
    _arrow(draw, (1160, 510), (1160, 365), color="#9AA9BF", width=5)
    image.save(path)


def create_scaling_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 800), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Scale the bottleneck without weakening governance",
        "API replicas and workers can grow independently, but every replica must share identity, database, storage, and execution rules.",
    )
    cards = [
        ((55, 175, 355, 390), "Measure first", "Queue age, execution duration, Oracle throttling, database connections, API latency, worker health, and storage growth.", BLUE_HEX),
        ((400, 175, 700, 390), "Scale API", "Add stateless FastAPI replicas behind one HTTPS origin while keeping one stable session secret.", BLUE_HEX),
        ((745, 175, 1045, 390), "Scale workers", "Add supervised workers; PostgreSQL row locks and SKIP LOCKED distribute different queue items safely.", ORANGE_HEX),
        ((1090, 175, 1345, 390), "Respect Oracle", "Limit concurrency by environment and operation so platform throughput does not become Oracle overload.", RED_HEX),
    ]
    for index, (bounds, heading, body, color) in enumerate(cards, 1):
        _card(draw, bounds, heading, body, color=color, label=str(index))
        if index < len(cards):
            _arrow(draw, (bounds[2] + 5, 280), (cards[index][0][0] - 5, 280), color="#9AA9BF", width=4)
    _card(draw, (170, 535, 1230, 700), "Invariant controls", "Same DATABASE_URL, compatible code and schema, stable WEB_SESSION_SECRET, consistent Oracle environment selection, shared runtime storage where needed, idempotent decisions, leases, and recovery-required quarantine.", color=TEAL_HEX, fill=f"#{TEAL_LIGHT}")
    image.save(path)


def create_roadmap(path: Path) -> None:
    image = Image.new("RGB", (1400, 860), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Improve operational maturity in deliberate stages",
        "The roadmap prioritizes production evidence and user trust before adding broader autonomous behavior.",
    )
    columns = [
        ((55, 160, 440, 735), "Now", "Harden deployment", ["Package API and worker services", "Centralize logs and alerts", "Define backup and retention", "Automate live smoke evidence", "Publish supported environment matrix"], BLUE_HEX, f"#{BLUE_LIGHT}"),
        ((505, 160, 890, 735), "Next", "Broaden quality controls", ["Browser end-to-end tests", "Prometheus or OpenTelemetry metrics", "Load and failure testing", "Automated catalog health monitoring", "Richer agent evaluation and replay"], ORANGE_HEX, f"#{ORANGE_LIGHT}"),
        ((955, 160, 1345, 735), "Later", "Expand safely", ["Additional schedule adapters", "External identity lifecycle integration", "Object storage for artifacts", "Task Manager synchronization when supported", "Policy-controlled multi-step agent plans"], TEAL_HEX, f"#{TEAL_LIGHT}"),
    ]
    for bounds, label, heading, items, color, fill in columns:
        x1, y1, x2, y2 = bounds
        draw.rounded_rectangle(bounds, radius=18, fill=fill, outline=color, width=3)
        draw.text((x1 + 24, y1 + 20), label.upper(), font=font(16, bold=True), fill=color)
        draw.text((x1 + 24, y1 + 56), heading, font=font(24, bold=True), fill=NAVY_HEX)
        y = y1 + 115
        for item in items:
            draw.ellipse((x1 + 26, y + 6, x1 + 38, y + 18), fill=color)
            for line in wrap(draw, item, font(18), x2 - x1 - 75):
                draw.text((x1 + 52, y), line, font=font(18), fill=MUTED_HEX)
                y += 25
            y += 27
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "extension": ASSET_DIR / "extension_flow.png",
        "testing": ASSET_DIR / "testing_pyramid.png",
        "release": ASSET_DIR / "release_flow.png",
        "topology": ASSET_DIR / "production_topology.png",
        "scaling": ASSET_DIR / "scaling_flow.png",
        "roadmap": ASSET_DIR / "roadmap.png",
    }
    create_extension_flow(assets["extension"])
    create_test_pyramid(assets["testing"])
    create_release_flow(assets["release"])
    create_topology(assets["topology"])
    create_scaling_flow(assets["scaling"])
    create_roadmap(assets["roadmap"])
    return assets


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "BISP SOLUTIONS  /  ENGINEERING TESTING DEPLOYMENT AND ROADMAP"
    set_run_font(paragraph.runs[0], size=8.5, color=MUTED, bold=True)
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.columns[0].width = Inches(5.7)
    table.columns[1].width = Inches(0.8)
    left = table.cell(0, 0).paragraphs[0]
    left.text = "Document 10  |  Engineering Testing Deployment and Roadmap"
    set_run_font(left.runs[0], size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(right)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:color"), LINE)
        borders.append(top)
        tc_pr.append(borders)


def add_cover(doc: Document) -> None:
    if LOGO.is_file():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions company logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(28)
    doc.add_paragraph("ENGINEERING AND DELIVERY GUIDE", style="Kicker")
    doc.add_paragraph(
        "Engineering Testing Deployment and Roadmap",
        style="Title",
    )
    doc.add_paragraph(
        "BISP Solutions Oracle EPM Automation Platform",
        style="Subtitle",
    )
    doc.add_paragraph(
        "A practical guide to extending the current platform, protecting contracts, testing without unnecessary Oracle calls, evaluating the governed Assistant, releasing database and application changes together, deploying API and workers, monitoring production, scaling safely, and prioritizing the remaining work.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [[
            "10 of the platform handbook",
            "10 September 2026",
            "Developers, technical leads, release engineers, platform operators, solution architects, security reviewers, and Service Administrators",
        ]],
        [2200, 2100, 5060],
    )
    add_note(
        doc,
        "Scope.",
        "This volume describes the current React 19, FastAPI, PostgreSQL, SQLAlchemy, Alembic, LangGraph, Gemini and Groq provider architecture; the durable execution worker; current CI and release gates; and the known improvement backlog. It documents the modern product only and excludes the retired interface.",
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        paragraph.add_run(
            "BISP Solutions  |  Stable contracts  |  Evidence before automation  |  Safe growth"
        ),
        size=10,
        color=MUTED,
        bold=True,
    )


def finish_document(doc: Document) -> None:
    compact_trailing_empty_paragraph(doc)
    for table in doc.tables:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            header_pr.append(header)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for section in doc.sections:
        for footer in (section.footer, section.even_page_footer):
            for table in footer.tables:
                header_pr = table.rows[0]._tr.get_or_add_trPr()
                if header_pr.find(qn("w:tblHeader")) is None:
                    header = OxmlElement("w:tblHeader")
                    header.set(qn("w:val"), "true")
                    header_pr.append(header)
    properties = doc.core_properties
    properties.title = "Engineering Testing Deployment and Roadmap"
    properties.subject = "Current developer extension testing release deployment monitoring scalability and roadmap guide"
    properties.author = "BISP Solutions"
    properties.keywords = (
        "Oracle EPM, FastAPI, React, PostgreSQL, LangGraph, testing, CI, "
        "deployment, monitoring, scalability, roadmap"
    )
    properties.comments = (
        "Current modern platform implementation as of 10 September 2026; "
        "retired interface excluded."
    )


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    patch_list_numbering(doc)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS VOLUME", style="Kicker")
    doc.add_paragraph("Move from source change to reliable production service", style="Heading 1")
    doc.add_paragraph(
        "This guide follows the lifecycle of a change. It begins with the code boundaries a developer extends, moves through automated and live testing, then explains release, deployment, monitoring, scaling, recovery, and the improvement roadmap. A reader does not need to be a programmer to understand why each control exists.",
        style="Lead",
    )
    add_table(
        doc,
        ["If you are...", "Read first", "What you should be able to do"],
        [
            ["Developer", "Chapters 2 to 8", "Add a capability without bypassing shared validation, permissions, durable execution, or evidence."],
            ["Tester or business reviewer", "Chapters 9 and 10", "Separate deterministic release checks from live Oracle and model-provider UAT."],
            ["Release engineer", "Chapters 11 to 13", "Promote code, PostgreSQL, LangGraph checkpoints, React assets, API, and workers in a safe order."],
            ["Platform operator", "Chapters 14 to 16", "Use health, logs, request IDs, queue state, Oracle evidence, and recovery procedures."],
            ["Product or architecture owner", "Chapters 17 and 18", "Scale the current design and prioritize remaining work by risk and user value."],
        ],
        [2300, 2100, 4960],
    )
    add_heading(doc, "Contents")
    add_table(
        doc,
        ["Part", "Chapters", "Subject"],
        [
            ["I", "1 to 4", "Current product architecture, repository map, extension boundaries, and business contracts"],
            ["II", "5 to 8", "Oracle adapters, API and React integration, Assistant capabilities, and data/schema change"],
            ["III", "9 to 10", "Automated testing, deterministic agent evaluation, and live UAT"],
            ["IV", "11 to 13", "Continuous integration, release engineering, deployment topology, and startup"],
            ["V", "14 to 16", "Monitoring, incidents, performance, recovery, and operational evidence"],
            ["VI", "17 to 18", "Scalability, current limitations, prioritized roadmap, and quick reference"],
        ],
        [1000, 1800, 6560],
    )

    add_chapter(
        doc,
        1,
        "Understand the engineering goal",
        "The platform should make Oracle Planning work easier without hiding risk. Engineering quality therefore means a simple user journey supported by explicit contracts, least privilege, durable execution, and explainable evidence.",
        new_page=False,
    )
    add_heading(doc, "What is being engineered")
    add_table(
        doc,
        ["Layer", "Current responsibility", "System of record"],
        [
            ["React application", "Business-facing workspaces, guided input, review, approvals, jobs, data review, scheduling, and administration.", "Browser state plus FastAPI responses"],
            ["FastAPI application", "Authentication, authorization, validation, API contracts, use-case composition, uploads, and health endpoints.", "Application rules in code"],
            ["PostgreSQL", "Identities, catalogs, schedules, queue state, workflows, Planning work, Assistant history, decisions, and checkpoints.", "Platform control state"],
            ["Durable worker", "Claims approved work, calls Oracle, monitors completion, writes evidence, cleans temporary files, and quarantines uncertainty.", "Execution queue and workflow history"],
            ["Oracle EPM", "Planning artifacts, metadata, cube data, saved jobs, Pipelines, Data Integrations, rules, variables, Data Maps, and Oracle job status.", "Planning business data and configuration"],
            ["AI providers", "Gemini or Groq produces language-model responses and tool calls within platform-controlled limits.", "No authority over permissions or approval"],
        ],
        [1900, 4660, 2800],
    )
    add_note(doc, "Primary rule.", "The user interface and Assistant may simplify a task, but they must not create a second Oracle configuration system or bypass the same validation, permission, approval, queue, and evidence path used by manual operation screens.")

    add_chapter(
        doc,
        2,
        "Read the repository as a set of responsibilities",
        "The project is organized so Oracle transport, business use cases, persistence, web delivery, and the Assistant can change independently while sharing stable models and services.",
    )
    add_heading(doc, "Current source map")
    add_table(
        doc,
        ["Location", "Purpose", "Typical change"],
        [
            ["app/models", "Typed domain and Oracle-facing data structures.", "Add or refine a business contract without importing FastAPI or React concerns."],
            ["app/clients and app/automation", "Planning REST and EPM Automate transport boundaries.", "Add a documented Oracle request or optional command adapter."],
            ["app/services", "Reusable Oracle operations, repositories, catalogs, reports, validation, identity, and notifications.", "Implement one focused capability behind a clear interface."],
            ["app/application", "Use cases spanning services, governance, durable execution, schedules, and evidence.", "Compose preflight, normalized payload, worker execution, and outcome policy."],
            ["app/infrastructure/database", "SQLAlchemy schema, engine, migration checks, and connection policy.", "Add a constrained table or repository dependency."],
            ["app/api/v1 and app/web", "Versioned API plus application composition, security middleware, uploads, and compatibility routes.", "Expose an authorized contract and map errors safely."],
            ["app/agent", "Provider adapters, LangGraph orchestration, tools, matching, checkpoints, persistence, and evaluation.", "Expose a least-privilege capability or improve governed conversation flow."],
            ["frontend/src", "React workspaces, typed API client, shared components, tests, and styles.", "Build a mouse-friendly task journey against API types."],
            ["alembic/versions", "Ordered PostgreSQL application-schema migrations through revision 0020.", "Change durable data shape without ad hoc database edits."],
            ["tests and evaluations", "Deterministic backend, frontend, security, queue, contract, and agent release evidence.", "Protect expected behavior before promotion."],
        ],
        [2200, 3500, 3660],
    )
    add_heading(doc, "Composition roots")
    add_bullets(
        doc,
        [
            "web_main.py loads validated settings, configures logging, and exposes the FastAPI application.",
            "worker.py composes schedule coordination and the durable execution worker separately from HTTP request handling.",
            "main.py remains a command-line composition root that reuses the same services; it is not a separate business implementation.",
            "frontend/src/main.tsx starts the modern React application, while the production build produces static assets for same-origin delivery or a trusted reverse proxy.",
        ],
    )

    add_chapter(
        doc,
        3,
        "Extend one capability through stable layers",
        "A new operation is complete only when its business contract, Oracle adapter, application use case, permission, durable execution, evidence, API, user experience, Assistant behavior, and tests agree.",
    )
    add_figure(
        doc,
        assets["extension"],
        "Figure 10.1 - End-to-end capability extension path.",
        width=6.65,
        alt_text="Five-step extension flow from typed business contract through Oracle service, application use case, API and React workflow, Assistant tool, and shared durable execution evidence.",
    )
    add_heading(doc, "Definition of done for a new operation")
    add_numbered(
        doc,
        [
            "Define typed inputs, allowed values, required fields, safe defaults, risk level, and the permission needed to discover and execute it.",
            "Implement the Oracle call in a focused REST or EPM Automate service, normalize Oracle responses, and raise typed application errors.",
            "Add live catalog or exact-name validation when Oracle supports it; never present stale registered artifacts as currently available.",
            "Create preflight and normalized execution payloads that can be serialized safely into PostgreSQL without passwords or opaque file handles.",
            "Route write-capable work through the durable queue and record actor, trigger source, effective Oracle username, steps, job IDs, counts, artifacts, and terminal result.",
            "Expose a versioned authorized API contract and build a guided React screen with clear selection, review, approval, progress, and failure states.",
            "If the Assistant needs the capability, expose a narrowly scoped tool, add input collection and approval behavior, and reuse the same execution service.",
            "Add backend tests, route and security tests, frontend interaction tests, agent cases where applicable, release notes, and live UAT instructions.",
        ],
    )
    add_code_block(
        doc,
        "from typing import Protocol\n\nclass OracleOperation(Protocol):\n    def preflight(self, request): ...\n    def execute(self, request): ...\n\n# React, API routes, schedules, and the Assistant depend on the\n# application contract; they do not construct Oracle HTTP calls directly.",
        "Code excerpt 10.1. The extension boundary expressed in plain Python.",
    )

    add_chapter(
        doc,
        4,
        "Design contracts for nontechnical users",
        "The most scalable interface is not a large form. It is a small sequence that discovers live choices, asks only for missing business inputs, shows the effect in ordinary language, and then monitors one approved execution.",
    )
    add_heading(doc, "User journey contract")
    add_table(
        doc,
        ["Stage", "User sees", "Platform guarantees"],
        [
            ["Discover", "Current cubes, jobs, Pipelines, integrations, rules, maps, files, dimensions, members, or variables as mouse-selectable choices.", "Live Oracle data is preferred; unavailable registered artifacts are hidden or clearly marked."],
            ["Prepare", "Only inputs required for the selected artifact and environment.", "Names, value types, files, periods, RTPs, scope, and role permissions are validated before approval."],
            ["Review", "Artifact, environment, business effect, files, periods, variable changes, clear behavior, and risk.", "The reviewed payload is normalized and safe to persist."],
            ["Approve", "One explicit approval for the exact action.", "Idempotent decisions prevent double-clicks or retries from creating duplicate submissions."],
            ["Run", "Queued, running, successful, failed, or recovery-required progress.", "A leased worker owns Oracle execution after the browser request ends."],
            ["Understand", "Business result, Oracle messages, counts, output files, error references, and next action.", "Evidence is redacted, correlated, and retained separately from transient UI state."],
        ],
        [1350, 4100, 3910],
    )
    add_heading(doc, "Contract rules")
    add_bullets(
        doc,
        [
            "Use exact Oracle identifiers internally while displaying helpful names and descriptions where available.",
            "Do not ask a user to type information the platform can discover safely from Oracle or an approved catalog.",
            "Do not invent defaults for destructive or environment-wide choices; use Oracle defaults only when that behavior is explicit.",
            "Separate similar business concepts such as Data Push through a Data Map, file-based Planning Data Import, and Data Integration.",
            "Keep source files and runtime values attached to the reviewed run so users do not re-enter them on a second screen.",
            "Return a request reference and execution ID that support teams can use without exposing secrets or internal paths.",
        ],
    )

    add_chapter(
        doc,
        5,
        "Add Oracle REST and EPM Automate adapters",
        "Oracle integration belongs behind reusable services. This protects the rest of the product from authentication details, endpoint differences, response shapes, command syntax, and deployment-specific behavior.",
    )
    add_heading(doc, "Current adapter responsibilities")
    add_table(
        doc,
        ["Adapter", "Use", "Engineering rule"],
        [
            ["EPMClient", "Reusable authenticated Planning REST session, JSON and binary requests, response handling, and transport exceptions.", "Centralize base URL, timeout, SSL verification, and authentication; never repeat requests blindly after an uncertain write."],
            ["JobService and JobMonitor", "Normalize submission responses, poll terminal status, and collect available diagnostics.", "A timeout is an uncertain outcome requiring reconciliation, not proof that Oracle rejected the job."],
            ["Focused REST services", "Metadata, data, integrations, Pipelines, Business Rules, Data Maps, variables, cubes, forms, files, and reports.", "Build only documented payload fields and return typed results."],
            ["EPMAutomateRunner", "Safe subprocess execution for operations explicitly configured to use EPM Automate.", "Use argument arrays without a command shell, enforce timeouts, and never log credential-bearing arguments."],
            ["EPMAutomateClient and services", "Encrypted password-file login, command execution, replacement uploads, and logout.", "Treat executable availability and encrypted password file as deployment prerequisites, not browser inputs."],
        ],
        [2100, 4300, 2960],
    )
    add_heading(doc, "Cloud and on-premises compatibility")
    add_para(doc, "Settings accept cloud, on-premises, or automatic deployment-mode resolution. Compatibility code may select documented endpoint variants, but capability discovery remains the deciding evidence. A feature must handle missing endpoints and different response shapes with a useful message instead of assuming every Planning environment exposes the same public REST resources.")
    add_note(doc, "Safe retry rule.", "Read-only discovery may be retried within bounded transport policy. A write that may have reached Oracle must be reconciled by Oracle job evidence before another submission is authorized.")

    add_chapter(
        doc,
        6,
        "Expose a versioned FastAPI contract",
        "FastAPI is the server boundary for the React product, approved Excel integration, and external clients. Pydantic schemas validate transport data while application services remain independent of HTTP.",
    )
    add_heading(doc, "API implementation rules")
    add_table(
        doc,
        ["Concern", "Current pattern", "Required extension behavior"],
        [
            ["Versioning", "The current product API is grouped under /api/v1; FastAPI also exposes OpenAPI from typed routes.", "Add new product contracts under the versioned router and preserve compatibility deliberately."],
            ["Authentication", "Signed browser sessions, optional Oracle OIDC, local recovery accounts, or scoped bearer tokens for approved external contracts.", "Resolve one active platform identity before processing protected input."],
            ["Authorization", "Method and path permissions plus service-level ownership and business checks.", "Enforce permissions on the server even when navigation hides the feature."],
            ["CSRF", "State-changing browser requests require the current session CSRF token.", "Use the shared API client rather than custom fetch logic that omits protection."],
            ["Validation", "Pydantic request schemas plus domain normalization and live preflight.", "Reject unsupported keys, unsafe files, unknown artifacts, invalid scopes, and stale selections before queuing."],
            ["Errors", "Safe user message, appropriate HTTP status, request ID, and redacted server log context.", "Do not return passwords, tokens, local paths, raw environment data, or unbounded Oracle payloads."],
        ],
        [1650, 3650, 4060],
    )
    add_code_block(
        doc,
        "@router.post(\"/operations/{operation_code}/runs\")\nasync def create_run(request: Request, payload: OperationRequest):\n    user = require_api_session(request)\n    require_permission(request, user)\n    reviewed = service.preflight(payload, actor=user)\n    return queue.submit(reviewed)",
        "Code excerpt 10.2. Representative route shape; authorization and application services remain explicit.",
    )

    add_chapter(
        doc,
        7,
        "Build the React workflow against shared types",
        "The modern frontend uses React 19, TypeScript 7, Vite 8, and a shared API client. Pages guide the user; they do not own Oracle business rules or durable execution state.",
    )
    add_heading(doc, "Frontend boundaries")
    add_table(
        doc,
        ["Frontend element", "Responsibility", "Must not do"],
        [
            ["Typed API client", "Send credentials and CSRF correctly, parse typed responses, and surface request references.", "Embed Oracle credentials or duplicate endpoint-specific error parsing in every component."],
            ["Workspace component", "Load live choices, capture business inputs, show guidance, and advance the user through prepare, review, run, and result.", "Decide permissions, fabricate catalog results, or infer Oracle success from HTTP acceptance."],
            ["Reusable selectors", "Provide cube, dimension, member, file, artifact, period, and variable selection with loading, empty, and error states.", "Force keyboard entry when a safe live choice exists."],
            ["Execution monitor", "Poll one execution ID and render queued, running, terminal, and recovery-required evidence.", "Keep the browser request open for the duration of the Oracle job."],
            ["Tests", "Validate rendering, user interaction, requests, error states, and accessible labels in jsdom.", "Depend on a live Oracle environment for routine component tests."],
        ],
        [1900, 3900, 3560],
    )
    add_heading(doc, "Frontend verification commands")
    add_code_block(
        doc,
        "Set-Location frontend\npnpm install --frozen-lockfile\npnpm test\npnpm build\n# or run both gates\npnpm verify",
        "Code excerpt 10.3. Current React test and production-build gate.",
    )
    add_note(doc, "Current test boundary.", "The repository contains Vitest and Testing Library coverage for the application shell, API client, access control, environment setup, Assistant, feedback, and scheduling. Browser-level end-to-end automation is a roadmap item, not a current guarantee.")

    add_chapter(
        doc,
        8,
        "Extend the governed EPM Assistant",
        "LangGraph coordinates conversation state, live inspection tools, clarification, input collection, approval, and execution handoff. Gemini and Groq are provider adapters; neither provider receives authority to bypass platform policy.",
    )
    add_heading(doc, "Assistant extension sequence")
    add_numbered(
        doc,
        [
            "Add or reuse an application capability that already enforces live validation and authorization.",
            "Define one narrowly scoped tool with a clear name, description, bounded schema, and safe result projection.",
            "Expose the tool only for intents and roles that need it; least-privilege tool selection reduces ambiguity and prompt-injection reach.",
            "Use deterministic matching to recommend live Oracle artifacts, but require user choice when the result is not clearly dominant.",
            "Extract obvious values from the conversation, then ask only for missing required fields through a structured input card.",
            "Create an action draft, run live preflight again, display the exact effect, and interrupt the graph for explicit approval.",
            "On approval, reserve an idempotent decision and submit the same normalized operation to the durable execution path.",
            "Persist readable conversation history separately from LangGraph checkpoints and retain execution evidence independently of the conversation.",
        ],
    )
    add_heading(doc, "Provider-neutral boundary")
    add_table(
        doc,
        ["Component", "Owns", "Does not own"],
        [
            ["GeminiAgentProvider or GroqAgentProvider", "Message conversion, tool definitions, provider SDK request, token limits, and provider error translation.", "Roles, permissions, Oracle credentials, approval policy, or execution."],
            ["AgentIntentRouter", "Deterministic intent and minimum tool exposure.", "Oracle artifact truth or final user decision."],
            ["AgentCapabilityGateway", "Allowlisted read tools and governed action preparation.", "Arbitrary Python, SQL, shell, or unrestricted HTTP access."],
            ["AgentGraphOrchestrator", "Nodes, interrupts, resume behavior, tool rounds, clarification, inputs, and approval transitions.", "Direct Oracle submission outside the approved application service."],
            ["SQLAgentRepository and PostgresSaver", "Readable product history and framework checkpoint state.", "Planning cube data or model-provider secrets."],
        ],
        [2500, 3450, 3410],
    )
    add_note(doc, "Model switchability.", "The provider protocol allows a future provider to replace Gemini or Groq without changing Oracle services, permissions, approval records, or the execution queue. A new adapter still requires provider-specific tests and live UAT.")

    add_chapter(
        doc,
        9,
        "Use a testing pyramid",
        "Fast deterministic tests protect the majority of behavior. PostgreSQL and contract tests protect integration boundaries. A smaller live UAT layer proves Oracle and model-provider behavior that mocks cannot guarantee.",
    )
    add_figure(
        doc,
        assets["testing"],
        "Figure 10.2 - Current testing pyramid and live UAT boundary.",
        width=6.65,
        alt_text="Testing pyramid with unit and service tests at the base, integration and contract tests, deterministic agent evaluation, and live Oracle and model provider UAT at the top.",
    )
    add_heading(doc, "Backend test domains")
    add_table(
        doc,
        ["Domain", "Examples protected today", "External dependency"],
        [
            ["Oracle services", "Payload construction, response normalization, file replacement, polling, diagnostics, timeouts, and engine-specific behavior.", "Mocked HTTP sessions or process runners"],
            ["Database and queue", "Schema architecture, repositories, transactions, claims, leases, heartbeats, terminal states, and recovery rules.", "Isolated test database compatibility; clean PostgreSQL migration in CI"],
            ["Security and identity", "Role model, sessions, CSRF, API tokens, Oracle password validation, OIDC and federated provisioning boundaries.", "Mocked identity sources"],
            ["Planning work", "Cycles, tasks, approvals, validation evidence, notifications, scheduling, and execution links.", "No live Oracle required"],
            ["Assistant", "Intent routing, tool exposure, matching, interrupts, input resume, approval, idempotency, scheduling, and provider error handling.", "Provider SDKs mocked for routine tests"],
            ["Reports and data review", "Grid selection, comparison tolerance, numeric normalization, workbook rendering, and artifact evidence.", "Synthetic form and data-slice responses"],
        ],
        [1800, 4950, 2610],
    )
    add_heading(doc, "Test commands")
    add_code_block(
        doc,
        "python -m pip install -r requirements-dev.txt\npython -m pytest\n\nSet-Location frontend\npnpm install --frozen-lockfile\npnpm verify",
        "Code excerpt 10.4. Local release-oriented test commands.",
    )
    add_note(doc, "Live-data safety.", "Routine automated tests must never call a customer Oracle environment or include real credentials, exported metadata, member data, uploaded files, or production prompt history. Use representative synthetic fixtures.")

    add_chapter(
        doc,
        10,
        "Evaluate the Assistant before release",
        "Agent quality needs two controls: deterministic evaluation of application-owned behavior and recorded live UAT for provider and Oracle behavior. Passing only one is not enough.",
    )
    add_heading(doc, "Deterministic release evaluation")
    add_table(
        doc,
        ["Property", "Current implementation"],
        [
            ["Suite", "evaluations/agent/release_v1.json with versioned case IDs, category, severity, kind, prompt, and expected result."],
            ["Supported case kinds", "intent_route and artifact_ranking."],
            ["What it protects", "Least-privilege intent routing, tool exposure, governed operation preparation, data review and evidence routing, prompt-injection boundaries, and conservative artifact ranking."],
            ["External calls", "None. It does not call Oracle, Gemini, or Groq and consumes no model tokens."],
            ["Threshold", "100 percent of critical cases and 100 percent of the current suite."],
            ["Evidence", "Machine-readable JSON report published by CI for thirty days."],
        ],
        [2400, 6960],
    )
    add_code_block(
        doc,
        "python -m app.agent.evaluation `\n  --suite evaluations/agent/release_v1.json `\n  --output reports/agent-evaluation.json `\n  --fail-on-threshold",
        "Code excerpt 10.5. Current deterministic Assistant release gate.",
    )
    add_heading(doc, "Live UAT matrix")
    add_table(
        doc,
        ["Area", "Minimum live checks"],
        [
            ["Providers", "Claimed Gemini and Groq model, invalid key, quota, timeout, oversized context, tool-call conversion, and readable error behavior."],
            ["Oracle environments", "Each supported cloud and on-premises version, authentication path, application selection, catalog discovery, job submission, polling, and failure evidence."],
            ["Governance", "Permission denial, cross-user conversation isolation, ambiguous artifacts, required inputs, approve, reject, cancel, retry, and interrupted resume."],
            ["Operations", "Representative Business Rule, Data Map, Pipeline, Data Integration, data import, metadata import, Cube Refresh, variables, data review, and reporting scenarios enabled for the release."],
            ["Security", "Prompt injection, hidden instructions in artifact text, unsafe file references, unsupported destructive requests, duplicate approvals, and secret redaction."],
        ],
        [2100, 7260],
    )

    add_chapter(
        doc,
        11,
        "Use continuous integration as a release gate",
        "The GitHub Actions workflow treats backend behavior, database installation, Agent evaluation, frontend tests, and the production build as one release contract.",
    )
    add_heading(doc, "Current CI jobs")
    add_table(
        doc,
        ["Job", "Environment", "Gate"],
        [
            ["Backend tests", "Ubuntu with Python 3.11 and 3.13", "Install development dependencies and run the complete pytest suite."],
            ["PostgreSQL migration smoke test", "PostgreSQL 16 with Python 3.13", "Upgrade a clean database to Alembic head, verify all heads, and initialize LangGraph checkpoints."],
            ["Agent release evaluation", "Python 3.13", "Run deterministic suite, fail on thresholds, and publish the JSON report even when the gate fails."],
            ["Frontend", "Node 22.13.1 and pnpm 11.19", "Install from the frozen lockfile, run Vitest, run strict TypeScript compilation, and create the Vite production bundle."],
        ],
        [2500, 2600, 4260],
    )
    add_heading(doc, "Branch and dependency policy")
    add_bullets(
        doc,
        [
            "Keep main protected and releasable; use short-lived feature branches and immutable release tags.",
            "Represent development, test, and production through configuration and secrets, not different long-lived source branches.",
            "Use requirements files and the pnpm lockfile as reviewed dependency inputs; never install an unreviewed production dependency during startup.",
            "Review dependency changes for license, security, Python and Node compatibility, provider SDK behavior, and transitive changes to serialization or networking.",
            "Do not place database passwords, Oracle credentials, AI keys, SMTP passwords, or EPM Automate encrypted-password files in workflow YAML or committed configuration.",
        ],
    )
    add_note(doc, "Current automation boundary.", "CI proves deterministic code, schema installation, and frontend build quality. It does not prove access to a customer's Oracle environment or a provider account; that remains recorded live UAT.")

    add_chapter(
        doc,
        12,
        "Release code and schemas together",
        "A release is not only Python code. It includes the React bundle, PostgreSQL application schema, LangGraph checkpoint schema, runtime configuration, API service, workers, and the behavior expected from Oracle integrations.",
    )
    add_figure(
        doc,
        assets["release"],
        "Figure 10.3 - Controlled release and deployment sequence.",
        width=6.65,
        alt_text="Five-step release flow covering review, continuous integration gates, database backup and stopped writers, schema and application deployment, and production verification with recorded evidence.",
    )
    add_heading(doc, "Deployment order")
    add_numbered(
        doc,
        [
            "Confirm the intended commit is on main, release checks passed, and the release is identified by an immutable tag.",
            "Back up PostgreSQL and preserve the approved secret-manager or environment configuration version.",
            "Stop workers from claiming new work; allow running Oracle operations to finish or record them for controlled reconciliation.",
            "Stop old API and worker services, install production dependencies from the tagged source, and keep generated runtime data outside the release directory.",
            "Run Alembic upgrade, verify all application-schema heads, and initialize or upgrade LangGraph checkpoint tables.",
            "Build the React bundle from the frozen lockfile.",
            "Start FastAPI in web mode and one or more separately supervised worker processes using the same database, runtime storage, Oracle context, and compatible code version.",
            "Verify liveness, readiness, administrator sign-in, environment selection, live catalog access, queue processing, and one approved low-risk smoke operation.",
        ],
    )
    add_code_block(
        doc,
        "python -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup\n\nSet-Location frontend\npnpm install --frozen-lockfile\npnpm build",
        "Code excerpt 10.6. Database, checkpoint, and frontend release preparation.",
    )
    add_heading(doc, "Rollback rule")
    add_para(doc, "Prefer a forward fix. If code must be rolled back, stop writers, prove the earlier code understands the current schema, and restore the pre-deployment PostgreSQL backup when a database rollback is necessary. Do not run an untested Alembic downgrade in production. Reconcile every Oracle job submitted near the failure boundary before authorizing another run.")

    add_chapter(
        doc,
        13,
        "Deploy API and workers as separate services",
        "Production separates fast browser requests from long-running Oracle operations. The API validates and queues work; supervised workers perform and monitor Oracle execution after the browser can safely disconnect.",
    )
    add_figure(
        doc,
        assets["topology"],
        "Figure 10.4 - Recommended current production topology.",
        width=6.65,
        alt_text="Production topology showing users through HTTPS to reverse proxy and FastAPI, PostgreSQL and runtime storage, separately supervised durable workers, and Oracle EPM.",
    )
    add_heading(doc, "Supported operating modes")
    add_table(
        doc,
        ["Mode", "Use", "Boundary"],
        [
            ["embedded", "Local development and simple testing where API and background execution share one process.", "Convenient but not the recommended production failure boundary."],
            ["web", "FastAPI request-serving process that validates and queues work.", "Does not claim durable Oracle work in the production split."],
            ["worker", "Dedicated worker process started through worker.py.", "Claims queue items, runs schedules, calls Oracle, monitors jobs, and persists outcomes."],
        ],
        [1600, 4000, 3760],
    )
    add_code_block(
        doc,
        "$env:EPM_EXECUTION_RUNTIME = \"web\"\npython -m uvicorn web_main:app --host 0.0.0.0 --port 8080\n\n# Start separately under a service supervisor\npython worker.py",
        "Code excerpt 10.7. Separate production API and worker entry points.",
    )
    add_heading(doc, "Deployment prerequisites")
    add_bullets(
        doc,
        [
            "One trusted HTTPS origin for the React bundle and API, or an exact tested CORS allowlist if separate origins become necessary.",
            "A stable WEB_SESSION_SECRET shared by API replicas and secure cookies enabled when browsers use HTTPS.",
            "One PostgreSQL database reachable by API and every worker, with encrypted transport and backup policy appropriate to the hosting environment.",
            "Consistent Oracle base URL, application selection, integration identity, deployment mode, SSL policy, and timeouts across the processes that cooperate on a run.",
            "Controlled runtime storage; multiple hosts need shared or object-backed storage for uploads and generated artifacts that another process may consume.",
            "A service supervisor or orchestrator that restarts crashed processes without automatically resubmitting recovery-required Oracle work.",
        ],
    )

    add_chapter(
        doc,
        14,
        "Monitor health logs queue and Oracle evidence",
        "Operational visibility begins with process health and request correlation, but a true business outcome comes from durable workflow and Oracle evidence rather than an HTTP success code or notification.",
    )
    add_heading(doc, "Health endpoints")
    add_table(
        doc,
        ["Endpoint", "Meaning", "Use"],
        [
            ["GET /health/live", "The API process can answer HTTP requests.", "Process liveness probe; repeated failure may justify restart."],
            ["GET /health/ready", "The API and PostgreSQL are available and the application can accept normal platform work.", "Load-balancer readiness; intentionally does not call Oracle."],
            ["Dashboard environment health", "Signed-in view of selected Oracle application and available platform checks.", "Operator diagnosis; Oracle outage should not force an API restart loop."],
        ],
        [2000, 4000, 3360],
    )
    add_heading(doc, "Current logging contract")
    add_para(doc, "The application writes timestamp, level, logger name, request correlation ID, and message to standard output. HTTP middleware records method, path, response status, duration, and the same X-Request-ID returned to the client. Production hosting should collect standard output into a protected searchable log service.")
    add_code_block(
        doc,
        "2026-09-10 12:15:04 | INFO | oracle_planning_automation.web |\nrequest=9f2a... | HTTP response: POST /api/v1/... status=202 duration_ms=84",
        "Code excerpt 10.8. Representative correlation-friendly log shape.",
    )
    add_heading(doc, "Minimum production signals")
    add_table(
        doc,
        ["Signal", "Question answered", "Alert example"],
        [
            ["API request rate latency and status", "Can users reach the product and which contracts are failing?", "Sustained 5xx rate or latency above service objective."],
            ["Database connectivity and pool pressure", "Can API and workers read and persist shared state?", "Connection timeouts, pool exhaustion, or schema-head mismatch."],
            ["Queue depth and oldest age", "Is approved work being claimed promptly?", "Queued work older than the expected operating window."],
            ["Worker heartbeat and lease expiry", "Are workers alive and maintaining ownership?", "No active workers or rising RECOVERY_REQUIRED records."],
            ["Oracle submission and duration", "Which environments, operations, or jobs are slow or failing?", "Repeated authentication, throttling, timeout, or terminal failure."],
            ["Assistant provider failures", "Are Gemini or Groq requests within token, quota, and timeout limits?", "Repeated provider errors or abnormal fallback volume."],
            ["Runtime storage and report growth", "Can uploads, evidence, and downloads complete safely?", "Capacity threshold or orphaned-file growth."],
        ],
        [2300, 4000, 3060],
    )
    add_note(doc, "Current gap.", "Structured request IDs and durable execution evidence are implemented. A Prometheus or OpenTelemetry metrics exporter, distributed tracing, dashboards, and alert definitions are not yet packaged in the repository and belong to the near-term roadmap.")

    add_chapter(
        doc,
        15,
        "Investigate and recover without duplicate Oracle work",
        "The most important incident rule is to separate a definite failure from an uncertain Oracle outcome. A timed-out request or expired worker lease may mean Oracle accepted the work even though the platform did not observe the result.",
    )
    add_heading(doc, "Evidence sequence")
    add_numbered(
        doc,
        [
            "Start with the user-visible Reference or X-Request-ID and locate the API log event.",
            "Find the execution ID, initiating user, trigger source, effective Oracle username, operation target, and queue record.",
            "Review ordered workflow steps, last successful boundary, Oracle job ID, returned messages, counts, generated artifacts, and safe terminal error.",
            "If the platform reports RECOVERY_REQUIRED or an Oracle submission timeout, inspect Oracle Job Console using the same time, user, job type, artifact, and application.",
            "Record whether Oracle ran, did not run, or remains uncertain. Never infer non-execution solely from a disconnected worker or HTTP error.",
            "Authorize a new run only after the earlier outcome is reconciled and the original cause, inputs, artifact state, and file references are revalidated.",
        ],
    )
    add_heading(doc, "Common incident decisions")
    add_table(
        doc,
        ["Condition", "Interpretation", "Action"],
        [
            ["HTTP 4xx before queue submission", "The platform rejected authentication, authorization, CSRF, schema, artifact, or input.", "Correct the request; no Oracle retry is needed unless separate evidence shows a submission."],
            ["Queue remains QUEUED", "No eligible worker has claimed the item or workers cannot reach shared state.", "Check worker health, database, storage, compatible release, and target concurrency policy."],
            ["Workflow FAILED with Oracle terminal status", "The platform observed a definite Oracle failure.", "Resolve the Oracle cause, validate changes, and create a new approved execution."],
            ["Worker lease expired", "Oracle outcome may be unknown.", "Quarantine as RECOVERY_REQUIRED and reconcile with Oracle before retry."],
            ["Notification delivery failed", "Communication failed; Oracle job status is independent.", "Use workflow and Oracle evidence as truth, then repair notification configuration."],
            ["Agent provider failed before approval", "No governed Oracle action was approved or queued.", "Review provider key, model, quota, context size, and network; retry the conversation request safely."],
        ],
        [2250, 3600, 3510],
    )
    add_note(doc, "Data preservation.", "Do not repair governance state with direct SQL updates or delete queue, workflow, approval, or decision records during an incident. Use a supported recovery service or a reviewed migration with backup and audit evidence.")

    add_chapter(
        doc,
        16,
        "Manage performance and capacity",
        "Most platform work is I/O-bound and depends on Oracle job duration. Performance tuning should reduce avoidable waiting and contention without increasing Oracle concurrency beyond the customer's safe operating limit.",
    )
    add_heading(doc, "Primary capacity controls")
    add_table(
        doc,
        ["Control", "Current behavior", "Tuning guidance"],
        [
            ["HTTP timeout", "EPM_REQUEST_TIMEOUT defaults to 30 seconds for individual REST requests.", "Increase only for a proven endpoint need; long jobs should submit and then poll rather than hold one request."],
            ["Oracle polling", "DEFAULT_POLL_INTERVAL defaults to 5 seconds and DEFAULT_JOB_TIMEOUT to 1800 seconds.", "Balance timely status with Oracle request volume; use operation-specific evidence before changing globally."],
            ["Worker polling", "EXECUTION_WORKER_POLL_INTERVAL defaults to 2 seconds.", "Reduce only when database load and empty-poll cost are understood."],
            ["Worker lease", "EXECUTION_LEASE_SECONDS defaults to 120 seconds and cannot be below 30 seconds.", "Keep the lease comfortably above heartbeat jitter and short process stalls."],
            ["Schedule polling", "SCHEDULE_POLL_INTERVAL defaults to 15 seconds.", "Match the precision users need; monthly and daily work rarely needs sub-second polling."],
            ["Database pool", "SQLAlchemy validates pooled connections and supports bounded pool plus overflow settings in the engine.", "Size for API replicas and workers together, leaving capacity for migrations and administration."],
            ["Agent context", "History, tool rounds, Groq input tokens, and completion tokens are bounded by settings.", "Prefer concise safe context and targeted tools over simply raising provider limits."],
            ["Uploads", "Browser uploads are size-limited and session-owned; temporary paths are validated under runtime storage.", "Monitor capacity, set hosting request limits, and move to shared or object storage before multi-host scaling."],
        ],
        [2050, 4000, 3310],
    )
    add_heading(doc, "Performance test plan")
    add_bullets(
        doc,
        [
            "Measure dashboard and catalog latency with Oracle available, slow, and unavailable.",
            "Submit concurrent read-only requests and approved queue items without exceeding Oracle environment concurrency limits.",
            "Measure queue claim time, database pool usage, worker throughput, job duration, and evidence-write latency.",
            "Exercise large but allowed file uploads, generated reports, data grids, and reconciliation exports while monitoring memory and storage.",
            "Terminate a worker during submission, polling, and cleanup to prove lease and recovery behavior.",
            "Run provider quota, token-limit, timeout, and network-failure cases without losing the pending governed state.",
        ],
    )

    add_chapter(
        doc,
        17,
        "Scale the current architecture safely",
        "The architecture supports separate API and worker scaling because durable state lives in PostgreSQL. Production scale still requires shared storage, connection budgeting, Oracle-aware concurrency, and compatible releases across every replica.",
    )
    add_figure(
        doc,
        assets["scaling"],
        "Figure 10.5 - Safe scaling sequence and invariants.",
        width=6.65,
        alt_text="Four-step scaling sequence: measure first, add stateless API replicas, add durable workers, respect Oracle capacity, while keeping database, secrets, storage, schema, and recovery controls consistent.",
    )
    add_heading(doc, "Scaling decisions")
    add_table(
        doc,
        ["Need", "Recommended response", "Required safeguard"],
        [
            ["More concurrent browser users", "Add FastAPI replicas behind one HTTPS load balancer and serve the same React build.", "Stable session secret, same database, compatible schema, exact origin policy, and shared environment selection."],
            ["More independent Oracle jobs", "Add supervised worker processes so SKIP LOCKED claims different queue rows.", "Environment and target concurrency limits, adequate database pool, shared runtime storage, and recovery monitoring."],
            ["Large reports and uploads", "Move runtime artifacts to shared storage or a provider-neutral object-storage adapter.", "Path isolation, signed download authorization, encryption, lifecycle policy, malware scanning as required, and evidence links."],
            ["High read traffic", "Cache safe discovery results briefly or precompute summaries where staleness is acceptable.", "Environment-scoped keys, explicit freshness, invalidation after synchronization, and no cached authorization decision."],
            ["Multiple Oracle environments", "Introduce an explicit tenant or environment boundary rather than reusing one global application context.", "Credential isolation, per-environment catalogs and queues, authorization scope, audit attribution, and capacity policy."],
        ],
        [2100, 3800, 3460],
    )
    add_note(doc, "Current topology limit.", "The implemented product selects one active Oracle environment and application context per deployment. Treat true multi-environment or multi-tenant operation as an architectural feature requiring explicit isolation, not a configuration shortcut.")

    add_chapter(
        doc,
        18,
        "Prioritize the remaining roadmap",
        "The next work should increase production trust and reduce support effort before broadening autonomous behavior. The current platform already has strong application controls; its largest gaps are deployment packaging, observability, browser-level testing, retention automation, and deeper live evaluation.",
    )
    add_figure(
        doc,
        assets["roadmap"],
        "Figure 10.6 - Prioritized improvement roadmap.",
        width=6.65,
        alt_text="Three-column roadmap showing immediate deployment hardening, next-stage quality and observability improvements, and later controlled expansion of schedules, identity lifecycle, storage, task synchronization, and agent plans.",
    )
    add_heading(doc, "Roadmap register")
    add_table(
        doc,
        ["Priority", "Improvement", "Why it matters", "Completion evidence"],
        [
            ["P0", "Production service packaging", "The repository has clear commands but no committed container, service-manager, or orchestration package for repeatable API and worker deployment.", "Reviewed deployment artifacts, health probes, non-root execution, secret injection, upgrade and rollback drill."],
            ["P0", "Central monitoring and alerts", "Structured logs exist, but operators need dashboards and proactive signals for API, database, queue, worker, Oracle, provider, and storage health.", "Metrics or traces, dashboards, alert thresholds, routing, runbooks, and incident test."],
            ["P0", "Retention backup and disaster recovery policy", "Important evidence is retained, but universal purge automation and environment-specific recovery objectives are not implemented.", "Approved data classes, retention jobs, legal hold, encrypted backup, restore test, RPO and RTO evidence."],
            ["P1", "Browser end-to-end testing", "Component tests do not prove routing, cookies, CSRF, uploads, downloads, responsive layout, and full user journeys in a real browser.", "Playwright or equivalent suite covering critical personas and governed operations in CI."],
            ["P1", "Live Oracle compatibility matrix", "Cloud and on-premises APIs vary by version and enabled feature.", "Versioned UAT results for each supported environment, authentication path, and operation."],
            ["P1", "Deeper Agent evaluation", "Current deterministic cases protect routing and ranking but not full provider conversations, resume replay, or long multi-step behavior.", "Expanded offline cases, replayable fixtures, provider contract tests, safety metrics, and reviewed live evaluation."],
            ["P1", "Catalog health automation", "Users should see only artifacts that still exist in the selected Oracle application.", "Scheduled read-only synchronization, staleness indicator, deactivation evidence, and administrator exceptions view."],
            ["P2", "Shared artifact storage", "Local runtime directories limit multi-host API and worker deployments.", "Provider-neutral storage interface, encryption, authorized downloads, lifecycle cleanup, and migration plan."],
            ["P2", "Additional schedule adapters", "Current unattended allowlist covers Oracle Pipelines and RTP registry synchronization.", "Risk-reviewed adapters with preflight, secret rejection, concurrency rules, execution evidence, pause and archive behavior."],
            ["P2", "Identity lifecycle automation", "Federated foundations and mappings exist, but enterprise joiner, mover, leaver integration needs deployment-specific connectors and policy.", "Approved directory connector, preview and apply controls, deprovision tests, audit, and emergency recovery access."],
            ["P3", "Oracle Task Manager synchronization", "Business users may benefit from Oracle-owned task visibility, but public API availability and semantics vary.", "Confirmed supported Oracle API, read-only proof of concept, ownership decision, mapping rules, and conflict policy."],
            ["P3", "Policy-controlled multi-step Assistant plans", "More autonomy can reduce manual work but raises compound risk and ambiguity.", "Explicit plan review, per-step policy, bounded tools, checkpoints, rollback limitations, evaluation, and human approval design."],
        ],
        [850, 2400, 3560, 2550],
    )
    add_heading(doc, "What should not be added yet")
    add_bullets(
        doc,
        [
            "Unrestricted natural-language execution without artifact validation and explicit approval.",
            "Automatic retry of timed-out, disconnected, or recovery-required Oracle writes.",
            "A second local configuration of Oracle Pipeline stages, Data Integration mappings, or saved job behavior that Oracle already owns.",
            "Generic administrator API tokens with broad platform authority before scope, rotation, client identity, and audit requirements are designed.",
            "Direct SQL maintenance screens that bypass services, constraints, archival behavior, and evidence.",
            "Multi-tenant deployment by merely changing environment variables between requests.",
        ],
    )

    add_page_break(doc)
    doc.add_paragraph("ENGINEERING QUICK REFERENCE", style="Kicker")
    doc.add_paragraph("Use the release and support checklist", style="Heading 1")
    add_heading(doc, "Local verification")
    add_code_block(
        doc,
        "python -m pip install -r requirements-dev.txt\npython -m pytest\npython -m app.agent.evaluation --suite evaluations/agent/release_v1.json --fail-on-threshold\n\nSet-Location frontend\npnpm install --frozen-lockfile\npnpm verify",
        "Quick reference 10.1. Deterministic application verification.",
    )
    add_heading(doc, "Database and startup")
    add_code_block(
        doc,
        "python -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup\n\n$env:EPM_EXECUTION_RUNTIME = \"web\"\npython -m uvicorn web_main:app --host 0.0.0.0 --port 8080\n# separately supervised\npython worker.py",
        "Quick reference 10.2. Current production preparation and process entry points.",
    )
    add_heading(doc, "Release evidence to retain")
    add_table(
        doc,
        ["Evidence", "Minimum record"],
        [
            ["Source", "Release tag, commit SHA, branch and pull request, reviewer, and dependency lock state."],
            ["Quality", "Backend matrix, migration smoke test, Agent evaluation report, frontend tests and build, plus live UAT result."],
            ["State protection", "PostgreSQL backup reference, old and new Alembic heads, LangGraph setup result, and restore readiness."],
            ["Deployment", "Start and end time, operator, API and worker version, environment, configuration version, and health result."],
            ["Smoke test", "Signed-in user, selected Oracle application, catalog check, execution ID, Oracle evidence, and outcome."],
            ["Decision", "Known issues, accepted risk, rollback or forward-fix decision, owner, and next review date."],
        ],
        [2300, 7060],
    )
    add_heading(doc, "Glossary")
    add_table(
        doc,
        ["Term", "Plain-language meaning"],
        [
            ["Adapter", "A small component that translates the platform contract to a specific external system or SDK."],
            ["Contract", "The agreed input, output, validation, permission, and behavior of a capability."],
            ["Composition root", "The entry point that creates and connects services for a running process."],
            ["Deterministic test", "A repeatable check whose result does not depend on a live model or Oracle environment."],
            ["Live UAT", "Recorded user acceptance testing against the real provider and Oracle deployment claimed as supported."],
            ["Release gate", "A required check that must pass before a version is promoted."],
            ["Liveness", "Proof that a process can answer; it does not prove every dependency is healthy."],
            ["Readiness", "Proof that a service can accept normal work using its required local dependencies."],
            ["Observability", "Logs, metrics, traces, and evidence that help operators understand current and past behavior."],
            ["Idempotency", "Protection that makes a repeated request reuse an earlier decision or result rather than duplicate work."],
            ["Recovery required", "A safety state used when Oracle may have accepted work but the platform cannot prove the terminal outcome."],
            ["RPO and RTO", "The approved maximum data-loss window and target time to restore service after a disaster."],
        ],
        [2400, 6960],
    )
    add_note(doc, "Handbook completion.", "Documents 1 through 10 together describe the current product, architecture, installation, Oracle operations, Planning lifecycle, governed Assistant, data review and reporting, administration and security, database governance, engineering, testing, deployment, and improvement roadmap.")

    finish_document(doc)
    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
