"""Build Document 05: Oracle EPM Automation Operations."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_DARK,
    BLUE_LIGHT,
    GOLD,
    GOLD_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    RED_LIGHT,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    add_bullets,
    add_callout,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    font,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    wrap,
)
from docs.builders.build_document_03 import (  # noqa: E402
    compact_trailing_empty_paragraph,
    configure_styles,
    patch_list_numbering,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-05"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / (
    "BISP_EPM_Automation_Document_05_Oracle_EPM_Automation_Operations.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
MUTED_HEX = f"#{MUTED}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"
RED_HEX = f"#{RED}"


def _arrow(draw: ImageDraw.ImageDraw, start, end, color=BLUE_HEX, width=5) -> None:
    draw.line((*start, *end), fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 17 * direction, y2 - 10), (x2 - 17 * direction, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 10, y2 - 17 * direction), (x2 + 10, y2 - 17 * direction)]
    draw.polygon(points, fill=color)


def _title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((55, 34), title, font=font(34, bold=True), fill=NAVY_HEX)
    draw.text((55, 82), subtitle, font=font(19), fill=MUTED_HEX)


def _card(draw: ImageDraw.ImageDraw, bounds, title: str, body: str, *, color: str, label: str = "") -> None:
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=18, fill="#FFFFFF", outline=color, width=3)
    y = y1 + 19
    if label:
        draw.text((x1 + 20, y), label.upper(), font=font(15, bold=True), fill=color)
        y += 27
    for line in wrap(draw, title, font(22, bold=True), x2 - x1 - 40):
        draw.text((x1 + 20, y), line, font=font(22, bold=True), fill=NAVY_HEX)
        y += 29
    y += 6
    for line in wrap(draw, body, font(16), x2 - x1 - 40):
        draw.text((x1 + 20, y), line, font=font(16), fill=MUTED_HEX)
        y += 22


def create_operation_map(path: Path) -> None:
    image = Image.new("RGB", (1400, 880), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Choose the smallest operation", "Start from the outcome. Oracle remains the owner of saved jobs and application design.")
    cards = [
        ((55, 145, 375, 325), "Calculate", "Business Rules", "Run a deployed rule and supply required runtime prompts.", BLUE_HEX),
        ((395, 145, 715, 325), "Move data", "Data Maps", "Push a governed source slice to its configured target.", ORANGE_HEX),
        ((735, 145, 1055, 325), "Load Planning data", "Planning Data Import", "Run a saved native Import Data job with an approved file.", TEAL_HEX),
        ((1075, 145, 1345, 325), "Load metadata", "Metadata Import", "Run a saved import and optionally refresh the cube.", RED_HEX),
        ((55, 365, 375, 545), "Integrate data", "Data Integration", "Run a configured integration for periods and modes.", TEAL_HEX),
        ((395, 365, 715, 545), "Run many stages", "Pipeline", "Inspect and execute one Oracle-owned orchestration.", ORANGE_HEX),
        ((735, 365, 1055, 545), "Synchronize metadata", "Cube Refresh", "Run an existing saved Refresh Database job.", RED_HEX),
        ((1075, 365, 1345, 545), "Set context", "Variables", "Maintain substitution definitions or user assignments.", BLUE_HEX),
        ((390, 600, 1010, 785), "Create an output", "Report Generation", "Inspect a registered definition or Planning form and generate a read-only Excel workbook.", TEAL_HEX),
    ]
    for bounds, title, label, body, color in cards:
        _card(draw, bounds, title, body, color=color, label=label)
    image.save(path)


def create_governance_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 700), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "One governed lifecycle for every operation", "The page may collect different inputs, but the safety and evidence pattern stays consistent.")
    stages = [
        ("1", "Prepare", "Choose a permitted service, current artifact, periods, values, or file."),
        ("2", "Review", "Read the exact target, Oracle-owned configuration, and impact."),
        ("3", "Approve", "Confirm the reviewed action. Nothing has been sent before this point."),
        ("4", "Queue & run", "A durable worker authenticates, submits, and monitors Oracle."),
        ("5", "Evidence", "Inspect steps, status, statistics, logs, files, user, and trigger source."),
    ]
    x_positions = [55, 325, 595, 865, 1135]
    colors = [BLUE_HEX, BLUE_HEX, ORANGE_HEX, ORANGE_HEX, TEAL_HEX]
    for index, ((number, title, body), x, color) in enumerate(zip(stages, x_positions, colors)):
        _card(draw, (x, 170, x + 215, 500), title, body, color=color, label=f"STEP {number}")
        if index < 4:
            _arrow(draw, (x + 220, 335), (x + 255, 335))
    draw.rounded_rectangle((185, 550, 1215, 635), radius=16, fill=f"#{TEAL_LIGHT}", outline=TEAL_HEX, width=2)
    draw.text((220, 575), "Leaving the browser does not cancel accepted work; Jobs & Activity retains the execution.", font=font(20, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_catalog_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 790), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Live discovery and durable registration", "Selectors show artifacts proven for the connected application, not stale names from another environment.")
    _card(draw, (55, 165, 420, 365), "Live saved-job discovery", "Business Rules, Data Maps, Import Metadata jobs, Import Data jobs, cubes, and available Cube Refresh jobs are read from Oracle.", color=BLUE_HEX, label="DIRECT")
    _card(draw, (510, 165, 890, 365), "Registered technical artifacts", "Pipelines and Data Integrations have durable records because public discovery differs across Oracle environments.", color=ORANGE_HEX, label="DURABLE")
    _card(draw, (980, 165, 1345, 365), "Environment identity", "The base URL and Planning application isolate registrations and verification state.", color=TEAL_HEX, label="SCOPE")
    _arrow(draw, (235, 400), (235, 475), color=BLUE_HEX)
    _arrow(draw, (700, 400), (700, 475), color=ORANGE_HEX)
    _arrow(draw, (1160, 400), (1160, 475), color=TEAL_HEX)
    statuses = [
        ((55, 495, 360, 690), "Verified", "Selectable and current for this Planning application.", TEAL_HEX),
        ((395, 495, 700, 690), "Pending", "Registered but not yet proven by a governed run or verification.", ORANGE_HEX),
        ((735, 495, 1040, 690), "Missing / hidden", "Definitive Oracle misses remove it from runnable selectors without erasing history.", RED_HEX),
        ((1075, 495, 1345, 690), "Unavailable", "Oracle could not be inspected; execution remains disabled until the environment is healthy.", BLUE_HEX),
    ]
    for bounds, title, body, color in statuses:
        _card(draw, bounds, title, body, color=color)
    image.save(path)


def create_file_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 750), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Three file-source choices", "Choose exactly one source. The saved Oracle job or integration still owns its technical mappings.")
    sources = [
        ((55, 170, 410, 410), "Configured file", "Use the filename already stored in the Oracle job, integration, or Pipeline definition. No upload is performed.", BLUE_HEX),
        ((520, 170, 875, 410), "Local upload", "Choose a file from the computer. A matching Oracle Inbox filename is replaced deliberately before execution.", ORANGE_HEX),
        ((985, 170, 1340, 410), "Existing Inbox file", "Choose a compatible live Oracle repository file. Lists are filtered by operation and extension.", TEAL_HEX),
    ]
    for bounds, title, body, color in sources:
        _card(draw, bounds, title, body, color=color)
        _arrow(draw, ((bounds[0] + bounds[2]) // 2, 430), ((bounds[0] + bounds[2]) // 2, 510), color=color)
    _card(draw, (385, 525, 1015, 675), "One reviewed file identity", "The review screen shows the selected source and filename before the request is queued. The execution log records upload replacement and Oracle submission evidence.", color=TEAL_HEX, label="GOVERNED INPUT")
    image.save(path)


def create_pipeline_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Pipeline execution uses the live Oracle definition", "The platform supplies runtime inputs; it does not rebuild the Pipeline stages.")
    _card(draw, (55, 160, 330, 405), "Verified code", "Select a Pipeline already verified for this environment, or let an administrator verify and register an exact code.", color=BLUE_HEX, label="1")
    _arrow(draw, (345, 280), (415, 280))
    _card(draw, (430, 160, 705, 405), "Inspect live", "Read display name, stages, job count, parallel behavior, variables, and stage file requirements.", color=BLUE_HEX, label="2")
    _arrow(draw, (720, 280), (790, 280))
    _card(draw, (805, 160, 1080, 405), "Resolve inputs", "Use Oracle defaults where valid; provide required runtime values and configured, uploaded, or Inbox files.", color=ORANGE_HEX, label="3")
    _arrow(draw, (1095, 280), (1165, 280))
    _card(draw, (1180, 160, 1345, 405), "Run once", "Submit and monitor the complete Pipeline as one governed execution.", color=TEAL_HEX, label="4")
    draw.rounded_rectangle((135, 505, 1265, 665), radius=20, fill="#FFFFFF", outline=ORANGE_HEX, width=3)
    draw.text((170, 535), "ORACLE OWNS", font=font(16, bold=True), fill=ORANGE_HEX)
    draw.text((170, 570), "Stages • job order • mappings • calculations • configured notifications", font=font(23, bold=True), fill=NAVY_HEX)
    draw.text((170, 615), "BISP validates inputs, stages replacement files, requests execution, monitors status, and retains evidence.", font=font(18), fill=MUTED_HEX)
    image.save(path)


def create_evidence_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 770), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "From queued request to retained evidence", "The execution record separates a browser experience from background Oracle work.")
    _card(draw, (55, 165, 350, 365), "Queued", "Request is durable and waiting for an available worker.", color=BLUE_HEX)
    _arrow(draw, (365, 265), (445, 265))
    _card(draw, (460, 165, 755, 365), "Running", "Worker validates, connects, submits Oracle work, and polls to terminal status.", color=ORANGE_HEX)
    _arrow(draw, (770, 265), (850, 265))
    _card(draw, (865, 165, 1160, 365), "Terminal", "Success, failed, or recovery-required is stored with step details.", color=TEAL_HEX)
    _arrow(draw, (1010, 390), (1010, 465), color=TEAL_HEX)
    _card(draw, (630, 485, 1345, 680), "Jobs & Activity", "Search and filter executions; inspect initiator, Oracle executor, trigger source, timestamps, step timeline, failure message, record statistics, generated files, and downloadable log.", color=TEAL_HEX, label="EVIDENCE")
    _card(draw, (55, 485, 535, 680), "Do not duplicate blindly", "The platform blocks an active duplicate for the same operation target. After a browser timeout, inspect the retained execution before trying again.", color=RED_HEX, label="SAFETY")
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "map": ASSET_DIR / "operation_map.png",
        "governance": ASSET_DIR / "governed_lifecycle.png",
        "catalog": ASSET_DIR / "catalog_lifecycle.png",
        "files": ASSET_DIR / "file_sources.png",
        "pipeline": ASSET_DIR / "pipeline_live_definition.png",
        "evidence": ASSET_DIR / "execution_evidence.png",
    }
    create_operation_map(assets["map"])
    create_governance_flow(assets["governance"])
    create_catalog_flow(assets["catalog"])
    create_file_flow(assets["files"])
    create_pipeline_flow(assets["pipeline"])
    create_evidence_flow(assets["evidence"])
    return assets


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break()


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    add_page_break(doc)
    doc.add_paragraph(f"CHAPTER {number:02d}", style="Kicker")
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "BISP SOLUTIONS  /  ORACLE EPM AUTOMATION OPERATIONS"
    set_run_font(paragraph.runs[0], size=8.5, color=MUTED, bold=True)
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.columns[0].width = Inches(5.7)
    table.columns[1].width = Inches(0.8)
    left = table.cell(0, 0).paragraphs[0]
    left.text = "Document 05  |  Oracle EPM Automation Operations"
    set_run_font(left.runs[0], size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(right)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:color"), LINE)
        borders.append(top)
        tc_pr.append(borders)


def add_cover(doc: Document) -> None:
    if LOGO.is_file():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions company logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(34)
    doc.add_paragraph("PLATFORM HANDBOOK", style="Kicker")
    doc.add_paragraph("Oracle EPM Automation\nOperations", style="Title")
    doc.add_paragraph("BISP Solutions Oracle EPM Automation Platform", style="Subtitle")
    doc.add_paragraph(
        "A current, operation-by-operation guide to live artifact discovery, governed inputs, file handling, explicit approval, durable Oracle execution, evidence, and troubleshooting.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [["05 of the platform handbook", "1 September 2026", "Planners, FP&A teams, consultants, administrators, reviewers and support teams"]],
        [2200, 2100, 5060],
    )
    add_callout(
        doc,
        "Verified current scope",
        "This volume was rebuilt from the current React runners, API contracts, services, role permissions, durable worker, and Oracle catalog behavior. The retired interface and earlier JSON-only assumptions are excluded. The audited operation-focused suite passed 67 tests before authoring.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("BISP Solutions  |  Safe, simple and observable Oracle Planning automation"), size=10, color=MUTED, bold=True)


def _finish_document(doc: Document) -> None:
    compact_trailing_empty_paragraph(doc)
    for table in doc.tables:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            header_pr.append(header)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for section in doc.sections:
        for footer in (section.footer, section.even_page_footer):
            for table in footer.tables:
                header_pr = table.rows[0]._tr.get_or_add_trPr()
                if header_pr.find(qn("w:tblHeader")) is None:
                    header = OxmlElement("w:tblHeader")
                    header.set(qn("w:val"), "true")
                    header_pr.append(header)
    properties = doc.core_properties
    properties.title = "Oracle EPM Automation Operations"
    properties.subject = "Current standalone Oracle Planning operation procedures for the BISP Solutions platform"
    properties.author = "BISP Solutions"
    properties.keywords = "Oracle EPM, Planning, operations, Pipeline, Data Integration, metadata, data load, business rules, data maps"
    properties.comments = "Documents the current React application and durable operation implementation as of 1 September 2026; retired UI excluded."


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    patch_list_numbering(doc)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS VOLUME", style="Kicker")
    doc.add_paragraph("Choose an outcome, then follow one safe path", style="Heading 1")
    doc.add_paragraph(
        "This is a practical operations manual, not a REST reference. Each chapter tells you what the service is for, what must already exist in Oracle, what the platform discovers, what the user supplies, what approval means, and what evidence proves the outcome.",
        style="Lead",
    )
    add_table(
        doc,
        ["Need", "Use", "Chapter"],
        [
            ["Calculate a Planning result", "Business Rules", "5"],
            ["Push data between configured cubes", "Data Maps", "6"],
            ["Load a native Planning data file", "Planning Data Import", "7"],
            ["Load members or properties", "Metadata Import", "8"],
            ["Run a Data Management / Data Integration load", "Data Integrations", "9"],
            ["Run a multi-stage technical process", "Pipelines", "10"],
            ["Synchronize metadata with the cube", "Planning Cube Refresh", "11"],
            ["Change shared or personal context", "Substitution / User Variables", "12–13"],
            ["Create a read-only Excel output", "Reports", "14"],
            ["Investigate any accepted request", "Jobs & Activity", "15"],
        ],
        [3600, 3600, 2160],
    )
    add_callout(doc, "Operations are for focused work", "Use a standalone operation when one service completes the task. Use an Oracle Pipeline when Oracle already owns the multi-stage process. Do not reproduce Pipeline stages manually in the platform.", fill=TEAL_LIGHT, accent=TEAL)
    add_heading(doc, "Contents")
    add_table(
        doc,
        ["Part", "Chapters", "Purpose"],
        [
            ["Foundation", "1–4", "Catalog, roles, governance and file-source rules."],
            ["Calculation and movement", "5–6", "Business Rules and Data Maps."],
            ["Data and metadata loading", "7–10", "Native import, metadata, Data Integration and Pipelines."],
            ["Administration and outputs", "11–14", "Cube Refresh, variables and reporting."],
            ["Evidence and runbooks", "15–16", "Monitoring, troubleshooting and quick-reference checklists."],
        ],
        [2100, 1800, 5460],
    )

    add_chapter(doc, 1, "Understand the current operation catalog", "The Operations page is a permission-filtered catalog of focused Oracle EPM services. It does not expose technical endpoints or services the signed-in role cannot use.")
    add_figure(doc, assets["map"], "Figure 5.1 — Current standalone operation map.", "Decision map linking a business outcome to Business Rules, Data Maps, data and metadata imports, Data Integration, Pipelines, Cube Refresh, variables, and report generation.", width=6.35)
    add_heading(doc, "Current services")
    add_table(
        doc,
        ["Service", "Risk", "What it changes"],
        [
            ["Business Rules", "Controlled", "Runs a deployed calculation with optional runtime prompts."],
            ["Data Maps", "Elevated", "Moves a configured source slice to a configured target; optional clear and overrides."],
            ["Planning Data Import", "Elevated", "Loads data through a saved native Import Data job."],
            ["Metadata Import", "Elevated", "Loads metadata through a saved job; may conditionally refresh the cube."],
            ["Data Integrations", "Elevated", "Runs a configured Data Integration for selected periods and modes."],
            ["Pipelines", "Elevated", "Runs a live multi-stage Oracle Data Integration Pipeline."],
            ["Planning Cube Refresh", "Elevated", "Runs an existing saved Refresh Database job."],
            ["Substitution Variables", "Elevated", "Updates or explicitly creates application- or cube-scoped definitions."],
            ["User Variables", "Controlled", "Changes one user's member assignment for an existing definition."],
            ["Report Generation", "Read only", "Reads Planning data and creates a downloadable Excel workbook."],
        ],
        [2500, 1500, 5360],
    )
    add_heading(doc, "Current role access")
    add_table(
        doc,
        ["Capability", "Service Admin", "Power User", "User", "Viewer"],
        [
            ["General Oracle operations", "Yes", "Yes", "No", "No"],
            ["Substitution Variables", "Yes", "No", "No", "No"],
            ["User Variables", "Self or authorized user", "Self", "Self", "No"],
            ["Report Generation", "Yes", "Yes", "Yes", "Yes"],
            ["Jobs & Activity", "Yes", "Yes", "No", "No"],
            ["Catalog synchronization / registration", "Yes", "No", "No", "No"],
        ],
        [2800, 1900, 1700, 1560, 1400],
    )
    add_callout(doc, "Two security layers", "The platform checks its role permissions first. Oracle then enforces the connected user's or execution identity's application permissions. Seeing a service does not grant access that Oracle itself denies.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 2, "Follow the governed runner", "Every operation uses the same safety sequence: prepare, review, approve, execute in the background, and inspect retained evidence.")
    add_figure(doc, assets["governance"], "Figure 5.2 — Shared governed lifecycle.", "Five-step flow from preparation and review through explicit approval, durable execution, and retained evidence.", width=6.35)
    add_heading(doc, "What happens at each stage")
    add_table(
        doc,
        ["Stage", "User responsibility", "Platform responsibility"],
        [
            ["Prepare", "Choose the correct environment, artifact and business inputs.", "Show only permitted services and current choices; validate required fields."],
            ["Review", "Read the exact target, file, period, mode, value and impact.", "Present a normalized summary; do not contact Oracle to change data."],
            ["Approve", "Confirm the reviewed action is authorized now.", "Create a durable queued record before background work begins."],
            ["Execute", "You may leave the page after acceptance.", "Worker validates again, authenticates, submits Oracle work and monitors it."],
            ["Result", "Review terminal status and business evidence.", "Retain steps, user identity, trigger, statistics, files, logs and failures."],
        ],
        [1600, 3500, 4260],
    )
    add_heading(doc, "Durable execution states")
    add_bullets(doc, [
        "Queued — accepted and waiting for a worker.",
        "Running — the worker is validating, connecting, submitting or monitoring Oracle.",
        "Success — every required workflow step completed.",
        "Failed — a validation, connection, Oracle job or verification step failed.",
        "Recovery required — the platform needs a deliberate, authorized recovery decision.",
    ])
    add_callout(doc, "Do not double-submit", "The platform blocks another active execution for the same operation target. If the browser times out or closes, open Jobs & Activity instead of immediately starting the same job again.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 3, "Understand live discovery, registration and synchronization", "The platform keeps selectors aligned with the connected Oracle base URL and Planning application so stale artifacts are not treated as current.")
    add_figure(doc, assets["catalog"], "Figure 5.3 — Oracle catalog lifecycle.", "Diagram of live discovery, durable Pipeline and Data Integration registration, environment scoping, and verified, pending, hidden, or unavailable states.", width=6.35)
    add_heading(doc, "Discovered live from Oracle")
    add_bullets(doc, [
        "Business Rule saved-job definitions.",
        "Data Map saved-job definitions.",
        "Import Metadata and Import Data jobs.",
        "Planning cubes / plan types where supported.",
        "Saved Cube Refresh jobs when Oracle exposes their exact names.",
    ])
    add_heading(doc, "Registered and verified")
    add_para(doc, "Pipelines and Data Integrations use durable registration because Oracle discovery is not equally complete in every Cloud and on-premises version. A registration belongs to one environment identity—the normalized Oracle base URL plus Planning application—not merely to a filename or browser session.")
    add_table(
        doc,
        ["State", "Selector behavior", "Meaning"],
        [
            ["Verified", "Visible and selectable", "The artifact was proven for the current application."],
            ["Pending", "Not treated as ordinarily runnable", "An exact Data Integration may await its first governed verification run."],
            ["Missing / inactive", "Hidden", "Oracle definitively reported that the artifact does not exist; history is preserved."],
            ["Oracle unavailable", "Execution disabled", "The platform cannot safely prove current Oracle context."],
        ],
        [1900, 2500, 4960],
    )
    add_heading(doc, "What the Sync control does")
    add_numbered(doc, [
        "Connects to the currently selected Planning application.",
        "Discovers supported saved jobs and cubes in one session.",
        "Verifies registered Pipeline codes and discovers or verifies Data Integrations where Oracle allows it.",
        "Marks current artifacts verified and hides definitively missing registrations from runnable dropdowns.",
        "Preserves historical registrations and executions instead of deleting audit evidence.",
    ])
    add_callout(doc, "Generic Cube Refresh names", "The label RefreshCube may be returned by Oracle as a generic API definition even when no saved job with that name exists. The current platform hides that generic label unless it can resolve a real saved Refresh Database job; the user may still enter an exact saved name when discovery is limited.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 4, "Choose files safely", "Data, metadata, Data Integration and Pipeline operations reuse one principle: choose exactly one approved source and show its identity again before approval.")
    add_figure(doc, assets["files"], "Figure 5.4 — Configured, local-upload and Oracle Inbox sources.", "Three file sources converge into one reviewed file identity before execution.", width=6.35)
    add_heading(doc, "Live Inbox selection")
    add_para(doc, "The common file picker lists compatible Oracle repository files newest first. Outbox, report and scheduler-output roots are excluded. If exactly one compatible file is returned and no prior choice exists, it can be selected automatically. Refresh re-reads Oracle; an exact manual reference remains available when discovery fails.")
    add_table(
        doc,
        ["Operation", "Visible compatible files", "Local upload guidance"],
        [
            ["Planning Data Import", ".csv, .txt, .zip", "CSV, TXT or ZIP"],
            ["Metadata Import", ".csv, .zip", "CSV or ZIP"],
            ["Data Integration", ".csv, .txt, .zip, .dat", "Use CSV, TXT or ZIP for the current local-upload path; DAT is safest as an existing/configured Oracle reference."],
            ["Pipeline", "Per live stage requirement", "Only extensions returned by the selected Pipeline requirement."],
        ],
        [2300, 2600, 4460],
    )
    add_heading(doc, "Replacement behavior")
    add_para(doc, "A local file is uploaded to the Oracle Inbox immediately before the approved job. If the same Inbox filename already exists, the platform replaces it and records whether a replacement occurred. This prevents the old-file problem while making the effect explicit in the log.")
    add_callout(doc, "A file extension is not a business classification", "The current filters prevent clearly incompatible output files from appearing, but they cannot know whether Revenue_January.csv belongs to one job and January_Sales.csv belongs to another. The user must still confirm the file's business content and saved Oracle job mapping.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 5, "Run a Business Rule", "Use Business Rules for an approved deployed Calculation Manager rule—not for creating or editing the rule itself.")
    add_heading(doc, "What must already exist in Oracle")
    add_bullets(doc, [
        "The rule is deployed and visible to the connected user.",
        "The rule's cube, calculation logic, security and defaults are configured in Oracle.",
        "Required Runtime Prompts have valid values or usable Oracle defaults.",
    ])
    add_heading(doc, "Current guided procedure")
    add_numbered(doc, [
        "Open Operations and choose Business Rules.",
        "Select the exact rule from the live Oracle catalog. An assistant handoff may prefill a verified choice, but the same review still applies.",
        "Review the synchronized Runtime Prompt definition when available. Enter values for prompts marked required without defaults; leave optional prompts empty to use Oracle defaults.",
        "If no synchronized definition exists, use the manual fallback only when you know the exact prompt names; leaving the list empty uses rule defaults.",
        "Review the rule name and prompt names/values, approve, and start.",
        "Inspect Oracle job status and the retained execution result.",
    ])
    add_heading(doc, "Runtime Prompt registry")
    add_para(doc, "Administrators can import Calculation Manager XML or an LCM ZIP to synchronize rule prompt definitions. The registry records labels, types, dimensions, scope, multiple-value behavior, defaults and limits. A failed import preserves the last-known-good registry rather than replacing it with incomplete data.")
    add_callout(doc, "No RTP guessing", "A rule can run without prompts, with Oracle defaults, or with required values. The platform does not invent a member. If a required value is unknown, stop and ask the rule owner.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 6, "Run a Data Map", "Use a Data Map to publish a configured Planning slice to its configured target. The Data Map definition remains in Oracle.")
    add_heading(doc, "What the platform lets you control")
    add_table(
        doc,
        ["Control", "Default", "Meaning"],
        [
            ["Data Map", "Required selection", "Exact live PLAN_TYPE_MAP job."],
            ["Clear target", "Off", "Requests Oracle to clear the target region before the push."],
            ["Member overrides", "None", "Overrides selected dimensions for this execution."],
            ["Exclusion overrides", "None", "Excludes selected members for this execution."],
        ],
        [2600, 1800, 4960],
    )
    add_numbered(doc, [
        "Select the live Data Map.",
        "Leave Clear target off unless the approved load strategy requires it.",
        "Optionally add each override dimension once with its exact member selection.",
        "Optionally add exclusions; incomplete or duplicate dimension entries are rejected.",
        "Review the target name, clear choice and override dimensions, then approve.",
        "Use Data Review or an approved report afterward when source-to-target business validation is required.",
    ])
    add_callout(doc, "Configuration boundary", "The platform does not change the Data Map's source cube, target cube, dimension mappings or default POV. Configure those in Oracle. Runtime overrides narrow or alter only the approved execution.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 7, "Run Planning Data Import", "Use native Planning import when a saved Import Data job already defines the cube, file format, dimensions, period columns and mappings.")
    add_heading(doc, "Prepare")
    add_numbered(doc, [
        "Select a live saved Import Data job.",
        "Optionally enter an error-output filename for rejected records when the Oracle job supports it.",
        "Choose one source: the file configured in the Oracle job, a local CSV/TXT/ZIP upload, or a compatible existing Inbox file.",
        "Confirm the chosen filename is the current business file and matches the saved job's format.",
    ])
    add_heading(doc, "What the saved job owns")
    add_bullets(doc, [
        "Target cube / plan type.",
        "Delimiter and import format.",
        "Dimension columns and mappings.",
        "Single-period, delimited-numeric or multi-column numeric behavior.",
        "Period-column handling and import behavior.",
    ])
    add_heading(doc, "Execution engine detail")
    add_para(doc, "Local-upload and existing-Inbox paths run through Planning REST. In the current implementation, choosing the file configured in the saved Import Data job uses EPM Automate because that route reliably executes the saved job without overriding its configured filename. The user still follows the same governed review and receives one durable result.")
    add_heading(doc, "Success evidence")
    add_para(doc, "When Oracle exposes record counters, the result shows records read, processed and rejected. If this environment does not expose counters, the job may still be successful; the log explicitly states that statistics were unavailable.")
    add_callout(doc, "Flexible files require matching Oracle configuration", "The Python service does not assume that dimensions are columns 1–5 or that the file has three periods. Different structures work when the selected saved job is configured for that exact file layout.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 8, "Run Metadata Import", "Use Metadata Import for approved structural changes through an existing saved Import Metadata job, with an optional dependent cube refresh.")
    add_heading(doc, "Prepare the job and source")
    add_numbered(doc, [
        "Select a live Import Metadata job.",
        "Choose exactly one source: configured Oracle job files, a local CSV/ZIP, or a compatible Oracle Inbox file.",
        "Optionally enter the error-output filename.",
        "Decide whether a saved Cube Refresh job should run only after a successful import.",
        "Review the structural impact, source and conditional refresh before approval.",
    ])
    add_heading(doc, "Conditional workflow")
    add_table(
        doc,
        ["Import result", "Refresh selected?", "Outcome"],
        [
            ["Success", "No", "Workflow ends after import evidence."],
            ["Success", "Yes", "The selected saved Cube Refresh job starts and is monitored."],
            ["Failure", "Either", "Workflow stops; refresh is not started."],
        ],
        [2100, 2100, 5160],
    )
    add_callout(doc, "Oracle owns metadata mapping", "The saved job defines dimensions, member properties, delimiters, validation, mappings and target cube. The platform does not infer a safe metadata mapping from the CSV filename.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 9, "Run a Data Integration", "Use Data Integration for a configured file-based integration with selected periods, import/export modes and one approved file source.")
    add_heading(doc, "Artifact selection and synchronization")
    add_para(doc, "Only verified active integrations for the current environment appear in the normal selector. An administrator can register an exact name when no safe read-only standalone lookup is available. That registration remains pending until Oracle accepts a governed verification run; definitively missing integrations become hidden without deleting their history.")
    add_heading(doc, "Period entry modes")
    add_table(
        doc,
        ["Mode", "Example", "When to use"],
        [
            ["Mapped period", "Jan-27", "The Data Integration period mapping expects a display-style mapped period."],
            ["Planning member", "Jan#FY27", "The integration expects a Planning period and year member pair."],
            ["Exact / advanced", "Environment-specific", "Use the exact start and end names already proven in Oracle."],
        ],
        [2300, 2200, 4860],
    )
    add_heading(doc, "Current import and export choices")
    add_para(doc, "Import choices are Replace, Append, Map and Validate, or No Import. Export choices are Merge, Replace, Accumulate, Subtract, No Export, or Check. These values are passed to the configured integration; they do not redefine its mappings or target application.")
    add_numbered(doc, [
        "Choose a verified integration.",
        "Choose the period naming mode, planning year and start/end period—or enter exact period names in advanced mode.",
        "Choose import and export modes approved for the business load.",
        "Choose the configured file, a compatible local upload, or an existing Inbox reference.",
        "Review all four input groups and approve.",
        "Monitor Oracle to terminal status and inspect record counters when returned.",
    ])
    add_callout(doc, "Multi-column formats are Oracle configuration", "A multi-column numeric integration and a delimited-numeric integration use the same platform runner. The chosen Oracle integration determines how dimensions and data columns are interpreted.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 10, "Run an Oracle Pipeline", "Use a Pipeline when Oracle already owns a multi-stage technical lifecycle. The platform inspects the live definition and asks only for unresolved run-time inputs.")
    add_figure(doc, assets["pipeline"], "Figure 5.5 — Live Pipeline inspection and execution.", "Flow from verified Pipeline code through live inspection and runtime input resolution to one governed execution, with Oracle retaining stage ownership.", width=6.35)
    add_heading(doc, "Select, verify and inspect")
    add_numbered(doc, [
        "Select a verified Pipeline. A missing or inactive registration is not shown as runnable.",
        "If the exact code is absent, a Service Administrator can Verify and register it; Oracle must confirm the code before normal use.",
        "Inspect the current Pipeline. The platform reads its display name, stages, job count, parallel behavior, runtime variables and file requirements.",
        "Supply required values; retain Oracle defaults when they are appropriate.",
        "Resolve each file requirement using its configured reference, a permitted local replacement, or another exact Oracle Inbox reference.",
        "Review stages, jobs, variables and file choices, then approve the whole Pipeline once.",
    ])
    add_heading(doc, "Current mouse-first behavior and exact references")
    add_para(doc, "Year and recognized period variables use dropdowns where the live definition exposes them. Other variables use their current Oracle defaults or exact values. Pipeline file requirements are stage-specific; the current runner displays the requirement and allowed extensions, but an alternative Oracle Inbox file is entered as the exact reference rather than selected through the shared Inbox list.")
    add_callout(doc, "Do not repeat Pipeline stages as platform operations", "If metadata import, data load, rules, Data Maps or notifications are already Pipeline stages, configure them once in Oracle and run the Pipeline. The platform should not add duplicate pre- or post-steps merely because equivalent standalone services exist.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 11, "Run Planning Cube Refresh", "Use Cube Refresh only after approved metadata changes or when the application's saved refresh process requires synchronization.")
    add_numbered(doc, [
        "Confirm that metadata changes are complete and active users can tolerate refresh impact.",
        "Choose a live saved Refresh Database job suggestion when Oracle exposes it, or enter the exact saved name shown in Oracle Planning.",
        "Review the application-wide impact and saved job name.",
        "Confirm the application is ready, approve, and start.",
        "Monitor Oracle to terminal status before allowing dependent work to continue.",
    ])
    add_heading(doc, "Why an exact-name fallback exists")
    add_para(doc, "Some Oracle environments expose only a generic CUBE_REFRESH definition rather than each saved job. The platform does not pretend that the generic label is executable. Exact-name execution remains available because Oracle can accept a real saved job even when public discovery is incomplete.")
    add_callout(doc, "Refresh can affect users", "A database refresh synchronizes Planning metadata with the underlying cube and can temporarily affect availability or performance. Schedule it deliberately; the platform executes the selected job but does not change its cube selections.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 12, "Maintain Substitution Variables", "Substitution Variables are shared Oracle definitions used by forms, rules, reports, jobs and Planning context. The current platform supports safe update and explicit creation.")
    add_heading(doc, "Update an existing definition")
    add_numbered(doc, [
        "Search the live variable catalog by name, value or scope and choose the exact definition.",
        "Review its current value and enter a different new value.",
        "Review the scope, name, old value and new value.",
        "Approve one change. The platform rechecks Oracle's current value before applying it.",
        "Oracle is reread after the update; an unconfirmed value fails the workflow.",
    ])
    add_heading(doc, "Create a new definition")
    add_numbered(doc, [
        "Choose ALL for application-wide scope or an exact discovered cube scope.",
        "Enter the variable name without a leading ampersand (&).",
        "Enter a non-empty initial value expected by the Oracle artifacts that will consume it.",
        "Confirm that the same name does not already exist in that scope.",
        "Review, approve, create and verify the new definition.",
    ])
    add_heading(doc, "Current validation boundary")
    add_para(doc, "The current service validates required scope/name/value, maximum lengths, duplicate identity and stale updates. It does not yet prove that a free-text value is a valid member for every form, rule or job that may consume the variable. Business validation still matters.")
    add_callout(doc, "Safe-update protection", "If another person changes the variable after the page was loaded, the approved update is rejected. Refresh the live catalog and review the new current value instead of overwriting it silently.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 13, "Maintain User Variables", "User Variables personalize one user's Planning context. The platform updates assignments for existing Oracle definitions; it does not create user-variable definitions.")
    add_heading(doc, "Current procedure")
    add_numbered(doc, [
        "The platform loads live definitions and current values for the signed-in Oracle user.",
        "A Service Administrator may enter another exact Oracle username and load that user's assignments; other permitted users update only themselves.",
        "Choose an existing user variable from the live dropdown.",
        "Enter the exact member for the displayed dimension and review the current versus new value.",
        "Approve. The platform rejects a stale current assignment and rereads Oracle to verify the new member.",
    ])
    add_heading(doc, "Compatibility fallback")
    add_para(doc, "If the Oracle version cannot return user-variable definitions directly, the service derives available definitions from current assigned values. If neither definitions nor assignments are available, the page explains that an administrator must configure Oracle or confirm REST support.")
    add_callout(doc, "Current usability boundary", "Variable selection is mouse-first, but the new member is currently an exact text value rather than a live member dropdown. Enter the member exactly as Oracle stores it. A future enhancement can reuse live dimension/member discovery for fully mouse-driven assignment.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 14, "Generate a read-only Excel report", "Report Generation reads Planning data and creates a retained workbook. It does not write values back to Oracle.")
    add_heading(doc, "Choose a source")
    add_table(
        doc,
        ["Source", "Best use", "Current behavior"],
        [
            ["Registered report", "Reusable governed layout or on-premises compatibility", "Uses a stored cube, POV, row and column definition."],
            ["Planning form", "Cloud environments that expose form layout/data through REST", "Uses the exact form name or ID and inspects its live layout."],
        ],
        [2300, 3100, 3960],
    )
    add_numbered(doc, [
        "Choose a registered report or enter an exact Planning form name/ID.",
        "Inspect the layout before continuing. The page shows row, column and page dimensions.",
        "Enter the workbook title and choose or enter each required POV member.",
        "Review source, cube when known, axes and POV; approve the read-only export.",
        "Download the generated .xlsx from the result or Jobs & Activity.",
    ])
    add_heading(doc, "Register a fallback definition")
    add_para(doc, "When direct form inspection is unavailable, an administrator can register a reusable definition: name, title, exact cube, optional fixed POV, at least one row dimension and at least one column dimension. Each dimension appears once, POV uses one member, and row/column member lists use exact names separated by a vertical bar.")
    add_callout(doc, "Registered report is not an Oracle object", "Registration stores the data-slice layout used by this platform. It does not create a Planning form or an Oracle Reports artifact. Oracle security and data access still apply during generation.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 15, "Monitor results and troubleshoot", "Every accepted operation produces durable execution evidence. Start diagnosis from the failed step and retained Oracle details, not from a browser symptom alone.")
    add_figure(doc, assets["evidence"], "Figure 5.6 — Durable execution and evidence.", "Queued request moves through background execution to a retained Jobs and Activity record, with duplicate prevention and troubleshooting evidence.", width=6.35)
    add_heading(doc, "What the result can show")
    add_bullets(doc, [
        "Execution identifier and workflow name.",
        "Queued, running, successful, failed or recovery-required status.",
        "Initiating platform user and display name.",
        "Oracle execution username when configured.",
        "Trigger source: manual, scheduled, API, Excel or AI agent.",
        "Step-by-step status, timestamps and safe retained details.",
        "Oracle job ID, status text and target artifact.",
        "Records read, processed and rejected when Oracle exposes counters.",
        "Generated Excel files and downloadable operation log.",
    ])
    add_heading(doc, "Common failures")
    add_table(
        doc,
        ["Message or symptom", "Meaning", "Next action"],
        [
            ["Artifact not found", "Name is stale, wrong for this application, or not visible to the Oracle identity.", "Synchronize; choose a verified artifact; confirm exact Oracle name and access."],
            ["Authentication failed", "Credentials, identity-domain behavior or Oracle access failed.", "Stop; verify environment and approved credentials. Do not repeatedly retry."],
            ["No jobId or status", "Oracle returned an unexpected submission shape.", "Inspect log and exact Oracle response; confirm endpoint/version compatibility."],
            ["Request timed out", "Browser or one request exceeded its window; the worker may still be running.", "Inspect Jobs & Activity before resubmitting."],
            ["Stale variable", "Oracle changed after review.", "Reload current values and approve a new review."],
            ["Record statistics unavailable", "This Oracle version did not expose counters.", "Use terminal job status, error file and log; do not treat missing counters alone as failure."],
            ["Recovery required", "An accepted execution needs deliberate recovery.", "Review evidence and use only the authorized recovery action offered."],
        ],
        [2500, 3350, 3510],
    )
    add_callout(doc, "Secrets are redacted", "Execution evidence sanitizes credential-shaped fields. Never place passwords, API keys, encrypted-password file contents, session cookies or complete environment files into operation notes, screenshots or support tickets.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 16, "Use the quick-reference runbooks", "These checklists support day-to-day execution, handover, testing and production readiness without requiring users to remember implementation details.")
    add_heading(doc, "Before every Oracle-changing operation")
    add_table(
        doc,
        ["Check", "Question"],
        [
            ["Environment", "Is the displayed Oracle base environment and Planning application correct?"],
            ["Permission", "Am I authorized for this service and for the underlying Oracle artifact?"],
            ["Artifact", "Is the selector live/verified and the exact name or code correct?"],
            ["Business context", "Are year, period, POV, scenario and runtime values correct?"],
            ["File", "Is the selected configured, local or Inbox file the intended current version?"],
            ["Impact", "Do I understand replace, clear, refresh, create or update behavior?"],
            ["Timing", "Are dependent work and active users ready for this operation now?"],
            ["Evidence", "Will I review the terminal result, counters, errors and downstream validation?"],
        ],
        [2200, 7160],
    )
    add_heading(doc, "After a successful load or push")
    add_numbered(doc, [
        "Confirm the execution reached terminal success.",
        "Review record statistics and rejected-record output when applicable.",
        "Open Data Review or an approved report for the target cube and business POV.",
        "Compare source and target with an appropriate tolerance when reconciliation is required.",
        "Complete or submit the linked Planning task only after business validation—not merely technical success.",
    ])
    add_heading(doc, "Operation ownership summary")
    add_table(
        doc,
        ["Platform supplies", "Oracle owns"],
        [
            ["Role-filtered access, live/verified selection and friendly inputs", "Cube design, metadata, forms, jobs, rules, maps, integrations and Pipelines"],
            ["Preflight, normalization, explicit review and approval", "Technical mappings, stage order, calculations, defaults and application security"],
            ["File staging, durable queue, monitoring and evidence", "Job execution and final Planning data/metadata state"],
        ],
        [4680, 4680],
    )
    add_heading(doc, "Current implementation notes")
    add_bullets(doc, [
        "This volume describes the current React application only; the retired interface is excluded.",
        "Oracle Cloud and on-premises compatibility can differ where Oracle REST endpoints expose different catalog or form capabilities; the UI uses explicit fallbacks rather than pretending unsupported discovery exists.",
        "Pipelines and Data Integrations are environment-scoped, registered and synchronized; Business Rules, Data Maps and saved import jobs are read live.",
        "The durable worker continues after the browser leaves and persists identity, trigger and step evidence in PostgreSQL-backed repositories.",
        "The audited operation-focused test set passed 67 tests on 1 September 2026.",
    ])
    add_callout(doc, "End of Document 05", "Use this volume for current standalone Oracle EPM operation procedures. Document 06 continues with the Planning Workspace and full business lifecycle: cycles, stages, assignments, approvals and scheduling.", fill=TEAL_LIGHT, accent=TEAL)

    _finish_document(doc)
    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
