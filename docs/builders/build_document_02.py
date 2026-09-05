"""Build Document 02: Solution Architecture and Technology Stack."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_DARK,
    BLUE_LIGHT,
    CONTENT_DXA,
    GOLD,
    GOLD_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    RED_LIGHT,
    SURFACE,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    add_bullets,
    add_callout,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    configure_styles,
    font,
    rgb,
    set_cell_fill,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    set_table_geometry,
    wrap,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-02"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / (
    "BISP_EPM_Automation_Document_02_Solution_Architecture_and_Technology_Stack.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

PAGE_BG = "#F8FAFD"
INK_HEX = f"#{INK}"
NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"
MUTED_HEX = f"#{MUTED}"
LINE_HEX = f"#{LINE}"


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str) -> None:
    """Draw one directional connector."""
    draw.line((start, end), fill=color, width=5)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 16 * direction, y2 - 10), (x2 - 16 * direction, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 10, y2 - 16 * direction), (x2 + 10, y2 - 16 * direction)]
    draw.polygon(points, fill=color)


def box(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    fill: str,
    outline: str,
    title_color: str = INK_HEX,
) -> None:
    """Draw a rounded architecture component."""
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=20, fill=fill, outline=outline, width=3)
    title_font = font(26, bold=True)
    body_font = font(18)
    title_lines = wrap(draw, title, title_font, x2 - x1 - 34)
    body_lines = wrap(draw, body, body_font, x2 - x1 - 34)
    y = y1 + 18
    for line in title_lines:
        draw.text((x1 + 17, y), line, font=title_font, fill=title_color)
        y += 32
    y += 7
    for line in body_lines:
        draw.text((x1 + 17, y), line, font=body_font, fill=MUTED_HEX)
        y += 24


def diagram_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((55, 35), title, font=font(35, bold=True), fill=NAVY_HEX)
    draw.text((55, 82), subtitle, font=font(20), fill=MUTED_HEX)


def create_system_context(path: Path) -> None:
    image = Image.new("RGB", (1500, 900), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(
        draw,
        "System context",
        "One governed platform between people, automation channels, Oracle EPM, and external providers",
    )
    box(draw, (60, 180, 330, 360), "People", "Planner, Power User, Service Administrator, Viewer", fill="#FFFFFF", outline=BLUE_HEX)
    box(draw, (60, 510, 330, 690), "Other entry points", "Approved schedule, Excel API token, CLI", fill="#FFFFFF", outline=TEAL_HEX)
    box(draw, (450, 150, 810, 340), "React web application", "Role-aware workspace, guided inputs, approvals, live monitoring", fill="#EDF3FF", outline=BLUE_HEX)
    box(draw, (450, 420, 810, 625), "FastAPI application", "Authentication, authorization, use cases, Oracle catalog, API v1", fill="#FFFFFF", outline=NAVY_HEX)
    box(draw, (450, 700, 810, 845), "Durable worker", "Claims queue jobs, calls Oracle, monitors status, stores evidence", fill="#E8F7F2", outline=TEAL_HEX)
    box(draw, (950, 135, 1390, 320), "Oracle EPM Planning", "REST v3, Interop files, jobs, rules, Data Maps, Pipelines, cubes and data", fill="#FFF2EC", outline=ORANGE_HEX)
    box(draw, (950, 385, 1390, 540), "PostgreSQL", "Users, tasks, artifacts, queue, history, audit, agent state", fill="#EDF3FF", outline=BLUE_HEX)
    box(draw, (950, 610, 1160, 800), "AI providers", "Gemini or Groq; model reasoning only", fill="#FFFFFF", outline=f"#{GOLD}")
    box(draw, (1190, 610, 1390, 800), "Notifications", "SMTP today; provider boundary for corporate delivery", fill="#FFFFFF", outline=f"#{GOLD}")
    arrow(draw, (330, 270), (450, 245), BLUE_HEX)
    arrow(draw, (330, 600), (450, 520), TEAL_HEX)
    arrow(draw, (630, 340), (630, 420), NAVY_HEX)
    arrow(draw, (810, 500), (950, 230), ORANGE_HEX)
    arrow(draw, (810, 530), (950, 465), BLUE_HEX)
    arrow(draw, (630, 625), (630, 700), TEAL_HEX)
    arrow(draw, (810, 770), (950, 500), TEAL_HEX)
    arrow(draw, (810, 500), (950, 700), f"#{GOLD}")
    arrow(draw, (810, 545), (1190, 700), f"#{GOLD}")
    image.save(path)


def create_layered_architecture(path: Path) -> None:
    image = Image.new("RGB", (1500, 930), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(draw, "Logical architecture", "Dependencies point inward toward reusable application rules")
    layers = [
        ("Experience layer", "React workspaces, Excel/API clients, CLI", "#EDF3FF", BLUE_HEX),
        ("Delivery layer", "FastAPI /api/v1 routes, sessions, CSRF, typed request/response schemas", "#FFFFFF", NAVY_HEX),
        ("Application layer", "Use cases, operation manager, Planning work, scheduling, approvals, agent service", "#E8F7F2", TEAL_HEX),
        ("Domain layer", "Typed models, permissions, statuses, validation rules, workflow contracts", "#FFFFFF", TEAL_HEX),
        ("Infrastructure layer", "SQLAlchemy repositories, PostgreSQL, Alembic, file storage, logging", "#FFF4D8", f"#{GOLD}"),
        ("Integration layer", "EPMClient, EPM Automate, Gemini/Groq adapters, SMTP provider", "#FFF2EC", ORANGE_HEX),
    ]
    y = 145
    for index, (title, body, fill, outline) in enumerate(layers, start=1):
        draw.rounded_rectangle((150, y, 1350, y + 105), radius=18, fill=fill, outline=outline, width=3)
        draw.rounded_rectangle((175, y + 20, 245, y + 85), radius=14, fill=outline)
        number = f"{index:02d}"
        bbox = draw.textbbox((0, 0), number, font=font(24, bold=True))
        draw.text((210 - (bbox[2] - bbox[0]) / 2, y + 37), number, font=font(24, bold=True), fill="#FFFFFF")
        draw.text((280, y + 18), title, font=font(25, bold=True), fill=NAVY_HEX)
        draw.text((280, y + 58), body, font=font(19), fill=MUTED_HEX)
        if index < len(layers):
            arrow(draw, (750, y + 105), (750, y + 127), MUTED_HEX)
        y += 125
    image.save(path)


def create_execution_flow(path: Path) -> None:
    image = Image.new("RGB", (1600, 760), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(draw, "Governed write path", "Every state-changing request follows the same controlled execution contract")
    steps = [
        ("1", "Discover", "Live artifact and capability"),
        ("2", "Collect", "Typed fields, files and context"),
        ("3", "Preflight", "Permissions and current Oracle state"),
        ("4", "Approve", "Exact target and effect"),
        ("5", "Queue", "Durable PostgreSQL record"),
        ("6", "Execute", "Worker calls Oracle"),
        ("7", "Monitor", "Oracle job to terminal status"),
        ("8", "Evidence", "History, logs, statistics, artifacts"),
    ]
    colors = [
        BLUE_HEX,
        BLUE_HEX,
        TEAL_HEX,
        f"#{GOLD}",
        NAVY_HEX,
        ORANGE_HEX,
        TEAL_HEX,
        f"#{BLUE_DARK}",
    ]
    x = 45
    y = 215
    width = 175
    for index, (number, title, body) in enumerate(steps):
        outline = colors[index]
        draw.rounded_rectangle((x, y, x + width, y + 265), radius=18, fill="#FFFFFF", outline=outline, width=3)
        draw.ellipse((x + 56, y + 20, x + 116, y + 80), fill=outline)
        bbox = draw.textbbox((0, 0), number, font=font(25, bold=True))
        draw.text((x + 86 - (bbox[2] - bbox[0]) / 2, y + 34), number, font=font(25, bold=True), fill="#FFFFFF")
        title_font = font(22, bold=True)
        bbox = draw.textbbox((0, 0), title, font=title_font)
        draw.text((x + width / 2 - (bbox[2] - bbox[0]) / 2, y + 100), title, font=title_font, fill=NAVY_HEX)
        lines = wrap(draw, body, font(17), width - 25)
        ty = y + 147
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font(17))
            draw.text((x + width / 2 - (bbox[2] - bbox[0]) / 2, ty), line, font=font(17), fill=MUTED_HEX)
            ty += 23
        if index < len(steps) - 1:
            arrow(draw, (x + width + 5, y + 132), (x + width + 25, y + 132), outline)
        x += 198
    draw.rounded_rectangle((300, 545, 1300, 675), radius=20, fill="#EDF3FF", outline=BLUE_HEX, width=2)
    draw.text((330, 570), "One shared contract", font=font(24, bold=True), fill=NAVY_HEX)
    draw.text((330, 610), "Web, schedules, Excel/API, CLI and EPM Assistant reuse the same application services and queue.", font=font(19), fill=MUTED_HEX)
    image.save(path)


def create_deployment(path: Path) -> None:
    image = Image.new("RGB", (1500, 900), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(draw, "Deployment topology", "The same queue contract supports a simple developer setup and a scalable production setup")
    draw.rounded_rectangle((45, 135, 715, 830), radius=25, fill="#FFFFFF", outline=BLUE_HEX, width=3)
    draw.text((80, 170), "Development", font=font(31, bold=True), fill=NAVY_HEX)
    draw.text((80, 215), "EPM_EXECUTION_RUNTIME=embedded", font=font(19), fill=BLUE_HEX)
    box(draw, (110, 285, 650, 410), "Vite development server", "React UI at port 5173; API proxy to FastAPI", fill="#EDF3FF", outline=BLUE_HEX)
    box(draw, (110, 455, 650, 590), "FastAPI + embedded worker", "web_main.py at port 8080; queue still written before execution", fill="#E8F7F2", outline=TEAL_HEX)
    box(draw, (110, 635, 365, 770), "PostgreSQL", "Shared durable state", fill="#FFFFFF", outline=NAVY_HEX)
    box(draw, (395, 635, 650, 770), "Oracle EPM", "Cloud or supported on-premises", fill="#FFF2EC", outline=ORANGE_HEX)
    arrow(draw, (380, 410), (380, 455), BLUE_HEX)
    arrow(draw, (300, 590), (240, 635), NAVY_HEX)
    arrow(draw, (480, 590), (520, 635), ORANGE_HEX)

    draw.rounded_rectangle((785, 135, 1455, 830), radius=25, fill="#FFFFFF", outline=TEAL_HEX, width=3)
    draw.text((820, 170), "Production", font=font(31, bold=True), fill=NAVY_HEX)
    draw.text((820, 215), "Web and worker are independently supervised", font=font(19), fill=TEAL_HEX)
    box(draw, (850, 285, 1390, 410), "HTTPS frontend + FastAPI", "Built React assets; FastAPI validates and commits work", fill="#EDF3FF", outline=BLUE_HEX)
    box(draw, (850, 455, 1390, 590), "One or more workers", "Lease-protected claims using FOR UPDATE SKIP LOCKED", fill="#E8F7F2", outline=TEAL_HEX)
    box(draw, (850, 635, 1105, 770), "PostgreSQL", "Queue, audit and business state", fill="#FFFFFF", outline=NAVY_HEX)
    box(draw, (1135, 635, 1390, 770), "Shared storage", "Uploads, logs and generated reports", fill="#FFF4D8", outline=f"#{GOLD}")
    arrow(draw, (1120, 410), (1120, 455), TEAL_HEX)
    arrow(draw, (1020, 590), (980, 635), NAVY_HEX)
    arrow(draw, (1240, 590), (1260, 635), f"#{GOLD}")
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "context": ASSET_DIR / "system-context.png",
        "layers": ASSET_DIR / "logical-architecture.png",
        "flow": ASSET_DIR / "governed-write-path.png",
        "deployment": ASSET_DIR / "deployment-topology.png",
    }
    create_system_context(assets["context"])
    create_layered_architecture(assets["layers"])
    create_execution_flow(assets["flow"])
    create_deployment(assets["deployment"])
    return assets


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    add_page_break(doc)
    kicker = doc.add_paragraph(style="Kicker")
    kicker_run = kicker.add_run(f"CHAPTER {number}")
    set_run_font(kicker_run, size=9, color=BLUE, bold=True)
    kicker.paragraph_format.space_before = Pt(4)
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_micro_spacer(doc: Document) -> None:
    """Force a small visual boundary between adjacent Word paragraphs."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(2)
    run = paragraph.add_run(" ")
    set_run_font(run, size=1, color=WHITE)


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.space_after = Pt(8 if level == 2 else 6)
    paragraph.paragraph_format.keep_with_next = True
    add_micro_spacer(doc)


def add_bullets(doc: Document, items: list[str]) -> None:
    """Add visually separated bullets for reliable Word/PDF rendering."""
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        run = paragraph.add_run(f"•\t{item}")
        set_run_font(run)


def add_code_block(doc: Document, code: str, caption: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_fill(cell, NAVY)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(code.splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=8.5, color=WHITE)
    if caption:
        p = doc.add_paragraph(caption, style="Caption")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(1)


def mark_table_header_row(table) -> None:
    """Mark the first row as a Word header for assistive technologies."""
    first_row = table.rows[0]
    tr_pr = first_row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def configure_header_footer(section) -> None:
    for header in (section.header, section.even_page_header):
        paragraph = header.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run("BISP SOLUTIONS  /  SOLUTION ARCHITECTURE")
        set_run_font(run, size=8.5, color=MUTED, bold=True)

    for footer in (section.footer, section.even_page_footer):
        table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
        set_table_geometry(table, [7000, 2360], indent_dxa=0)
        left = table.cell(0, 0).paragraphs[0]
        left.text = "Document 02  |  Solution Architecture and Technology Stack"
        set_run_font(left.runs[0], size=8.5, color=MUTED)
        add_page_number(table.cell(0, 1).paragraphs[0])
        for cell in table.rows[0].cells:
            set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            top = OxmlElement("w:top")
            top.set(qn("w:val"), "single")
            top.set(qn("w:sz"), "6")
            top.set(qn("w:color"), LINE)
            borders.append(top)
            tc_pr.append(borders)


def add_cover(doc: Document, assets: dict[str, Path]) -> None:
    if LOGO.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(38)
    doc.add_paragraph("ENTERPRISE TECHNICAL GUIDE", style="Kicker")
    doc.add_paragraph("Solution Architecture\nand Technology Stack", style="Title")
    doc.add_paragraph(
        "BISP Solutions Oracle EPM Automation Platform",
        style="Subtitle",
    )
    doc.add_paragraph(
        "A beginner-friendly explanation of how the modern React experience, FastAPI services, PostgreSQL data platform, durable worker, Oracle EPM integrations, and governed LangGraph assistant work together.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [["02 of the platform handbook", "24 August 2026", "Business, functional, technical, security and operations teams"]],
        [2200, 2100, 5060],
    )
    add_callout(
        doc,
        "Scope statement",
        "This guide documents the current modern React/FastAPI platform and the latest implemented architecture. Retired server-rendered UI components and superseded experimental flows are intentionally excluded.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BISP Solutions  |  Production-oriented Oracle EPM automation")
    set_run_font(run, size=10, color=MUTED, bold=True)


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc, assets)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS GUIDE", style="Kicker")
    doc.add_paragraph("Architecture before implementation detail", style="Heading 1")
    doc.add_paragraph(
        "Start with the system diagrams and plain-language explanations. Technical readers can then use the stack tables, source-tree map, code excerpts, and deployment guidance as an implementation reference.",
        style="Lead",
    )
    add_heading(doc, "What this document answers")
    add_bullets(doc, [
        "Which systems and people participate in the platform?",
        "Where do presentation, business rules, persistence, Oracle communication, and AI orchestration live?",
        "How does a request move safely from a screen or conversation to Oracle EPM?",
        "Which technologies are used today, and why were they selected?",
        "How do security, migrations, monitoring, testing, deployment, and future extensions work?",
    ])
    add_heading(doc, "Contents")
    add_table(
        doc,
        ["Chapter", "Focus"],
        [
            ["1", "Architecture at a glance"],
            ["2", "Design principles and system boundaries"],
            ["3", "Modern frontend architecture"],
            ["4", "FastAPI backend and clean application structure"],
            ["5", "Oracle EPM integration architecture"],
            ["6", "PostgreSQL persistence and Alembic migrations"],
            ["7", "Durable execution, workers, scheduling and recovery"],
            ["8", "Security, access control and federated identity"],
            ["9", "LangGraph EPM Assistant architecture"],
            ["10", "Cross-cutting platform services"],
            ["11", "Testing and engineering quality"],
            ["12", "Deployment topology and operational model"],
            ["13", "Repository map and extension patterns"],
        ],
        [1450, 7910],
    )
    add_callout(
        doc,
        "Important distinction",
        "The platform login identity and Oracle Planning service credentials solve different problems. Platform identity controls who may use features. The server-managed Oracle connection controls what the configured Oracle account can access.",
        fill=GOLD_LIGHT,
        accent=GOLD,
    )

    add_chapter(doc, 1, "Architecture at a glance", "The product is a governed control layer around Oracle EPM, not a replacement for Oracle's calculation, metadata, integration, or security engines.")
    add_figure(
        doc,
        assets["context"],
        "Figure 1. Current system context for the modern platform.",
        "System context connecting users and other entry points to the React application, FastAPI backend, durable worker, Oracle EPM, PostgreSQL, AI providers, and notifications.",
        width=6.45,
    )
    add_heading(doc, "The simplest mental model")
    add_numbered(doc, [
        "A user starts from the React application, an approved schedule, Excel/API integration, the CLI, or the EPM Assistant.",
        "FastAPI authenticates the caller, checks permissions, validates the requested artifact and inputs, and records approved work.",
        "PostgreSQL holds durable business state, queue state, audit evidence, catalog data, and agent state.",
        "A worker executes long-running work through Oracle Planning REST or, where deliberately configured, EPM Automate.",
        "The platform monitors Oracle status and returns the same evidence regardless of how the request started.",
    ])
    ownership = doc.add_paragraph("Ownership boundaries", style="Heading 2")
    ownership.paragraph_format.page_break_before = True
    ownership.paragraph_format.line_spacing = 1.1
    ownership.paragraph_format.space_after = Pt(8)
    add_table(
        doc,
        ["System", "Primary responsibility", "Does not own"],
        [
            ["Oracle EPM Planning", "Application design, cubes, dimensions, rules, jobs, Pipelines, Data Maps, integrations and data", "Platform navigation, cross-channel queueing, platform audit and AI governance"],
            ["BISP platform", "Business-friendly access, discovery, validation, approval, execution control, monitoring and evidence", "Reimplementation of Oracle calculation or integration engines"],
            ["Model provider", "Natural-language interpretation and tool selection within an allow-list", "Oracle credentials, authorization policy, final approval or direct arbitrary REST access"],
        ],
        [2200, 3900, 3260],
    )

    add_chapter(doc, 2, "Design principles and system boundaries", "The architecture is deliberately modular so one feature can evolve without duplicating Oracle logic or coupling the product to one interface.")
    add_figure(
        doc,
        assets["layers"],
        "Figure 2. Logical layers and dependency direction.",
        "Six logical layers: experience, delivery, application, domain, infrastructure, and external integration.",
        width=6.35,
    )
    add_heading(doc, "Core principles")
    add_table(
        doc,
        ["Principle", "How the implementation applies it"],
        [
            ["One source of business truth", "React, schedules, Excel/API, CLI and AI call shared application services instead of maintaining separate Oracle logic."],
            ["Oracle owns Oracle configuration", "Saved jobs, rules, Pipelines, integrations and Data Maps remain defined in Oracle; the platform discovers and executes them."],
            ["Read before write", "Live catalogs, input discovery and preflight happen before a state-changing request can be approved."],
            ["Durability before convenience", "Approved work is committed to PostgreSQL before a worker performs long Oracle activity."],
            ["Provider-neutral boundaries", "Oracle REST/EPM Automate, AI providers and notification providers are isolated behind replaceable adapters."],
            ["Fail visibly", "Typed exceptions, step history, Oracle diagnostics and recovery states replace silent retries and ambiguous failures."],
        ],
        [2600, 6760],
    )
    add_heading(doc, "Clean architecture in practical language")
    add_para(doc, "The React UI knows how to present choices, but it does not know how to construct Oracle job payloads. FastAPI routes know how to receive and authorize requests, but they delegate business behavior. Application services coordinate use cases. Domain models describe validated inputs and states. Infrastructure repositories and external clients perform database and network work.")
    add_callout(doc, "Extension rule", "A new channel or AI provider should reuse existing application services. A new Oracle operation should add a typed model, focused service, application orchestration, API contract and tests - not logic inside a React component.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 3, "Modern frontend architecture", "The supported user experience is a React and TypeScript single-page application designed for planners, consultants, administrators and executives.")
    add_heading(doc, "Current frontend stack")
    add_table(
        doc,
        ["Technology", "Current constraint", "Why it is used"],
        [
            ["React", "19.2.8", "Composable workspaces, predictable state-driven UI and reusable guided-operation components."],
            ["TypeScript", "7.0.2", "Compile-time contracts for API payloads, operation inputs, users, cycles, jobs and agent messages."],
            ["Vite", "8.2.1", "Fast development server, production bundling and API proxying."],
            ["CSS", "Application-owned responsive design system", "A tailored BISP enterprise experience without a generic admin-template dependency."],
            ["Vitest + Testing Library", "Vitest 4.1.10; Testing Library 16.3.2", "Behavior-focused component tests in a jsdom environment."],
            ["pnpm", "11.16.0", "Deterministic JavaScript dependency management through the lockfile."],
        ],
        [2100, 1900, 5360],
    )
    add_heading(doc, "Supported modern workspaces")
    add_bullets(doc, [
        "Home, My Work, Planning Cycles, Approvals and Notifications for the business lifecycle.",
        "Operations for independent Oracle services, including rules, imports, integrations, Pipelines, Data Maps, variables and Cube Refresh.",
        "Data Review for live cube, dimension and member selection, grids, comparison and validation evidence.",
        "Reports, Jobs & Activity, Scheduling, Access Control and EPM Assistant.",
    ])
    add_heading(doc, "Frontend/backend boundary")
    add_para(doc, "The browser calls the typed /api/v1 contract. It never receives the Oracle password. Security-sensitive validation, permission checks, artifact verification and queue submission always remain server-side. The UI provides progressive disclosure, confirmations, loading states, actionable errors and live polling; it does not become the authority for execution.")
    add_code_block(
        doc,
        "function versionedApiPath(path: string): string {\n  if (!path.startsWith(\"/api/\") || path.startsWith(\"/api/v1/\")) return path;\n  return `/api/v1/${path.slice(\"/api/\".length)}`;\n}\n\nasync function request<T>(path: string, options?: RequestInit): Promise<T> {\n  // Centralized browser request handling returns a typed result.\n}",
        "Code excerpt 1. The frontend centralizes versioned API access instead of scattering HTTP details across screens.",
    )
    add_callout(doc, "Migration-friendly UI", "Because business behavior is behind API contracts, React can be reorganized or even replaced later without rewriting Oracle services, queueing, audit, security or agent governance.")

    add_chapter(doc, 4, "FastAPI backend and clean application structure", "FastAPI is the server boundary; Python application services remain the reusable core.")
    add_heading(doc, "Current backend stack")
    add_table(
        doc,
        ["Technology", "Purpose"],
        [
            ["Python 3.11+", "Typed implementation language for Oracle integration, workflows, reporting, worker and AI orchestration."],
            ["FastAPI 0.115+", "Versioned JSON API, authentication endpoints, upload/download boundaries and application composition."],
            ["Uvicorn 0.30+", "ASGI runtime for the FastAPI service."],
            ["Pydantic (through FastAPI)", "Request parsing and deterministic validation at the HTTP boundary."],
            ["python-dotenv", "Environment-specific configuration without hardcoded secrets."],
            ["Python logging", "Centralized operational, security, Oracle request and worker diagnostics."],
            ["openpyxl", "Professional Excel report and validation workbook generation."],
        ],
        [2600, 6760],
    )
    add_heading(doc, "Versioned API contract")
    add_para(doc, "The modern frontend uses /api/v1. Compatibility aliases can remain temporarily for older integrations, but new functionality belongs in the versioned contract. Routes enforce the signed session and ask the application layer for business results.")
    add_code_block(
        doc,
        "@router.get(\"/operations\", response_model=OperationsResponse)\nasync def standalone_operations(request: Request) -> OperationsResponse:\n    require_api_session(request)\n    user = current_user(request)\n    operations = [\n        operation for operation in request.app.state.operation_catalog.definitions()\n        if user.has_permission(_operation_permission(operation.code))\n    ]\n    return OperationsResponse(status=\"success\", operations=operations)",
        "Code excerpt 2. The server filters the Operations catalog using the current user's permissions.",
    )
    add_heading(doc, "Module responsibilities")
    add_table(
        doc,
        ["Package", "Responsibility"],
        [
            ["app/api and app/web", "HTTP contracts, sessions, CSRF, uploads, redirects and composition."],
            ["app/application", "Use-case coordination for operations, Planning work, execution, scheduling, reports and variables."],
            ["app/models", "Typed domain records, statuses, requests and results."],
            ["app/services", "Focused Oracle operations, catalogs, repositories, workflows, rendering and notifications."],
            ["app/clients and app/automation", "Reusable Oracle REST session and safe EPM Automate command execution."],
            ["app/infrastructure", "Database engine, tables and migration checks."],
            ["app/agent and app/identity", "Governed LangGraph assistant and local/federated identity adapters."],
        ],
        [2900, 6460],
    )

    add_chapter(doc, 5, "Oracle EPM integration architecture", "Oracle public APIs and configured Oracle artifacts remain the authoritative execution surface.")
    add_heading(doc, "REST-first shared client")
    add_para(doc, "EPMClient creates one requests.Session, applies HTTP Basic Authentication, common headers, SSL and timeout policy, logs request metadata, normalizes URLs, and translates transport or API failures into platform exceptions. Future services use this client rather than opening their own sessions.")
    add_code_block(
        doc,
        "self._session = session or requests.Session()\nself._session.auth = HTTPBasicAuth(\n    settings.epm_username,\n    settings.require_rest_password(),\n)\nself._session.headers.update({\n    \"Accept\": \"application/json\",\n    \"User-Agent\": \"oracle-planning-automation/0.1.0\",\n})",
        "Code excerpt 3. Authentication and connection reuse are centralized in EPMClient.",
    )
    add_heading(doc, "Supported Oracle communication paths")
    add_table(
        doc,
        ["Path", "When used", "Architecture note"],
        [
            ["Planning REST v3", "Primary engine for authentication, jobs, rules, maps, data, forms, variables and monitoring", "Preferred where Oracle exposes a supported endpoint."],
            ["Interop/file APIs", "Inbox listing, upload, download, replacement and logs", "File concerns are centralized and purpose-filtered in the UI."],
            ["EPM Automate", "Optional alternative for selected imports, integrations, rules, Pipelines and Data Maps", "Executed as argument lists without an interactive shell; encrypted password file supported."],
            ["Oracle OIDC", "Human platform sign-in when Oracle Cloud Identity is configured", "Separate from REST service credentials used for automation."],
        ],
        [2100, 3900, 3360],
    )
    add_heading(doc, "Cloud and on-premises compatibility")
    add_para(doc, "Settings can infer Cloud hostnames or accept an explicit cloud/on-premises deployment mode. Capability detection and supported fallbacks handle REST differences; the platform does not pretend an unavailable Oracle endpoint exists. Registered definitions remain useful for features that older releases cannot discover directly.")
    add_heading(doc, "Unified Oracle artifact catalog")
    add_para(doc, "Environment-scoped catalog records unify live rules, Data Maps, data and metadata jobs, Cube Refresh jobs, cubes, registered Pipelines and Data Integrations. Synchronization verifies current Oracle state, hides stale entries from selectors and preserves historical evidence instead of deleting audit records after a temporary outage or environment switch.")
    add_callout(doc, "Oracle remains authoritative", "The platform can create a safe registration record for artifacts that need local metadata, but it does not create arbitrary Oracle jobs, rules, Pipeline stages or mappings through undocumented APIs.", fill=ORANGE_LIGHT, accent=ORANGE)

    add_chapter(doc, 6, "PostgreSQL persistence and Alembic migrations", "Production state is relational, transactional and versioned; SQLite is retained only as an isolated automated-test compatibility path.")
    add_heading(doc, "Database technology")
    add_table(
        doc,
        ["Component", "Role"],
        [
            ["PostgreSQL", "Production system of record and concurrency foundation."],
            ["SQLAlchemy 2.x Core", "Explicit SQL tables, transactions, connection pooling and repository queries."],
            ["psycopg 3", "PostgreSQL driver, including binary package for straightforward deployment."],
            ["Alembic 1.16+", "Ordered, reviewable and reversible application-schema migrations."],
            ["LangGraph PostgreSQL checkpointer", "Framework-owned graph checkpoints in tables separate from Alembic-managed business tables."],
        ],
        [2900, 6460],
    )
    add_heading(doc, "Current application table domains")
    add_table(
        doc,
        ["Domain", "Representative tables", "Purpose"],
        [
            ["Access and identity", "platform_users, roles, permissions, authentication_events, identity_providers, external identities/entitlements, mappings and sync runs", "Local access, Oracle identity synchronization, role mapping and sign-in evidence."],
            ["Oracle catalog", "oracle_artifacts, business_rule_rtp_sync_runs, RTP definitions and parameters", "Environment-scoped artifact lifecycle and structured Business Rule prompts."],
            ["Planning lifecycle", "planning_cycles, stages, tasks, dependencies, executions, validations, approvals and notifications", "Business process ownership, readiness, review and decisions."],
            ["Execution", "workflow_runs, workflow_steps, execution_queue, process_schedules", "Durable work, step evidence, leases, recovery and recurring triggers."],
            ["Agent", "agent_conversations, messages, tool activities, action drafts and action decisions", "Auditable conversations, governed proposals and human decisions."],
            ["API integration", "api_tokens", "Hashed, scoped, revocable credentials for Excel and approved external clients."],
        ],
        [2100, 3800, 3460],
    )
    add_heading(doc, "Migration policy")
    add_numbered(doc, [
        "A developer adds an Alembic revision for every production schema change.",
        "Deployment applies revisions with alembic upgrade head before the web or worker service starts.",
        "Startup compares the installed database head with the code head and refuses to run when they differ.",
        "Repositories use the schema; they do not create or alter production tables at runtime.",
        "LangGraph's idempotent checkpoint setup runs separately after Alembic when installing or upgrading the graph framework.",
    ])
    add_code_block(
        doc,
        "alembic upgrade head\nalembic current --check-heads\npython -m app.agent.checkpoint_setup",
        "Code excerpt 4. Production database and LangGraph checkpoint initialization are intentionally separate.",
    )
    add_callout(doc, "Latest schema snapshot", "The codebase currently includes Alembic revisions through 0016_scoped_business_rule_rtps, which adds environment/cube-aware runtime-prompt registry support.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 7, "Durable execution, workers, scheduling and recovery", "Long Oracle work is decoupled from the browser request while preserving exactly which user, channel, artifact and inputs started it.")
    add_figure(
        doc,
        assets["flow"],
        "Figure 3. Shared governed write path for all supported channels.",
        "Eight-step governed flow from discovery and typed input through approval, durable queueing, worker execution, monitoring, and evidence.",
        width=6.45,
    )
    add_heading(doc, "Why PostgreSQL leases are used")
    add_para(doc, "Workers claim queued rows transactionally. PostgreSQL FOR UPDATE SKIP LOCKED permits multiple workers to process different jobs without claiming the same item. Each running job carries a lease owner, expiry and heartbeat. This is safer and more scalable than keeping work only in a web-process thread.")
    add_code_block(
        doc,
        "statement = (\n    select(execution_queue)\n    .where(status == \"QUEUED\", available_at <= now)\n    .order_by(priority, created_at)\n    .limit(1)\n    .with_for_update(skip_locked=True)\n)\n# The selected row is atomically changed to RUNNING with a lease.",
        "Code excerpt 5. Simplified lease-protected claim pattern used by the durable worker.",
    )
    add_heading(doc, "Failure and recovery semantics")
    add_bullets(doc, [
        "A normal Oracle failure becomes FAILED with its diagnostics and terminal evidence.",
        "A worker that disappears after Oracle may have accepted work creates an uncertain outcome.",
        "Expired running leases are quarantined as RECOVERY_REQUIRED rather than automatically retried.",
        "An authorized user reviews the Oracle Job Console before starting a linked recovery execution.",
        "Schedules use atomic claims so one due occurrence is submitted once even when multiple processes are active.",
    ])
    add_heading(doc, "Trigger attribution")
    add_para(doc, "Every execution records whether it originated from MANUAL, SCHEDULED, API, EXCEL or AI_AGENT. This gives auditors one common history without losing the business channel that initiated the request.")

    add_chapter(doc, 8, "Security, access control and federated identity", "Authorization is enforced in the backend and is independent from what a browser chooses to display.")
    add_heading(doc, "Four current business roles")
    add_table(
        doc,
        ["Role", "Typical capability boundary"],
        [
            ["Service Administrator", "All platform permissions, including users, cycles, schedules, catalogs, variables, design and operations."],
            ["Power User", "Approved processes and operations, own/user-variable work, Data Review, reports, history and Assistant; no user administration or process design."],
            ["User", "Assigned Planning work, approved process runs, own user variables, Data Review, reports and Assistant."],
            ["Viewer", "Reports and read-only Assistant guidance without operational controls."],
        ],
        [2300, 7060],
    )
    add_heading(doc, "Authentication options")
    add_bullets(doc, [
        "Local platform accounts use versioned scrypt password hashing, random salts, failed-login throttling and a timed lockout.",
        "Oracle Cloud federated sign-in uses OIDC discovery and Authorization Code flow with state, nonce and PKCE S256.",
        "The OIDC callback verifies token/UserInfo subjects and signs in only an already approved synchronized platform identity.",
        "Oracle identity synchronization, entitlement mapping and platform provisioning use preview-and-apply checksums to prevent stale administrative changes.",
    ])
    add_heading(doc, "Web and API controls")
    add_table(
        doc,
        ["Control", "Implementation intent"],
        [
            ["Signed server sessions", "The browser receives a session, not Oracle credentials."],
            ["CSRF validation", "State-changing browser requests must carry the server-issued token."],
            ["Secure-cookie configuration", "Production HTTPS deployments can require secure cookies."],
            ["Scoped API tokens", "Excel and approved clients use revocable, expiring tokens stored only as hashes."],
            ["Permission-aware APIs", "Route and application checks prevent UI manipulation from granting access."],
            ["Evidence redaction", "Credential-shaped fields are removed before history is returned to a browser."],
        ],
        [2700, 6660],
    )
    add_callout(doc, "Separation of duties", "Oracle Cloud OIDC proves a human identity to the platform. Basic Authentication in EPMClient authenticates the server's configured automation connection to Oracle Planning. Neither replaces the other.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 9, "LangGraph EPM Assistant architecture", "The Assistant is a governed orchestration channel over the same platform services, not an autonomous Oracle administrator.")
    add_heading(doc, "Current AI stack")
    add_table(
        doc,
        ["Component", "Current implementation", "Reason"],
        [
            ["Graph orchestrator", "LangGraph 1.2.7+", "Explicit model, tool and approval nodes with durable PostgreSQL checkpoints."],
            ["Model providers", "Google Gemini or Groq", "Provider adapters allow configuration changes without rewriting business tools or governance."],
            ["Provider SDKs", "google-genai 2.13+ and groq 1.6+", "Current supported client methods and structured tool calls."],
            ["Application audit", "PostgreSQL agent tables", "Conversations, tool activity, drafts and decisions are separate from framework checkpoints."],
            ["Capability gateway", "Permission-aware allow-list", "The model can request only predefined read or preparation tools."],
        ],
        [2100, 3300, 3960],
    )
    add_heading(doc, "Graph flow")
    add_numbered(doc, [
        "The service loads the signed-in user, recent conversation history and allowed tool names.",
        "Deterministic intent helpers resolve known operation patterns before relying on model judgment where possible.",
        "The model may request an allow-listed inspection or preparation tool; the gateway rechecks permissions.",
        "Live Oracle catalogs and explainable artifact matching prevent invented names from reaching approval.",
        "Structured cards collect exact periods, files, modes, RTPs, variables, members and other required values.",
        "LangGraph interrupts at clarification, input or approval boundaries and persists resumable state.",
        "Only an explicit UI approval can submit a validated operation or multi-step standalone flow to the durable queue.",
        "The worker executes it with AI_AGENT attribution using the same monitoring and evidence path as manual work.",
    ])
    add_code_block(
        doc,
        "builder.add_node(\"model\", self._model_node)\nbuilder.add_node(\"approval\", self._approval_node)\nbuilder.add_node(\"tools\", self._tool_node)\nbuilder.add_edge(START, \"model\")\nbuilder.add_conditional_edges(\n    \"model\", self._route_after_model,\n    {\"approval\": \"approval\", \"tools\": \"tools\", \"end\": END},\n)\nbuilder.add_edge(\"approval\", \"tools\")\nbuilder.add_edge(\"tools\", \"model\")",
        "Code excerpt 6. The current LangGraph topology makes tool use and approval explicit.",
    )
    add_heading(doc, "Current tool families")
    add_bullets(doc, [
        "Environment, operation catalog, recent history and execution evidence inspection.",
        "Live cube, dimension and member discovery; read-only slice review and comparison.",
        "Live Oracle artifact listing and task-based matching for rules, Data Maps, Pipelines, integrations, imports and variables.",
        "Single-operation preparation and approved multi-operation standalone flows.",
    ])
    add_heading(doc, "Guardrails")
    add_bullets(doc, [
        "Conversation text is not approval; the Assistant cannot approve its own proposal.",
        "The tool allow-list is intersected with the signed-in user's permissions.",
        "Artifacts and guided inputs are revalidated immediately before queue submission.",
        "The model cannot generate arbitrary Oracle REST calls, delete Oracle objects, schedule work or silently retry a write.",
        "Groq history and large tool results are compacted to respect account token limits while preserving the latest request and deterministic selection state.",
    ])
    add_callout(doc, "Provider independence", "Switching between Gemini and Groq changes the model adapter and environment configuration. The LangGraph state machine, tools, permission checks, approval UI, operation manager and Oracle services remain unchanged.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 10, "Cross-cutting platform services", "Several shared capabilities keep Oracle operations consistent regardless of feature or entry point.")
    add_table(
        doc,
        ["Capability", "Architecture"],
        [
            ["Configuration", "A frozen Settings dataclass loads .env and process variables, validates ranges and provides deployment-mode helpers."],
            ["Logging and exceptions", "Central logging plus a typed EPM exception hierarchy separates configuration, authentication, transport, operation, timeout and recovery failures."],
            ["Job monitoring", "Shared polling normalizes Oracle status, timeouts, child jobs, diagnostics and record statistics."],
            ["Files", "Server-side upload receipts, Oracle Inbox discovery, exact replacement behavior, controlled shared runtime storage and protected downloads."],
            ["Data Review", "Live cube/dimension/member discovery creates read-only grids, comparisons, tolerances, quality rules and Planning-task validation evidence."],
            ["Reporting", "Registered data-slice or supported form exports are rendered into downloadable Excel workbooks through openpyxl."],
            ["Notifications", "Operation outcomes generate provider-neutral events; SMTP is implemented and delivery failure cannot rewrite Oracle success/failure."],
            ["Catalog sync", "A last-known-good environment-scoped artifact registry separates temporary discovery failure from confirmed stale state."],
        ],
        [2500, 6860],
    )
    add_heading(doc, "Why Data Review and reporting remain separate")
    add_para(doc, "Data Review is an interactive, read-only analysis surface with live selections and reconciliation. Reporting creates a durable business artifact, usually Excel, from an approved definition or supported form export. Both share Oracle data services but serve different user outcomes.")
    add_heading(doc, "Why the platform keeps typed operation inputs")
    add_para(doc, "Business Rules, imports, Data Integrations, Pipelines, Data Maps, Cube Refresh, substitution variables and user variables each have different required fields and risk. Dedicated dataclasses and validators preserve those differences while the common operation manager standardizes approval, queueing and monitoring.")

    add_chapter(doc, 11, "Testing and engineering quality", "The test strategy protects business behavior without requiring a live Oracle environment in the automated suite.")
    add_heading(doc, "Current quality stack")
    add_table(
        doc,
        ["Area", "Tools and approach"],
        [
            ["Python", "pytest 8.x with mocked requests sessions, process runners, repositories and application services."],
            ["Frontend", "Vitest, React Testing Library and jsdom for role-aware screens, workflows and feedback states."],
            ["Type safety", "Python type hints and TypeScript API/domain contracts."],
            ["Database", "Migration-architecture tests, repository tests and test-only isolated SQLite compatibility where appropriate."],
            ["Agent", "Intent, matching, graph interrupts/resume, provider adapters, approvals, permissions and deterministic operation regression tests."],
            ["Oracle integration", "Payload, URL, timeout, status, replacement, job diagnostics and command construction tests without live execution."],
        ],
        [2200, 7160],
    )
    add_heading(doc, "Representative verification commands")
    add_code_block(
        doc,
        "python -m pytest\n\ncd frontend\npnpm test\npnpm build",
        "Code excerpt 7. Backend tests, frontend behavior tests and a production TypeScript/Vite build form the basic release gate.",
    )
    add_page_break(doc)
    add_heading(doc, "Engineering safeguards")
    add_bullets(doc, [
        "Small operation-specific services avoid one unmaintainable Oracle client.",
        "Dependency injection and test factories isolate network, storage and provider behavior.",
        "Startup schema checks prevent code from running against an unexpected database revision.",
        "No automated test is expected to call a real Oracle environment or installed EPM Automate executable.",
        "Backward-compatible API aliases can be retired only after clients migrate to /api/v1.",
    ])

    add_chapter(doc, 12, "Deployment topology and operational model", "Development keeps setup simple; production separates web and worker responsibilities for reliability and scale.")
    add_figure(
        doc,
        assets["deployment"],
        "Figure 4. Development and production deployment patterns.",
        "Comparison of an embedded development worker with independently supervised production web and worker services.",
        width=6.35,
    )
    add_heading(doc, "Development startup")
    add_code_block(doc, "python web_main.py\n\ncd frontend\npnpm dev", "Code excerpt 8. Embedded development mode still writes the durable queue before local execution.")
    add_heading(doc, "Production startup")
    add_code_block(doc, "# API service\npython web_main.py\n\n# Separately supervised execution service\npython worker.py", "Code excerpt 9. Production runs FastAPI and the durable worker as separate processes.")
    add_heading(doc, "Production requirements")
    add_bullets(doc, [
        "Set EPM_EXECUTION_RUNTIME=web for the API process and run one or more worker processes.",
        "Use the same DATABASE_URL and controlled RUNTIME_DATA_DIR across services.",
        "Build the React application with pnpm build and serve frontend/dist over HTTPS.",
        "Route /api, authentication, downloads and compatibility entry points to FastAPI.",
        "Apply Alembic and LangGraph checkpoint setup before application startup.",
        "Use a stable WEB_SESSION_SECRET, secure cookies, TLS and managed secret injection.",
        "Supervise web and workers, centralize logs and alert on FAILED or RECOVERY_REQUIRED executions.",
    ])
    add_callout(doc, "Horizontal worker scale", "Multiple workers can run safely because PostgreSQL row locking prevents duplicate claims. Scaling does not remove the need to respect Oracle concurrency limits and job dependencies.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 13, "Repository map and extension patterns", "The source tree reflects the architectural boundaries described in this guide.")
    add_heading(doc, "Current high-level structure")
    add_code_block(
        doc,
        "frontend/                 React + TypeScript modern UI\napp/api/v1/               Versioned frontend API\napp/web/                  FastAPI composition, sessions and uploads\napp/application/          Reusable use cases and execution coordination\napp/models/               Typed domain models and statuses\napp/services/             Oracle operations, catalogs and repositories\napp/clients/              Shared Planning REST client\napp/automation/           Safe EPM Automate adapter\napp/infrastructure/       SQLAlchemy engine, tables and migration checks\napp/agent/                LangGraph, providers, tools and agent persistence\napp/identity/             Oracle identity sync and OIDC adapter\nalembic/versions/         Ordered production schema changes\ntests/                    Backend regression suite\nweb_main.py               FastAPI entry point\nworker.py                 Durable execution/schedule worker",
        "Code excerpt 10. The current source tree follows explicit responsibility boundaries.",
    )
    add_heading(doc, "How to add a new Oracle operation")
    add_numbered(doc, [
        "Confirm that Oracle exposes a supported REST or EPM Automate capability and document prerequisites.",
        "Create typed request/result models and focused validation rules.",
        "Implement a service that uses EPMClient or the EPM Automate adapter and shared JobMonitor behavior.",
        "Add an application-level input type and operation-manager execution path.",
        "Expose a permission-protected /api/v1 contract and a reusable React runner.",
        "Add catalog discovery, evidence, notification and file handling only when the operation needs them.",
        "Write unit, application, route and UI tests before adding an Assistant capability.",
        "If the Assistant may use it, add a narrow tool/preparation definition and retain explicit approval for writes.",
    ])
    add_heading(doc, "How to add another AI provider")
    add_numbered(doc, [
        "Implement the AgentProvider protocol for model generation and provider-native tool-result continuation.",
        "Normalize provider tool calls into AgentToolCall and never execute them inside the adapter.",
        "Add validated environment configuration and keep provider secrets server-side.",
        "Run the shared graph, capability, approval and persistence tests plus provider-specific correlation tests.",
    ])
    add_heading(doc, "Architecture decision summary")
    add_table(
        doc,
        ["Decision", "Outcome"],
        [
            ["React is the only supported presentation layer", "No duplicate modern/legacy UI behavior."],
            ["FastAPI owns the backend contract", "All clients share authentication, authorization and application services."],
            ["PostgreSQL is mandatory in normal runtime", "Reliable transactions, queues, locks, audit and migrations."],
            ["Oracle Pipeline owns repeatable Oracle technical sequences", "The platform avoids duplicating configuration that belongs in Oracle."],
            ["LangGraph governs the Assistant", "Durable state, explicit interrupts and provider-neutral tool orchestration."],
            ["Human approval is a platform event", "A chat response alone can never execute a state-changing Oracle action."],
        ],
        [3500, 5860],
    )
    add_callout(doc, "Final takeaway", "The architecture is designed so convenience layers can change while the governed core remains stable: typed use cases, server-side permissions, live Oracle verification, durable PostgreSQL execution, explicit approval and auditable evidence.", fill=BLUE_LIGHT, accent=BLUE)

    add_page_break(doc)
    doc.add_paragraph("IMPLEMENTATION REFERENCE", style="Kicker")
    doc.add_paragraph("Technology stack at a glance", style="Heading 1")
    doc.add_paragraph("Current direct dependency constraints from the project manifests.", style="Lead")
    add_table(
        doc,
        ["Layer", "Technologies"],
        [
            ["Frontend", "React 19.2.8, React DOM 19.2.8, TypeScript 7.0.2, Vite 8.2.1, pnpm 11.16.0"],
            ["Frontend testing", "Vitest 4.1.10, Testing Library React 16.3.2, jsdom 30.0.1"],
            ["Backend", "Python 3.11+, FastAPI 0.115+, Uvicorn 0.30+, requests 2.32+, python-dotenv 1.0+"],
            ["Persistence", "PostgreSQL, SQLAlchemy 2.x, Alembic 1.16+, psycopg 3.2+"],
            ["AI agent", "LangGraph 1.2.7+, langgraph-checkpoint-postgres 3.1+, google-genai 2.13+, groq 1.6+"],
            ["Identity and HTTP", "Authlib 1.7.2+, httpx 0.28+, itsdangerous 2.2+"],
            ["Documents and time", "openpyxl 3.1+, tzdata 2025.2+"],
            ["Testing", "pytest 8.2+ plus mocked external boundaries"],
        ],
        [2400, 6960],
    )
    add_heading(doc, "Configuration groups")
    add_bullets(doc, [
        "Oracle connection and deployment mode: EPM_BASE_URL, EPM_USERNAME, EPM_PASSWORD, APPLICATION_NAME and EPM_DEPLOYMENT_MODE.",
        "Database and runtime: DATABASE_URL, RUNTIME_DATA_DIR, EPM_EXECUTION_RUNTIME, worker polling and lease values.",
        "Web and identity: WEB_SESSION_SECRET, WEB_SECURE_COOKIES, WEB_FRONTEND_URL and optional Oracle OIDC settings.",
        "Agent: AGENT_PROVIDER, AGENT_ORCHESTRATOR=langgraph, AGENT_MODEL, provider API key, history/tool limits and Groq token budgets.",
        "Oracle engines and catalogs: per-operation REST/EPM Automate defaults and registered catalog files where needed.",
        "Notifications: provider, SMTP connection, sender, recipients and TLS/SSL settings.",
    ])
    add_page_break(doc)
    add_callout(doc, "Secret handling", "Production secrets belong in the deployment platform's secret manager or injected environment, never in source control. The .env file is a local development convenience and must remain uncommitted.", fill=RED_LIGHT, accent=RED)

    doc.settings.odd_and_even_pages_header_footer = True

    properties = doc.core_properties
    properties.title = "BISP Oracle EPM Automation Platform - Solution Architecture and Technology Stack"
    properties.subject = "Document 02 - Current modern platform architecture"
    properties.author = "BISP Solutions"
    properties.keywords = "Oracle EPM, Planning, architecture, FastAPI, React, PostgreSQL, LangGraph, BISP Solutions"
    properties.comments = "Documents the current modern React/FastAPI platform as of 24 August 2026."
    for table in doc.tables:
        mark_table_header_row(table)
    for section in doc.sections:
        for part in (section.header, section.even_page_header, section.footer, section.even_page_footer):
            for table in part.tables:
                mark_table_header_row(table)
    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
