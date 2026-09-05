"""Build Document 09: Administration, Security, Database, and Governance."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_LIGHT,
    GOLD_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    RED_LIGHT,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    add_bullets,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    font,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    wrap,
)
from docs.builders.build_document_03 import (  # noqa: E402
    add_code_block,
    compact_trailing_empty_paragraph,
    configure_styles,
    patch_list_numbering,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-09"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = (
    OUTPUT_DIR
    / "BISP_EPM_Automation_Document_09_Administration_Security_Database_and_Governance.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
MUTED_HEX = f"#{MUTED}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"
RED_HEX = f"#{RED}"


def _arrow(draw: ImageDraw.ImageDraw, start, end, *, color=BLUE_HEX, width=5) -> None:
    draw.line((*start, *end), fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [
            (x2, y2),
            (x2 - 17 * direction, y2 - 10),
            (x2 - 17 * direction, y2 + 10),
        ]
    else:
        direction = 1 if y2 > y1 else -1
        points = [
            (x2, y2),
            (x2 - 10, y2 - 17 * direction),
            (x2 + 10, y2 - 17 * direction),
        ]
    draw.polygon(points, fill=color)


def _title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((55, 34), title, font=font(34, bold=True), fill=NAVY_HEX)
    y = 82
    for line in wrap(draw, subtitle, font(19), 1280):
        draw.text((55, y), line, font=font(19), fill=MUTED_HEX)
        y += 26


def _card(
    draw: ImageDraw.ImageDraw,
    bounds,
    title: str,
    body: str,
    *,
    color: str,
    label: str = "",
    fill: str = "#FFFFFF",
) -> None:
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=18, fill=fill, outline=color, width=3)
    y = y1 + 18
    if label:
        draw.text((x1 + 20, y), label.upper(), font=font(15, bold=True), fill=color)
        y += 28
    for line in wrap(draw, title, font(22, bold=True), x2 - x1 - 40):
        draw.text((x1 + 20, y), line, font=font(22, bold=True), fill=NAVY_HEX)
        y += 29
    y += 5
    for raw_line in body.split("\n"):
        for line in wrap(draw, raw_line, font(16), x2 - x1 - 40):
            draw.text((x1 + 20, y), line, font=font(16), fill=MUTED_HEX)
            y += 22
        y += 3


def create_governance_layers(path: Path) -> None:
    image = Image.new("RGB", (1400, 820), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Five layers protect every governed action",
        "Business convenience sits above controls that remain active regardless of whether work begins in the UI, a schedule, Excel, or the Assistant.",
    )
    layers = [
        ((120, 160, 1280, 255), "1  Identity", "Who is requesting the action?", BLUE_HEX, f"#{BLUE_LIGHT}"),
        ((170, 275, 1230, 370), "2  Authorization", "Does the current platform role permit this capability?", BLUE_HEX, "#FFFFFF"),
        ((220, 390, 1180, 485), "3  Validation and approval", "Are the environment, artifact, inputs, impact, and decision explicit?", ORANGE_HEX, f"#{ORANGE_LIGHT}"),
        ((270, 505, 1130, 600), "4  Durable execution", "Can one leased worker perform the approved Oracle request without duplicate submission?", TEAL_HEX, f"#{TEAL_LIGHT}"),
        ((320, 620, 1080, 715), "5  Evidence and recovery", "Can an administrator prove the outcome and avoid unsafe automatic retries?", NAVY_HEX, "#FFFFFF"),
    ]
    for bounds, heading, body, color, fill in layers:
        x1, y1, x2, y2 = bounds
        draw.rounded_rectangle(bounds, radius=18, fill=fill, outline=color, width=3)
        draw.text((x1 + 25, y1 + 18), heading, font=font(21, bold=True), fill=NAVY_HEX)
        draw.text((x1 + 325, y1 + 22), body, font=font(17), fill=MUTED_HEX)
    image.save(path)


def create_identity_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 800), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Identity and access lifecycle",
        "Authentication proves identity; synchronized entitlements and platform role mappings decide what the person may do.",
    )
    boxes = [
        ((55, 165, 300, 360), "Identity source", "Local account, Oracle credential validation, or Oracle Cloud OIDC.", BLUE_HEX, f"#{BLUE_LIGHT}"),
        ((390, 165, 635, 360), "Platform profile", "Active user record, linked external subject, and safe identity metadata.", BLUE_HEX, "#FFFFFF"),
        ((725, 165, 970, 360), "Role assignment", "Service Administrator, Power User, User, or Viewer.", ORANGE_HEX, f"#{ORANGE_LIGHT}"),
        ((1060, 165, 1345, 360), "Permission check", "The API authorizes the exact endpoint before any work is prepared or submitted.", TEAL_HEX, f"#{TEAL_LIGHT}"),
    ]
    for i, (bounds, title, body, color, fill) in enumerate(boxes):
        _card(draw, bounds, title, body, color=color, label=str(i + 1), fill=fill)
        if i < len(boxes) - 1:
            _arrow(draw, (bounds[2] + 8, 260), (boxes[i + 1][0][0] - 8, 260), color="#9AA9BF", width=4)
    _card(
        draw,
        (110, 500, 640, 700),
        "Administrator controls",
        "Create or deactivate local users, assign approved roles, synchronize external identities, review mappings, and protect the final active administrator.",
        color=BLUE_HEX,
        fill="#FFFFFF",
    )
    _card(
        draw,
        (760, 500, 1290, 700),
        "Oracle remains authoritative",
        "A platform role never grants Oracle access that the effective Oracle execution identity does not already possess.",
        color=TEAL_HEX,
        fill=f"#{TEAL_LIGHT}",
    )
    image.save(path)


def create_database_domains(path: Path) -> None:
    image = Image.new("RGB", (1400, 900), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "PostgreSQL stores control state, not Planning data",
        "Thirty-nine application tables are organized into a small set of business domains; Oracle EPM remains the system of record for Planning artifacts and cube data.",
    )
    domains = [
        ((55, 160, 420, 365), "Identity and access", "Users, roles, permissions, login events, external identities, entitlements, mappings, sync runs, API tokens, and selected Oracle application.", BLUE_HEX, f"#{BLUE_LIGHT}"),
        ((515, 160, 880, 365), "Oracle catalog and RTPs", "Environment-scoped artifact verification plus Business Rule runtime-prompt definitions and sync evidence.", ORANGE_HEX, f"#{ORANGE_LIGHT}"),
        ((975, 160, 1345, 365), "Planning lifecycle", "Cycles, stages, tasks, dependencies, task executions, validations, approvals, and user notifications.", TEAL_HEX, f"#{TEAL_LIGHT}"),
        ((55, 465, 420, 670), "Execution and scheduling", "Workflow runs and steps, queue leases, process definitions, profiles, schedules, and schedule occurrences.", TEAL_HEX, "#FFFFFF"),
        ((515, 465, 880, 670), "EPM Assistant", "Conversations, messages, tool activity, action drafts, and immutable approval decisions.", BLUE_HEX, "#FFFFFF"),
        ((975, 465, 1345, 670), "LangGraph checkpoints", "Framework-owned checkpoint tables created by LangGraph setup, separate from the Alembic application schema.", NAVY_HEX, "#FFFFFF"),
    ]
    for bounds, title, body, color, fill in domains:
        _card(draw, bounds, title, body, color=color, fill=fill)
    draw.rounded_rectangle((170, 760, 1230, 835), radius=18, fill="#FFFFFF", outline="#C5D1E2", width=2)
    draw.text((215, 784), "Files, passwords, Planning cube data, and Oracle job configuration are not copied into these tables.", font=font(19, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_migration_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Controlled database release flow",
        "The application refuses to run against an older schema, preventing code and database structures from drifting apart.",
    )
    steps = [
        (55, "Back up", "Take a recoverable PostgreSQL backup and record the current revision.", BLUE_HEX),
        (315, "Stop writers", "Pause API workers or coordinate a maintenance window before structural change.", ORANGE_HEX),
        (575, "Apply Alembic", "Run ordered migrations through the code's current head: 0020_execution_identity.", ORANGE_HEX),
        (835, "Set up LangGraph", "Run the checkpoint setup command for framework-owned agent tables.", BLUE_HEX),
        (1095, "Verify and start", "Confirm current heads, start API and workers, then perform read-only checks.", TEAL_HEX),
    ]
    for index, (x, heading, body, color) in enumerate(steps, 1):
        _card(draw, (x, 175, x + 225, 500), heading, body, color=color, label=f"STEP {index}")
        if index < len(steps):
            _arrow(draw, (x + 229, 335), (x + 252, 335), color="#9AA9BF", width=4)
    draw.text((160, 620), "Startup gate", font=font(17, bold=True), fill=RED_HEX)
    draw.text((300, 617), "Installed Alembic heads must exactly equal the code's heads; otherwise FastAPI and worker startup stop with a targeted error.", font=font(19, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_audit_chain(path: Path) -> None:
    image = Image.new("RGB", (1400, 820), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "One action leaves a connected evidence chain",
        "The evidence answers who requested the work, what was approved, which worker executed it, what Oracle returned, and how the final result was communicated.",
    )
    steps = [
        ((55, 165, 310, 360), "Actor and trigger", "User identity plus MANUAL, SCHEDULED, API, EXCEL, or AI_AGENT source.", BLUE_HEX),
        ((390, 165, 645, 360), "Reviewed decision", "Normalized inputs, artifact, checksum, approve or reject, and decision time.", ORANGE_HEX),
        ((725, 165, 980, 360), "Queue ownership", "Execution ID, target, status, worker lease, heartbeat, and attempt count.", TEAL_HEX),
        ((1060, 165, 1345, 360), "Oracle evidence", "Workflow steps, job IDs, returned status, counts, files, messages, and safe errors.", NAVY_HEX),
    ]
    for i, (bounds, heading, body, color) in enumerate(steps):
        _card(draw, bounds, heading, body, color=color, label=str(i + 1))
        if i < len(steps) - 1:
            _arrow(draw, (bounds[2] + 8, 260), (steps[i + 1][0][0] - 8, 260), color="#9AA9BF", width=4)
    _card(draw, (190, 505, 1210, 705), "Administrator review", "Jobs and Activity joins the durable workflow record with step evidence. Authentication events, schedule occurrences, Planning task links, approvals, notifications, and agent decisions provide the surrounding context when applicable.", color=BLUE_HEX, label="CORRELATE BY EXECUTION ID", fill=f"#{BLUE_LIGHT}")
    image.save(path)


def create_worker_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 820), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Schedules and workers share one durable control path",
        "Unattended work is allowlisted, persisted, claimed once, monitored, and quarantined when Oracle outcome is uncertain.",
    )
    _card(draw, (55, 160, 330, 360), "Schedule occurrence", "Timezone-aware one-time, daily, weekly, or monthly occurrence; secrets are rejected from its JSON configuration.", color=BLUE_HEX, fill=f"#{BLUE_LIGHT}")
    _card(draw, (410, 160, 685, 360), "Durable queue", "QUEUED item with one target key; active duplicate targets are blocked.", color=ORANGE_HEX, fill=f"#{ORANGE_LIGHT}")
    _card(draw, (765, 160, 1040, 360), "Leased worker", "Atomic claim, heartbeat, effective Oracle identity, and terminal persistence.", color=TEAL_HEX, fill=f"#{TEAL_LIGHT}")
    _card(draw, (1120, 160, 1345, 360), "Oracle EPM", "The approved REST or EPM Automate operation runs under the configured execution identity.", color=NAVY_HEX)
    for x1, x2 in [(338, 402), (693, 757), (1048, 1112)]:
        _arrow(draw, (x1, 260), (x2, 260), color="#9AA9BF", width=4)
    _card(draw, (100, 520, 610, 705), "Normal terminal result", "SUCCESS or FAILED is written to the queue and workflow records. Schedule run status and notification delivery are updated independently.", color=TEAL_HEX, fill=f"#{TEAL_LIGHT}")
    _card(draw, (790, 520, 1300, 705), "Expired worker lease", "RECOVERY_REQUIRED. The platform does not automatically repeat the Oracle write; an administrator first reconciles with Oracle Job Console.", color=RED_HEX, fill=f"#{RED_LIGHT}")
    _arrow(draw, (900, 375), (1050, 505), color=RED_HEX)
    _arrow(draw, (900, 375), (350, 505), color=TEAL_HEX)
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "governance": ASSET_DIR / "governance_layers.png",
        "identity": ASSET_DIR / "identity_lifecycle.png",
        "database": ASSET_DIR / "database_domains.png",
        "migration": ASSET_DIR / "migration_workflow.png",
        "audit": ASSET_DIR / "governed_audit_chain.png",
        "worker": ASSET_DIR / "schedule_worker_flow.png",
    }
    create_governance_layers(assets["governance"])
    create_identity_flow(assets["identity"])
    create_database_domains(assets["database"])
    create_migration_flow(assets["migration"])
    create_audit_chain(assets["audit"])
    create_worker_flow(assets["worker"])
    return assets


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    kicker = doc.add_paragraph(f"CHAPTER {number:02d}", style="Kicker")
    kicker.paragraph_format.page_break_before = True
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_note(doc: Document, lead: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(7)
    set_run_font(paragraph.add_run(f"{lead} "), bold=True, color=NAVY)
    set_run_font(paragraph.add_run(text), color=INK)


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "BISP SOLUTIONS  /  ADMINISTRATION SECURITY DATABASE AND GOVERNANCE"
    set_run_font(paragraph.runs[0], size=8.5, color=MUTED, bold=True)
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.columns[0].width = Inches(5.7)
    table.columns[1].width = Inches(0.8)
    left = table.cell(0, 0).paragraphs[0]
    left.text = "Document 09  |  Administration Security Database and Governance"
    set_run_font(left.runs[0], size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(right)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:color"), LINE)
        borders.append(top)
        tc_pr.append(borders)


def add_cover(doc: Document) -> None:
    if LOGO.is_file():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions company logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(28)
    doc.add_paragraph("PLATFORM ADMINISTRATION GUIDE", style="Kicker")
    doc.add_paragraph(
        "Administration Security Database and Governance",
        style="Title",
    )
    doc.add_paragraph(
        "BISP Solutions Oracle EPM Automation Platform",
        style="Subtitle",
    )
    doc.add_paragraph(
        "A practical guide to administering identities and roles, protecting browser and API access, understanding the PostgreSQL schema, applying controlled migrations, supervising schedules and workers, preserving evidence, and preparing the current platform for production governance.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [[
            "09 of the platform handbook",
            "5 September 2026",
            "Service Administrators, database administrators, security reviewers, support teams, Oracle EPM administrators, and developers",
        ]],
        [2200, 2100, 5060],
    )
    add_note(
        doc,
        "Scope.",
        "This volume describes the current React and FastAPI platform, PostgreSQL application schema through Alembic revision 0020_execution_identity, LangGraph checkpoint storage, durable execution worker, current four-role access model, Oracle identity options, and present security boundaries. It excludes the retired interface.",
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        paragraph.add_run(
            "BISP Solutions  |  Least privilege  |  Durable evidence  |  Controlled change"
        ),
        size=10,
        color=MUTED,
        bold=True,
    )


def _finish_document(doc: Document) -> None:
    compact_trailing_empty_paragraph(doc)
    for table in doc.tables:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            header_pr.append(header)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for section in doc.sections:
        for footer in (section.footer, section.even_page_footer):
            for table in footer.tables:
                header_pr = table.rows[0]._tr.get_or_add_trPr()
                if header_pr.find(qn("w:tblHeader")) is None:
                    header = OxmlElement("w:tblHeader")
                    header.set(qn("w:val"), "true")
                    header_pr.append(header)
    properties = doc.core_properties
    properties.title = "Administration Security Database and Governance"
    properties.subject = "Current administration security PostgreSQL migration audit and governance guide"
    properties.author = "BISP Solutions"
    properties.keywords = (
        "Oracle EPM, administration, security, RBAC, PostgreSQL, SQLAlchemy, "
        "Alembic, governance, audit, worker, scheduling"
    )
    properties.comments = (
        "Current modern platform implementation as of 5 September 2026; "
        "retired interface excluded."
    )


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    patch_list_numbering(doc)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS VOLUME", style="Kicker")
    doc.add_paragraph("Administer the platform as one controlled system", style="Heading 1")
    doc.add_paragraph(
        "The platform is not secured by one setting. Safe operation depends on connected controls across identity, authorization, Oracle access, database state, execution ownership, audit evidence, and operational recovery. This guide explains those controls in the order an administrator normally uses them.",
        style="Lead",
    )
    add_table(
        doc,
        ["If you are...", "Read first", "Primary responsibility"],
        [
            ["Service Administrator", "Chapters 1-6 and 12-16", "Users, roles, identity integration, application selection, schedules, evidence, and operating policy."],
            ["Database administrator", "Chapters 7-11 and 15", "PostgreSQL availability, migrations, backups, restore tests, access, capacity, and release evidence."],
            ["Security reviewer", "Chapters 2-6 and 14-16", "Least privilege, sessions, CSRF, API tokens, secrets, logs, data boundaries, and production hardening."],
            ["Application support", "Chapters 8 and 12-16", "Table ownership, execution correlation, queue health, notifications, recovery, and support evidence."],
            ["Developer or consultant", "Chapters 7-13", "Repositories, transactions, schema evolution, agent state, scheduling, queue behavior, and safe extension."],
        ],
        [2200, 1900, 5260],
    )
    add_heading(doc, "Contents")
    add_table(
        doc,
        ["Part", "Chapters", "What it explains"],
        [
            ["Administration model", "1-3", "Governance layers, roles, permissions, account lifecycle, and administrator safeguards."],
            ["Identity and security", "4-6", "Authentication, sessions, OIDC, API tokens, same-origin web design, CSRF, headers, secrets, and Oracle boundaries."],
            ["Database control", "7-11", "PostgreSQL purpose, all application table domains, SQLAlchemy repositories, Alembic, and LangGraph checkpoints."],
            ["Operations governance", "12-14", "Schedules, queue leases, audit evidence, logging, notifications, and worker supervision."],
            ["Production stewardship", "15-16", "Retention, backups, recovery, hardening, periodic reviews, limitations, and handoff checklist."],
        ],
        [2100, 1500, 5760],
    )

    add_chapter(
        doc,
        1,
        "Understand the administration model",
        "Administration connects business ownership to technical controls. No single administrator should treat the web screen, PostgreSQL, or Oracle EPM as an isolated system.",
    )
    add_figure(
        doc,
        assets["governance"],
        "Figure 9.1 - Five layers protecting a governed platform action.",
        "Five stacked governance layers: identity, authorization, validation and approval, durable execution, and evidence and recovery.",
        width=6.35,
    )
    add_heading(doc, "Systems of record")
    add_table(
        doc,
        ["System", "Owns", "Does not own"],
        [
            ["Oracle EPM Planning", "Planning cube data, deployed artifacts, saved jobs, Oracle permissions, and Oracle job status.", "Platform users, schedules, approval decisions, or local workflow history."],
            ["BISP platform", "Platform identities, permissions, cycle work, synchronized catalogs, schedules, approvals, queue state, and audit evidence.", "The authoritative Planning cube, Oracle artifact configuration, or Oracle identity policy."],
            ["Language model provider", "Model inference for understanding and explanation during an Assistant request.", "Authorization, approval, Oracle credentials, execution ownership, or final evidence."],
        ],
        [2000, 3700, 3660],
    )
    add_heading(doc, "Administrator ownership")
    add_bullets(doc, [
        "Service Administrator: platform users, roles, Planning cycles, schedules, catalogs, and governed operation access.",
        "Oracle EPM administrator: Oracle users and groups, application access, deployed artifacts, saved jobs, cubes, metadata, and maintenance windows.",
        "Database administrator: PostgreSQL availability, least-privilege login, backup, restore, capacity, migration window, and revision evidence.",
        "Platform operations: API and worker supervision, HTTPS routing, secret injection, shared storage, monitoring, logs, and recovery coordination.",
        "Security owner: identity policy, access reviews, API token policy, log protection, key rotation, and exception approval.",
    ])
    add_note(doc, "Main principle.", "A platform role can remove or expose a platform capability, but it cannot grant Oracle privileges. Oracle still evaluates the effective execution identity on every Oracle request.")

    add_chapter(
        doc,
        2,
        "Use the four role access model",
        "The current platform uses four stable business-facing roles. Permissions are seeded from code and enforced by FastAPI rather than relying on hidden navigation alone.",
    )
    add_heading(doc, "Role summary")
    add_table(
        doc,
        ["Role", "Designed for", "Administrative boundary"],
        [
            ["Service Administrator", "Platform administration and all governed work.", "Owns all twelve current permissions, including users, design, schedules, variables, catalogs, and execution."],
            ["Power User", "Consultants and process owners who run approved work and investigate outcomes.", "Cannot manage users, design processes/cycles, manage schedules/catalogs, or update application substitution variables."],
            ["User", "Planners completing assigned work and authorized reviews.", "Can run approved processes, update own user variables, review data, generate reports, and use the Assistant; cannot run arbitrary standalone operations."],
            ["Viewer", "Read-oriented consumers of reports and Assistant guidance.", "Has report generation and Assistant access only; the Assistant exposes only tools allowed for this role."],
        ],
        [2100, 3300, 3960],
    )
    add_heading(doc, "Permission matrix")
    add_table(
        doc,
        ["Permission", "Service Admin", "Power User", "User", "Viewer"],
        [
            ["process.run", "Yes", "Yes", "Yes", "No"],
            ["process.design", "Yes", "No", "No", "No"],
            ["schedule.manage", "Yes", "No", "No", "No"],
            ["operation.execute", "Yes", "Yes", "No", "No"],
            ["variable.update", "Yes", "No", "No", "No"],
            ["user_variable.update", "Yes", "Yes", "Yes", "No"],
            ["data.review", "Yes", "Yes", "Yes", "No"],
            ["report.generate", "Yes", "Yes", "Yes", "Yes"],
            ["history.view", "Yes", "Yes", "No", "No"],
            ["user.manage", "Yes", "No", "No", "No"],
            ["catalog.manage", "Yes", "No", "No", "No"],
            ["agent.use", "Yes", "Yes", "Yes", "Yes"],
        ],
        [3200, 1540, 1540, 1540, 1540],
    )
    add_heading(doc, "How enforcement works")
    add_numbered(doc, [
        "The signed session or bearer token resolves one active platform user.",
        "The server maps the HTTP method and protected path to one or more required permissions.",
        "The request is rejected with HTTP 403 when the user has none of the accepted permissions.",
        "The service layer validates environment, artifact, inputs, ownership, and other business rules.",
        "Oracle independently enforces the permissions of the effective Oracle execution account.",
    ])
    add_code_block(
        doc,
        "# Stable permission codes used by the current API\nPROCESS_RUN = \"process.run\"\nOPERATION_EXECUTE = \"operation.execute\"\nVARIABLE_UPDATE = \"variable.update\"\nUSER_MANAGE = \"user.manage\"\nCATALOG_MANAGE = \"catalog.manage\"\nAGENT_USE = \"agent.use\"",
        "Code excerpt 9.1. Representative server-enforced permission codes.",
    )

    add_chapter(
        doc,
        3,
        "Manage platform users safely",
        "The user lifecycle is explicit: bootstrap once, create or link identities, assign the smallest suitable role, review activity, and deactivate access without destroying evidence.",
    )
    add_figure(
        doc,
        assets["identity"],
        "Figure 9.2 - Identity and authorization lifecycle.",
        "Four-step identity lifecycle from identity source through platform profile and role assignment to API permission check, with administrator and Oracle authority boundaries.",
        width=6.35,
    )
    add_heading(doc, "First administrator bootstrap")
    add_numbered(doc, [
        "Start with an empty, migrated PostgreSQL database.",
        "Open the application and create the first Service Administrator using a unique username, display name, optional email, and passphrase of at least twelve characters.",
        "A PostgreSQL advisory transaction lock ensures that concurrent first-run requests cannot create multiple bootstrap administrators.",
        "The user, Service Administrator assignment, and BOOTSTRAP_ADMIN_CREATED event are committed atomically.",
        "After the first record exists, the bootstrap endpoint cannot be used again.",
    ])
    add_heading(doc, "Local account controls")
    add_table(
        doc,
        ["Control", "Current behavior", "Administrator action"],
        [
            ["Password storage", "Versioned scrypt hash with a random salt; plaintext is never retained.", "Use unique passphrases and a controlled reset process."],
            ["Minimum length", "12 characters; maximum 256 characters.", "Apply stronger organizational policy through training and identity-provider controls where required."],
            ["Failed sign-in", "After 5 failed attempts, the account is locked for 15 minutes.", "Investigate repeated failures through authentication events; do not weaken the threshold casually."],
            ["Deactivation", "Inactive users cannot authenticate and their sessions cease resolving as active users.", "Deactivate departed users; preserve history for attribution."],
            ["Final administrator", "The last active Service Administrator cannot be deactivated or have that role removed.", "Create and test a second administrator before changing the final administrator."],
        ],
        [2100, 3900, 3360],
    )
    add_heading(doc, "Routine access review")
    add_bullets(doc, [
        "Review active users, role assignments, authentication source, and last sign-in.",
        "Remove unnecessary elevated roles and avoid using Service Administrator for ordinary daily work.",
        "Confirm disabled or departed Oracle identities no longer have a usable linked platform profile.",
        "Review bearer tokens separately; deactivating the owner immediately makes owned tokens unusable.",
        "Record who approved exceptional access and when it must be reviewed again outside the application if no expiry workflow exists.",
    ])

    add_chapter(
        doc,
        4,
        "Understand authentication and sessions",
        "The platform supports local authentication and optional Oracle-backed identity paths while keeping browser session data separate from Oracle integration credentials.",
    )
    add_heading(doc, "Authentication paths")
    add_table(
        doc,
        ["Path", "What proves identity", "Important boundary"],
        [
            ["Local platform account", "Username and scrypt-verified local passphrase.", "This identity is independent of the Oracle integration account."],
            ["Oracle credential validation", "User-supplied Oracle credentials are validated against Planning for sign-in.", "Entered Oracle passwords are not persisted; repeated failures are throttled."],
            ["Oracle Cloud OIDC", "Validated issuer, state, nonce, authorization response, and immutable OIDC subject.", "The subject must map to an approved synchronized and linked platform profile."],
        ],
        [2200, 3600, 3560],
    )
    add_heading(doc, "Signed browser session")
    add_table(
        doc,
        ["Session property", "Current behavior"],
        [
            ["Cookie name", "bisp_epm_session"],
            ["Lifetime", "Eight hours"],
            ["SameSite", "Lax"],
            ["Secure flag", "Controlled by WEB_SECURE_COOKIES; production behind HTTPS should set it to true."],
            ["Rotation", "Successful authentication clears the prior session and creates a new session identifier and CSRF token."],
            ["Contents", "Safe identifiers such as user ID, session ID, CSRF token, and authentication method; never Oracle passwords or API keys."],
            ["Signing secret", "WEB_SESSION_SECRET. If missing, an ephemeral value is generated and all browser sessions reset at restart."],
        ],
        [2400, 6960],
    )
    add_note(doc, "Cookie limitation.", "Starlette's session cookie is signed against tampering, not a place for confidential payloads. Keep only safe session identifiers in it and protect the signing secret consistently across all API instances.")
    add_heading(doc, "CSRF protection")
    add_para(doc, "Every state-changing browser request using POST, PUT, PATCH, or DELETE must carry the session's X-CSRF-Token. Missing or mismatched tokens receive HTTP 403 before the requested mutation is performed.")
    add_heading(doc, "Oracle sign-in throttling")
    add_para(doc, "The platform separately limits repeated Oracle credential failures: three recent failures for a username or fifteen from an IP address within the fifteen-minute window trigger throttling. Authentication events retain the safe outcome and client address; entered passwords are not recorded.")

    add_chapter(
        doc,
        5,
        "Protect APIs and browser traffic",
        "The current deployment is designed around a same-origin browser boundary or a development proxy. Cross-origin access is not enabled by a permissive global CORS policy.",
    )
    add_heading(doc, "Same-origin and CORS design")
    add_para(doc, "FastAPI does not currently install CORSMiddleware. In local development, Vite proxies /api, /app, /auth, /login, /logout, /setup, and /static to FastAPI, so the browser still sees one frontend origin. In production, serve the built React application and API from one HTTPS origin or use a reverse proxy that presents them as one origin.")
    add_note(doc, "If separate origins become necessary.", "Add a deliberate FastAPI CORS allowlist for exact trusted frontend origins, enable credentials only when required, test cookies and CSRF together, and never use a wildcard origin for authenticated browser traffic.")
    add_heading(doc, "Security headers")
    add_table(
        doc,
        ["Header", "Current policy", "Purpose"],
        [
            ["Content-Security-Policy", "default self; images self/data; styles, scripts, fonts, connections, forms self; no framing.", "Restricts browser resource and connection origins."],
            ["X-Frame-Options", "DENY", "Prevents the interface from being framed."],
            ["X-Content-Type-Options", "nosniff", "Prevents MIME-type guessing."],
            ["Referrer-Policy", "strict-origin-when-cross-origin", "Limits referrer detail sent across origins."],
            ["Permissions-Policy", "Camera, microphone, and geolocation disabled.", "Removes unused browser capabilities."],
            ["Cache-Control", "no-store for HTML", "Reduces retention of authenticated page content in caches."],
            ["X-Request-ID", "Generated or accepted when it matches the safe pattern.", "Connects browser errors, API responses, and server logs."],
        ],
        [2300, 3900, 3160],
    )
    add_heading(doc, "Scoped bearer API tokens")
    add_table(
        doc,
        ["Control", "Current behavior"],
        [
            ["Supported scopes", "pipeline.read, pipeline.run, and execution.read for Excel and controlled external clients."],
            ["Secret handling", "A random token is shown once at creation; only SHA-256 hash and a short prefix are stored."],
            ["Expiry and revocation", "Optional expiry plus irreversible revocation timestamp."],
            ["Ongoing authorization", "Each request checks scope, token state, owner activity, and the owner's current platform permission."],
            ["Usage evidence", "Last-used time and IP are updated at a bounded interval to avoid excessive writes."],
        ],
        [2500, 6860],
    )
    add_note(doc, "API boundary.", "Current API tokens are intentionally limited to the Excel Pipeline contract. Do not reuse them as general administrator tokens or expose browser session cookies to macros or scripts.")

    add_chapter(
        doc,
        6,
        "Manage secrets and Oracle boundaries",
        "Secrets enter the server process through deployment configuration. They must not be stored in PostgreSQL JSON, React source, screenshots, logs, schedules, or generated documents.",
    )
    add_heading(doc, "Secret inventory")
    add_table(
        doc,
        ["Secret", "Used by", "Storage rule"],
        [
            ["PostgreSQL password", "DATABASE_URL used by API, workers, migrations, and LangGraph checkpoints.", "Production secret manager or protected service configuration; URL-encode reserved characters."],
            ["Oracle integration password", "FastAPI and worker Oracle clients.", "Server-side environment only; never returned to the browser or stored in platform tables."],
            ["WEB_SESSION_SECRET", "Signed browser sessions.", "Long random value shared consistently by every API instance; rotate with a planned session reset."],
            ["OIDC client secret", "Oracle Cloud authorization-code exchange.", "Server-side secret store; pair with exact issuer and redirect URI."],
            ["Gemini or Groq key", "Configured Assistant provider.", "Server-side secret store; restrict access and monitor provider quota."],
            ["SMTP password", "Optional email notification provider.", "Server-side secret store; never place in schedule configuration."],
            ["EPM Automate .epw", "Operations explicitly configured for the EPM Automate engine.", "Encrypted file with restricted filesystem ACL; reference by path only."],
        ],
        [2100, 3500, 3760],
    )
    add_heading(doc, "Schedule secret rejection")
    add_para(doc, "The schedule service recursively rejects JSON keys containing password, secret, token, API key, credential, or private key. Schedules store only safe target and runtime configuration; workers obtain credentials from their own deployment environment.")
    add_heading(doc, "Database and log redaction boundary")
    add_bullets(doc, [
        "Authentication records store outcomes, usernames, user references, client IP, timestamps, and safe details - not passwords.",
        "Workflow and Assistant evidence is projected through safe allowlists and bounded error messages before presentation.",
        "Agent-facing execution evidence excludes filesystem paths, upload tokens, credentials, and unrelated internal step data.",
        "Logging uses structured context and request IDs, but operators must still avoid adding raw request bodies or secret environment values.",
        "Production database connections should require encrypted transport where the database is remote; configure the approved PostgreSQL SSL mode and certificates in DATABASE_URL.",
    ])

    add_chapter(
        doc,
        7,
        "Understand the PostgreSQL data boundary",
        "PostgreSQL is the durable control database for the platform. It does not replace Oracle Planning and it does not contain a copied Planning cube.",
    )
    add_figure(
        doc,
        assets["database"],
        "Figure 9.3 - Current PostgreSQL storage domains.",
        "Six PostgreSQL domains: identity and access, Oracle catalog and runtime prompts, Planning lifecycle, execution and scheduling, EPM Assistant, and separate LangGraph checkpoints.",
        width=6.35,
    )
    add_heading(doc, "Technology choices")
    add_table(
        doc,
        ["Component", "Current use", "Reason"],
        [
            ["PostgreSQL", "Production application and LangGraph state.", "Transactions, constraints, JSONB, concurrent row locking, durability, backup, and operational tooling."],
            ["SQLAlchemy Core", "Canonical schema, queries, repositories, transactions, and connection pooling.", "Explicit SQL-oriented control without coupling the application to one web framework."],
            ["Alembic", "Ordered application-schema revisions.", "Repeatable installation and upgrade history with a startup compatibility check."],
            ["psycopg and psycopg_pool", "PostgreSQL driver and LangGraph checkpoint connection pool.", "Supported native PostgreSQL connectivity for both application and framework state."],
        ],
        [2100, 3400, 3860],
    )
    add_heading(doc, "Connection behavior")
    add_bullets(doc, [
        "The production runtime accepts PostgreSQL only. SQLite paths are supported solely as isolated automated-test compatibility hooks.",
        "The SQLAlchemy pool checks connections before use, keeps ten pooled connections, allows twenty overflow connections, waits thirty seconds for a connection, and recycles connections after thirty minutes.",
        "API and every worker must use the same DATABASE_URL so queue, history, identity, schedule, and Assistant records remain coherent.",
        "Repository writes use explicit transaction contexts. Related records such as a bootstrap user and role assignment either commit together or roll back together.",
        "PostgreSQL JSONB stores flexible safe metadata; relationships and high-value business states remain constrained columns and foreign keys.",
    ])
    add_note(doc, "Not stored.", "Oracle cube data, full form extracts, source-file contents, user-entered Oracle passwords, and language-model API keys are not intended for the application tables. Temporary uploads and generated files live in controlled runtime storage with separate lifecycle rules.")

    add_chapter(
        doc,
        8,
        "Know the database table responsibilities",
        "The current Alembic-managed schema contains thirty-nine application tables. They are grouped below by responsibility so administrators can query and protect them without treating every table as an isolated feature.",
    )
    add_heading(doc, "Identity environment and access tables")
    add_table(
        doc,
        ["Table", "Responsibility"],
        [
            ["platform_users", "Local or linked platform profiles, active state, password hash when local, lockout counters, and login timestamps."],
            ["platform_roles", "Four seeded platform roles and their business descriptions."],
            ["platform_role_permissions", "Many-to-many role-to-permission assignments seeded from current code."],
            ["platform_user_roles", "User role assignments, assignment time, and assigning administrator."],
            ["authentication_events", "Successful and failed login, logout, bootstrap, user administration, Oracle credential, and external authentication audit events."],
            ["oracle_environment_settings", "Per-base-URL deployment mode, discovered applications, selected application, selection source, error, timestamp, and selecting user."],
            ["identity_providers", "Safe provider metadata for LOCAL, ORACLE_CLOUD, OIDC, or LDAP adapters; secrets are excluded."],
            ["external_identities", "Synchronized provider subjects, optional linked platform user, active state, safe profile fields, and last-seen/authenticated timestamps."],
            ["external_entitlements", "Synchronized application roles, granular roles, and groups exposed by the provider."],
            ["external_identity_entitlements", "Direct or inherited entitlement observations for each external identity."],
            ["identity_role_mappings", "Administrator-approved mapping from one external entitlement to one platform role."],
            ["identity_sync_runs", "Status and counts for each directory synchronization, including bounded error evidence."],
            ["api_tokens", "One-time-issued external-client token hashes, prefixes, scopes, expiry, usage, and revocation metadata."],
        ],
        [3000, 6360],
    )
    add_heading(doc, "Oracle catalog and runtime prompt tables")
    add_table(
        doc,
        ["Table", "Responsibility"],
        [
            ["oracle_artifacts", "Environment-scoped live and registered Pipelines, Data Integrations, Business Rules, Data Maps, import jobs, Cube Refresh jobs, and cubes with verification state."],
            ["business_rule_rtp_sync_runs", "Source checksum, parser version, counts, warnings, and outcome for Calc Manager RTP registry synchronization."],
            ["business_rule_rtp_definitions", "One active rule definition per environment and normalized name, including cube and source lineage."],
            ["business_rule_rtp_parameters", "Ordered prompt name, label, type, dimension, default, required/hidden/multiple flags, scope, limits, and source metadata."],
        ],
        [3000, 6360],
    )
    add_heading(doc, "Execution process and scheduling tables")
    add_table(
        doc,
        ["Table", "Responsibility"],
        [
            ["workflow_runs", "Durable top-level execution identity, workflow name, status, actor snapshot, trigger source, effective Oracle username, and terminal error."],
            ["workflow_steps", "Ordered execution steps with status, timestamps, safe details, and step error."],
            ["execution_queue", "Serializable Oracle work, priority, active state, worker lease, heartbeat, attempts, completion, and recovery-required evidence."],
            ["automation_schedules", "Current generic allowlisted recurrences for Oracle Pipelines and RTP registry synchronization."],
            ["automation_schedule_runs", "Every claimed occurrence, resolved safe payload, execution link, status, and error."],
            ["planning_processes", "Stable process code used by the retained process-definition subsystem."],
            ["planning_process_versions", "Versioned process and cycle definitions with draft or active status."],
            ["process_run_profiles", "Archived-or-active reusable process input profiles such as year, periods, variables, and file references."],
            ["process_schedules", "Earlier process-specific recurrence records retained for compatibility and history; generic automation_schedules is the current schedule model for allowlisted unattended automation."],
        ],
        [3000, 6360],
    )
    add_note(doc, "Current workflow direction.", "Planning-cycle creation and Oracle-owned Pipelines are the supported business workflow focus. The retained process-definition tables remain part of the installed schema, but the retired interface and discontinued custom process-creation experience are not part of this handbook.")

    add_chapter(
        doc,
        9,
        "Understand Planning and Assistant tables",
        "Planning lifecycle records connect assignments, validations, approvals, and Oracle executions. Assistant records add conversation context and immutable approval evidence without creating a separate execution system.",
    )
    add_heading(doc, "Planning lifecycle tables")
    add_table(
        doc,
        ["Table", "Responsibility"],
        [
            ["planning_cycles", "Cycle identity, type, process code, scenario/year/period context, dates, status, creator, completion, and archive state."],
            ["planning_cycle_stages", "Ordered business stages, dates, and completion state inside one cycle."],
            ["planning_tasks", "Assigned work, priority, context, action routing, configuration, due time, and completion state."],
            ["planning_task_dependencies", "Task prerequisites within a cycle."],
            ["planning_task_executions", "Every execution attempt linked to a task, including actor, attempt number, status, and error."],
            ["planning_task_validations", "Summary-only validation evidence: cubes, selection, criteria, checked/matched/exception counts, performer, and warning acknowledgement."],
            ["planning_approvals", "Submission and decision chain linking cycle, source task, approval task, optional validation, decision actor, time, and comment."],
            ["user_notifications", "In-application notification title, severity, message, action link, source reference, creation time, and read time."],
        ],
        [3000, 6360],
    )
    add_heading(doc, "EPM Assistant tables")
    add_table(
        doc,
        ["Table", "Responsibility"],
        [
            ["agent_conversations", "User-owned conversation identity, title, provider, model, and timestamps."],
            ["agent_messages", "Ordered USER and ASSISTANT messages for a conversation."],
            ["agent_tool_activities", "Allowed tool name, bounded arguments, status, summary, and time for explainable inspection activity."],
            ["agent_action_drafts", "Prepared operation, artifact, risk, route, normalized input schema and values, preflight checks, and status."],
            ["agent_action_decisions", "Idempotent approve/reject request, actor snapshot, payload checksum and snapshot, outcome, execution link, failure, and timestamps."],
        ],
        [3000, 6360],
    )
    add_heading(doc, "Deletion and archival behavior")
    add_table(
        doc,
        ["Record", "Current behavior", "Governance meaning"],
        [
            ["Agent conversation", "The owning user can delete it; child messages, tool activity, drafts, and linked agent decision rows follow configured foreign-key behavior.", "Do not treat conversations as the only execution audit; workflow execution evidence remains separate."],
            ["Schedule", "Archived rather than physically deleted; run history is retained.", "Inactive definitions disappear from normal lists without erasing occurrences."],
            ["Run profile", "Archived rather than physically deleted.", "Historical references remain explainable."],
            ["Planning cycle", "Archive column exists and active queries exclude archived records.", "Define an approved lifecycle policy before exposing broader archival controls."],
            ["Workflow execution", "No automatic purge is currently implemented.", "Retention must be governed at deployment level before records are removed."],
        ],
        [2100, 3800, 3460],
    )

    add_chapter(
        doc,
        10,
        "Use SQLAlchemy repositories and transactions",
        "Schema definitions describe data shape. Repositories own SQL and transaction boundaries so web routes, workers, schedulers, and the Assistant do not manipulate tables ad hoc.",
    )
    add_heading(doc, "Data access structure")
    add_table(
        doc,
        ["Layer", "Responsibility", "Rule"],
        [
            ["SQLAlchemy schema", "Tables, columns, PostgreSQL JSONB, foreign keys, unique constraints, check constraints, indexes, and naming conventions.", "One canonical metadata object is used by Alembic and repositories."],
            ["Repository", "Read and write a coherent aggregate such as users, workflows, schedules, planning work, or agent state.", "No Oracle calls and no HTTP-specific behavior."],
            ["Application service", "Validate business inputs, combine repositories and adapters, and enforce workflow policy.", "Use explicit outcomes and domain exceptions."],
            ["API or worker", "Authenticate/authorize, translate transport data, invoke a use case, and present or persist the result.", "Never bypass repositories to change governance state."],
        ],
        [1900, 4300, 3160],
    )
    add_heading(doc, "Transaction examples")
    add_bullets(doc, [
        "Bootstrap commits the first user, administrator role, and authentication event together under an advisory lock.",
        "Worker claim uses PostgreSQL row locking with SKIP LOCKED so multiple workers can safely claim different queue rows.",
        "Agent approval reservation relies on a unique request ID so double-clicks and retries reuse the recorded decision rather than submitting Oracle twice.",
        "Native PostgreSQL upsert supports idempotent role seeding and synchronized catalog updates.",
        "Foreign keys use deliberate delete behavior: CASCADE for owned child records, SET NULL when attribution may outlive an account, and RESTRICT where referenced system records must remain.",
    ])
    add_code_block(
        doc,
        "with database.begin() as connection:\n    connection.execute(update_statement)\n    connection.execute(audit_insert)\n# Both statements commit together; an exception rolls both back.",
        "Code excerpt 9.2. Explicit transaction boundary used by repository services.",
    )
    add_heading(doc, "Safe read-only administration queries")
    add_code_block(
        doc,
        "-- Recent workflow outcomes\nSELECT execution_id, workflow_name, status, trigger_source,\n       initiated_by_username, oracle_execution_username, started_at, completed_at\nFROM workflow_runs\nORDER BY started_at DESC\nLIMIT 50;\n\n-- Due or recently triggered schedules\nSELECT schedule_id, name, target_type, target_key, is_enabled,\n       next_run_at, last_outcome, last_execution_id\nFROM automation_schedules\nWHERE archived_at IS NULL\nORDER BY next_run_at NULLS LAST;",
        "Code excerpt 9.3. Read-only operational queries. Use a read-only database role for routine reporting.",
    )
    add_note(doc, "Do not repair by SQL.", "Direct UPDATE or DELETE statements can bypass application validation, audit creation, role safeguards, archive semantics, and checksum rules. Use the supported application service or a reviewed migration for changes.")

    add_chapter(
        doc,
        11,
        "Apply Alembic and LangGraph schema changes",
        "Application tables and LangGraph checkpoint tables have different owners and upgrade commands. A production release must manage both explicitly.",
    )
    add_figure(
        doc,
        assets["migration"],
        "Figure 9.4 - Controlled PostgreSQL release and migration sequence.",
        "Five-step database release sequence from backup through stopping writers, Alembic migration, LangGraph checkpoint setup, and verified restart.",
        width=6.35,
    )
    add_heading(doc, "Current Alembic history")
    add_table(
        doc,
        ["Revision range", "Capability introduced"],
        [
            ["0001-0004", "Initial platform schema, Planning cycle/task engine, external API tokens, approvals, and notifications."],
            ["0005-0008", "Four-role access model, task execution attempts, validation evidence, and environment artifact registry."],
            ["0009-0012", "Durable execution queue, unified Oracle catalog, idempotent agent decisions, and standalone-flow queue support."],
            ["0013-0016", "Federated identity foundation, Oracle OIDC subject binding, Business Rule RTP registry, and scoped RTP metadata."],
            ["0017-0020", "Generic automation schedules, synchronous scheduled RTP sync, environment application selection, and effective Oracle execution identity."],
        ],
        [2500, 6860],
    )
    add_heading(doc, "Installation and upgrade commands")
    add_code_block(
        doc,
        "python -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup",
        "Code excerpt 9.4. Apply the current application schema, verify every Alembic head, then prepare framework-owned LangGraph checkpoint tables.",
    )
    add_heading(doc, "Startup compatibility gate")
    add_para(doc, "FastAPI and worker.py inspect PostgreSQL before composing services. If the installed Alembic heads do not exactly match the code's heads, startup stops and instructs the administrator to run alembic upgrade head. This prevents new code from processing Oracle work against old tables.")
    add_heading(doc, "Migration operating rules")
    add_bullets(doc, [
        "Back up PostgreSQL and test restoration before a production schema change.",
        "Review generated migrations; autogeneration is a starting point, not production approval.",
        "Coordinate API and worker shutdown so no process writes with a mixed schema version.",
        "Prefer forward-compatible fixes. A database downgrade requires a tested data-preservation and restoration plan.",
        "Record the code tag, previous revision, new revision, operator, timestamps, backup reference, and verification result.",
        "Run LangGraph checkpoint setup after changing the LangGraph checkpoint package or deploying to a new database.",
    ])

    add_chapter(
        doc,
        12,
        "Govern Assistant state schedules and queue records",
        "These three durable subsystems have different lifecycles but share PostgreSQL so a user request can be traced from context to decision to execution.",
    )
    add_heading(doc, "LangGraph checkpoints")
    add_para(doc, "The application-owned agent tables store readable product history. LangGraph's PostgresSaver separately stores graph checkpoints required to pause for a choice, input, or approval and resume the exact conversation state. The setup command creates or upgrades those framework-owned tables; they are not declared in Alembic's canonical application schema.")
    add_table(
        doc,
        ["State", "Owner", "Purpose"],
        [
            ["Conversation, messages, tool activity, drafts, decisions", "BISP application schema and repositories", "User-visible history, governed proposal state, approval idempotency, and audit."],
            ["Graph checkpoints, pending writes, serialized state", "LangGraph PostgresSaver", "Reliable interrupt and resume for the orchestration engine."],
            ["Workflow runs, steps, and queue", "Shared execution framework", "Oracle execution ownership, status, evidence, and recovery."],
        ],
        [3000, 2700, 3660],
    )
    add_heading(doc, "Current unattended schedule allowlist")
    add_table(
        doc,
        ["Target", "Execution behavior", "Governance"],
        [
            ["Oracle Pipeline", "Creates a durable operation execution using registered, environment-valid Pipeline context.", "Service Administrator manages recurrence; concurrency policy skips when the same target is active."],
            ["Business Rule RTP registry sync", "Runs controlled Calc Manager export/import synchronization and records a completed synchronous occurrence.", "No Oracle business calculation is launched; generated export is cleaned up best-effort."],
        ],
        [2200, 3900, 3260],
    )
    add_heading(doc, "Schedule controls")
    add_bullets(doc, [
        "Frequencies are one-time, daily, weekly, or monthly using an explicit IANA timezone and a local first-run date/time.",
        "Input policy is ORACLE_DEFAULTS, FIXED, or DYNAMIC as supported by the target adapter.",
        "The current concurrency policy is SKIP_IF_ACTIVE; overlapping work is not queued blindly.",
        "Misfire behavior is RUN_ONCE or SKIP, and every claimed occurrence receives its own run record.",
        "Archiving hides a schedule from active lists while retaining schedule-run and execution evidence.",
        "Credentials and secret-like configuration keys are rejected before schedule persistence.",
    ])

    add_chapter(
        doc,
        13,
        "Supervise durable execution and recovery",
        "The queue prevents a browser, scheduler, Excel macro, or Assistant request from owning long-running Oracle work after submission.",
    )
    add_figure(
        doc,
        assets["worker"],
        "Figure 9.5 - Durable schedule and worker control flow.",
        "Schedule occurrence enters a durable queue, a worker claims and heartbeats the item, Oracle executes, and the platform records either a terminal result or recovery-required state after lease expiry.",
        width=6.35,
    )
    add_heading(doc, "Queue states")
    add_table(
        doc,
        ["State", "Meaning", "Administrator response"],
        [
            ["QUEUED", "Validated work awaits an eligible worker.", "Confirm at least one worker uses the same database and storage configuration."],
            ["RUNNING", "One worker owns the lease and sends heartbeats.", "Monitor workflow steps and Oracle job evidence; do not submit a duplicate target."],
            ["SUCCESS", "Worker completed the operation and persisted terminal evidence.", "Review returned counts or artifacts where business validation is required."],
            ["FAILED", "Worker observed a definite failure and persisted a safe error.", "Correct the cause, revalidate inputs and artifact, then authorize a new run."],
            ["RECOVERY_REQUIRED", "The worker lease expired while Oracle outcome may be unknown.", "Reconcile with Oracle Job Console and platform evidence before any retry."],
        ],
        [1800, 3900, 3660],
    )
    add_heading(doc, "Worker deployment")
    add_code_block(
        doc,
        "# Production API process\nEPM_EXECUTION_RUNTIME=web\npython -m uvicorn web_main:app --host 0.0.0.0 --port 8080\n\n# One or more supervised worker processes\npython worker.py",
        "Code excerpt 9.5. Production separates request handling from Oracle execution.",
    )
    add_bullets(doc, [
        "Every worker needs the same DATABASE_URL, Oracle environment configuration, runtime storage, and operation dependencies.",
        "The queue uses row locks and SKIP LOCKED so several workers can claim different jobs safely.",
        "A worker lease lasts at least thirty seconds and is refreshed at a bounded heartbeat interval.",
        "Upload cleanup refuses paths outside the controlled runtime upload root.",
        "Worker supervision belongs to the deployment platform; restart policies must not automatically resubmit RECOVERY_REQUIRED Oracle writes.",
    ])

    add_chapter(
        doc,
        14,
        "Use audit evidence logging and notifications",
        "A supportable operation needs more than a success banner. The platform retains connected evidence while limiting secret and path exposure.",
    )
    add_figure(
        doc,
        assets["audit"],
        "Figure 9.6 - Governed operation audit and evidence chain.",
        "Connected evidence chain from actor and trigger through reviewed decision, queue ownership, Oracle evidence, and administrator correlation by execution ID.",
        width=6.35,
    )
    add_heading(doc, "Execution evidence")
    add_table(
        doc,
        ["Question", "Current evidence"],
        [
            ["Who requested it?", "Initiating username and display snapshot plus trigger source: MANUAL, SCHEDULED, API, EXCEL, or AI_AGENT."],
            ["Who executed in Oracle?", "Effective Oracle execution username stored on the workflow run."],
            ["What was requested?", "Workflow or operation name, target, normalized payload, artifact, and Assistant decision snapshot when applicable."],
            ["Was it approved?", "Planning approval row or immutable Agent action decision with checksum, actor, time, and outcome."],
            ["What happened?", "Queue status, workflow status, ordered step status, timestamps, Oracle job IDs, statistics, messages, output artifacts, and safe errors."],
            ["How was it triggered?", "Manual UI, scheduler, scoped API/Excel token, or AI Agent attribution."],
        ],
        [2600, 6760],
    )
    add_heading(doc, "Logging")
    add_para(doc, "The application logger writes timestamp, level, logger name, request correlation ID, and message to standard output. Request middleware logs method, path, response status, and duration, and returns X-Request-ID to the client. Production should collect stdout through the hosting platform into a protected searchable log service.")
    add_table(
        doc,
        ["Log practice", "Recommendation"],
        [
            ["Correlation", "Start with the X-Request-ID shown in the UI or response, then correlate execution ID, schedule run ID, Oracle job ID, and timestamps."],
            ["Confidentiality", "Do not log Authorization headers, cookies, CSRF tokens, passwords, API keys, .epw content, full environment dumps, or unreviewed request bodies."],
            ["Integrity", "Restrict modification and deletion, synchronize system clocks, and retain platform plus infrastructure logs according to approved policy."],
            ["Alerting", "Alert on repeated authentication failures, worker unavailability, growing queue age, RECOVERY_REQUIRED, migration mismatch, and repeated provider failures."],
        ],
        [2200, 7160],
    )
    add_heading(doc, "Notifications")
    add_para(doc, "In-application user notifications are persisted in PostgreSQL. Optional email uses a provider-neutral notification service with the current SMTP adapter. TLS and SSL are mutually exclusive, configuration is validated at startup, and delivery failure is logged without changing the true task or execution outcome.")
    add_note(doc, "Operational meaning.", "A successful notification does not prove an Oracle job succeeded, and a failed email does not convert a successful Oracle job into a failed job. Always treat workflow and Oracle evidence as authoritative.")

    add_chapter(
        doc,
        15,
        "Define retention backup and recovery",
        "The application preserves important history but does not currently implement a universal database purge policy. Production retention must be an explicit organizational decision.",
    )
    add_heading(doc, "Current lifecycle behavior")
    add_table(
        doc,
        ["Data class", "Current behavior", "Production policy needed"],
        [
            ["Temporary browser uploads", "Maximum 100 MB; session-owned; expire after eight hours; removed after terminal execution or sign-out where the cleanup path runs.", "Monitor runtime storage and use a startup/host cleanup procedure for orphaned files after abrupt process loss."],
            ["Generated reports and logs", "Stored under configured runtime directories and referenced by execution evidence.", "Set retention by business, audit, and privacy need; back up only when required."],
            ["Schedules and profiles", "Archived instead of deleted so history remains explainable.", "Define who may archive and how long history is retained."],
            ["Workflow and Planning evidence", "Persisted without a built-in automatic purge.", "Approve retention periods, legal holds, export needs, and deletion procedure before production."],
            ["Authentication events", "Persisted for security investigation without an automatic purge.", "Align with security-monitoring and privacy policy."],
            ["Assistant conversations", "User-owned deletion is supported; execution evidence is separate.", "Define conversation retention, privacy expectations, and checkpoint cleanup coordination."],
            ["Oracle artifacts and RTP registry", "Updated or deactivated by synchronization; verification history and errors remain in current records/sync runs.", "Schedule synchronization and define stale-record review."],
        ],
        [2300, 3900, 3160],
    )
    add_heading(doc, "PostgreSQL backup minimum")
    add_numbered(doc, [
        "Use the database platform's supported consistent backup mechanism and encrypt backup storage.",
        "Include application tables, Alembic version state, and LangGraph checkpoint tables in the same recoverable database protection plan.",
        "Record database version, application code tag, Alembic heads, checkpoint package version, timestamp, and backup identifier.",
        "Test restoration into an isolated environment; a backup is not proven until it can be restored and validated.",
        "After restore, verify Alembic heads and checkpoint setup before starting workers.",
        "Reconcile any Oracle jobs that were running at the backup or failure boundary before retrying them.",
    ])
    add_heading(doc, "Recovery order")
    add_table(
        doc,
        ["Order", "Recovery check"],
        [
            ["1", "PostgreSQL is restored, reachable, and protected from unintended writers."],
            ["2", "Code release matches the restored schema, or an approved forward migration is applied."],
            ["3", "LangGraph checkpoint setup succeeds and agent state is readable."],
            ["4", "Runtime storage and report/log paths are restored or intentionally empty according to policy."],
            ["5", "API starts in web mode and health/readiness checks pass."],
            ["6", "Oracle configuration and selected application are verified using read-only checks."],
            ["7", "Queue rows and RUNNING/RECOVERY_REQUIRED executions are reconciled with Oracle."],
            ["8", "Workers and schedules resume only after the reconciliation decision is recorded."],
        ],
        [1200, 8160],
    )

    add_chapter(
        doc,
        16,
        "Operate a production governance checklist",
        "Production readiness means every control has an owner, evidence, and a tested response when it fails.",
    )
    add_heading(doc, "Before production access")
    add_bullets(doc, [
        "Serve one trusted HTTPS origin; set WEB_SECURE_COOKIES=true and keep a stable WEB_SESSION_SECRET across API instances.",
        "Use a dedicated least-privilege PostgreSQL login, encrypted remote database transport, protected backups, capacity monitoring, and tested restoration.",
        "Use a dedicated Oracle integration identity with only the permissions required by approved operations; record the effective identity in execution evidence.",
        "Create at least two controlled Service Administrator accounts and verify the last-administrator safeguard.",
        "Complete an initial and periodic role review; separate daily operational accounts from elevated administration.",
        "Store database, Oracle, OIDC, model-provider, SMTP, and EPM Automate secrets in an approved secret manager or protected service configuration.",
        "Run Alembic and LangGraph setup as explicit deployment steps; block startup on schema drift.",
        "Run API and workers under service supervision with the same database and approved shared storage.",
        "Centralize logs, preserve X-Request-ID, synchronize clocks, and alert on authentication abuse, queue age, worker health, and recovery-required executions.",
        "Define retention, archival, purge, legal hold, conversation privacy, report/log lifecycle, and backup policies before broad use.",
        "Test one read-only Oracle request, one low-risk governed execution, one schedule occurrence, one worker failure/recovery exercise, and one backup restoration.",
    ])
    add_heading(doc, "Periodic administrator review")
    add_table(
        doc,
        ["Frequency", "Review"],
        [
            ["Daily or each operating window", "API/worker health, overdue queue items, failed schedules, RECOVERY_REQUIRED, Oracle environment health, and storage capacity."],
            ["Weekly", "Authentication anomalies, inactive or stale users, notification failures, catalog synchronization errors, and repeated operation failures."],
            ["Monthly", "Role and token review, Oracle integration access, backup success, restore sampling, database growth, retention exceptions, and unresolved recovery cases."],
            ["Each release", "Code tag, tests, migration review, backup, Alembic heads, LangGraph setup, frontend build, security configuration, smoke tests, and release evidence."],
            ["At least annually", "Threat model, identity-provider mappings, disaster recovery exercise, retention policy, privileged access, supplier/provider review, and penetration testing plan."],
        ],
        [2100, 7260],
    )
    add_heading(doc, "Known current boundaries")
    add_bullets(doc, [
        "There is no built-in universal database retention or purge engine; policy and controlled implementation are still required.",
        "CORS is not enabled globally; deployments should preserve same-origin behavior unless an exact allowlist is deliberately implemented and tested.",
        "SMTP is the only current external email adapter, and notification delivery is intentionally independent of execution success.",
        "API tokens currently cover the Pipeline/Excel integration contract rather than general administrative automation.",
        "Application logging writes to stdout; centralized storage, alerting, tamper protection, and retention are deployment responsibilities.",
        "Database at-rest encryption, PostgreSQL TLS, host hardening, vulnerability management, WAF/rate limiting, and disaster recovery are infrastructure controls, not automatically provided by FastAPI.",
        "Platform RBAC does not replace Oracle EPM security; the effective Oracle identity must remain least privileged.",
    ])
    add_note(doc, "Final control.", "When an Oracle outcome is uncertain, stop and reconcile. Convenience, scheduling, and AI assistance never justify automatically repeating a write that Oracle may already have accepted.")

    add_chapter(
        doc,
        17,
        "Use the administration quick reference",
        "This final reference turns the guide into a repeatable release and incident checklist.",
    )
    add_heading(doc, "Database release commands")
    add_code_block(
        doc,
        "python -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup",
        "Quick reference 9.1. Current application and LangGraph schema preparation.",
    )
    add_heading(doc, "Health and evidence sequence")
    add_numbered(doc, [
        "GET /health/live confirms the API process can serve requests.",
        "GET /health/ready confirms the API and PostgreSQL are available.",
        "Sign in and review the selected Oracle application and deployment mode.",
        "Synchronize or inspect one read-only Oracle catalog appropriate to the role.",
        "Submit one approved low-risk execution and record its execution ID.",
        "Verify queue state, workflow steps, Oracle evidence, initiating user, trigger source, and effective Oracle username.",
        "Confirm notifications and logs correlate to the same event without becoming the source of truth.",
    ])
    add_heading(doc, "Incident evidence to capture")
    add_table(
        doc,
        ["Capture", "Example"],
        [
            ["Release", "Application version/tag, Alembic heads, LangGraph package/checkpoint setup state."],
            ["Request", "X-Request-ID, endpoint, time, user, role, authentication method, and client IP where appropriate."],
            ["Execution", "Execution ID, trigger source, target, queue status, lease owner, workflow status, step, and Oracle job ID."],
            ["Environment", "Oracle base URL classification, selected application, deployment mode, and effective Oracle execution username - never the password."],
            ["Failure", "First safe error, Oracle messages/statistics, worker log excerpt, recovery state, and any user-visible reference."],
            ["Decision", "Approval or rejection evidence, retry authorization, Oracle reconciliation result, owner, and timestamps."],
        ],
        [2200, 7160],
    )
    add_heading(doc, "Glossary")
    add_table(
        doc,
        ["Term", "Plain-language meaning"],
        [
            ["Authentication", "Proof of who a user or external client is."],
            ["Authorization", "Decision about what that identity may do."],
            ["RBAC", "Role-based access control: permissions are assigned through named roles."],
            ["CSRF", "Protection against another site causing an authenticated browser to submit an unwanted change."],
            ["CORS", "Browser policy controlling whether one web origin may call another origin."],
            ["Migration", "A versioned, ordered database structure change."],
            ["Checkpoint", "Durable LangGraph state that lets an interrupted Assistant workflow resume."],
            ["Lease", "Time-limited worker ownership of one queue item."],
            ["Idempotency", "Protection that makes a repeated request reuse a recorded decision instead of causing duplicate work."],
            ["Recovery required", "A terminal safety state used when Oracle may have accepted work but the worker outcome is uncertain."],
            ["Trigger source", "The origin of work: manual, scheduled, API, Excel, or AI Agent."],
            ["System of record", "The authoritative owner of a kind of data or configuration."],
        ],
        [2400, 6960],
    )
    add_note(doc, "Next handbook volume.", "Document 10 continues with developer extension, automated testing, release engineering, deployment topology, monitoring, scalability, and the future improvement roadmap.")

    _finish_document(doc)
    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
