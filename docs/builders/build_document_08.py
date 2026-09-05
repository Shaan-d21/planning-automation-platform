"""Build Document 08: EPM Assistant and Governed AI Agent."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_DARK,
    BLUE_LIGHT,
    GOLD,
    GOLD_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    RED_LIGHT,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    add_bullets,
    add_callout,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    font,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    wrap,
)
from docs.builders.build_document_03 import (  # noqa: E402
    add_code_block,
    compact_trailing_empty_paragraph,
    configure_styles,
    patch_list_numbering,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-08"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / "BISP_EPM_Automation_Document_08_EPM_Assistant_and_Governed_AI_Agent.docx"
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
MUTED_HEX = f"#{MUTED}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"
RED_HEX = f"#{RED}"


def _arrow(draw: ImageDraw.ImageDraw, start, end, color=BLUE_HEX, width=5) -> None:
    draw.line((*start, *end), fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 17 * direction, y2 - 10), (x2 - 17 * direction, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 10, y2 - 17 * direction), (x2 + 10, y2 - 17 * direction)]
    draw.polygon(points, fill=color)


def _title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((55, 34), title, font=font(34, bold=True), fill=NAVY_HEX)
    y = 82
    for line in wrap(draw, subtitle, font(19), 1280):
        draw.text((55, y), line, font=font(19), fill=MUTED_HEX)
        y += 26


def _card(draw: ImageDraw.ImageDraw, bounds, title: str, body: str, *, color: str, label: str = "", fill: str = "#FFFFFF") -> None:
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=18, fill=fill, outline=color, width=3)
    y = y1 + 18
    if label:
        draw.text((x1 + 20, y), label.upper(), font=font(15, bold=True), fill=color)
        y += 28
    for line in wrap(draw, title, font(22, bold=True), x2 - x1 - 40):
        draw.text((x1 + 20, y), line, font=font(22, bold=True), fill=NAVY_HEX)
        y += 29
    y += 5
    for line in wrap(draw, body, font(16), x2 - x1 - 40):
        draw.text((x1 + 20, y), line, font=font(16), fill=MUTED_HEX)
        y += 22


def create_assistant_journey(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "A request becomes governed work", "The conversational experience removes navigation and typing, but never removes validation, authority, or evidence.")
    steps = [
        (55, "Ask", "Describe the business outcome in ordinary language.", BLUE_HEX),
        (315, "Understand", "Classify intent and expose only the tools needed for this request.", BLUE_HEX),
        (575, "Resolve", "Inspect live Oracle context, recommend artifacts, and collect exact inputs.", ORANGE_HEX),
        (835, "Approve", "Show the operation, artifact, risk, and reviewed values before one explicit decision.", ORANGE_HEX),
        (1095, "Monitor", "Queue the shared worker and return live status plus retained evidence.", TEAL_HEX),
    ]
    for index, (x, title, body, color) in enumerate(steps, 1):
        _card(draw, (x, 175, x + 225, 485), title, body, color=color, label=f"STEP {index}")
        if index < len(steps):
            _arrow(draw, (x + 229, 330), (x + 252, 330), color="#9AA9BF", width=4)
    draw.rounded_rectangle((150, 555, 1250, 680), radius=20, fill=f"#{TEAL_LIGHT}", outline=TEAL_HEX, width=3)
    draw.text((185, 580), "THE IMPORTANT DIVISION", font=font(16, bold=True), fill=TEAL_HEX)
    draw.text((185, 620), "AI interprets and explains. The platform validates, authorizes, executes, and proves.", font=font(23, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_architecture(path: Path) -> None:
    image = Image.new("RGB", (1400, 890), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Current agent architecture", "A provider-neutral LangGraph workflow sits behind the React workspace and reuses the platform's governed Oracle execution services.")
    _card(draw, (55, 150, 325, 330), "React workspace", "Conversations, Markdown answers, structured choice/input cards, approvals, and live execution status.", color=BLUE_HEX, label="USER EXPERIENCE", fill=f"#{BLUE_LIGHT}")
    _arrow(draw, (335, 240), (440, 240))
    _card(draw, (450, 140, 950, 345), "FastAPI agent service", "User-scoped conversations, deterministic intent router, LangGraph orchestration, capability gateway, and approval endpoints.", color=BLUE_HEX, label="CONTROL PLANE")
    _arrow(draw, (960, 240), (1065, 240))
    _card(draw, (1075, 150, 1345, 330), "Model provider", "Gemini or Groq through one provider contract. The model does not call Oracle directly.", color=ORANGE_HEX, label="INTERPRETATION", fill=f"#{ORANGE_LIGHT}")
    _card(draw, (55, 470, 390, 690), "PostgreSQL", "Conversations, messages, tool activity, drafts, decisions, execution evidence, and durable LangGraph checkpoints.", color=TEAL_HEX, label="DURABLE STATE", fill=f"#{TEAL_LIGHT}")
    _card(draw, (530, 470, 870, 690), "Shared execution worker", "Consumes the durable queue, submits governed operations, monitors Oracle jobs, and retains results.", color=TEAL_HEX, label="EXECUTION", fill=f"#{TEAL_LIGHT}")
    _card(draw, (1010, 470, 1345, 690), "Oracle EPM", "Live catalogs, cubes, dimensions, members, jobs, integrations, rules, variables, data, and returned evidence.", color=NAVY_HEX, label="SYSTEM OF RECORD")
    _arrow(draw, (700, 350), (700, 455), color="#9AA9BF")
    _arrow(draw, (950, 580), (995, 580), color=TEAL_HEX)
    _arrow(draw, (515, 580), (405, 580), color=TEAL_HEX)
    draw.rounded_rectangle((210, 755, 1190, 830), radius=18, fill="#FFFFFF", outline="#C5D1E2", width=2)
    draw.text((245, 778), "One execution framework serves the Assistant, Operations, schedules, and lifecycle work - not a separate hidden automation path.", font=font(19, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_trust_boundary(path: Path) -> None:
    image = Image.new("RGB", (1400, 800), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "The trust boundary", "Useful language intelligence is kept separate from the controls that protect Oracle EPM.")
    _card(draw, (55, 155, 650, 660), "AI model", "Understands the request\n\nExplains capabilities\n\nSuggests likely Oracle artifacts\n\nSummarizes read-only findings\n\nAsks for missing business context", color=BLUE_HEX, label="MAY INTERPRET")
    _card(draw, (750, 155, 1345, 660), "Deterministic platform", "Checks the signed-in user's permission\n\nDiscovers current Oracle artifacts\n\nNormalizes and validates exact inputs\n\nFreezes choices around approval\n\nQueues the worker and records evidence", color=TEAL_HEX, label="MUST CONTROL", fill=f"#{TEAL_LIGHT}")
    draw.rounded_rectangle((560, 310, 840, 500), radius=22, fill=f"#{GOLD_LIGHT}", outline=ORANGE_HEX, width=4)
    draw.text((610, 340), "NO DIRECT", font=font(18, bold=True), fill=ORANGE_HEX)
    draw.text((590, 380), "MODEL-TO-ORACLE", font=font(20, bold=True), fill=NAVY_HEX)
    draw.text((650, 420), "PATH", font=font(20, bold=True), fill=NAVY_HEX)
    draw.text((586, 460), "Explicit approval is required", font=font(15, bold=True), fill=RED_HEX)
    draw.text((210, 720), "A fluent answer is not evidence. The platform's validated card, Oracle response, execution record, and step evidence are the source of truth.", font=font(19, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_artifact_resolution(path: Path) -> None:
    image = Image.new("RGB", (1400, 830), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "How the Assistant chooses an Oracle artifact", "Exact matches are preferred; explainable recommendations reduce effort; ambiguity always returns to the user.")
    _card(draw, (55, 165, 310, 355), "Read request", "Extract purpose and any exact artifact name from the current conversation.", color=BLUE_HEX, label="1")
    _arrow(draw, (320, 260), (410, 260))
    _card(draw, (420, 165, 675, 355), "Inspect catalog", "Use artifacts valid for the currently connected Planning environment.", color=BLUE_HEX, label="2")
    _arrow(draw, (685, 260), (775, 260))
    _card(draw, (785, 165, 1040, 355), "Score matches", "Compare task words with names; label strong or possible matches with a reason.", color=ORANGE_HEX, label="3")
    _arrow(draw, (1050, 260), (1140, 260))
    _card(draw, (1150, 165, 1345, 355), "Decide", "Auto-select only when the choice is safe and unambiguous.", color=TEAL_HEX, label="4", fill=f"#{TEAL_LIGHT}")
    branches = [
        ((70, 495, 395, 715), "Exact or one safe match", "Prefill the artifact and continue to its required input card.", TEAL_HEX, f"#{TEAL_LIGHT}"),
        ((535, 495, 860, 715), "Several plausible matches", "Show recommendations and the complete live list in a One choice needed card.", ORANGE_HEX, f"#{GOLD_LIGHT}"),
        ((1000, 495, 1330, 715), "Not in local catalog", "For Pipeline or Data Integration, synchronize with Oracle and allow authorized exact registration.", RED_HEX, f"#{RED_LIGHT}"),
    ]
    for bounds, title, body, color, fill in branches:
        _card(draw, bounds, title, body, color=color, fill=fill)
    image.save(path)


def create_approval_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 780), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "One explicit approval, one monitored run", "Directly supported actions no longer bounce through a second governed screen after the Assistant has collected and validated the inputs.")
    steps = [
        ((55, 160, 330, 380), "Reviewed proposal", "Exact operation, artifact, risk, effect, and normalized inputs.", BLUE_HEX),
        ((405, 160, 680, 380), "Approve and run", "The user selects the dedicated control. A chat message such as yes is not executable approval.", ORANGE_HEX),
        ((755, 160, 1030, 380), "Reserve decision", "A unique request record prevents a double click or replay from submitting twice.", ORANGE_HEX),
        ((1105, 160, 1345, 380), "Queue work", "The shared durable execution manager accepts the governed operation.", TEAL_HEX),
    ]
    for i, (bounds, title, body, color) in enumerate(steps):
        _card(draw, bounds, title, body, color=color, label=str(i + 1))
        if i < len(steps) - 1:
            _arrow(draw, (bounds[2] + 8, 270), (steps[i + 1][0][0] - 8, 270), color="#9AA9BF", width=4)
    _arrow(draw, (1225, 395), (1225, 485), color=TEAL_HEX)
    _card(draw, (880, 500, 1345, 690), "Live execution card", "Queued, running, success, or failure; execution ID; step details; Oracle counts when available; link to Jobs & Activity.", color=TEAL_HEX, label="EVIDENCE", fill=f"#{TEAL_LIGHT}")
    _card(draw, (55, 500, 680, 690), "Reject or cancel", "Rejecting a proposal records the decision without starting Oracle. Cancelling an input or choice card safely ends that pending interaction.", color=RED_HEX, label="SAFE EXIT", fill=f"#{RED_LIGHT}")
    image.save(path)


def create_orchestration_choice(path: Path) -> None:
    image = Image.new("RGB", (1400, 800), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Pipeline or standalone flow?", "The Assistant prefers an existing Oracle Pipeline for an end-to-end lifecycle; it can prepare a platform-managed sequence when no suitable Pipeline exists.")
    _card(draw, (475, 145, 925, 300), "Multi-step request", "Example: load metadata, refresh the cube, load data, calculate, and publish.", color=BLUE_HEX, label="USER OUTCOME")
    _arrow(draw, (700, 310), (700, 390))
    draw.rounded_rectangle((480, 405, 920, 485), radius=18, fill=f"#{GOLD_LIGHT}", outline=ORANGE_HEX, width=3)
    draw.text((535, 430), "Is a suitable registered Oracle Pipeline available?", font=font(19, bold=True), fill=NAVY_HEX)
    _arrow(draw, (500, 495), (330, 565), color=TEAL_HEX)
    _arrow(draw, (900, 495), (1070, 565), color=ORANGE_HEX)
    _card(draw, (55, 575, 580, 745), "Use the Oracle Pipeline", "One Oracle-governed orchestration with its configured stages, variables, file requirements, and notifications.", color=TEAL_HEX, label="PREFERRED WHEN IT FITS", fill=f"#{TEAL_LIGHT}")
    _card(draw, (820, 575, 1345, 745), "Use a standalone flow", "Configure 2-12 supported operations, approve once, run sequentially, and stop at the first failed step.", color=ORANGE_HEX, label="WHEN NO PIPELINE FITS", fill=f"#{ORANGE_LIGHT}")
    image.save(path)


def create_provider_model(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Provider-neutral by design", "The agent workflow, tools, approval model, and Oracle integration remain stable when the configured language-model provider changes.")
    _card(draw, (55, 180, 390, 500), "LangGraph workflow", "Conversation state, tool rounds, durable interrupts, approval resume, and deterministic shortcuts.", color=BLUE_HEX, label="STABLE ORCHESTRATION", fill=f"#{BLUE_LIGHT}")
    _arrow(draw, (405, 340), (520, 340))
    _card(draw, (530, 205, 870, 475), "Provider contract", "One model step receives messages and tool schemas, then returns text and correlated tool calls.", color=NAVY_HEX, label="ADAPTER BOUNDARY")
    _arrow(draw, (885, 295), (1000, 245), color=ORANGE_HEX)
    _arrow(draw, (885, 385), (1000, 445), color=ORANGE_HEX)
    _card(draw, (1010, 160, 1345, 330), "Google Gemini", "Current default example: gemini-3.5-flash-lite.", color=ORANGE_HEX, label="OPTION A", fill=f"#{ORANGE_LIGHT}")
    _card(draw, (1010, 395, 1345, 565), "Groq", "Current example: openai/gpt-oss-120b with bounded prompts.", color=TEAL_HEX, label="OPTION B", fill=f"#{TEAL_LIGHT}")
    draw.rounded_rectangle((200, 615, 1200, 700), radius=18, fill="#FFFFFF", outline="#C5D1E2", width=2)
    draw.text((245, 642), "Switch configuration, restart the API, and begin a new conversation. Oracle controls and evidence do not change.", font=font(19, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "journey": ASSET_DIR / "assistant_journey.png",
        "architecture": ASSET_DIR / "agent_architecture.png",
        "trust": ASSET_DIR / "trust_boundary.png",
        "artifacts": ASSET_DIR / "artifact_resolution.png",
        "approval": ASSET_DIR / "approval_execution.png",
        "orchestration": ASSET_DIR / "orchestration_choice.png",
        "providers": ASSET_DIR / "provider_abstraction.png",
    }
    create_assistant_journey(assets["journey"])
    create_architecture(assets["architecture"])
    create_trust_boundary(assets["trust"])
    create_artifact_resolution(assets["artifacts"])
    create_approval_flow(assets["approval"])
    create_orchestration_choice(assets["orchestration"])
    create_provider_model(assets["providers"])
    return assets


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    kicker = doc.add_paragraph(f"CHAPTER {number:02d}", style="Kicker")
    kicker.paragraph_format.page_break_before = True
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "BISP SOLUTIONS  /  EPM ASSISTANT AND GOVERNED AI AGENT"
    set_run_font(paragraph.runs[0], size=8.5, color=MUTED, bold=True)
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.columns[0].width = Inches(5.7)
    table.columns[1].width = Inches(0.8)
    left = table.cell(0, 0).paragraphs[0]
    left.text = "Document 08  |  EPM Assistant and Governed AI Agent"
    set_run_font(left.runs[0], size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(right)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:color"), LINE)
        borders.append(top)
        tc_pr.append(borders)


def add_cover(doc: Document) -> None:
    if LOGO.is_file():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions company logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(30)
    doc.add_paragraph("PLATFORM HANDBOOK", style="Kicker")
    doc.add_paragraph("EPM Assistant and\nGoverned AI Agent", style="Title")
    doc.add_paragraph("BISP Solutions Oracle EPM Automation Platform", style="Subtitle")
    doc.add_paragraph(
        "A beginner-friendly guide to conversational Oracle EPM assistance: live inspection, explainable artifact recommendations, structured inputs, explicit approvals, monitored execution, provider configuration, security controls, and practical limitations.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [["08 of the platform handbook", "3 September 2026", "Planners, business users, FP&A leaders, Oracle EPM consultants, administrators, security reviewers, and support teams"]],
        [2200, 2100, 5060],
    )
    add_callout(
        doc,
        "Current implementation only",
        "This volume reflects the current React Assistant workspace, FastAPI agent APIs, deterministic intent and capability controls, LangGraph orchestration, Gemini and Groq adapters, PostgreSQL persistence, durable execution worker, and live Oracle EPM services. It excludes the retired interface and the earlier process-creation experiment.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraph.add_run("BISP Solutions  |  Conversational convenience with governed control"), size=10, color=MUTED, bold=True)


def _finish_document(doc: Document) -> None:
    compact_trailing_empty_paragraph(doc)
    for table in doc.tables:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            header_pr.append(header)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for section in doc.sections:
        for footer in (section.footer, section.even_page_footer):
            for table in footer.tables:
                header_pr = table.rows[0]._tr.get_or_add_trPr()
                if header_pr.find(qn("w:tblHeader")) is None:
                    header = OxmlElement("w:tblHeader")
                    header.set(qn("w:val"), "true")
                    header_pr.append(header)
    properties = doc.core_properties
    properties.title = "EPM Assistant and Governed AI Agent"
    properties.subject = "Current conversational Oracle EPM assistance, approval, execution, and safety guide"
    properties.author = "BISP Solutions"
    properties.keywords = "Oracle EPM, Planning, AI Agent, EPM Assistant, LangGraph, Gemini, Groq, governed automation"
    properties.comments = "Documents the current React, FastAPI, LangGraph, PostgreSQL, execution-worker, and Oracle EPM agent implementation as of 3 September 2026; retired UI excluded."


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    patch_list_numbering(doc)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS VOLUME", style="Kicker")
    doc.add_paragraph("Follow a request from conversation to evidence", style="Heading 1")
    doc.add_paragraph(
        "The Assistant is the platform's natural-language entry point. This guide explains what a business user sees, what the AI contributes, what deterministic controls protect Oracle, and where an administrator configures or tests the service.",
        style="Lead",
    )
    add_table(
        doc,
        ["If you are...", "Start with", "Primary goal"],
        [
            ["A planner or business user", "Chapters 1, 4-8", "Ask for work, select the correct artifact, provide guided inputs, approve, and monitor."],
            ["An FP&A process owner", "Chapters 6-9", "Understand operations, orchestration choices, approval, evidence, and recovery."],
            ["A Service Administrator", "Chapters 2, 9-12", "Understand permissions, persistence, provider setup, testing, support, and limits."],
            ["A security or audit reviewer", "Chapters 3, 8-10", "Verify trust boundaries, explicit authorization, immutable decisions, and retained evidence."],
            ["A developer or consultant", "Chapters 2, 5, 10-12", "Understand LangGraph, provider adapters, capability tools, configuration, and release gates."],
        ],
        [2400, 1800, 5160],
    )
    add_heading(doc, "Contents")
    add_table(
        doc,
        ["Part", "Chapters", "What it explains"],
        [
            ["Purpose and design", "1-3", "User value, architecture, trust boundary, permissions, and least privilege."],
            ["Guided work", "4-7", "Conversation UX, read-only tools, artifact matching, and structured operation inputs."],
            ["Governed execution", "8-9", "One-click approval, durable queuing, monitoring, Pipeline choice, standalone flows, and recovery."],
            ["Operate the agent", "10-12", "Persistence, privacy, provider configuration, testing, troubleshooting, boundaries, and glossary."],
        ],
        [2000, 1600, 5760],
    )

    add_chapter(doc, 1, "Understand what the Assistant is", "The Assistant turns a business outcome into the smallest safe next action; it is not an unrestricted robot with direct Oracle access.")
    add_figure(doc, assets["journey"], "Figure 8.1 - From conversational request to governed evidence.", "Five-step journey from a natural-language request through classification, live artifact and input resolution, explicit approval, and monitored evidence, with a clear division between AI interpretation and platform control.", width=6.35)
    add_heading(doc, "Why it exists")
    add_bullets(doc, [
        "Give business users one place to ask what the platform can do, inspect live EPM context, prepare supported operations, and follow execution.",
        "Reduce navigation, technical terminology, exact-name typing, and repeated re-entry of information already stated in the conversation.",
        "Recommend likely Oracle artifacts for a described task without forcing a user to search hundreds of rules, maps, Pipelines, or integrations manually.",
        "Keep the same role, permission, approval, execution, monitoring, and audit controls used elsewhere in the platform.",
        "Explain limitations honestly and route unsupported work to the appropriate governed workspace rather than pretending it was completed.",
    ])
    add_heading(doc, "Three kinds of assistance")
    add_table(doc, ["Mode", "Example", "Result"], [
        ["Explain", "What can this platform automate?", "A readable answer derived from the current capability catalog and permitted tools."],
        ["Inspect", "Show the latest failed execution or review a Forecast slice.", "Live or retained read-only evidence, constrained by role and Oracle access."],
        ["Prepare and run", "Run the travel-expense rule or monthly Pipeline.", "Exact artifact and inputs, one explicit approval, then monitored execution."],
    ], [2000, 3300, 4060])
    add_callout(doc, "Not a second security model", "The Assistant cannot grant access or bypass Oracle. It only makes existing platform capabilities easier to discover and use.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 2, "See how the architecture works", "The current design separates presentation, interpretation, deterministic control, persistence, execution, and Oracle connectivity.")
    add_figure(doc, assets["architecture"], "Figure 8.2 - Current EPM Assistant architecture.", "React connects to a FastAPI agent service using deterministic routing and LangGraph; a provider adapter supplies language intelligence, PostgreSQL stores durable state, the shared worker runs governed jobs, and Oracle EPM remains the system of record.", width=6.35)
    add_heading(doc, "Component responsibilities")
    add_table(doc, ["Component", "Responsibility", "Why it is separate"], [
        ["React Assistant workspace", "Conversation list, messages, structured cards, explicit approval, execution monitoring, and navigation.", "Keeps the user experience replaceable without moving business controls into the browser."],
        ["FastAPI agent API", "Authenticates the user, owns conversations, routes intent, exposes allowed capabilities, validates resumes, and handles approvals.", "All authoritative checks stay server-side."],
        ["LangGraph orchestrator", "Coordinates model turns, tool calls, durable interruptions, approval waits, and resume behavior.", "Long-running human-in-the-loop work can pause safely."],
        ["Provider adapters", "Translate the common agent contract to Gemini or Groq requests and tool results.", "The model vendor can change without rewriting the workflow."],
        ["Capability gateway", "Offers a narrow set of read-only and preparation tools; never arbitrary REST.", "The model can ask only for actions the platform has defined."],
        ["PostgreSQL", "Stores application records and durable graph checkpoints in managed schemas.", "Conversation and approval state survives process restarts."],
        ["Execution manager and worker", "Queues, submits, monitors, and records Oracle operations.", "The Assistant reuses the platform's trusted execution path."],
        ["Oracle EPM", "Owns Planning artifacts, security, data, jobs, and final operation responses.", "Oracle remains the source of truth."],
    ], [2100, 4300, 2960])
    add_heading(doc, "What LangGraph adds")
    add_bullets(doc, [
        "A clear model -> tools -> model loop with a maximum number of tool rounds.",
        "Durable interrupt and resume for clarification, structured input, and explicit approval.",
        "Checkpointed conversation workflow keyed by authenticated user, conversation, and connected Oracle environment.",
        "Deterministic shortcuts for common intents so the safest route does not depend on the model inventing the next tool call.",
    ])
    add_callout(doc, "Environment-aware state", "The connected Oracle base URL and application name form part of the agent environment key. When the environment changes, stale contextual assumptions and pending choices cannot silently carry into the new application.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 3, "Understand the safety and permission model", "The language model may interpret and recommend; deterministic application code must decide what is allowed and what Oracle receives.")
    add_figure(doc, assets["trust"], "Figure 8.3 - AI interpretation versus deterministic control.", "The AI model explains and recommends on one side; the deterministic platform enforces permissions, live discovery, input validation, approval, queuing, and evidence on the other, with no direct model-to-Oracle path.", width=6.35)
    add_heading(doc, "Non-negotiable controls")
    add_bullets(doc, [
        "Conversation text is never treated as execution approval. The user must select the dedicated Approve and run control.",
        "The server rechecks the exact permission at preparation, approval, and execution boundaries.",
        "The model cannot call an arbitrary URL, submit arbitrary REST, invent an operation, or execute a tool outside the capability gateway.",
        "Live artifact choices are frozen while the graph is interrupted. A stale, changed, or mismatched choice is rejected on resume.",
        "The Assistant never claims success unless the shared execution system and Oracle evidence confirm it.",
        "Credentials, API keys, session tokens, passwords, and uploaded file contents are excluded from model context and audit snapshots.",
    ])
    add_heading(doc, "Role-aware experience")
    add_table(doc, ["Platform role", "Assistant experience"], [
        ["Service Administrator", "All Assistant capabilities plus administrative operations, variable management, catalog management, Data Review, history, and reports."],
        ["Power User", "Governed business operations, own or permitted user-variable changes, Data Review, reports, history, and Assistant guidance."],
        ["User", "Planning work, Data Review and reports when permitted, own user-variable updates, and Assistant guidance; no general elevated operation execution."],
        ["Viewer", "Read-only Assistant guidance and report generation only; no operations, execution history, or Data Review tools."],
    ], [2350, 7010])
    add_para(doc, "Every role can open the Assistant, but the available tools shrink to the signed-in user's actual permissions. The model never receives a tool that the user is not allowed to use.")
    add_heading(doc, "Least-privilege intent routing")
    add_table(doc, ["Detected intent", "Typical permitted tool group"], [
        ["Platform discovery", "Environment summary and operation catalog."],
        ["History review", "Recent executions and exact execution evidence, only with history permission."],
        ["Data Review", "Cubes, dimensions, members, live slice, and comparison, only with Data Review permission."],
        ["Operation preparation", "Artifact discovery and the relevant governed preparation tools, only with the required operation or variable permission."],
        ["General guidance", "A minimal informational tool set; no unnecessary execution tools."],
    ], [2650, 6710])
    add_callout(doc, "Language is not authority", "A helpful recommendation can still be wrong. The validated structured card, current Oracle catalog, explicit user decision, and execution evidence take precedence over the prose response.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 4, "Work in the Assistant", "The workspace combines ordinary conversation with structured controls whenever a choice, value, or authorization must be exact.")
    add_heading(doc, "Start a request")
    add_numbered(doc, [
        "Open EPM Assistant. The header identifies the configured provider and model.",
        "Start a new conversation for a new business objective or after changing provider configuration.",
        "Use a suggested prompt such as Explore, Review data, Activity, Environment, Prepare, Quality, or Guidance, or describe the outcome in your own words.",
        "Review the Assistant's explanation and any live context it retrieved.",
        "Complete the structured card when the Assistant needs an exact artifact, operation input, or approval.",
    ])
    add_heading(doc, "Why cards replace free text")
    add_table(doc, ["Card", "What it prevents", "What the user does"], [
        ["One choice needed", "Guessing among similar Oracle artifacts.", "Select a recommendation or choose from the complete current list."],
        ["Review live Oracle context", "Misspelled cubes, dimensions, members, variables, jobs, periods, or file names.", "Use dropdowns, member buttons, and exact compatible choices."],
        ["Provide run inputs", "Unstructured key-value text and unsupported fields.", "Choose required operation-specific values and files."],
        ["Your approval is required", "Accidental or implied execution.", "Review the exact effect, then Approve and run or Reject proposal."],
        ["Live execution", "Unverifiable success statements.", "Follow queued/running/final status and open Jobs & Activity."],
    ], [2400, 3250, 3710])
    add_heading(doc, "Conversation behavior")
    add_bullets(doc, [
        "A conversation belongs to one authenticated user. Another user cannot load it by guessing its identifier.",
        "The first request becomes a concise conversation title, and the current implementation accepts messages up to 4,000 characters.",
        "Only a bounded number of recent messages are supplied to the model; older application records remain in PostgreSQL but do not grow every model request forever.",
        "The composer is disabled while a choice, input, or approval card is pending. Complete or cancel that card before starting a different branch.",
        "Responses render headings, lists, code, and quotes for readability. Structured business choices use platform cards rather than Markdown tables generated by the model.",
    ])
    add_callout(doc, "Context is useful, not permanent authority", "The Assistant can remember the selected artifact or business context within a conversation, but a later governed action still revalidates current permissions, inputs, and Oracle state.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 5, "Use read-only investigation tools", "The Assistant can answer with current platform facts when the user's role permits access, instead of relying only on general language-model knowledge.")
    add_heading(doc, "Current read-only capabilities")
    add_table(doc, ["Capability", "What it can retrieve", "Important boundary"], [
        ["Environment summary", "Configured Planning application and safe connection context.", "No password, API key, or session token."],
        ["Operation catalog", "Available platform operations, purpose, category, and risk level.", "Availability is role-filtered."],
        ["Recent execution history", "Latest runs and statuses within a bounded limit.", "Requires history permission."],
        ["Execution evidence", "Exact or latest execution, steps, failures, and Oracle counts when returned.", "Unavailable counters remain unavailable; they are never invented."],
        ["Planning cubes", "Cubes visible through the live Data Review service.", "Oracle release and configured identity can limit discovery."],
        ["Cube dimensions", "Dimensions exposed for the selected cube.", "Exact-name fallback may be needed on older endpoints."],
        ["Dimension member search", "Bounded live member results, up to 100 per request.", "The Assistant cannot invent missing members."],
        ["Review data slice", "One exact read-only POV, row, and column selection.", "No data writeback."],
        ["Compare data slices", "Matching source and target intersections with tolerance and bounded mismatch evidence.", "Layout alignment is required."],
        ["Operation artifacts", "Current rules, maps, Pipelines, integrations, jobs, and registered definitions for preparation.", "Only current-environment choices are eligible."],
    ], [2200, 4100, 3060])
    add_heading(doc, "Data Review inside the conversation")
    add_bullets(doc, [
        "Choose a live cube, place every dimension in POV, rows, or columns, and select members with mouse-first controls.",
        "Display the returned grid inline with row, column, cell, and missing metrics, plus truncation notice when applicable.",
        "Export the slice to Excel or open the full Data Review workspace for spreadsheet-style analysis.",
        "Compare source and target slices and preview up to 100 mismatches in the conversation, with an Excel evidence option.",
        "Persist validated selection context, not the financial cell values. A previous live slice must be read again to obtain current values.",
    ])
    add_callout(doc, "Live Oracle limits still apply", "Cloud and on-premises releases do not expose identical discovery APIs. When metadata cannot be discovered, the Assistant may request an exact known name, but Oracle's live data or operation endpoint remains the final validator.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 6, "Resolve the correct Oracle artifact", "The Assistant combines exact-name recognition, deterministic recommendations, complete live choices, and controlled catalog recovery.")
    add_figure(doc, assets["artifacts"], "Figure 8.4 - Explainable artifact resolution.", "Flow from request and current catalog through deterministic matching to exact selection, a One choice needed card, or Pipeline and Data Integration synchronization and registration.", width=6.35)
    add_heading(doc, "How recommendations work")
    add_bullets(doc, [
        "Exact names in the request are recognized despite harmless differences in case, punctuation, spaces, or underscores.",
        "When the user describes a purpose, deterministic matching compares meaningful task words with artifact names, including CamelCase terms.",
        "Common words are ignored, exact and partial word overlap is scored, and each result includes a short reason and Strong match or Possible match label.",
        "The platform returns at most five recommendations while keeping the complete current list available in the card.",
        "A generic or ambiguous request is never auto-selected simply to save a click.",
    ])
    add_heading(doc, "What happens in common cases")
    add_table(doc, ["Situation", "User experience"], [
        ["Exact current artifact named", "The card is prefilled and moves to required inputs or approval."],
        ["One clearly safe current match", "The platform can auto-select it and still shows the exact artifact in the final review."],
        ["Several plausible matches", "Recommended matches appear first; the user selects one or searches the complete list."],
        ["No good match", "The full list remains available, but the Assistant does not guess."],
        ["Pipeline or Data Integration missing from the registered catalog", "The user can synchronize with Oracle. An authorized catalog manager may register an exact identifier when discovery supports or requires it."],
        ["Registered item is stale for this environment", "It is hidden or marked unavailable rather than treated as executable; the final run verifies again."],
    ], [3200, 6160])
    add_callout(doc, "Recommendations do not execute", "Selecting a recommendation only resolves the artifact. The operation still requires validated inputs and an explicit approval before it enters the queue.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 7, "Provide trusted operation inputs", "Every operation has a server-owned schema and a purpose-built card, so the user does not need to invent technical key-value syntax.")
    add_heading(doc, "Current guided operations")
    add_table(doc, ["Operation", "Guided input behavior"], [
        ["Business Rule", "Choose Calculation Manager defaults or provide synchronized runtime prompts. Required prompt names and known defaults are displayed; exact manual names are a compatibility fallback."],
        ["Data Map", "Choose the exact map, make the clear-target decision explicit, and optionally supply member or exclusion overrides."],
        ["Pipeline", "Select live variables, including modeled year and period choices, and satisfy every file requirement using the configured Oracle file, an existing Inbox file, or a local upload."],
        ["Data Integration", "Choose mapped, Planning, or exact periods, import/export modes, and configured, Inbox, or uploaded input file."],
        ["Planning Data Import", "Use the job's configured file, choose an existing Inbox file, or upload locally; optionally request an error output file."],
        ["Metadata Import", "Use configured, Inbox, or uploaded metadata; optionally request error output and run an exact saved Cube Refresh job after successful import."],
        ["Cube Refresh", "Choose an exact current saved refresh job and review its application-wide impact. A generic RefreshCube label is not treated as a real saved job unless Oracle confirms it."],
        ["Substitution Variable", "Update an existing variable or explicitly create one; choose ALL or a live cube scope. Exact scope, name, and value from the request are prefilled when complete."],
        ["User Variable", "Choose the definition, dimension, target user allowed by role, and new member; the current assignment is reread before write."],
        ["Standalone flow", "Configure 2-12 supported steps sequentially and review one combined plan before approval."],
    ], [2500, 6860])
    add_heading(doc, "File choices")
    add_bullets(doc, [
        "Use configured file keeps the file already owned by the Oracle job or integration.",
        "Choose from Inbox lists available server-side files so the user can select with the mouse.",
        "Upload from this computer stores the file temporarily until approval; it does not silently upload to Oracle when merely selected.",
        "Extension, ownership, operation schema, and temporary upload token are validated on the server.",
    ])
    add_heading(doc, "Variable protections")
    add_bullets(doc, [
        "Substitution-variable names are nonempty, at most 80 characters, and must not begin with an ampersand; values are nonempty and at most 255 characters.",
        "Updating an existing substitution variable includes the expected current value. If another user changed it meanwhile, the operation stops rather than overwriting silently.",
        "User-variable values and identifiers are nonempty and bounded; the current assignment and resulting Oracle value are confirmed.",
        "A universal semantic member check is not possible for every substitution variable because some valid variables contain dates, text, expressions, or application-specific tokens. Oracle acceptance and returned value remain authoritative.",
    ])
    add_callout(doc, "Schema-owned inputs", "The server rejects unsupported or unexpected fields. This protects the platform from stale cards, model-invented parameters, and inputs copied from a different operation.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 8, "Approve and run safely", "For directly supported actions, the current experience has one final approval and then immediately enters the shared monitored execution workflow.")
    add_figure(doc, assets["approval"], "Figure 8.5 - Explicit approval and durable execution.", "A reviewed proposal moves through the dedicated Approve and run action, immutable decision reservation, durable queue, and live evidence; rejection safely ends the proposal without Oracle execution.", width=6.35)
    add_heading(doc, "What the approval card shows")
    add_bullets(doc, [
        "The operation category and exact Oracle artifact or variable definition.",
        "The risk level and a plain-language description of what approval will do.",
        "Every reviewed runtime value, period, mode, file choice, refresh choice, override, or target user that affects the run.",
        "A clear warning when the action is elevated, application-wide, data-changing, or otherwise consequential.",
    ])
    add_heading(doc, "Why one click does not mean uncontrolled")
    add_numbered(doc, [
        "The user selects Approve and run, not merely a conversational yes.",
        "The API verifies the pending graph state, authenticated user, environment, operation permission, artifact, and normalized inputs.",
        "An immutable decision record with a unique request ID is reserved before submission.",
        "Repeated clicks or retries return the existing processing or submitted result instead of starting another Oracle operation.",
        "The shared execution manager writes the durable queue entry and the worker performs the operation.",
        "The Assistant displays the same execution status and retained evidence available in Jobs & Activity.",
    ])
    add_heading(doc, "Decision evidence")
    add_table(doc, ["Recorded item", "Purpose"], [
        ["Conversation, request, user, operation, and artifact", "Answers who approved what and from which interaction."],
        ["Decision and timestamps", "Distinguishes approval from rejection and records when it occurred and finalized."],
        ["Sanitized payload snapshot and checksum", "Proves the reviewed input context without retaining secrets or raw file content."],
        ["Outcome and execution ID", "Links the human decision to the queued or completed execution."],
    ], [3100, 6260])
    add_callout(doc, "No silent write retry", "If a governed write fails, the Assistant does not quietly submit it again. Recovery requires renewed review and explicit authority.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 9, "Choose orchestration and recover from failure", "The Assistant can match an existing Oracle Pipeline or build a bounded platform-managed sequence without confusing the two models.")
    add_figure(doc, assets["orchestration"], "Figure 8.6 - Oracle Pipeline versus standalone flow.", "A multi-step request is matched first to a suitable registered Oracle Pipeline; when none fits, a platform-managed sequence of 2-12 supported operations can be configured and approved.", width=6.35)
    add_heading(doc, "Prefer an Oracle Pipeline when")
    add_bullets(doc, [
        "The end-to-end lifecycle is already configured and governed in Oracle Data Integration.",
        "Its stages, variables, file requirements, mappings, notifications, and operational ownership match the requested business outcome.",
        "The organization wants one Oracle-owned orchestration and one Oracle job trail.",
    ])
    add_heading(doc, "Use a platform-managed standalone flow when")
    add_bullets(doc, [
        "No registered Pipeline safely matches the requested sequence.",
        "Every step is already a supported governed platform operation.",
        "The user wants one combined review while keeping each underlying Oracle action explicit.",
        "The flow contains between 2 and 12 steps and accepts stop-on-failure behavior.",
    ])
    add_heading(doc, "Execution and recovery")
    add_table(doc, ["Behavior", "Current rule"], [
        ["Order", "Standalone steps run sequentially; the next step starts only after the previous one succeeds."],
        ["Failure", "The flow stops after the first failed step. Completed steps are not automatically rerun."],
        ["Metadata refresh", "An optional Cube Refresh after Metadata Import runs only if the import succeeded."],
        ["Recovery", "A new linked execution starts from the failed step only after renewed review and approval; original evidence remains unchanged."],
        ["Temporary files", "The user must reselect temporary local uploads for recovery because they are not treated as permanent reusable inputs."],
    ], [2600, 6760])
    add_callout(doc, "Automation without duplication", "The Assistant does not rebuild technical work that belongs inside a well-designed Oracle Pipeline. It helps the user find, parameterize, approve, and monitor that Pipeline; standalone orchestration is the controlled fallback.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 10, "Protect data, state, and audit evidence", "The agent persists enough state to be reliable while separating application records, framework checkpoints, transient files, and Oracle business data.")
    add_heading(doc, "What PostgreSQL stores")
    add_table(doc, ["Record area", "Examples", "Lifecycle"], [
        ["Conversations", "Conversation owner, title, provider context, creation and update timestamps.", "User-scoped; deleting the conversation cascades owned message history."],
        ["Messages", "User and Assistant content plus structured response metadata.", "Durable application record with bounded model-visible history."],
        ["Tool activity", "Allowed capability name, status, sanitized inputs, result summary, and timing.", "Operational and diagnostic evidence."],
        ["Action drafts and decisions", "Proposed operation, normalized inputs, explicit decision, actor, checksum, outcome, and execution link.", "Governance evidence; decisions are treated as immutable events."],
        ["LangGraph checkpoints", "Workflow state and pending interrupt data in a framework-owned schema.", "Survives API restart; managed separately from Alembic tables."],
        ["Executions", "Durable queue state, steps, logs, artifacts, and Oracle evidence used by the whole platform.", "Retained according to platform operations policy."],
    ], [2350, 4300, 2710])
    add_heading(doc, "Privacy and information flow")
    add_bullets(doc, [
        "The platform does not provide Oracle passwords, API keys, browser session tokens, or raw uploaded file content to the language model.",
        "The configured provider may process the user's conversation, system instructions, permitted tool schemas, and the result summaries or metadata needed to answer the request. Organizations should choose a provider and retention policy appropriate for their data classification.",
        "Data Review grids use structured platform cards. The selected layout can be retained as context, while live financial cell values are reread from Oracle and are not stored as prior-grid context.",
        "Audit-safe snapshots remove credential-like values and ephemeral upload tokens before persistence.",
        "Oracle data security and the configured Oracle service identity still determine which artifacts and values can be discovered or executed.",
    ])
    add_heading(doc, "Checkpoint setup")
    add_para(doc, "Production uses the PostgreSQL LangGraph saver. The in-memory saver is reserved for tests and Path-based temporary databases; it is not durable deployment storage.")
    add_code_block(doc, "python -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup", "Code excerpt 1. Apply application migrations and initialize or upgrade LangGraph checkpoint tables during deployment.")
    add_callout(doc, "Two migration owners", "Alembic owns the application's business tables. LangGraph's checkpoint setup owns its framework tables. Run both upgrade paths instead of attempting to hand-edit either schema.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 11, "Configure providers and runtime", "Gemini and Groq are current provider options behind one interface; switching changes language-model service, not governance or Oracle execution.")
    add_figure(doc, assets["providers"], "Figure 8.7 - Provider-neutral orchestration.", "A stable LangGraph workflow and provider contract connect to either Google Gemini or Groq, with the same platform tools, permissions, approvals, and Oracle execution path.", width=6.35)
    add_heading(doc, "Google Gemini example")
    add_code_block(doc, "AGENT_PROVIDER=gemini\nAGENT_ORCHESTRATOR=langgraph\nAGENT_MODEL=gemini-3.5-flash-lite\nGEMINI_API_KEY=replace-with-your-key\nAGENT_MAX_TOOL_ROUNDS=4\nAGENT_HISTORY_MESSAGES=20\nLANGGRAPH_STRICT_MSGPACK=true", "Code excerpt 2. Current Gemini configuration example; model availability and quotas are controlled by the provider.")
    add_heading(doc, "Groq example")
    add_code_block(doc, "AGENT_PROVIDER=groq\nAGENT_ORCHESTRATOR=langgraph\nAGENT_MODEL=openai/gpt-oss-120b\nGROQ_API_KEY=replace-with-your-key\nGROQ_MAX_INPUT_TOKENS=2500\nGROQ_MAX_COMPLETION_TOKENS=384\nAGENT_MAX_TOOL_ROUNDS=4\nAGENT_HISTORY_MESSAGES=20", "Code excerpt 3. Current Groq configuration example with bounded input and completion budgets.")
    add_heading(doc, "Switch providers")
    add_numbered(doc, [
        "Stop the API service so it will not retain the old provider client.",
        "Change AGENT_PROVIDER, AGENT_MODEL, and the relevant provider API key in the deployment secret or environment configuration.",
        "Keep AGENT_ORCHESTRATOR=langgraph for the current supported workflow. The legacy value exists only as a compatibility fallback.",
        "Restart the API and begin a new conversation so old provider-native message state is not mixed with the new provider.",
        "Run a read-only environment request and one governed non-production test before enabling business use.",
    ])
    add_heading(doc, "Provider-specific protections")
    add_table(doc, ["Provider", "Current adapter behavior"], [
        ["Gemini", "Preserves complete provider-native content, including thought signatures required for tool-call continuation."],
        ["Groq", "Compacts history under a configured token budget, limits model-visible catalog previews while the platform retains full choices, and returns readable messages for token, quota, network, and credential failures."],
        ["Both", "Receive the same provider-neutral messages and tool schemas and cannot bypass the capability gateway or approval service."],
    ], [2100, 7260])
    add_callout(doc, "Keys and credits are external", "The platform does not create provider credits. A valid key, supported tool-calling model, network access, quota, and provider availability are required. Changing to another key under the same organization or service tier may not change its rate limit.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 12, "Test, troubleshoot, and respect the boundaries", "A governed agent needs deterministic release gates and live user acceptance testing because language services and Oracle environments both vary.")
    add_heading(doc, "Deterministic release gate")
    add_para(doc, "The repository includes a synthetic agent evaluation suite that does not call a language model or Oracle. The current release_v1 suite contains 13 cases and requires 100% overall and 100% critical pass rates in CI.")
    add_bullets(doc, [
        "General routing and platform discovery.",
        "Execution statistics and evidence retrieval.",
        "Business Rule, Pipeline, Data Map, and multi-operation preparation.",
        "Data Review and contextual refinement, including explicit operation intent overriding prior review context.",
        "Prompt injection remaining governed and unsupported destructive requests staying blocked.",
        "Exact renamed rule, compact rule name, and generic requests that must not guess.",
    ])
    add_code_block(doc, "python -m app.agent.evaluation --suite evaluations/agent/release_v1.json --fail-on-threshold", "Code excerpt 4. Run the deterministic agent release gate locally; CI also publishes the evaluation report.")
    add_heading(doc, "Live acceptance checklist")
    add_bullets(doc, [
        "Gemini and Groq availability, correct model name, key, quota, latency, token limits, and tool-call continuation.",
        "Cross-user conversation isolation and role-specific tool visibility.",
        "Current Oracle catalog discovery, ambiguous recommendations, exact-name selection, Pipeline and Data Integration synchronization, and stale-item handling.",
        "Required runtime prompts, year and period values, Inbox files, configured files, local uploads, and expiry behavior.",
        "Approve, reject, cancel, double-click replay, API restart during interruption, failure evidence, and recovery from the failed step.",
        "Oracle Cloud and on-premises compatibility, including metadata-discovery fallback without false success.",
    ])
    add_heading(doc, "Common problems")
    add_table(doc, ["Symptom", "Meaning", "Action"], [
        ["Provider could not complete the request", "Invalid key, unsupported model, quota, network, or provider-side error.", "Check server logs and provider account; retry a new conversation after correcting configuration."],
        ["Groq request too large or TPM exceeded", "The organization and service tier rejected the requested token volume.", "Use a new conversation, reduce input/history budgets, or choose an appropriate model or tier."],
        ["One choice needed after a clear request", "The exact current artifact was not safely resolved or the catalog is stale.", "Review recommendations; for Pipeline or Data Integration, synchronize with Oracle."],
        ["Unsupported action inputs", "A stale or mismatched card attempted fields outside the current schema.", "Cancel the card, refresh the application, and prepare the operation again."],
        ["Approval no longer valid", "User, environment, permission, artifact, catalog, or pending graph state changed.", "Start a fresh reviewed proposal; do not force-resume stale state."],
        ["Oracle operation failed", "The governed request reached Oracle but Oracle rejected it or a monitored job failed.", "Open execution evidence, correct the Oracle-side cause or input, and explicitly approve a new run or recovery."],
    ], [2350, 3400, 3610])
    add_heading(doc, "What the Assistant cannot do")
    add_bullets(doc, [
        "Create, redesign, edit, or delete Oracle Planning applications, dimensions, forms, Business Rules, Data Maps, Pipelines, integrations, jobs, users, or security assignments.",
        "Approve its own proposal, treat yes in chat as authority, or run an operation that the signed-in user cannot run through the platform.",
        "Use arbitrary REST endpoints, execute unregistered code, or silently retry a write operation.",
        "Guarantee that a language-model response is correct, that a provider remains free, or that a model, quota, or API behavior will remain unchanged.",
        "Bypass Oracle release differences, missing public endpoints, artifact security, Planning data security, or invalid source data.",
        "Schedule an operation directly or generate a governed report workbook entirely from free text; scheduling and report generation use their dedicated workspaces.",
        "Replace accountable business review, segregation of duties, Oracle administration, or production change management.",
    ])
    add_heading(doc, "Glossary")
    add_table(doc, ["Term", "Plain-language meaning"], [
        ["AI Agent", "The complete governed workflow that combines a language model, deterministic routing, platform tools, human decisions, and monitored execution."],
        ["Language model", "The external service that interprets natural language and drafts explanations or tool requests; Gemini and Groq are current provider options."],
        ["LangGraph", "The stateful orchestration library that coordinates model, tool, interrupt, approval, and resume steps."],
        ["Capability gateway", "The server-side allowlist of narrowly defined read-only and preparation tools available to the agent."],
        ["Deterministic", "Controlled by ordinary application rules and code rather than by the model's probabilistic language generation."],
        ["Artifact", "A named Oracle or platform object used by an operation, such as a Business Rule, Data Map, Pipeline, integration, import job, or refresh job."],
        ["Clarification", "A structured choice that resolves ambiguity before inputs or approval."],
        ["Interrupt / resume", "A durable pause while the user selects a choice, supplies inputs, or makes an approval decision, followed by continuation of the same graph state."],
        ["Action draft", "A normalized proposed operation and its reviewed inputs before execution authority is granted."],
        ["Approval decision", "The explicit, immutable user event that approves or rejects a specific reviewed request."],
        ["Idempotency", "Protection that makes repeated submission of the same approved request return the same result instead of running twice."],
        ["Checkpoint", "Durable LangGraph state that allows an interrupted conversation workflow to survive a process restart."],
        ["Standalone flow", "A platform-managed sequential plan of 2-12 supported operations used when no suitable Oracle Pipeline exists."],
        ["Execution evidence", "Retained status, step details, failure information, Oracle counts when available, timestamps, and produced artifacts for a run."],
    ], [2600, 6760])
    _finish_document(doc)
    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
