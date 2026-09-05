"""Build Document 03: Installation, Configuration, and Startup."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_DARK,
    BLUE_LIGHT,
    CONTENT_DXA,
    GOLD,
    GOLD_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    RED_LIGHT,
    SURFACE,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    add_bullets,
    add_callout,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    font,
    set_cell_fill,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    set_table_geometry,
    wrap,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-03"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / (
    "BISP_EPM_Automation_Document_03_Installation_Configuration_and_Startup.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

PAGE_BG = "#F8FAFD"
INK_HEX = f"#{INK}"
NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"
MUTED_HEX = f"#{MUTED}"
LINE_HEX = f"#{LINE}"


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    color: str,
    *,
    width: int = 5,
) -> None:
    """Draw one directional connector with an arrowhead."""
    draw.line((start, end), fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [
            (x2, y2),
            (x2 - 16 * direction, y2 - 10),
            (x2 - 16 * direction, y2 + 10),
        ]
    else:
        direction = 1 if y2 > y1 else -1
        points = [
            (x2, y2),
            (x2 - 10, y2 - 16 * direction),
            (x2 + 10, y2 - 16 * direction),
        ]
    draw.polygon(points, fill=color)


def diagram_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((55, 34), title, font=font(35, bold=True), fill=NAVY_HEX)
    draw.text((55, 82), subtitle, font=font(20), fill=MUTED_HEX)


def box(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    fill: str,
    outline: str,
    number: str | None = None,
) -> None:
    """Draw one rounded process or component card."""
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=20, fill=fill, outline=outline, width=3)
    left = x1 + 22
    if number:
        draw.ellipse((x1 + 20, y1 + 18, x1 + 78, y1 + 76), fill=outline)
        bbox = draw.textbbox((0, 0), number, font=font(23, bold=True))
        draw.text(
            (x1 + 49 - (bbox[2] - bbox[0]) / 2, y1 + 31),
            number,
            font=font(23, bold=True),
            fill="#FFFFFF",
        )
        left = x1 + 95
    y = y1 + 17
    for line in wrap(draw, title, font(24, bold=True), x2 - left - 18):
        draw.text((left, y), line, font=font(24, bold=True), fill=NAVY_HEX)
        y += 31
    y += 7
    if number:
        y = max(y, y1 + 90)
    for line in wrap(draw, body, font(18), x2 - x1 - 42):
        draw.text((x1 + 22, y), line, font=font(18), fill=MUTED_HEX)
        y += 24


def create_installation_path(path: Path) -> None:
    image = Image.new("RGB", (1600, 880), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(
        draw,
        "Installation path",
        "A predictable sequence turns a clean machine into a verified Oracle EPM Automation environment",
    )
    steps = [
        ("1", "Prepare", "Python, Node.js, pnpm, PostgreSQL, Git and Oracle network access", BLUE_HEX),
        ("2", "Install", "Create the Python virtual environment; install backend and frontend dependencies", BLUE_HEX),
        ("3", "Configure", "Copy .env.example, add secrets safely and select deployment options", f"#{GOLD}"),
        ("4", "Migrate", "Apply Alembic application migrations and LangGraph checkpoint setup", NAVY_HEX),
        ("5", "Start", "Run FastAPI and React locally, or API plus durable workers in production", TEAL_HEX),
        ("6", "Verify", "Health probes, first administrator, Oracle application and a read-only smoke test", ORANGE_HEX),
    ]
    x_positions = [65, 565, 1065, 65, 565, 1065]
    y_positions = [165, 165, 165, 485, 485, 485]
    for index, ((number, title, body, color), x, y) in enumerate(
        zip(steps, x_positions, y_positions, strict=True)
    ):
        box(
            draw,
            (x, y, x + 420, y + 215),
            title,
            body,
            fill="#FFFFFF",
            outline=color,
            number=number,
        )
        if index == 0:
            arrow(draw, (485, 272), (555, 272), BLUE_HEX)
        elif index == 1:
            arrow(draw, (985, 272), (1055, 272), BLUE_HEX)
        elif index == 2:
            arrow(draw, (1275, 380), (1275, 470), NAVY_HEX)
        elif index == 3:
            arrow(draw, (485, 592), (555, 592), NAVY_HEX)
        elif index == 4:
            arrow(draw, (985, 592), (1055, 592), TEAL_HEX)
    image.save(path)


def create_configuration_boundary(path: Path) -> None:
    image = Image.new("RGB", (1500, 920), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(
        draw,
        "Configuration and secret boundary",
        "The browser receives safe session data; infrastructure and Oracle secrets stay on the server",
    )
    box(
        draw,
        (70, 180, 420, 420),
        "React browser",
        "UI state, signed session cookie, CSRF token and typed API responses. No Oracle or database password.",
        fill="#EDF3FF",
        outline=BLUE_HEX,
    )
    box(
        draw,
        (570, 145, 980, 455),
        "FastAPI + worker",
        "Loads validated Settings from process environment or .env. Owns authentication, permissions, queues, Oracle calls and redaction.",
        fill="#FFFFFF",
        outline=NAVY_HEX,
    )
    box(
        draw,
        (1130, 180, 1430, 420),
        "Oracle EPM",
        "Planning REST and optional EPM Automate. Oracle remains the system of record for Planning data and artifacts.",
        fill="#FFF2EC",
        outline=ORANGE_HEX,
    )
    box(
        draw,
        (210, 595, 600, 825),
        "PostgreSQL",
        "Platform identities, catalog state, queue, history, audit, schedules and agent state. No Oracle passwords.",
        fill="#E8F7F2",
        outline=TEAL_HEX,
    )
    box(
        draw,
        (820, 595, 1250, 825),
        "External providers",
        "Gemini or Groq API key, optional SMTP credentials and optional Oracle OIDC client secret.",
        fill="#FFF4D8",
        outline=f"#{GOLD}",
    )
    arrow(draw, (420, 300), (560, 300), BLUE_HEX)
    arrow(draw, (980, 300), (1120, 300), ORANGE_HEX)
    arrow(draw, (700, 455), (500, 585), TEAL_HEX)
    arrow(draw, (850, 455), (1000, 585), f"#{GOLD}")
    draw.rounded_rectangle((470, 500, 1030, 560), radius=18, fill="#FDEDEC", outline=f"#{RED}", width=2)
    draw.text((500, 517), "Secrets never belong in Git, React source, screenshots or logs.", font=font(20, bold=True), fill=f"#{RED}")
    image.save(path)


def create_startup_modes(path: Path) -> None:
    image = Image.new("RGB", (1550, 930), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(
        draw,
        "Choose the correct startup mode",
        "Development minimizes processes; production separates request handling from Oracle execution",
    )
    draw.rounded_rectangle((45, 145, 750, 855), radius=25, fill="#FFFFFF", outline=BLUE_HEX, width=3)
    draw.text((80, 175), "Local development", font=font(31, bold=True), fill=NAVY_HEX)
    draw.text((80, 220), "EPM_EXECUTION_RUNTIME=embedded", font=font(19, bold=True), fill=BLUE_HEX)
    box(draw, (105, 285, 690, 420), "Terminal 1: FastAPI", "python web_main.py — starts API and an embedded queue worker", fill="#E8F7F2", outline=TEAL_HEX)
    box(draw, (105, 475, 690, 610), "Terminal 2: React", "cd frontend; pnpm dev — starts Vite at 127.0.0.1:5173", fill="#EDF3FF", outline=BLUE_HEX)
    box(draw, (105, 665, 385, 795), "PostgreSQL", "Required durable state", fill="#FFFFFF", outline=NAVY_HEX)
    box(draw, (410, 665, 690, 795), "Oracle EPM", "Cloud or supported on-premises", fill="#FFF2EC", outline=ORANGE_HEX)
    arrow(draw, (397, 420), (397, 465), TEAL_HEX)
    arrow(draw, (285, 610), (245, 655), NAVY_HEX)
    arrow(draw, (520, 610), (550, 655), ORANGE_HEX)

    draw.rounded_rectangle((800, 145, 1505, 855), radius=25, fill="#FFFFFF", outline=TEAL_HEX, width=3)
    draw.text((835, 175), "Production", font=font(31, bold=True), fill=NAVY_HEX)
    draw.text((835, 220), "EPM_EXECUTION_RUNTIME=web", font=font(19, bold=True), fill=TEAL_HEX)
    box(draw, (860, 285, 1445, 420), "HTTPS frontend + API", "Built React assets and FastAPI; validates and commits requests", fill="#EDF3FF", outline=BLUE_HEX)
    box(draw, (860, 475, 1445, 610), "One or more workers", "python worker.py — claims queue items, calls Oracle and polls schedules", fill="#E8F7F2", outline=TEAL_HEX)
    box(draw, (860, 665, 1140, 795), "PostgreSQL", "Shared database and locks", fill="#FFFFFF", outline=NAVY_HEX)
    box(draw, (1165, 665, 1445, 795), "Shared storage", "Uploads, logs and reports", fill="#FFF4D8", outline=f"#{GOLD}")
    arrow(draw, (1152, 420), (1152, 465), TEAL_HEX)
    arrow(draw, (1035, 610), (1000, 655), NAVY_HEX)
    arrow(draw, (1270, 610), (1300, 655), f"#{GOLD}")
    image.save(path)


def create_troubleshooting_tree(path: Path) -> None:
    image = Image.new("RGB", (1550, 900), PAGE_BG)
    draw = ImageDraw.Draw(image)
    diagram_title(
        draw,
        "Startup troubleshooting path",
        "Follow the first failing boundary instead of changing several settings at once",
    )
    box(draw, (520, 140, 1030, 270), "Application did not become ready", "Start with the first error in the FastAPI or worker log.", fill="#FDEDEC", outline=f"#{RED}")
    candidates = [
        ((70, 390, 410, 585), "Configuration", "Missing or invalid environment value", BLUE_HEX),
        ((460, 390, 800, 585), "PostgreSQL", "Connection, credentials or migration head", NAVY_HEX),
        ((850, 390, 1190, 585), "Frontend", "Vite proxy, build or public origin", TEAL_HEX),
        ((1240, 390, 1480, 585), "Oracle", "URL, identity, DNS, TLS or application access", ORANGE_HEX),
    ]
    centers = []
    for bounds, title, body, color in candidates:
        box(draw, bounds, title, body, fill="#FFFFFF", outline=color)
        centers.append(((bounds[0] + bounds[2]) // 2, bounds[1]))
    for center in centers:
        arrow(draw, (775, 270), center, MUTED_HEX, width=4)
    draw.rounded_rectangle((145, 690, 1405, 825), radius=22, fill="#EDF3FF", outline=BLUE_HEX, width=2)
    draw.text((185, 718), "Verify in order", font=font(24, bold=True), fill=NAVY_HEX)
    draw.text((185, 760), "1. /health/live   2. /health/ready   3. browser sign-in   4. authenticated Oracle health   5. one read-only catalog", font=font(20), fill=MUTED_HEX)
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "installation": ASSET_DIR / "installation-path.png",
        "configuration": ASSET_DIR / "configuration-boundary.png",
        "startup": ASSET_DIR / "startup-modes.png",
        "troubleshooting": ASSET_DIR / "troubleshooting-path.png",
    }
    create_installation_path(assets["installation"])
    create_configuration_boundary(assets["configuration"])
    create_startup_modes(assets["startup"])
    create_troubleshooting_tree(assets["troubleshooting"])
    return assets


def configure_styles(doc: Document) -> None:
    """Apply the compact_reference_guide preset and named brand overrides."""
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = __import__("docx").shared.RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE_DARK, 18, 10),
        ("Heading 2", 13, BLUE_DARK, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = __import__("docx").shared.RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(10)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = __import__("docx").shared.RGBColor.from_string(BLUE_DARK)
    subtitle.paragraph_format.space_after = Pt(12)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = __import__("docx").shared.RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = __import__("docx").shared.RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    for name in ("Kicker", "Lead", "Small", "Table Text", "Table Header"):
        if name not in styles:
            styles.add_style(name, 1)

    kicker = styles["Kicker"]
    kicker.font.name = "Calibri"
    kicker._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    kicker._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    kicker.font.size = Pt(9)
    kicker.font.bold = True
    kicker.font.color.rgb = __import__("docx").shared.RGBColor.from_string(BLUE)
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(3)
    kicker.paragraph_format.keep_with_next = True

    lead = styles["Lead"]
    lead.font.name = "Calibri"
    lead._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    lead._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    lead.font.size = Pt(12.5)
    lead.font.color.rgb = __import__("docx").shared.RGBColor.from_string(MUTED)
    lead.paragraph_format.space_after = Pt(12)
    lead.paragraph_format.line_spacing = 1.2

    small = styles["Small"]
    small.font.name = "Calibri"
    small._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    small._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    small.font.size = Pt(9)
    small.font.color.rgb = __import__("docx").shared.RGBColor.from_string(MUTED)
    small.paragraph_format.space_after = Pt(4)

    table_text = styles["Table Text"]
    table_text.font.name = "Calibri"
    table_text._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_text._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_text.font.size = Pt(9.5)
    table_text.font.color.rgb = __import__("docx").shared.RGBColor.from_string(INK)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.1

    table_header = styles["Table Header"]
    table_header.font.name = "Calibri"
    table_header._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_header._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_header.font.size = Pt(9.5)
    table_header.font.bold = True
    table_header.font.color.rgb = __import__("docx").shared.RGBColor.from_string(WHITE)
    table_header.paragraph_format.space_after = Pt(0)


def patch_list_numbering(doc: Document) -> None:
    """Encode compact-reference list indentation and spacing in numbering XML."""
    numbering = doc.part.numbering_part.element
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        num_pr = style._element.pPr.numPr
        if num_pr is None or num_pr.numId is None:
            continue
        num_id = str(num_pr.numId.val)
        abstract_id = None
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == num_id:
                node = num.find(qn("w:abstractNumId"))
                abstract_id = node.get(qn("w:val")) if node is not None else None
                break
        if abstract_id is None:
            continue
        for abstract in numbering.findall(qn("w:abstractNum")):
            if abstract.get(qn("w:abstractNumId")) != abstract_id:
                continue
            level = abstract.find(qn("w:lvl"))
            if level is None:
                break
            p_pr = level.find(qn("w:pPr"))
            if p_pr is None:
                p_pr = OxmlElement("w:pPr")
                level.append(p_pr)
            tabs = p_pr.find(qn("w:tabs"))
            if tabs is None:
                tabs = OxmlElement("w:tabs")
                p_pr.append(tabs)
            for child in list(tabs):
                tabs.remove(child)
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), "540")
            tabs.append(tab)
            ind = p_pr.find(qn("w:ind"))
            if ind is None:
                ind = OxmlElement("w:ind")
                p_pr.append(ind)
            ind.set(qn("w:left"), "540")
            ind.set(qn("w:hanging"), "270")
            spacing = p_pr.find(qn("w:spacing"))
            if spacing is None:
                spacing = OxmlElement("w:spacing")
                p_pr.append(spacing)
            spacing.set(qn("w:after"), "80")
            spacing.set(qn("w:line"), "300")
            spacing.set(qn("w:lineRule"), "auto")
            break


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def compact_trailing_empty_paragraph(doc: Document) -> bool:
    """Keep Word's required paragraph after a table from creating a blank page."""
    if not doc.paragraphs or not doc.paragraphs[-1].text.strip():
        paragraph = doc.paragraphs[-1] if doc.paragraphs else doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(1)
        marker = paragraph.add_run("\u200b")
        marker.font.size = Pt(1)
        return True
    return False


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    trailing_empty = compact_trailing_empty_paragraph(doc)
    if trailing_empty:
        doc.paragraphs[-1].paragraph_format.page_break_before = True
    kicker = doc.add_paragraph(style="Kicker")
    if not trailing_empty:
        kicker.paragraph_format.page_break_before = True
    kicker_run = kicker.add_run(f"CHAPTER {number}")
    set_run_font(kicker_run, size=9, color=BLUE, bold=True)
    kicker.paragraph_format.space_before = Pt(4)
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True


def add_code_block(doc: Document, code: str, caption: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_fill(cell, NAVY)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(code.splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=8.3, color=WHITE)
    if caption:
        p = doc.add_paragraph(caption, style="Caption")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)


def mark_table_header_row(table) -> None:
    first_row = table.rows[0]
    tr_pr = first_row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def configure_header_footer(section) -> None:
    for header in (section.header, section.even_page_header):
        paragraph = header.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run("BISP SOLUTIONS  /  INSTALLATION & STARTUP")
        set_run_font(run, size=8.5, color=MUTED, bold=True)

    for footer in (section.footer, section.even_page_footer):
        table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
        set_table_geometry(table, [7000, 2360], indent_dxa=0)
        left = table.cell(0, 0).paragraphs[0]
        left.text = "Document 03  |  Installation, Configuration, and Startup"
        set_run_font(left.runs[0], size=8.5, color=MUTED)
        add_page_number(table.cell(0, 1).paragraphs[0])
        for cell in table.rows[0].cells:
            set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            top = OxmlElement("w:top")
            top.set(qn("w:val"), "single")
            top.set(qn("w:sz"), "6")
            top.set(qn("w:color"), LINE)
            borders.append(top)
            tc_pr.append(borders)


def add_cover(doc: Document) -> None:
    if LOGO.is_file():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions company logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(38)
    doc.add_paragraph("ENTERPRISE INSTALLATION GUIDE", style="Kicker")
    doc.add_paragraph("Installation, Configuration,\nand Startup", style="Title")
    doc.add_paragraph("BISP Solutions Oracle EPM Automation Platform", style="Subtitle")
    doc.add_paragraph(
        "A beginner-friendly, production-oriented guide to preparing PostgreSQL, connecting Oracle EPM, configuring identity and optional providers, starting the modern React/FastAPI platform, and verifying a safe deployment.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [["03 of the platform handbook", "29 August 2026", "Administrators, consultants, developers, security and operations teams"]],
        [2200, 2100, 5060],
    )
    add_callout(
        doc,
        "Scope statement",
        "This guide documents the current React frontend, FastAPI backend, PostgreSQL database, Alembic migrations, durable worker, Oracle EPM connection, optional EPM Automate, identity options, notifications, and LangGraph AI providers. The retired UI is not part of this installation.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("BISP Solutions  |  Current production-oriented baseline")
    set_run_font(run, size=10, color=MUTED, bold=True)


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    patch_list_numbering(doc)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS GUIDE", style="Kicker")
    doc.add_paragraph("Choose the path that matches your responsibility", style="Heading 1")
    doc.add_paragraph(
        "A local developer can follow the chapters in order. A production administrator should begin with prerequisites, database preparation, configuration and production deployment. Feature owners can go directly to the optional-service chapters after the core platform is healthy.",
        style="Lead",
    )
    add_table(
        doc,
        ["If you are...", "Read first", "Expected result"],
        [
            ["A first-time developer", "Chapters 1–8", "A working local React and FastAPI environment with embedded execution."],
            ["A database administrator", "Chapters 2, 4 and 12", "A PostgreSQL database at the current Alembic head with a backup and upgrade plan."],
            ["A production platform administrator", "Chapters 2–6 and 9–12", "A secure HTTPS deployment with separately supervised API and workers."],
            ["An Oracle EPM consultant", "Chapters 2, 5, 7 and 10", "A verified Oracle application connection and the correct optional operation engines."],
            ["An AI/notification feature owner", "Chapter 7", "A provider configured only after the core platform is healthy."],
        ],
        [2200, 2300, 4860],
    )
    add_heading(doc, "Installation outcome")
    add_bullets(doc, [
        "PostgreSQL is reachable, versioned and ready for application and LangGraph state.",
        "FastAPI starts without a configuration or migration error.",
        "The React application opens and creates a signed platform session.",
        "The first Service Administrator account is created once, then normal sign-in takes over.",
        "The selected Oracle Planning application is verified with live, permission-aware access.",
        "Queued work is executed by the embedded worker in development or by a separate worker in production.",
    ])
    add_callout(
        doc,
        "Important",
        "Do not configure every optional feature on the first attempt. Bring up PostgreSQL, FastAPI, React and the Oracle connection first; then add AI, email or EPM Automate one provider at a time.",
        fill=GOLD_LIGHT,
        accent=GOLD,
    )

    add_chapter(doc, 1, "What is being installed", "The platform is one product composed of independently deployable web, execution, storage and integration responsibilities.")
    add_figure(
        doc,
        assets["installation"],
        "Figure 1. Recommended installation sequence.",
        "Six-step installation path from prerequisites through verification.",
        width=6.35,
    )
    add_heading(doc, "Core components")
    add_table(
        doc,
        ["Component", "Purpose", "Required?"],
        [
            ["React frontend", "Business-user interface, guided forms, approvals, monitoring and the EPM Assistant workspace.", "Yes"],
            ["FastAPI backend", "Authentication, permissions, API v1, validation, Oracle catalogs, queue submission and protected downloads.", "Yes"],
            ["Durable worker", "Claims queue records, calls Oracle, monitors jobs, polls schedules and stores evidence.", "Embedded in development; separate in production"],
            ["PostgreSQL", "System of record for platform identity, configuration, schedules, execution history, audit and agent state.", "Yes"],
            ["Oracle EPM Planning", "System of record for Planning applications, data, metadata, rules, jobs, integrations, Pipelines and Data Maps.", "Yes"],
            ["EPM Automate", "Alternative engine for operations explicitly configured to use it.", "Optional"],
            ["Gemini or Groq", "Language-model provider behind the governed LangGraph assistant.", "Optional"],
            ["SMTP", "Success and failure email delivery.", "Optional"],
        ],
        [1800, 5460, 2100],
    )
    add_heading(doc, "What PostgreSQL does not replace")
    add_para(doc, "PostgreSQL does not store Planning cube data, metadata files, Oracle passwords, AI API keys, report workbooks or uploaded file contents. Oracle Planning remains authoritative for the EPM application; the database records how the automation platform safely coordinates work around it.")
    add_heading(doc, "Current runtime boundary")
    add_bullets(doc, [
        "Normal runtime requires PostgreSQL. SQLite remains only an isolated test compatibility path.",
        "The modern React application is the only supported presentation layer.",
        "FastAPI is the only backend contract; browser, Excel, schedules and the Assistant reuse it and the same application services.",
        "Production table creation and alteration belong only to Alembic migrations.",
        "Oracle writes are accepted into a durable queue before a worker performs them.",
    ])

    add_chapter(doc, 2, "Prerequisites and access", "Verify software, network and Oracle permissions before troubleshooting application code.")
    add_heading(doc, "Supported baseline")
    add_table(
        doc,
        ["Dependency", "Supported baseline", "Why it is needed"],
        [
            ["Operating system", "Windows development is documented; production may use a supported Windows or Linux service host.", "Runs the Python, Node and PostgreSQL clients."],
            ["Python", "3.11 through 3.13", "FastAPI, Oracle clients, workers, migrations and agent orchestration."],
            ["Node.js", "22.13 or a supported newer LTS; package manifest also accepts 24.x", "Builds and serves the React development experience."],
            ["pnpm", "11.x", "Reproducible frontend dependency and build management."],
            ["PostgreSQL", "16 or a compatible supported release", "Transactions, locks, queue, audit and durable application state."],
            ["Git", "Current supported client", "Retrieves and versions the source code."],
            ["Oracle EPM", "HTTPS Planning environment reachable from backend and workers", "Live application and automation target."],
        ],
        [1800, 3240, 4320],
    )
    add_heading(doc, "Oracle prerequisites")
    add_bullets(doc, [
        "A backend integration identity with the Planning permissions required by the operations that workers will execute.",
        "At least one accessible Planning application, or permission to discover and select one.",
        "Saved jobs and deployed Oracle artifacts already configured for the operations you intend to run.",
        "Outbound HTTPS access to the Oracle EPM hostname from both the API host and every worker host.",
        "For Cloud Basic Authentication, the username format required by the identity domain, which may be username or identitydomain.username.",
    ])
    add_heading(doc, "Network and filesystem checklist")
    add_table(
        doc,
        ["Flow", "Typical port/path", "Check"],
        [
            ["Browser → frontend", "HTTPS 443 in production; 5173 in local development", "User can open the public origin."],
            ["Frontend → FastAPI", "HTTPS 443 through proxy; 8080 locally", "Routes for /api and authentication reach FastAPI."],
            ["FastAPI/worker → PostgreSQL", "TCP 5432 by default", "Both services can authenticate to the same database."],
            ["FastAPI/worker → Oracle EPM", "HTTPS 443; configured HTTP/HTTPS for supported on-premises", "DNS, TLS and firewall allow the configured host."],
            ["FastAPI/worker → AI or SMTP", "Provider HTTPS or SMTP 587/465", "Required only when the feature is enabled."],
            ["API/worker filesystem", "RUNTIME_DATA_DIR and report directory", "Service identities can read/write controlled shared storage."],
        ],
        [2300, 2600, 4460],
    )
    add_callout(doc, "Least privilege", "Do not use a personal Oracle password as a permanent worker secret. Use an approved integration identity, grant only the needed Oracle roles, and keep the platform's Service Administrator role separate from Oracle data access.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 3, "Install the source and dependencies", "Use isolated and locked dependency environments so one machine change does not unexpectedly alter the platform.")
    add_heading(doc, "Create the Python environment")
    add_code_block(
        doc,
        "Set-Location \"D:\\path\\to\\Planning Automation\"\npython -m venv .venv\n.\\.venv\\Scripts\\Activate.ps1\npython -m pip install --upgrade pip\npython -m pip install -r requirements.txt",
        "Code excerpt 1. Backend runtime installation from the project root.",
    )
    add_para(doc, "Use requirements-dev.txt instead when the machine will run the backend test suite. It includes the production dependencies plus pytest and the HTTP test client.")
    add_code_block(doc, "python -m pip install -r requirements-dev.txt", "Code excerpt 2. Development and test dependencies.")
    add_heading(doc, "Install the React dependencies")
    add_code_block(
        doc,
        "Set-Location frontend\npnpm install --frozen-lockfile\npnpm verify\nSet-Location ..",
        "Code excerpt 3. Locked frontend install followed by tests and a production build.",
    )
    add_heading(doc, "What should remain local")
    add_bullets(doc, [
        ".env, encrypted .epw files, virtual environments and node_modules are intentionally excluded from Git.",
        "Runtime folders such as var, logs and reports are deployment data, not source code.",
        "frontend/dist is a generated build artifact and should be recreated from the tagged source release.",
        "Do not copy a developer's .env into production; create environment-specific secrets and paths.",
    ])
    add_callout(doc, "Windows PowerShell", "If script execution policy prevents virtual-environment activation, use the organization's approved PowerShell policy or invoke .venv\\Scripts\\python.exe directly. Do not disable security controls globally just to activate the environment.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 4, "Prepare PostgreSQL and migrations", "PostgreSQL must exist and match the code's Alembic head before either FastAPI or the worker is allowed to start.")
    add_heading(doc, "Create an empty database")
    add_para(doc, "A PostgreSQL administrator can create a dedicated login and database with pgAdmin or an approved SQL client. Replace the placeholder with a strong secret and store it in the deployment secret manager.")
    add_code_block(
        doc,
        "CREATE ROLE epm_automation LOGIN PASSWORD '<strong-password>';\nCREATE DATABASE epm_automation\n    OWNER epm_automation\n    ENCODING 'UTF8';",
        "Code excerpt 4. Example PostgreSQL role and database creation.",
    )
    add_heading(doc, "Configure the connection URL")
    add_code_block(
        doc,
        "DATABASE_URL=postgresql+psycopg://epm_automation:<url-encoded-password>@localhost:5432/epm_automation",
        "Code excerpt 5. SQLAlchemy/psycopg PostgreSQL connection URL.",
    )
    add_bullets(doc, [
        "URL-encode reserved password characters such as @, :, /, # and %.",
        "Keep the same DATABASE_URL for the API and every worker in one deployment.",
        "Run Alembic from the project root so alembic/env.py loads the project-root .env.",
        "An operating-system DATABASE_URL overrides the value in .env; remove or update a stale process variable before retrying.",
    ])
    add_heading(doc, "Apply the two schema families")
    add_numbered(doc, [
        "Run Alembic to create or upgrade platform-owned tables.",
        "Check that the installed revision matches every current Alembic head.",
        "Run the LangGraph checkpoint setup for framework-owned checkpoint tables.",
    ])
    add_code_block(
        doc,
        "python -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup",
        "Code excerpt 6. Current database initialization sequence. The application schema head is 0020_execution_identity in this release.",
    )
    add_heading(doc, "Why startup refuses an old schema")
    add_para(doc, "Both web_main.py and worker.py inspect PostgreSQL before composing application services. If the installed revision does not equal the code's head, startup stops with a targeted message. This avoids processing Oracle work with code and tables from different releases.")
    add_callout(doc, "No SQLite conversion", "This release starts with a clean PostgreSQL database. SQLite is not a production fallback and no earlier SQLite records are copied. Future upgrades are applied as ordered Alembic migrations to the same PostgreSQL database.", fill=BLUE_LIGHT, accent=BLUE)
    add_heading(doc, "Migration rules")
    add_bullets(doc, [
        "Back up PostgreSQL before a production schema change.",
        "Stop API and workers during a coordinated migration window so no process uses mixed code and schema versions.",
        "Review generated migration scripts before release; Alembic autogeneration is a starting point, not an automatic production approval.",
        "Prefer forward-compatible fixes. A production downgrade requires a tested data-preservation and restoration plan.",
        "Record the revision before and after deployment as release evidence.",
    ])

    add_chapter(doc, 5, "Configure the environment", "Settings are validated once at startup so configuration errors appear before an Oracle job is queued.")
    add_figure(
        doc,
        assets["configuration"],
        "Figure 2. Server-side configuration and secret boundary.",
        "React receives safe session data while FastAPI and workers retain database, Oracle, AI and notification secrets.",
        width=6.35,
    )
    add_heading(doc, "Create the local .env")
    add_code_block(
        doc,
        "Copy-Item .env.example .env\n\n# Then edit .env with environment-specific values.\n# Never commit .env.",
        "Code excerpt 7. Start from the complete safe template.",
    )
    add_heading(doc, "Minimum working configuration")
    add_code_block(
        doc,
        "EPM_BASE_URL=https://your-epm-host.example.com\nEPM_INTEGRATION_USERNAME=service.account\nEPM_INTEGRATION_PASSWORD=<secret>\nAPPLICATION_NAME=\n\nDATABASE_URL=postgresql+psycopg://epm_automation:<encoded-secret>@localhost:5432/epm_automation\nRUNTIME_DATA_DIR=var\n\nWEB_SESSION_SECRET=<long-random-secret>\nWEB_SECURE_COOKIES=false\nWEB_FRONTEND_URL=http://127.0.0.1:5173\nEPM_EXECUTION_RUNTIME=embedded",
        "Code excerpt 8. Minimum local configuration; leave APPLICATION_NAME blank to use supported discovery when available.",
    )
    add_heading(doc, "Generate a session secret")
    add_code_block(
        doc,
        "python -c \"import secrets; print(secrets.token_urlsafe(48))\"",
        "Code excerpt 9. Generate a random value and store it only in the environment secret store or local .env.",
    )
    add_heading(doc, "Configuration precedence")
    add_numbered(doc, [
        "The process environment is checked first.",
        "The project-root .env supplies only values that are not already present.",
        "The Settings class normalizes URLs, parses numbers and booleans, validates allowed values and resolves relative paths from the project root.",
        "Invalid configuration raises ConfigurationError and prevents startup.",
    ])
    add_heading(doc, "Oracle URL formats")
    add_table(
        doc,
        ["Deployment", "EPM_BASE_URL example", "Rule"],
        [
            ["Oracle EPM Cloud", "https://epm-test-example.epm.us-ashburn-1.ocs.oraclecloud.com", "Use the environment root; do not append /epmcloud, /HyperionPlanning or /rest/v3."],
            ["Supported on-premises", "http://planning-host:9000/HyperionPlanning", "Include the deployed Planning context required by that server."],
            ["Reverse proxy/custom Cloud host", "https://planning.company.example", "Set EPM_DEPLOYMENT_MODE=cloud when auto-detection cannot identify the Cloud host."],
        ],
        [1900, 3860, 3600],
    )
    add_heading(doc, "Core configuration groups")
    add_table(
        doc,
        ["Group", "Important settings", "Guidance"],
        [
            ["Oracle connection", "EPM_BASE_URL, EPM_INTEGRATION_USERNAME, EPM_INTEGRATION_PASSWORD, APPLICATION_NAME", "The legacy EPM_USERNAME/EPM_PASSWORD aliases remain migration-only compatibility inputs."],
            ["Oracle behavior", "EPM_DEPLOYMENT_MODE, EPM_REQUEST_TIMEOUT, EPM_VERIFY_SSL, DEFAULT_POLL_INTERVAL, DEFAULT_JOB_TIMEOUT", "Keep SSL verification enabled except in a controlled diagnostic environment with an approved certificate plan."],
            ["Database/storage", "DATABASE_URL, RUNTIME_DATA_DIR, REPORT_OUTPUT_DIR", "API and workers must use the same database and controlled storage."],
            ["Execution", "EPM_EXECUTION_RUNTIME, EXECUTION_WORKER_POLL_INTERVAL, EXECUTION_LEASE_SECONDS, SCHEDULE_POLL_INTERVAL", "Lease must be at least 30 seconds; production API uses web mode."],
            ["Web", "WEB_SESSION_SECRET, WEB_SECURE_COOKIES, WEB_FRONTEND_URL", "Use HTTPS, secure cookies and a stable secret in production."],
            ["Catalogs/defaults", "PIPELINE_CATALOG_FILE, DATA_INTEGRATION_CATALOG_FILE, REPORT_CATALOG_FILE and DEFAULT_* values", "Defaults improve convenience; Oracle live discovery and preflight remain authoritative."],
        ],
        [1800, 3680, 3880],
    )
    add_callout(doc, "Secret rule", "Do not place database passwords, Oracle passwords, AI keys, SMTP credentials or OIDC client secrets in React variables, source files, screenshots, logs or committed configuration. Production should inject them from the hosting platform's secret manager.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 6, "Identity, sessions and first administrator", "The platform authenticates its own users separately from the backend identity that executes Oracle jobs.")
    add_heading(doc, "Three identities to understand")
    add_table(
        doc,
        ["Identity", "Purpose", "Stored secret?"],
        [
            ["Platform user", "Signs into the BISP application and receives one or more platform roles.", "Local password is scrypt-hashed; federated profiles may be passwordless."],
            ["Oracle human identity", "Optionally proves the user's identity through Oracle Cloud OIDC or Oracle credential validation.", "OIDC tokens are session-scoped; entered Oracle password is not stored."],
            ["Backend integration identity", "Used by FastAPI/workers for REST or EPM Automate execution.", "Secret is supplied by server environment; never returned to browser."],
        ],
        [2300, 4160, 2900],
    )
    add_heading(doc, "First-run bootstrap")
    add_numbered(doc, [
        "Start PostgreSQL, apply migrations and start FastAPI plus the React frontend.",
        "Open the React URL. The safe /api/v1/bootstrap response reports that initial setup is required.",
        "Create the first Service Administrator with a username, display name, optional email and a passphrase of at least 12 characters.",
        "The database transaction creates the account and assigns the Service Administrator role exactly once.",
        "Use Access Control to create or synchronize subsequent users and assign the approved four-role model.",
    ])
    add_table(
        doc,
        ["Role", "Primary responsibility"],
        [
            ["Service Administrator", "Platform access, users, cycles, schedules, catalogs and all governed operations."],
            ["Power User", "Approved operations, processes, data review, reports, history, user variables and Assistant."],
            ["User", "Assigned Planning work, approved processes, data review, reports, own user variables and Assistant."],
            ["Viewer", "Reports and read-only Assistant guidance without operational execution controls."],
        ],
        [2200, 7160],
    )
    add_heading(doc, "Optional Oracle Cloud OIDC")
    add_code_block(
        doc,
        "IDENTITY_PROVIDER=oracle_cloud\nORACLE_IDENTITY_ISSUER_URL=https://<identity-domain>.identity.oraclecloud.com\nORACLE_IDENTITY_CLIENT_ID=<client-id>\nORACLE_IDENTITY_CLIENT_SECRET=<confidential-client-secret>\nORACLE_IDENTITY_REDIRECT_URI=https://automation.example.com/auth/oracle/callback",
        "Code excerpt 10. Optional federated sign-in; localhost HTTP is accepted only for local development.",
    )
    add_bullets(doc, [
        "The redirect URI must point to this automation platform, not to the Oracle EPM application host.",
        "Register the exact callback URI in the Oracle identity application.",
        "Federated sign-in does not automatically grant a platform role; an approved linked platform identity is still required.",
        "Local recovery sign-in remains available for controlled administration.",
    ])
    add_callout(doc, "Session continuity", "Keep WEB_SESSION_SECRET stable across restarts and instances. Changing it invalidates active browser sessions. Set WEB_SECURE_COOKIES=true when users access the product over HTTPS.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 7, "Configure optional providers", "Enable one optional integration only after the core database, web and Oracle connection checks pass.")
    add_heading(doc, "EPM Assistant: Gemini")
    add_code_block(
        doc,
        "AGENT_PROVIDER=gemini\nAGENT_ORCHESTRATOR=langgraph\nAGENT_MODEL=gemini-3.5-flash-lite\nGEMINI_API_KEY=<server-side-key>\nAGENT_MAX_TOOL_ROUNDS=4\nAGENT_HISTORY_MESSAGES=20\nLANGGRAPH_STRICT_MSGPACK=true",
        "Code excerpt 11. Gemini provider configuration.",
    )
    add_heading(doc, "EPM Assistant: Groq")
    add_code_block(
        doc,
        "AGENT_PROVIDER=groq\nAGENT_ORCHESTRATOR=langgraph\nAGENT_MODEL=openai/gpt-oss-120b\nGROQ_API_KEY=<server-side-key>\nGROQ_MAX_INPUT_TOKENS=2500\nGROQ_MAX_COMPLETION_TOKENS=384",
        "Code excerpt 12. Groq token budgets keep requests below common account limits.",
    )
    add_para(doc, "After switching provider or model, restart FastAPI and begin a new Assistant conversation. The LangGraph workflow, allowed tools, approvals, permissions and worker path do not change.")
    add_heading(doc, "Email notifications")
    add_code_block(
        doc,
        "EMAIL_NOTIFICATIONS_ENABLED=true\nEMAIL_PROVIDER=smtp\nEMAIL_SMTP_HOST=smtp.gmail.com\nEMAIL_SMTP_PORT=587\nEMAIL_SMTP_USERNAME=<test-account>\nEMAIL_SMTP_PASSWORD=<app-password>\nEMAIL_FROM=<sender>\nEMAIL_TO=<recipient>\nEMAIL_USE_TLS=true\nEMAIL_USE_SSL=false",
        "Code excerpt 13. Development SMTP example. TLS and SSL cannot both be true.",
    )
    add_heading(doc, "EPM Automate")
    add_code_block(
        doc,
        "EPM_AUTOMATE_EXECUTABLE=C:\\Program Files\\Oracle\\EPM Automate\\bin\\epmautomate.bat\nEPM_AUTOMATE_PASSWORD_FILE=C:\\Secure\\EPMPassword.epw\nEPM_AUTOMATE_COMMAND_TIMEOUT=1800",
        "Code excerpt 14. EPM Automate requires an encrypted .epw file and is needed only for operations configured with the epmautomate engine.",
    )
    add_table(
        doc,
        ["Optional feature", "Enable when", "Leave disabled when"],
        [
            ["AI provider", "The Assistant will be tested and provider terms/limits are approved.", "Core automation is being commissioned."],
            ["SMTP", "A verified development or corporate mail account is available.", "Delivery is not yet approved; operations still record success/failure."],
            ["EPM Automate", "A selected operation explicitly uses the epmautomate engine.", "Every enabled operation uses REST."],
            ["OIDC", "The public HTTPS callback and Oracle IAM application are ready.", "Local platform authentication is sufficient for the current phase."],
        ],
        [1900, 4160, 3300],
    )

    add_chapter(doc, 8, "Start a local development environment", "Development uses two terminals: FastAPI includes an embedded worker, while Vite serves the React experience.")
    add_figure(
        doc,
        assets["startup"],
        "Figure 3. Development and production startup modes.",
        "Local development runs FastAPI with an embedded worker and a Vite frontend; production separates the API and durable worker.",
        width=6.35,
    )
    add_heading(doc, "Start in the correct order")
    add_numbered(doc, [
        "Start PostgreSQL and confirm the configured database is reachable.",
        "Activate the Python virtual environment from the project root.",
        "Apply Alembic migrations and LangGraph checkpoint setup after a new install or upgrade.",
        "Start FastAPI with EPM_EXECUTION_RUNTIME=embedded.",
        "Start Vite from the frontend directory.",
        "Open the React origin and complete first-run setup or sign in.",
    ])
    add_page_break(doc)
    add_code_block(
        doc,
        "# Terminal 1 — project root\n.\\.venv\\Scripts\\Activate.ps1\npython web_main.py\n\n# Terminal 2 — project root\nSet-Location frontend\npnpm dev",
        "Code excerpt 15. Local startup. worker.py is not required in embedded mode.",
    )
    add_heading(doc, "Local addresses")
    add_table(
        doc,
        ["Address", "Purpose", "Expected result"],
        [
            ["http://127.0.0.1:5173", "React business-user interface", "Login or one-time administrator setup."],
            ["http://127.0.0.1:8080/health/live", "FastAPI process liveness", "HTTP 200 and status alive."],
            ["http://127.0.0.1:8080/health/ready", "FastAPI and PostgreSQL readiness", "HTTP 200 and database available."],
            ["/api/v1/health after sign-in", "Live Oracle connection check", "Authenticated environment and application result."],
        ],
        [3300, 2700, 3360],
    )
    add_heading(doc, "When to restart")
    add_bullets(doc, [
        "Restart FastAPI after changing .env, Oracle application selection, identity configuration or AI provider.",
        "Restart the Vite server after changing its proxy/host configuration.",
        "Begin a new Assistant conversation after changing AI provider or model.",
        "Do not change the active Oracle environment for running API/worker processes; restart both after a controlled selection change.",
    ])

    add_chapter(doc, 9, "Deploy for production", "Production separates the web tier from execution workers so requests, schedules and Oracle jobs remain observable and recoverable.")
    add_heading(doc, "Production build and schema")
    add_code_block(
        doc,
        "python -m pip install -r requirements.txt\npython -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup\n\nSet-Location frontend\npnpm install --frozen-lockfile\npnpm build\nSet-Location ..",
        "Code excerpt 16. Deploy a tagged release, current schemas and a reproducible React bundle.",
    )
    add_heading(doc, "Start separately supervised services")
    add_code_block(
        doc,
        "# API service environment\nEPM_EXECUTION_RUNTIME=web\npython -m uvicorn web_main:app --host 0.0.0.0 --port 8080\n\n# Worker service using the same DATABASE_URL and RUNTIME_DATA_DIR\npython worker.py",
        "Code excerpt 17. Production needs both the API service and at least one worker.",
    )
    add_heading(doc, "Do both commands run every time?")
    add_table(
        doc,
        ["Environment", "FastAPI", "worker.py", "Frontend"],
        [
            ["Local development", "Yes; web_main.py with embedded runtime", "No", "pnpm dev in a second terminal"],
            ["Production", "Yes; web runtime", "Yes; one or more supervised workers", "Serve frontend/dist over HTTPS or let FastAPI serve the built entry point"],
            ["API-only maintenance", "May start for diagnostics", "Keep stopped while migrations or controlled recovery are underway", "Optional"],
        ],
        [1900, 2800, 2200, 2460],
    )
    add_heading(doc, "Reverse proxy and routing")
    add_bullets(doc, [
        "Serve the public React origin over HTTPS and set WEB_FRONTEND_URL to that origin.",
        "Route /api, /auth, /login, /logout, /app, /setup and protected download/static compatibility paths to FastAPI.",
        "If WEB_FRONTEND_URL is omitted and frontend/dist exists, FastAPI can serve the built React entry point.",
        "Do not expose PostgreSQL or worker control ports to browsers.",
        "Use /health/live for process liveness and /health/ready for API/database readiness; readiness intentionally does not call Oracle.",
    ])
    add_heading(doc, "Shared production resources")
    add_bullets(doc, [
        "Every API and worker instance uses the same DATABASE_URL.",
        "Use controlled shared storage for RUNTIME_DATA_DIR, uploads, logs and report artifacts when services run on different hosts.",
        "Keep WEB_SESSION_SECRET identical across API instances.",
        "Workers use PostgreSQL row locking so multiple workers can claim different jobs without duplicate execution.",
        "Respect Oracle job concurrency, integration dependencies and customer maintenance windows even when worker capacity increases.",
    ])
    add_callout(doc, "Recovery behavior", "If a worker disappears during an Oracle write, the lease becomes RECOVERY_REQUIRED. The platform does not automatically repeat the operation because Oracle may already have accepted it. Review Oracle Job Console and platform evidence before retrying.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 10, "Verify the first start", "A successful process start is not enough; verify each boundary from infrastructure to a safe Oracle read.")
    add_heading(doc, "Commissioning sequence")
    add_numbered(doc, [
        "Confirm /health/live returns HTTP 200.",
        "Confirm /health/ready returns HTTP 200 with the database component available.",
        "Open the React application and complete bootstrap or normal sign-in.",
        "Open Dashboard environment health and confirm the configured Oracle URL, deployment mode and selected application.",
        "If multiple applications are discovered, select one as Service Administrator and restart API and workers once.",
        "Use Operations catalog synchronization to inspect current Oracle artifacts without running a write.",
        "Open Data Review or another read-only catalog appropriate to the Oracle version and permissions.",
        "Only then test one low-risk governed operation with explicit review and approval.",
    ])
    add_heading(doc, "Expected evidence")
    add_table(
        doc,
        ["Check", "Healthy evidence", "Where to look"],
        [
            ["Database", "Alembic current reports the current head; readiness says available.", "Deployment log and /health/ready"],
            ["Frontend/API", "React loads safe bootstrap state; API responses carry a request reference.", "Browser and FastAPI log"],
            ["Identity", "Signed-in role and permitted navigation are correct.", "Profile menu and Access Control"],
            ["Oracle", "Application and catalogs resolve without exposing credentials.", "Dashboard environment health and Operations"],
            ["Queue/worker", "Approved work moves from queued to running to terminal status.", "Jobs & Activity and worker log"],
            ["Optional providers", "Assistant or SMTP succeeds without changing core Oracle status semantics.", "Provider-specific test and execution evidence"],
        ],
        [1900, 4200, 3260],
    )
    add_chapter(doc, 11, "Troubleshooting", "Diagnose one boundary at a time and preserve the first meaningful error message.")
    add_figure(
        doc,
        assets["troubleshooting"],
        "Figure 4. Recommended startup troubleshooting order.",
        "Decision path separating configuration, PostgreSQL, frontend and Oracle failures before ordered health checks.",
        width=6.35,
    )
    add_table(
        doc,
        ["Symptom", "Likely cause", "Corrective action"],
        [
            ["DATABASE_URL is required although .env contains it", "Command was run outside the project root, the file is not named .env, or a process variable is overriding it.", "Run from the repository root; check Get-Item .env and $env:DATABASE_URL; update or remove the stale process value."],
            ["PostgreSQL password authentication failed", "The server is reachable but the URL password does not match the role.", "Reset the role password, update the secret and URL-encode reserved characters."],
            ["PostgreSQL schema is not current", "Code was upgraded without applying every migration.", "Stop workers; run alembic upgrade head and current --check-heads; then restart."],
            ["React says backend unavailable", "FastAPI is stopped, Vite proxy target is wrong, or production proxy routes are incomplete.", "Verify port 8080 and /health/ready; inspect vite.config.ts or the HTTPS reverse proxy."],
            ["The React frontend has not been built", "WEB_FRONTEND_URL is empty and frontend/dist does not exist.", "Run pnpm build or configure the separate public React origin."],
            ["Browser session resets after restart", "WEB_SESSION_SECRET was missing or changed.", "Configure one stable long random secret across API instances."],
            ["Oracle authentication failed", "Username format, integration password, identity domain or application permission is incorrect.", "Validate the account directly in Oracle and use the required username format; avoid changing unrelated API code."],
            ["Oracle host cannot be resolved", "DNS, VPN or local network path is unavailable.", "Resolve the hostname with network tooling, test port 443 and involve network administrators for a permanent fix."],
            ["Oracle returns 404", "Base URL/context or endpoint capability differs between Cloud and on-premises versions.", "Verify EPM_BASE_URL and EPM_DEPLOYMENT_MODE; use platform fallbacks or supported live capabilities."],
            ["Execution remains queued", "No production worker is running or worker settings point to another database.", "Start worker.py and verify identical DATABASE_URL and RUNTIME_DATA_DIR."],
            ["Execution becomes RECOVERY_REQUIRED", "Worker lease expired during uncertain Oracle work.", "Inspect Oracle Job Console and execution evidence before an authorized retry."],
            ["Gemini/Groq request fails", "Missing key, unsupported model/tool protocol, provider quota or token limit.", "Check provider, model, account limits and server network; use configured Groq token budgets and restart with a new conversation."],
            ["Email configuration blocks startup", "Notifications are enabled with incomplete SMTP fields or both TLS and SSL.", "Supply host, sender, recipients and matching credentials; choose TLS or SSL, not both."],
            ["EPM Automate cannot start", "Executable or encrypted .epw path is wrong.", "Use the full executable path, an existing .epw file and only enable the engine where required."],
        ],
        [2100, 3320, 3940],
    )
    add_callout(doc, "Avoid repeated writes", "Do not use a failing production write as a connectivity test. Confirm health, authentication and a read-only catalog first. A transport timeout does not prove that Oracle rejected an earlier request.", fill=RED_LIGHT, accent=RED)
    add_heading(doc, "Logging and correlation")
    add_para(doc, "When an API or browser request fails, capture the X-Request-ID response value or the Reference displayed by the UI. Search the API log for request=<value> to correlate status, duration and the safe exception summary. Never add secrets to a support screenshot.")

    add_chapter(doc, 12, "Upgrade, rollback and operations", "Treat code, PostgreSQL schema, LangGraph checkpoints, React build, API and workers as one versioned release unit.")
    add_heading(doc, "Pre-deployment checklist")
    add_numbered(doc, [
        "Confirm the intended immutable Git tag and successful backend, frontend and migration CI checks.",
        "Create a PostgreSQL backup with the organization's approved tool.",
        "Preserve the current secret-manager version and runtime configuration.",
        "Stop workers from claiming new jobs and allow running Oracle jobs to finish or enter controlled recovery.",
        "Deploy code and production dependencies.",
        "Apply Alembic, check the current head and run LangGraph checkpoint setup.",
        "Build the React bundle and start the API in web mode.",
        "Start one or more workers and repeat the commissioning checks.",
    ])
    add_heading(doc, "Release verification commands")
    add_code_block(
        doc,
        "python -m pip install -r requirements-dev.txt\npython -m pytest\npython -m alembic current --check-heads\npython -m app.agent.evaluation --fail-on-threshold\n\nSet-Location frontend\npnpm install --frozen-lockfile\npnpm verify",
        "Code excerpt 18. Current backend, migration, agent and frontend release gates.",
    )
    add_heading(doc, "Rollback policy")
    add_bullets(doc, [
        "Prefer a forward fix when it safely preserves data and audit history.",
        "Before rolling back code, confirm the previous release understands the current database schema.",
        "Restore the pre-deployment PostgreSQL backup when a database rollback is required; do not improvise an Alembic downgrade in production.",
        "Deploy the previous immutable code tag together with its matching React bundle.",
        "Reconcile every Oracle job accepted before rollback against Oracle Job Console and platform evidence.",
    ])
    add_heading(doc, "Operational ownership")
    add_table(
        doc,
        ["Owner", "Ongoing responsibility"],
        [
            ["Platform administrator", "Secrets, HTTPS, sessions, service supervision, health probes, logs and access reviews."],
            ["Database administrator", "PostgreSQL backup, availability, capacity, recovery and approved migration window."],
            ["Oracle EPM administrator", "Integration account, application access, saved jobs, Oracle artifacts and environment maintenance."],
            ["Application support", "Release checks, catalog synchronization, queue health, RECOVERY_REQUIRED triage and evidence."],
            ["AI/provider owner", "Approved model, keys, quotas, provider terms, evaluation gate and live-provider UAT."],
        ],
        [2400, 6960],
    )
    add_heading(doc, "Release evidence to retain")
    add_bullets(doc, [
        "Git tag and commit SHA for the deployed source.",
        "Backend, frontend, migration and agent evaluation results.",
        "PostgreSQL backup reference and Alembic revision before and after deployment.",
        "Deployment operator, start/end time, API health and worker health.",
        "Known issues, recovery decisions and any Oracle jobs reconciled during the change window.",
    ])

    add_page_break(doc)
    doc.add_paragraph("QUICK REFERENCE", style="Kicker")
    doc.add_paragraph("Daily command and configuration reference", style="Heading 1")
    doc.add_paragraph("Use this section after the first successful installation.", style="Lead")
    add_heading(doc, "Local startup")
    add_code_block(doc, "# Terminal 1\npython web_main.py\n\n# Terminal 2\ncd frontend\npnpm dev", "Quick reference 1. Development uses the embedded worker.")
    add_heading(doc, "Production startup")
    add_code_block(doc, "# API, EPM_EXECUTION_RUNTIME=web\npython -m uvicorn web_main:app --host 0.0.0.0 --port 8080\n\n# Worker\npython worker.py", "Quick reference 2. Both processes are required for production execution and schedules.")
    add_heading(doc, "Migration check")
    add_code_block(doc, "python -m alembic upgrade head\npython -m alembic current --check-heads\npython -m app.agent.checkpoint_setup", "Quick reference 3. Run during installation and deployment upgrades.")
    add_heading(doc, "Production readiness checklist")
    add_bullets(doc, [
        "Tagged source release and reproducible frontend build",
        "Current Alembic head and LangGraph checkpoint schema",
        "Stable WEB_SESSION_SECRET and WEB_SECURE_COOKIES=true behind HTTPS",
        "API and workers share DATABASE_URL and controlled storage",
        "Oracle integration identity uses approved least privilege",
        "Health probes and service supervision configured",
        "Database backup and rollback evidence recorded",
        "No .env, .epw, logs, var, reports, node_modules or dist copied from a developer workstation",
        "Optional AI, SMTP, OIDC and EPM Automate features enabled only after individual verification",
        "One read-only Oracle smoke test and one approved low-risk queue/worker test completed",
    ])
    closing_kicker = doc.add_paragraph("INSTALLATION COMPLETE", style="Kicker")
    closing_kicker.paragraph_format.page_break_before = True
    doc.add_paragraph("Ready for a controlled first run", style="Heading 1")
    doc.add_paragraph(
        "The environment is ready when infrastructure health, platform identity, the Oracle application and the queue/worker path have all been verified independently.",
        style="Lead",
    )
    add_heading(doc, "Before handing the platform to users")
    add_bullets(doc, [
        "Record the production URLs, support owner and release evidence.",
        "Confirm the four platform roles and the Oracle integration account's least-privilege access.",
        "Demonstrate one read-only inspection and one approved low-risk operation from request through evidence.",
        "Provide the next handbook volume that matches the audience: user experience, Oracle operations, Planning lifecycle, Data Review or EPM Assistant.",
    ])
    add_callout(doc, "Final takeaway", "A reliable installation is not just a running web page. It is a versioned PostgreSQL schema, a secure server configuration, a verified Oracle application, a governed queue, an active worker, a tested React/API boundary and evidence that every optional provider can fail without bypassing platform controls.", fill=BLUE_LIGHT, accent=BLUE)

    doc.settings.odd_and_even_pages_header_footer = True
    properties = doc.core_properties
    properties.title = "BISP Oracle EPM Automation Platform - Installation, Configuration, and Startup"
    properties.subject = "Document 03 - Current installation and production startup guide"
    properties.author = "BISP Solutions"
    properties.keywords = "Oracle EPM, Planning, installation, configuration, FastAPI, React, PostgreSQL, Alembic, LangGraph, worker"
    properties.comments = "Documents the current modern platform installation baseline as of 29 August 2026."
    for table in doc.tables:
        mark_table_header_row(table)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for section in doc.sections:
        for part in (section.header, section.even_page_header, section.footer, section.even_page_footer):
            for table in part.tables:
                mark_table_header_row(table)
    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
