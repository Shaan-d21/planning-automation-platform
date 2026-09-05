"""Build Document 04: Modern Web Application User Guide."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_DARK,
    BLUE_LIGHT,
    CONTENT_DXA,
    GOLD,
    GOLD_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    RED_LIGHT,
    SURFACE,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    add_bullets,
    add_callout,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    font,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    set_table_geometry,
    wrap,
)
from docs.builders.build_document_03 import (  # noqa: E402
    compact_trailing_empty_paragraph,
    configure_styles,
    patch_list_numbering,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-04"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / (
    "BISP_EPM_Automation_Document_04_Modern_Web_Application_User_Guide.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
MUTED_HEX = f"#{MUTED}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"


def _arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    color: str = BLUE_HEX,
    width: int = 5,
) -> None:
    draw.line((start, end), fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 17 * direction, y2 - 10), (x2 - 17 * direction, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 10, y2 - 17 * direction), (x2 + 10, y2 - 17 * direction)]
    draw.polygon(points, fill=color)


def _text_block(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    fill: str,
    outline: str,
    label: str | None = None,
) -> None:
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=20, fill=fill, outline=outline, width=3)
    y = y1 + 20
    if label:
        draw.text((x1 + 22, y), label.upper(), font=font(16, bold=True), fill=outline)
        y += 28
    for line in wrap(draw, title, font(23, bold=True), x2 - x1 - 44):
        draw.text((x1 + 22, y), line, font=font(23, bold=True), fill=NAVY_HEX)
        y += 30
    y += 7
    for line in wrap(draw, body, font(17), x2 - x1 - 44):
        draw.text((x1 + 22, y), line, font=font(17), fill=MUTED_HEX)
        y += 23


def _diagram_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((55, 34), title, font=font(34, bold=True), fill=NAVY_HEX)
    draw.text((55, 82), subtitle, font=font(19), fill=MUTED_HEX)


def create_interface_anatomy(path: Path) -> None:
    image = Image.new("RGB", (1400, 820), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _diagram_title(draw, "The workspace at a glance", "Four stable areas help a user understand where they are and what to do next.")
    draw.rounded_rectangle((55, 135, 1345, 760), radius=24, fill="#FFFFFF", outline=f"#{LINE}", width=3)
    draw.rounded_rectangle((70, 150, 345, 745), radius=18, fill=NAVY_HEX)
    draw.text((98, 180), "BISP EPM AUTOMATION", font=font(22, bold=True), fill="#FFFFFF")
    draw.text((98, 222), "ROLE-AWARE SIDEBAR", font=font(16, bold=True), fill="#9DB7FF")
    for index, item in enumerate(("Home", "My Work", "Planning", "Automation", "Analysis", "Administration")):
        y = 285 + index * 64
        fill = "#2859D8" if index == 0 else "#13233F"
        draw.rounded_rectangle((92, y, 320, y + 46), radius=10, fill=fill)
        draw.text((112, y + 12), item, font=font(18, bold=index == 0), fill="#FFFFFF")
    draw.rounded_rectangle((365, 150, 1330, 250), radius=16, fill="#FFFFFF", outline=f"#{LINE}", width=2)
    draw.text((400, 178), "Oracle EPM  /  Current page", font=font(20, bold=True), fill=NAVY_HEX)
    draw.text((1010, 178), "Alerts   Refresh   User", font=font(18, bold=True), fill=BLUE_HEX)
    draw.rounded_rectangle((390, 285, 1295, 690), radius=18, fill="#F8FAFD", outline=f"#{LINE}", width=2)
    draw.text((425, 320), "PAGE PURPOSE", font=font(16, bold=True), fill=BLUE_HEX)
    draw.text((425, 355), "Clear heading and plain-language description", font=font(28, bold=True), fill=NAVY_HEX)
    draw.text((425, 407), "Cards, filters, progress, and the next safe action appear here.", font=font(19), fill=MUTED_HEX)
    _text_block(draw, (425, 475, 690, 640), "Context", "Current cycle, Oracle application, or selected service.", fill="#FFFFFF", outline=BLUE_HEX)
    _text_block(draw, (715, 475, 980, 640), "Decision", "Only choices permitted for the signed-in role.", fill="#FFFFFF", outline=TEAL_HEX)
    _text_block(draw, (1005, 475, 1260, 640), "Outcome", "Success, warning, failure, and evidence.", fill="#FFFFFF", outline=ORANGE_HEX)
    image.save(path)


def create_destination_map(path: Path) -> None:
    image = Image.new("RGB", (1400, 900), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _diagram_title(draw, "Where should I go?", "Start with the business objective; the platform directs you to the right workspace.")
    cards = [
        ((55, 145, 440, 320), "Complete assigned work", "MY WORK", "Open tasks, follow readiness, submit work or launch its linked operation.", BLUE_HEX),
        ((505, 145, 890, 320), "Run one Oracle service", "OPERATIONS", "Choose a Business Rule, Pipeline, Data Integration, import, Data Map, variable, or refresh.", ORANGE_HEX),
        ((955, 145, 1340, 320), "Review live Planning data", "DATA REVIEW", "Choose a cube, place dimensions, select members, load the read-only grid, and export it.", TEAL_HEX),
        ((55, 365, 440, 540), "Approve submitted work", "APPROVALS", "Review the task evidence and comment before approving or returning it.", TEAL_HEX),
        ((505, 365, 890, 540), "Investigate execution", "JOBS & ACTIVITY", "Inspect queued, running, successful, failed, or recovery-required executions.", BLUE_HEX),
        ((955, 365, 1340, 540), "Generate an output", "REPORTS", "Choose the registered report and business point of view, then download Excel output.", ORANGE_HEX),
        ((55, 585, 440, 760), "Ask for guided help", "EPM ASSISTANT", "Describe the goal in business language. The assistant may inspect, recommend, or prepare a governed action.", BLUE_HEX),
        ((505, 585, 890, 760), "Plan recurring work", "SCHEDULES", "Administrators schedule approved operations; execution remains monitored and auditable.", ORANGE_HEX),
        ((955, 585, 1340, 760), "Administer the platform", "PLANNING CYCLES / ACCESS", "Administrators create cycles, manage users, and map approved Oracle identities.", TEAL_HEX),
    ]
    for bounds, title, label, body, color in cards:
        _text_block(draw, bounds, title, body, fill="#FFFFFF", outline=color, label=label)
    image.save(path)


def create_governed_journey(path: Path) -> None:
    image = Image.new("RGB", (1400, 700), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _diagram_title(draw, "A governed Oracle operation", "Preparation is separate from approval; every accepted action produces monitored evidence.")
    stages = [
        ("1", "Choose", "Select a permitted service and live Oracle artifact."),
        ("2", "Prepare", "Provide only the required periods, files, values, or prompts."),
        ("3", "Review", "Confirm the exact application, artifact, inputs, and impact."),
        ("4", "Approve", "Explicitly authorize the reviewed operation."),
        ("5", "Monitor", "Watch queue, Oracle job status, results, and evidence."),
    ]
    x_positions = [55, 325, 595, 865, 1135]
    colors = [BLUE_HEX, BLUE_HEX, ORANGE_HEX, ORANGE_HEX, TEAL_HEX]
    for index, ((number, title, body), x, color) in enumerate(zip(stages, x_positions, colors)):
        _text_block(draw, (x, 180, x + 215, 500), title, body, fill="#FFFFFF", outline=color, label=f"STEP {number}")
        if index < len(stages) - 1:
            _arrow(draw, (x + 220, 340), (x + 255, 340), color=BLUE_HEX)
    draw.rounded_rectangle((180, 555, 1220, 635), radius=16, fill=f"#{TEAL_LIGHT}", outline=TEAL_HEX, width=2)
    draw.text((215, 578), "Nothing reaches Oracle until the required review and approval are complete.", font=font(21, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_status_lifecycle(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _diagram_title(draw, "Two status systems, one clear story", "Task status describes business work; execution status describes the underlying automation.")
    draw.text((60, 150), "PLANNING TASK", font=font(17, bold=True), fill=BLUE_HEX)
    task_items = [("Ready", "Can be started"), ("In progress", "Owner is working"), ("Submitted", "Awaiting review"), ("Completed", "Business work accepted")]
    exec_items = [("Queued", "Waiting for a worker"), ("Running", "Oracle work in progress"), ("Succeeded", "Terminal success"), ("Failed", "Review details"), ("Recovery required", "Safe attention needed")]
    x = 60
    for index, (title, body) in enumerate(task_items):
        _text_block(draw, (x, 185, x + 280, 345), title, body, fill="#FFFFFF", outline=TEAL_HEX if index == 3 else BLUE_HEX)
        if index < len(task_items) - 1:
            _arrow(draw, (x + 285, 265), (x + 335, 265))
        x += 340
    draw.text((60, 420), "AUTOMATION EXECUTION", font=font(17, bold=True), fill=ORANGE_HEX)
    x = 60
    for index, (title, body) in enumerate(exec_items):
        color = TEAL_HEX if title == "Succeeded" else (f"#{RED}" if title in {"Failed", "Recovery required"} else ORANGE_HEX)
        _text_block(draw, (x, 455, x + 230, 635), title, body, fill="#FFFFFF", outline=color)
        if index < len(exec_items) - 1:
            _arrow(draw, (x + 235, 545), (x + 270, 545), color=ORANGE_HEX)
        x += 275
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "interface": ASSET_DIR / "interface_anatomy.png",
        "destinations": ASSET_DIR / "destination_map.png",
        "journey": ASSET_DIR / "governed_operation_journey.png",
        "statuses": ASSET_DIR / "status_lifecycle.png",
    }
    create_interface_anatomy(assets["interface"])
    create_destination_map(assets["destinations"])
    create_governed_journey(assets["journey"])
    create_status_lifecycle(assets["statuses"])
    return assets


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    trailing_empty = compact_trailing_empty_paragraph(doc)
    if trailing_empty and number != 1:
        doc.paragraphs[-1].paragraph_format.page_break_before = True
    kicker = doc.add_paragraph(style="Kicker")
    if not trailing_empty and number != 1:
        kicker.paragraph_format.page_break_before = True
    set_run_font(kicker.add_run(f"CHAPTER {number}"), size=9, color=BLUE, bold=True)
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True


def configure_header_footer(section) -> None:
    for header in (section.header, section.even_page_header):
        paragraph = header.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run("BISP SOLUTIONS  /  MODERN WEB APPLICATION USER GUIDE"), size=8.5, color=MUTED, bold=True)

    for footer in (section.footer, section.even_page_footer):
        table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
        set_table_geometry(table, [7000, 2360], indent_dxa=0)
        left = table.cell(0, 0).paragraphs[0]
        left.text = "Document 04  |  Modern Web Application User Guide"
        set_run_font(left.runs[0], size=8.5, color=MUTED)
        add_page_number(table.cell(0, 1).paragraphs[0])
        for cell in table.rows[0].cells:
            set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            top = OxmlElement("w:top")
            top.set(qn("w:val"), "single")
            top.set(qn("w:sz"), "6")
            top.set(qn("w:color"), LINE)
            borders.append(top)
            tc_pr.append(borders)


def add_cover(doc: Document) -> None:
    if LOGO.is_file():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions company logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(38)
    doc.add_paragraph("END-USER HANDBOOK", style="Kicker")
    doc.add_paragraph("Modern Web Application\nUser Guide", style="Title")
    doc.add_paragraph("BISP Solutions Oracle EPM Automation Platform", style="Subtitle")
    doc.add_paragraph(
        "A plain-language guide to signing in, navigating the role-aware workspace, completing Planning work, approving governed actions, monitoring Oracle execution, and finding the right feature without technical training.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [["04 of the platform handbook", "31 August 2026", "Planners, reviewers, consultants, administrators, finance leaders and first-time users"]],
        [2200, 2100, 5060],
    )
    add_callout(
        doc,
        "Current product scope",
        "This guide documents only the current React web application and its current role-based pages. The retired user interface is intentionally excluded. Detailed Oracle operation procedures, lifecycle administration, Data Review, reporting, and AI-agent design continue in later handbook volumes.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraph.add_run("BISP Solutions  |  Beginner-friendly, governed, enterprise Planning automation"), size=10, color=MUTED, bold=True)


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    patch_list_numbering(doc)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS GUIDE", style="Kicker")
    doc.add_paragraph("Begin with your objective—not the technology", style="Heading 1")
    doc.add_paragraph(
        "The application is organized around normal Planning responsibilities. You do not need to know a REST endpoint, job identifier, or database table. Start from the page that describes the outcome you need, follow the visible steps, and stop if the Oracle context is not the one you expect.",
        style="Lead",
    )
    add_table(
        doc,
        ["If you need to...", "Begin here", "Then read"],
        [
            ["Complete work assigned to you", "Home or My Work", "Chapters 5 and 6"],
            ["Run one approved Oracle service", "Operations", "Chapters 8 and 9"],
            ["Review or approve submitted work", "Approvals", "Chapter 7"],
            ["Inspect live Planning values", "Data Review", "Chapter 11"],
            ["Investigate an automation result", "Jobs & Activity", "Chapter 9"],
            ["Manage cycles, schedules, or access", "Administration pages", "Chapter 10"],
            ["Ask for guided help", "EPM Assistant", "Chapter 11"],
        ],
        [3300, 2800, 3260],
    )
    add_callout(
        doc,
        "A useful safety habit",
        "Before any action that changes Oracle, read the connected application name, the artifact name, the business period, and the reviewed input summary. If one item is wrong, return to the inputs instead of approving.",
        fill=GOLD_LIGHT,
        accent=GOLD,
    )
    add_heading(doc, "Visual language used in the application")
    add_table(
        doc,
        ["Visual cue", "Meaning", "What you should do"],
        [
            ["Blue primary button", "The recommended next step", "Select it only after the page inputs are correct."],
            ["Green success state", "The step or execution completed", "Review the evidence or continue to the next task."],
            ["Amber warning", "Attention or explicit approval is needed", "Read the impact and make a deliberate decision."],
            ["Red error", "The request could not complete", "Read the plain-language message; do not repeatedly resubmit."],
        ],
        [1800, 3300, 4260],
    )

    add_chapter(doc, 1, "What the web application does", "The platform brings Planning work, governed Oracle automation, business review, and operational evidence into one role-aware workspace.")
    add_heading(doc, "A business workspace over Oracle EPM")
    add_para(doc, "Oracle Planning remains the system that owns application metadata, data, forms, rules, integrations, Pipelines, Data Maps, and saved jobs. The BISP platform provides a simpler operating layer: it discovers current Oracle artifacts, presents only permitted actions, collects inputs, requires the right review, submits work, monitors results, and preserves evidence.")
    add_bullets(doc, [
        "Business users see their work and the pages relevant to their role.",
        "Consultants can run approved services without remembering REST payloads or command syntax.",
        "Reviewers can approve or return Planning responsibilities with a comment.",
        "Administrators can manage cycles, schedules, platform access, and Oracle identity mappings.",
        "Every execution has a durable record that distinguishes who requested it, how it was triggered, and what happened.",
    ])
    add_heading(doc, "What the web application does not replace")
    add_para(doc, "The platform does not redesign a Planning cube, create every Oracle artifact, or bypass Oracle security. Application design remains in Oracle. The platform automates the safe use of approved artifacts and makes the lifecycle easier to operate.")
    add_callout(doc, "Simple mental model", "Oracle defines what the Planning application can do. BISP controls who can request it, how inputs are reviewed, when it runs, and where the result is monitored.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 2, "Sign in and establish context", "Authentication protects both platform access and the Oracle environment used for live discovery and execution.")
    add_heading(doc, "Sign-in choices")
    add_table(
        doc,
        ["Sign-in method", "When it appears", "What it means"],
        [
            ["Platform sign-in", "Always available as the controlled local or recovery path", "Uses a platform account, assigned role, and secure session."],
            ["Oracle credentials", "When enabled by configuration", "Validates the user against the connected Oracle environment and maps access to the platform."],
            ["Oracle Cloud Identity", "When federation is configured", "Redirects to the organization identity provider and returns to the platform."],
            ["Create Platform Administrator", "Only before the first account exists", "Bootstraps the first administrator; it is not shown after setup."],
        ],
        [2200, 3100, 4060],
    )
    add_numbered(doc, [
        "Open the web application URL supplied by the administrator.",
        "Choose the sign-in method available for your environment.",
        "Enter only your own approved credentials and complete identity-provider steps if shown.",
        "Confirm the displayed user name, persona, application name, and deployment mode after sign-in.",
    ])
    add_heading(doc, "Session and security behavior")
    add_bullets(doc, [
        "The top bar identifies the signed-in user and business persona.",
        "Server-authorized navigation is returned after authentication; hidden pages are not merely cosmetic.",
        "A CSRF-protected session is used for actions that change state.",
        "Repeated failed platform sign-ins can cause a temporary lockout.",
        "Use Sign out when leaving a shared workstation.",
    ])
    add_callout(doc, "If the wrong application is shown", "Do not run an operation. A Service Administrator should review Environment Setup and select the intended Oracle Planning application.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 3, "Navigate the role-aware workspace", "The same structure is used across pages so a first-time user can quickly find context, actions, and outcomes.")
    add_figure(doc, assets["interface"], "Figure 4.1 — Current React application anatomy.", "Diagram of the BISP application showing a role-aware sidebar, top context bar, main work area, and context-decision-outcome cards.", width=6.35)
    add_heading(doc, "Sidebar")
    add_para(doc, "The sidebar is built from permissions returned by the backend. It groups destinations by purpose: everyday workspace, Planning, automation, analysis, and administration. A user sees only permitted destinations.")
    add_heading(doc, "Top bar")
    add_bullets(doc, [
        "Breadcrumb: confirms the current page.",
        "Notifications: opens meaningful Planning events, not raw technical logs.",
        "Refresh: retrieves the latest workspace state.",
        "User identity: shows the signed-in person and persona.",
        "Sign out: closes the authenticated session.",
    ])
    add_heading(doc, "Responsive and accessible behavior")
    add_para(doc, "On smaller screens, the navigation becomes a menu with a page scrim. Buttons and form controls retain accessible names, loading states are announced, confirmation dialogs identify their purpose, and error screens provide a visible retry action.")
    add_heading(doc, "Sidebar groups")
    add_table(
        doc,
        ["Group", "Purpose"],
        [
            ["Workspace", "Home, personal tasks, notifications, and the EPM Assistant."],
            ["Planning", "Approvals and live Data Review for permitted users."],
            ["Automation", "Standalone Operations and administrator-managed Schedules."],
            ["Analysis", "Reports plus execution history for users who can monitor it."],
            ["Administration", "Planning-cycle design and platform Access Control."],
        ],
        [2400, 6960],
    )
    add_callout(doc, "Refresh with purpose", "Use the application Refresh control after another person or Oracle process has changed shared state. Do not refresh repeatedly while a button already shows a spinner or a job is being polled.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 4, "Understand roles and page access", "Four business roles keep the experience simple while backend permissions enforce what each person may do.")
    add_table(
        doc,
        ["Role", "Typical responsibility", "Core experience"],
        [
            ["Service Administrator", "Platform, cycle, schedule, catalog, access, and governed-operation administration", "All current pages and permissions."],
            ["Power User", "FP&A or consultant execution, data review, monitoring, and reporting", "Runs permitted operations, reviews data, sees jobs, and uses the assistant."],
            ["User", "Planner completing assigned work and approved processes", "My Work, approvals, Data Review, reports, user-variable actions, and assistant guidance."],
            ["Viewer", "Executive or consumer of approved information", "Home, My Work, notifications, reports, and read-only assistant guidance."],
        ],
        [1800, 3500, 4060],
    )
    add_heading(doc, "Current navigation matrix")
    add_table(
        doc,
        ["Page", "Admin", "Power User", "User", "Viewer"],
        [
            ["Home / My Work / Notifications", "Yes", "Yes", "Yes", "Yes"],
            ["Approvals", "Yes", "Yes", "Yes", "No"],
            ["Data Review", "Yes", "Yes", "Yes", "No"],
            ["Operations", "Yes", "Yes", "Limited to permitted services", "No"],
            ["Schedules", "Yes", "No", "No", "No"],
            ["Reports", "Yes", "Yes", "Yes", "Yes"],
            ["Jobs & Activity", "Yes", "Yes", "No", "No"],
            ["EPM Assistant", "Yes", "Yes", "Yes", "Read-only guidance"],
            ["Planning Cycles / Access Control", "Yes", "No", "No", "No"],
        ],
        [3100, 1500, 1700, 1760, 1300],
    )
    add_chapter(doc, 5, "Use Home as the daily starting point", "Home translates the signed-in persona and current Planning state into the most useful next actions.")
    add_heading(doc, "What Home answers")
    add_bullets(doc, [
        "What is my current Planning priority?",
        "Which cycle is active and when is it due?",
        "What recent activity or approved reporting is relevant to me?",
        "Which permitted workspace should I open next?",
        "For managers, what requires review and how are teams progressing?",
    ])
    add_heading(doc, "Home varies by persona")
    add_table(
        doc,
        ["Persona", "Home emphasis"],
        [
            ["Service Administrator", "Cycle health, exceptions, administration, executions, and approvals."],
            ["Power User / FP&A", "Ready tasks, operational activity, data validation, and governed Oracle actions."],
            ["User / Planner", "Assigned tasks, due work, current cycle, and approved next steps."],
            ["Viewer / Executive", "Approved results, current business context, and reports; operational controls are hidden."],
        ],
        [2500, 6860],
    )
    add_heading(doc, "Empty states are useful")
    add_para(doc, "Messages such as “No active cycle” or “No recent activity” are not errors. They explain why a card is empty and what responsible action is available. For example, only an administrator can open a prepared cycle, while a Viewer can still open approved reports.")

    add_chapter(doc, 6, "Complete work in My Work", "My Work is the business task list: it brings assignments, readiness, deadlines, linked operations, and submission into one place.")
    add_heading(doc, "Read a task card")
    add_table(
        doc,
        ["Field", "Meaning"],
        [
            ["Status", "Where the task is in its business lifecycle."],
            ["Readiness", "Whether prerequisites allow work to begin."],
            ["Priority", "Normal or higher attention for sequencing work."],
            ["Cycle and stage", "The Planning process and business stage that own the task."],
            ["Period / scenario / entity", "Business context attached by the cycle definition, when relevant."],
            ["Due date", "The expected completion point."],
            ["Open / Start task / Run operation", "The next permitted action for that task type."],
        ],
        [2800, 6560],
    )
    add_heading(doc, "Normal task journey")
    add_numbered(doc, [
        "Open My Work and begin with a Ready task whose dependencies are complete.",
        "Read the description, business context, and expected evidence.",
        "Start the task. If it is linked to automation, use Run operation and complete the governed runner.",
        "Review the result or validation evidence returned to the task.",
        "Mark the task complete or submit it for approval, according to the configured workflow.",
    ])
    add_heading(doc, "When a task is not ready")
    add_para(doc, "Do not work around the dependency. The platform calculates readiness from the cycle and prior work. Open the task to understand the blocking prerequisite, or contact the cycle owner when the upstream responsibility is overdue.")
    add_para(doc, "Operation-backed tasks: a linked operation opens the same governed service, carries task context when supported, and returns execution evidence to the task.", bold_lead="Operation-backed tasks:")

    add_chapter(doc, 7, "Use Approvals and Notifications", "Approvals control business acceptance; Notifications keep important decisions and exceptions visible without mixing them with technical logs.")
    add_heading(doc, "Approvals workspace")
    add_numbered(doc, [
        "Open the submitted task and read the task description, cycle context, submitter, and available evidence.",
        "Confirm that the result satisfies the business responsibility—not merely that a technical job succeeded.",
        "Add a useful comment.",
        "Approve when the work is acceptable, or Return when the owner must correct it.",
    ])
    add_para(doc, "A returned task goes back to the responsible user with the reviewer comment. Approval and return decisions are recorded for auditability.")
    add_heading(doc, "Notifications workspace")
    add_bullets(doc, [
        "Shows assignments, approval decisions, deadlines, and operational exceptions relevant to the user.",
        "Displays an unread count in the page and top navigation.",
        "Supports an Unread only filter.",
        "Lets the user mark individual items as read.",
        "Keeps detailed execution diagnostics in Jobs & Activity rather than cluttering the business inbox.",
    ])
    add_callout(doc, "Approval is a business decision", "A green Oracle job result can support approval, but the reviewer should also confirm scope, period, totals, and the intended business outcome.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 8, "Run a standalone operation", "Operations is the single place to run one approved Oracle EPM service when a complete Planning cycle or Pipeline is not required.")
    add_figure(doc, assets["destinations"], "Figure 4.2 — Objective-to-workspace map.", "Nine cards map common business objectives to My Work, Operations, Data Review, Approvals, Jobs and Activity, Reports, EPM Assistant, Schedules, and Administration.", width=6.35)
    add_heading(doc, "Current service families")
    add_bullets(doc, [
        "Calculation: Business Rules with optional runtime prompt values.",
        "Data loading: native Planning Data Import and Data Integration.",
        "Metadata: saved Metadata Import jobs with file handling and optional governed follow-up.",
        "Data movement: Data Maps with reviewed clear and override controls.",
        "Orchestration: Oracle Data Integration Pipelines and their live variables or file requirements.",
        "Administration: Cube Refresh, substitution variables, and user variables according to permission.",
    ])
    add_heading(doc, "Connected Oracle catalog")
    add_para(doc, "The catalog area identifies the connected Planning application and the synchronized availability of supported artifacts. Search and category filters reduce the list. Registered artifacts that are no longer available are not treated as runnable live choices; administrators can synchronize the catalog and review unavailable registrations where supported.")
    add_figure(doc, assets["journey"], "Figure 4.3 — Governed operation journey.", "Five-stage process from choosing a service through preparation, review, approval, and monitored evidence.", width=6.35)
    add_callout(doc, "Use the configured Oracle artifact", "Import formats, mappings, target cubes, Pipeline stages, Business Rule logic, and saved job definitions remain owned by Oracle. The runner gathers only the runtime choices needed for the selected operation.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 9, "Monitor Jobs & Activity", "The monitoring workspace explains what the platform requested, what Oracle returned, and whether safe human attention is required.")
    add_heading(doc, "Execution states")
    add_figure(doc, assets["statuses"], "Figure 4.4 — Business task and automation execution states.", "Two parallel lifecycles show Planning task states and automation execution states including recovery required.", width=6.35)
    add_table(
        doc,
        ["Execution state", "Interpretation", "User action"],
        [
            ["Queued", "Accepted and waiting for an available worker", "Wait; avoid submitting a duplicate."],
            ["Running", "A worker is executing or monitoring Oracle", "Open details if needed; the page polls active work."],
            ["Succeeded", "All required steps reached a successful terminal result", "Review evidence and complete the business task."],
            ["Failed", "A step ended with an error", "Read the failed step, Oracle message, and available log evidence."],
            ["Recovery required", "The platform cannot safely infer whether to resume or retry", "Use the governed recovery action only after reviewing completed and pending steps."],
        ],
        [1900, 4200, 3260],
    )
    add_heading(doc, "Investigate without creating more risk")
    add_numbered(doc, [
        "Confirm the execution identifier, service, requester, trigger source, and start time.",
        "Open the step timeline and find the first failed or uncertain step.",
        "Read Oracle status, response detail, load statistics, and downloadable evidence where available.",
        "Correct the underlying problem before retrying.",
        "Use a presented recovery or retry action only when the platform identifies it as safe and permitted.",
    ])
    add_heading(doc, "Evidence worth checking")
    add_bullets(doc, [
        "Oracle job identifier and terminal status.",
        "Submitted file name, business period, and selected artifact.",
        "Load statistics, rejected records, downloaded logs, or generated outputs.",
        "The user and trigger source that started the request.",
    ])
    add_callout(doc, "Do not equate browser timeout with Oracle failure", "The worker continues independently of the browser. Reopen Jobs & Activity and inspect the durable execution before deciding whether another request is needed.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 10, "Use schedules and administration pages", "Administrative pages configure recurring execution, Planning cycles, platform users, and identity mappings without exposing these controls to everyday users.")
    add_heading(doc, "Schedules")
    add_para(doc, "A Service Administrator can schedule supported approved operations. A schedule stores the operation, reviewed input set, recurrence, status, and execution context. Each scheduled occurrence creates its own execution record and is processed by the durable worker.")
    add_bullets(doc, [
        "Use a schedule for repeatable technical work with stable, reviewed inputs.",
        "Pause a schedule before changing the underlying Oracle artifact or environment.",
        "Use Jobs & Activity to inspect each occurrence; the schedule itself is not the execution result.",
        "Avoid scheduling an operation whose required file or period cannot be resolved reliably at runtime.",
    ])
    add_heading(doc, "Planning Cycles")
    add_para(doc, "Planning Cycles model business stages, dates, assignments, readiness, task behavior, approvals, and optional linked operations. They are different from Oracle Pipelines: a Pipeline orchestrates technical Oracle stages, while a Planning cycle coordinates people and business responsibilities around those technical services.")
    add_heading(doc, "Access Control")
    add_bullets(doc, [
        "Create or update platform users and assign exactly one current business role.",
        "Activate or deactivate access without deleting audit history.",
        "Reset local recovery credentials where permitted.",
        "Inspect Oracle identity users, groups, and roles without storing passwords.",
        "Map approved Oracle entitlements to platform roles and provision reviewed identities.",
        "The last active Service Administrator is protected from accidental removal.",
    ])
    add_callout(doc, "Separation of duties", "Administration defines access and reusable workflow. Reviewers accept business work. Workers execute durable automation. Oracle still enforces the connected application's own permissions.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 11, "Use Data Review, Reports, and EPM Assistant", "These workspaces help users understand Planning information and take guided action while preserving role, review, and approval boundaries.")
    add_heading(doc, "Data Review")
    add_para(doc, "Data Review creates a read-only, spreadsheet-like grid directly from a live Planning cube. The user selects a live cube, places each discovered dimension into POV, rows, or columns, chooses members from Oracle metadata, and loads the exact data intersection. Results can be searched, paged, refreshed, inspected cell by cell, and exported to Excel.")
    add_bullets(doc, [
        "POV dimensions use one fixed member.",
        "Row and column dimensions can use multiple selected members.",
        "Every cube dimension must be placed exactly once.",
        "The grid is read-only; it does not save values back to Oracle.",
        "Assigned validation tasks can restore a suggested intersection and retain reusable context.",
    ])
    add_heading(doc, "Reports")
    add_para(doc, "Reports creates a downloadable workbook from a registered reporting definition and its required point of view. The page is available to all four roles, but it returns only output permitted by the connected Planning application and the platform configuration.")
    add_page_break(doc)
    add_heading(doc, "EPM Assistant")
    add_para(doc, "The assistant accepts business-language requests, inspects the platform's allowed capabilities, retrieves current Oracle choices through approved tools, and can recommend or prepare a governed action. It never receives permission merely because a user asks in natural language.")
    add_table(
        doc,
        ["Assistant behavior", "What the user should expect"],
        [
            ["Explain or inspect", "A readable answer grounded in current platform or Oracle context."],
            ["Recommend an artifact", "Likely Business Rule, Data Map, Pipeline, Data Integration, or import match with a choice when ambiguous."],
            ["Ask for missing input", "A structured card with dropdowns, files, dates, or values rather than an unsafe guess."],
            ["Prepare, approve, and run", "An exact review card appears first; only explicit authorization and permission checks can move it to durable execution."],
        ],
        [2900, 6460],
    )
    add_heading(doc, "What the assistant never does")
    add_bullets(doc, [
        "Bypass the signed-in user's platform or Oracle permissions.",
        "Treat an ambiguous artifact suggestion as final without confirmation.",
        "Change Oracle merely because an action was described in conversation.",
        "Replace durable execution monitoring with a chat response.",
    ])
    add_callout(doc, "Assistant boundary", "Treat the assistant as a governed copilot, not an autonomous Oracle administrator. Validate its suggested artifact and inputs before approving any change.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 12, "Handle errors, empty states, and support", "The interface separates recoverable user choices, Oracle failures, access restrictions, and platform availability so the next step is understandable.")
    add_heading(doc, "Common situations")
    add_table(
        doc,
        ["What you see", "Likely meaning", "Recommended response"],
        [
            ["No services or artifacts match", "Search/filter is too narrow, permission does not allow them, or the live catalog has none", "Clear the search; confirm environment and role; ask an administrator to synchronize."],
            ["Oracle artifact unavailable", "A saved registration no longer matches the connected application", "Choose a verified live artifact or ask an administrator to review the registration."],
            ["Authentication or environment error", "Oracle credentials, identity mapping, application selection, or network access failed", "Stop execution and contact the administrator with the exact message and time."],
            ["Request timed out", "The browser or Oracle request exceeded its response window", "Inspect Jobs & Activity before resubmitting."],
            ["Approval required", "The operation is prepared but not authorized", "Review the exact inputs and approve only if correct."],
            ["Recovery required", "The platform needs a deliberate decision before continuing", "Review step history and use only the offered governed recovery path."],
        ],
        [2500, 3300, 3560],
    )
    add_heading(doc, "Useful support information")
    add_bullets(doc, [
        "Exact time and timezone of the issue.",
        "Signed-in user and displayed role—never the password.",
        "Connected application name and deployment mode.",
        "Page, service, artifact, period, and selected file name.",
        "Execution identifier, failed step, Oracle status, and concise error message.",
        "Whether the request was manual, scheduled, API, Excel, or AI-assisted.",
    ])
    add_para(doc, "Never include secrets: do not copy passwords, API keys, encrypted-password files, session cookies, or complete environment files into screenshots, emails, tickets, or assistant messages.", bold_lead="Never include secrets:")

    add_chapter(doc, 13, "First-day walkthrough", "This short path lets a new user learn the application without starting an unnecessary Oracle change.")
    add_numbered(doc, [
        "Sign in and confirm your name, persona, connected application, and deployment mode.",
        "Open Home and read the current cycle, priority, and suggested next action.",
        "Open My Work and inspect one task without changing its status.",
        "Open Notifications and use the Unread only filter.",
        "If permitted, open Data Review, select a live cube, place dimensions, choose members, and load a small read-only grid.",
        "Open Reports and review the available registered outputs; generate only an approved test POV.",
        "If permitted, open Operations and inspect the service catalog without approving an action.",
        "Open Jobs & Activity and inspect a prior completed execution.",
        "Ask EPM Assistant a read-only question such as “What can I do in this platform?” and review its answer.",
        "Sign out.",
    ])
    add_heading(doc, "Quick decision checklist before an Oracle-changing action")
    add_table(
        doc,
        ["Check", "Question"],
        [
            ["Environment", "Is this the intended Oracle application and deployment?"],
            ["Artifact", "Is this the exact live Rule, Pipeline, integration, Data Map, import job, or refresh job?"],
            ["Business context", "Are year, period, scenario, scope, and point of view correct?"],
            ["File", "Is the chosen local, Inbox, or configured file the intended current version?"],
            ["Impact", "Do I understand replace, clear, refresh, create, or update behavior?"],
            ["Approval", "Am I authorized and ready to start this operation now?"],
            ["Monitoring", "Will I remain responsible for checking the terminal result and evidence?"],
        ],
        [2100, 7260],
    )
    add_callout(doc, "You are ready", "A new user who can answer these seven questions can safely use the guided runners. Document 05 continues with detailed procedures for every Oracle EPM automation operation.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 14, "Glossary and quick reference", "Use these terms when discussing the platform with Planning, finance, operations, security, and support teams.")
    add_table(
        doc,
        ["Term", "Plain-language meaning"],
        [
            ["Artifact", "A named Oracle object such as a Business Rule, Data Map, Pipeline, integration, report, or saved job."],
            ["Catalog", "The platform's synchronized view of registered and live Oracle artifacts."],
            ["Governed action", "A prepared operation that requires valid inputs, permission, review, and explicit approval."],
            ["Planning cycle", "A business calendar of stages, responsibilities, dependencies, assignments, and approvals."],
            ["Pipeline", "An Oracle Data Integration orchestration containing technical stages and runtime variables."],
            ["POV", "Point of view: the fixed members that identify the business slice being reviewed or reported."],
            ["Execution", "The durable record of one automation request and all of its steps."],
            ["Worker", "The background service that claims queued work and continues independently of the browser."],
            ["Evidence", "Status, response, statistics, logs, files, timestamps, and other facts retained for review."],
            ["Trigger source", "How work began: manual, scheduled, API, Excel, or AI agent."],
            ["Recovery required", "A protected state where the platform needs a reviewed human decision before proceeding."],
        ],
        [2400, 6960],
    )
    add_page_break(doc)
    add_heading(doc, "Handbook continuation")
    add_bullets(doc, [
        "Document 05 — Oracle EPM Automation Operations: detailed procedures for every standalone service.",
        "Document 06 — Planning Workspace and Lifecycle: cycles, stages, tasks, approvals, scheduling, and end-to-end operation.",
        "Document 07 — Data Review and Reporting: cube discovery, grid design, comparison, export, and report generation.",
        "Document 08 — EPM Assistant and AI Agent: providers, LangGraph workflow, tools, approvals, limitations, and responsible use.",
    ])
    add_callout(doc, "End of Document 04", "Use this guide for the current web application experience. Continue with Document 05 when you are ready for detailed, operation-by-operation Oracle EPM procedures.", fill=BLUE_LIGHT, accent=BLUE)

    compact_trailing_empty_paragraph(doc)
    for table in doc.tables:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            header_pr.append(header)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for section in doc.sections:
        for footer in (section.footer, section.even_page_footer):
            for table in footer.tables:
                header_pr = table.rows[0]._tr.get_or_add_trPr()
                if header_pr.find(qn("w:tblHeader")) is None:
                    header = OxmlElement("w:tblHeader")
                    header.set(qn("w:val"), "true")
                    header_pr.append(header)

    properties = doc.core_properties
    properties.title = "Modern Web Application User Guide"
    properties.subject = "BISP Solutions Oracle EPM Automation Platform end-user handbook"
    properties.author = "BISP Solutions"
    properties.keywords = "Oracle EPM, Planning, user guide, React, governed automation, BISP Solutions"
    properties.comments = "Documents the current React web application only; retired UI excluded."

    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
