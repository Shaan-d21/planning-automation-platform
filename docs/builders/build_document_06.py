"""Build Document 06: Planning Workspace and Lifecycle."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.builders.build_document_01 import (  # noqa: E402
    BLUE,
    BLUE_DARK,
    BLUE_LIGHT,
    GOLD,
    GOLD_LIGHT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_LIGHT,
    RED,
    RED_LIGHT,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    add_bullets,
    add_callout,
    add_figure,
    add_numbered,
    add_page_number,
    add_para,
    add_table,
    configure_page,
    font,
    set_cell_margins,
    set_picture_alt_text,
    set_run_font,
    wrap,
)
from docs.builders.build_document_03 import (  # noqa: E402
    compact_trailing_empty_paragraph,
    configure_styles,
    patch_list_numbering,
)


OUTPUT_DIR = ROOT / "outputs" / "documentation" / "document-06"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / (
    "BISP_EPM_Automation_Document_06_Planning_Workspace_and_Lifecycle.docx"
)
LOGO = ROOT / "app" / "web" / "static" / "images" / "bisp-logo.png"

NAVY_HEX = f"#{NAVY}"
BLUE_HEX = f"#{BLUE}"
MUTED_HEX = f"#{MUTED}"
TEAL_HEX = f"#{TEAL}"
ORANGE_HEX = f"#{ORANGE}"
RED_HEX = f"#{RED}"


def _arrow(draw: ImageDraw.ImageDraw, start, end, color=BLUE_HEX, width=5) -> None:
    draw.line((*start, *end), fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [
            (x2, y2),
            (x2 - 17 * direction, y2 - 10),
            (x2 - 17 * direction, y2 + 10),
        ]
    else:
        direction = 1 if y2 > y1 else -1
        points = [
            (x2, y2),
            (x2 - 10, y2 - 17 * direction),
            (x2 + 10, y2 - 17 * direction),
        ]
    draw.polygon(points, fill=color)


def _title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((55, 34), title, font=font(34, bold=True), fill=NAVY_HEX)
    y = 82
    for line in wrap(draw, subtitle, font(19), 1280):
        draw.text((55, y), line, font=font(19), fill=MUTED_HEX)
        y += 26


def _card(
    draw: ImageDraw.ImageDraw,
    bounds,
    title: str,
    body: str,
    *,
    color: str,
    label: str = "",
    fill: str = "#FFFFFF",
) -> None:
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=18, fill=fill, outline=color, width=3)
    y = y1 + 18
    if label:
        draw.text((x1 + 20, y), label.upper(), font=font(15, bold=True), fill=color)
        y += 27
    for line in wrap(draw, title, font(22, bold=True), x2 - x1 - 40):
        draw.text((x1 + 20, y), line, font=font(22, bold=True), fill=NAVY_HEX)
        y += 29
    y += 5
    for line in wrap(draw, body, font(16), x2 - x1 - 40):
        draw.text((x1 + 20, y), line, font=font(16), fill=MUTED_HEX)
        y += 22


def create_layer_model(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(
        draw,
        "Business lifecycle and technical execution",
        "The Planning Workspace coordinates people and dates. Oracle remains the owner of saved technical definitions.",
    )
    _card(
        draw,
        (55, 165, 640, 430),
        "Planning Cycle",
        "Names the business period, stages, accountable owners, dependencies, due dates, approval responsibilities, and progress.",
        color=BLUE_HEX,
        label="BISP BUSINESS LAYER",
        fill=f"#{BLUE_LIGHT}",
    )
    _card(
        draw,
        (760, 165, 1345, 430),
        "Oracle technical artifacts",
        "Pipelines, rules, integrations, Data Maps, import jobs, refresh jobs, forms, reports, mappings, and security remain configured in Oracle.",
        color=ORANGE_HEX,
        label="ORACLE EXECUTION LAYER",
        fill=f"#{ORANGE_LIGHT}",
    )
    _arrow(draw, (640, 295), (750, 295), color=TEAL_HEX)
    _arrow(draw, (760, 345), (650, 345), color=TEAL_HEX)
    draw.text((663, 260), "launch", font=font(15, bold=True), fill=TEAL_HEX)
    draw.text((670, 366), "evidence", font=font(15, bold=True), fill=TEAL_HEX)
    _card(
        draw,
        (250, 515, 1150, 690),
        "One user experience",
        "My Work sends the user to the correct governed screen. A successful retained execution completes the linked task; a failure blocks it for a controlled retry.",
        color=TEAL_HEX,
        label="CONNECTED BY TASK CONTEXT",
    )
    image.save(path)


def create_cycle_hierarchy(path: Path) -> None:
    image = Image.new("RGB", (1400, 820), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "The Planning Cycle hierarchy", "One dated cycle contains ordered business stages; each stage contains accountable responsibilities.")
    _card(draw, (430, 145, 970, 285), "March Forecast FY27", "Forecast | FY27 | actual through Feb | forecast starts Mar", color=BLUE_HEX, label="CYCLE")
    stage_x = [70, 365, 660, 955]
    stage_names = ["Actuals &\nReconciliation", "Planning Input", "Review &\nApproval", "Publish &\nReporting"]
    colors = [BLUE_HEX, BLUE_HEX, ORANGE_HEX, TEAL_HEX]
    for index, (x, name, color) in enumerate(zip(stage_x, stage_names, colors), start=1):
        _arrow(draw, (700, 300), (x + 185, 365), color="#9AA9BF", width=3)
        _card(draw, (x, 375, x + 280, 505), name, f"Business checkpoint {index}", color=color, label=f"STAGE {index}")
        _arrow(draw, (x + 140, 515), (x + 140, 565), color=color)
        task = ["Reconcile actual data", "Complete input\nSubmit plan", "Review submission", "Publish results"][index - 1]
        _card(draw, (x + 15, 575, x + 265, 740), task, "Assigned to one user or one role.", color=color, label="RESPONSIBILITY")
    draw.text((55, 775), "Dependencies can connect responsibilities within or across stages; later work waits until every selected prerequisite is completed.", font=font(18, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_readiness_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 730), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Readiness is calculated, not guessed", "A task's stored status and its unfinished prerequisites determine what the user can do next.")
    _card(draw, (55, 170, 340, 405), "Waiting", "At least one selected prerequisite is not completed. The Start or Run action is unavailable.", color=MUTED_HEX, label="COMPUTED")
    _arrow(draw, (355, 285), (445, 285), color=MUTED_HEX)
    _card(draw, (460, 170, 745, 405), "Ready", "All prerequisites are complete. The assigned user or role can start, open, submit, or run the task.", color=BLUE_HEX, label="COMPUTED")
    _arrow(draw, (760, 285), (850, 285), color=BLUE_HEX)
    _card(draw, (865, 170, 1150, 405), "In progress", "A manual task has started, or an Oracle execution has been queued or is running.", color=ORANGE_HEX, label="STORED")
    _arrow(draw, (1005, 425), (1005, 500), color=ORANGE_HEX)
    _card(draw, (760, 520, 1060, 665), "Completed", "Manual confirmation, accepted validation, approved review, or successful Oracle execution.", color=TEAL_HEX, label="TERMINAL")
    _card(draw, (1100, 520, 1345, 665), "Blocked", "The user marked a problem or the latest Oracle execution failed. Retry or resume after correction.", color=RED_HEX, label="ATTENTION")
    _arrow(draw, (1150, 285), (1220, 505), color=RED_HEX)
    _card(draw, (55, 500, 650, 665), "Important rule", "Cancelled work is terminal for stage progress, but only Completed satisfies another task's dependency. Design dependencies carefully.", color=ORANGE_HEX, label="DEPENDENCY SAFETY", fill=f"#{GOLD_LIGHT}")
    image.save(path)


def create_task_routing(path: Path) -> None:
    image = Image.new("RGB", (1400, 850), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Where should the user work?", "The task action chooses a safe destination; it does not recreate Oracle configuration inside the cycle.")
    groups = [
        ((55, 155, 430, 365), "Business work", "Checklist task in My Work; start, resume, or mark complete.", BLUE_HEX, "MANUAL"),
        ((510, 155, 885, 365), "Data Review", "Open a read-only grid. Run quality or comparison validation before completion.", TEAL_HEX, "EVIDENCE"),
        ((965, 155, 1340, 365), "Approval", "Submit for review or decide an assigned submission in Approvals.", ORANGE_HEX, "GOVERNANCE"),
        ((55, 455, 430, 685), "Oracle operation", "Pipeline, rule, integration, Data Map, imports, Cube Refresh, or substitution variable update.", RED_HEX, "EXECUTION"),
        ((510, 455, 885, 685), "Report", "Open Reports to generate or review an authorized business output.", TEAL_HEX, "READ-ONLY"),
        ((965, 455, 1340, 685), "Current boundary", "The cycle selects the operation type. The governed runner selects the current Oracle artifact and run inputs.", BLUE_HEX, "NO DUPLICATION"),
    ]
    for bounds, title, body, color, label in groups:
        _card(draw, bounds, title, body, color=color, label=label)
    draw.rounded_rectangle((245, 745, 1155, 815), radius=15, fill=f"#{BLUE_LIGHT}", outline=BLUE_HEX, width=2)
    draw.text((285, 768), "Every executable task carries its task ID so execution evidence can be linked back to the business responsibility.", font=font(18, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_approval_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Approval is a controlled loop", "The reviewer sees business context and the latest validation evidence linked to the submitted task.")
    stages = [
        (55, "Validate", "Complete required checks. Warnings need acknowledgement; failures cannot proceed.", TEAL_HEX),
        (325, "Submit", "The submit task completes and creates a pending request for each dependent review task.", BLUE_HEX),
        (595, "Review", "Reviewer sees cycle, entity, scenario/period, submitter, and validation evidence.", ORANGE_HEX),
        (865, "Decide", "Approve, or return with a required explanation of what must change.", ORANGE_HEX),
        (1135, "Continue", "Approval completes reviewer work. Return reopens submission and waits for resubmission.", TEAL_HEX),
    ]
    for index, (x, title, body, color) in enumerate(stages):
        _card(draw, (x, 165, x + 215, 500), title, body, color=color, label=str(index + 1))
        if index < len(stages) - 1:
            _arrow(draw, (x + 220, 335), (x + 255, 335), color=BLUE_HEX)
    draw.rounded_rectangle((205, 560, 1195, 680), radius=18, fill=f"#{GOLD_LIGHT}", outline=ORANGE_HEX, width=2)
    draw.text((245, 585), "RETURN PATH", font=font(15, bold=True), fill=ORANGE_HEX)
    draw.text((245, 620), "Submitted task -> In progress | Review task -> Not started | Planner receives the review comment", font=font(20, bold=True), fill=NAVY_HEX)
    image.save(path)


def create_monthly_timeline(path: Path) -> None:
    image = Image.new("RGB", (1400, 770), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Example monthly forecast lifecycle", "A practical D-2 to close sequence using current Planning Workspace capabilities.")
    y = 270
    draw.line((100, y, 1300, y), fill="#AAB6C7", width=6)
    events = [
        (120, "D-2", "Prepare", "Confirm actuals, files, Oracle artifacts, owners, and dates.", BLUE_HEX),
        (350, "D-1", "Reconcile", "Run loads or Pipeline, then validate source data.", TEAL_HEX),
        (590, "D0-D2", "Plan", "Planners update assumptions and run assigned calculations.", BLUE_HEX),
        (830, "D2-D3", "Review", "Validate, submit, return corrections, and approve.", ORANGE_HEX),
        (1070, "D3-D4", "Publish", "Push reporting data, generate reports, and inspect evidence.", TEAL_HEX),
        (1280, "END", "Close", "All stage responsibilities complete; cycle reaches 100%.", NAVY_HEX),
    ]
    for index, (x, day, title, body, color) in enumerate(events):
        draw.ellipse((x - 16, y - 16, x + 16, y + 16), fill=color, outline="#FFFFFF", width=4)
        top = 330 if index % 2 == 0 else 470
        _arrow(draw, (x, y + 20), (x, top - 10), color=color, width=3)
        left = max(35, min(x - 100, 1165))
        _card(draw, (left, top, left + 205, top + 205), title, body, color=color, label=day)
    image.save(path)


def create_schedule_flow(path: Path) -> None:
    image = Image.new("RGB", (1400, 780), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "Unattended scheduling remains governed", "The scheduler stores recurrence and safe inputs, then revalidates the current Oracle target at every occurrence.")
    stages = [
        (55, "Configure", "Choose Pipeline or RTP registry sync, input policy, first run, frequency, timezone, and missed-run behavior.", BLUE_HEX),
        (330, "Live review", "Inspect the current Pipeline or confirm Oracle can generate the Calculation Manager snapshot.", BLUE_HEX),
        (605, "Claim once", "A durable occurrence is claimed atomically so multiple workers cannot submit it twice.", ORANGE_HEX),
        (880, "Revalidate", "Check target, required values, Oracle Inbox references, and environment identity again.", ORANGE_HEX),
        (1155, "Execute", "Submit a Pipeline to the worker or complete the RTP registry synchronization and retain history.", TEAL_HEX),
    ]
    for index, (x, title, body, color) in enumerate(stages):
        _card(draw, (x, 165, x + 210, 505), title, body, color=color, label=str(index + 1))
        if index < len(stages) - 1:
            _arrow(draw, (x + 215, 335), (x + 260, 335), color=TEAL_HEX)
    _card(draw, (120, 570, 650, 710), "Misfire: Run once or Skip", "Choose what happens when the service recovers after an occurrence was missed.", color=ORANGE_HEX, label="RECOVERY")
    _card(draw, (750, 570, 1280, 710), "Concurrency: Skip if active", "The same target is not submitted while an active execution already exists.", color=RED_HEX, label="DUPLICATE SAFETY")
    image.save(path)


def create_evidence_chain(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    _title(draw, "One evidence chain from responsibility to result", "Business progress is not separated from technical proof.")
    _card(draw, (55, 165, 315, 420), "Planning task", "Owner, due date, priority, dependencies, action type, and business context.", color=BLUE_HEX, label="WHY")
    _arrow(draw, (330, 290), (415, 290))
    _card(draw, (430, 165, 690, 420), "Execution attempt", "Task ID, execution ID, attempt number, initiator, and queued/running/success/failed state.", color=ORANGE_HEX, label="HOW")
    _arrow(draw, (705, 290), (790, 290))
    _card(draw, (805, 165, 1065, 420), "Oracle evidence", "Step timeline, job status, logs, files, statistics, errors, and timestamps.", color=TEAL_HEX, label="PROOF")
    _arrow(draw, (1080, 290), (1165, 290))
    _card(draw, (1180, 165, 1345, 420), "Progress", "Task, stage, and cycle update automatically.", color=TEAL_HEX, label="OUTCOME")
    _card(draw, (185, 520, 645, 680), "Failed attempt", "Task becomes Blocked. Correct the cause, then use Retry operation; previous attempts remain visible.", color=RED_HEX, label="EXCEPTION")
    _card(draw, (755, 520, 1215, 680), "Successful attempt", "Latest task completes. Stage and cycle progress recalculate from retained task states.", color=TEAL_HEX, label="SUCCESS")
    image.save(path)


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "layers": ASSET_DIR / "business_technical_layers.png",
        "hierarchy": ASSET_DIR / "cycle_hierarchy.png",
        "readiness": ASSET_DIR / "task_readiness.png",
        "routing": ASSET_DIR / "task_action_routing.png",
        "approval": ASSET_DIR / "approval_loop.png",
        "timeline": ASSET_DIR / "monthly_forecast_timeline.png",
        "schedule": ASSET_DIR / "schedule_flow.png",
        "evidence": ASSET_DIR / "evidence_chain.png",
    }
    create_layer_model(assets["layers"])
    create_cycle_hierarchy(assets["hierarchy"])
    create_readiness_flow(assets["readiness"])
    create_task_routing(assets["routing"])
    create_approval_flow(assets["approval"])
    create_monthly_timeline(assets["timeline"])
    create_schedule_flow(assets["schedule"])
    create_evidence_chain(assets["evidence"])
    return assets


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break()


def add_hard_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_chapter(doc: Document, number: int, title: str, lead: str) -> None:
    add_page_break(doc)
    doc.add_paragraph(f"CHAPTER {number:02d}", style="Kicker")
    doc.add_paragraph(title, style="Heading 1")
    doc.add_paragraph(lead, style="Lead")


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "BISP SOLUTIONS  /  PLANNING WORKSPACE AND LIFECYCLE"
    set_run_font(paragraph.runs[0], size=8.5, color=MUTED, bold=True)
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.columns[0].width = Inches(5.7)
    table.columns[1].width = Inches(0.8)
    left = table.cell(0, 0).paragraphs[0]
    left.text = "Document 06  |  Planning Workspace and Lifecycle"
    set_run_font(left.runs[0], size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(right)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:color"), LINE)
        borders.append(top)
        tc_pr.append(borders)


def add_cover(doc: Document) -> None:
    if LOGO.is_file():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.75))
        set_picture_alt_text(run, "BISP Solutions company logo")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(34)
    doc.add_paragraph("PLATFORM HANDBOOK", style="Kicker")
    doc.add_paragraph("Planning Workspace\nand Lifecycle", style="Title")
    doc.add_paragraph("BISP Solutions Oracle EPM Automation Platform", style="Subtitle")
    doc.add_paragraph(
        "A beginner-friendly guide to opening a Planning cycle, organizing stages and responsibilities, completing dependent work, validating data, running governed Oracle operations, approving submissions, scheduling unattended work, and closing with retained evidence.",
        style="Lead",
    )
    add_table(
        doc,
        ["Document", "Implementation snapshot", "Audience"],
        [["06 of the platform handbook", "2 September 2026", "Planners, FP&A managers, Oracle EPM consultants, administrators, reviewers, and support teams"]],
        [2200, 2100, 5060],
    )
    add_callout(
        doc,
        "Current implementation only",
        "This volume was rebuilt from the current React Planning Workspace, FastAPI API contracts, PostgreSQL lifecycle records, task execution links, Data Review evidence, approvals, notifications, and automation schedules. It excludes the retired interface and the older custom Process Designer flow.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        p.add_run("BISP Solutions  |  Business clarity with governed Oracle execution"),
        size=10,
        color=MUTED,
        bold=True,
    )


def _finish_document(doc: Document) -> None:
    compact_trailing_empty_paragraph(doc)
    for table in doc.tables:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            header_pr.append(header)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for section in doc.sections:
        for footer in (section.footer, section.even_page_footer):
            for table in footer.tables:
                header_pr = table.rows[0]._tr.get_or_add_trPr()
                if header_pr.find(qn("w:tblHeader")) is None:
                    header = OxmlElement("w:tblHeader")
                    header.set(qn("w:val"), "true")
                    header_pr.append(header)
    properties = doc.core_properties
    properties.title = "Planning Workspace and Lifecycle"
    properties.subject = "Current business lifecycle and scheduling guide for the BISP Solutions Oracle EPM Automation Platform"
    properties.author = "BISP Solutions"
    properties.keywords = "Oracle EPM, Planning, lifecycle, Planning Workspace, approvals, tasks, scheduling, governance"
    properties.comments = "Documents the current React application and PostgreSQL-backed lifecycle implementation as of 2 September 2026; retired UI excluded."


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = create_assets()
    doc = Document()
    configure_styles(doc)
    patch_list_numbering(doc)
    for section in doc.sections:
        configure_page(section, first_page=True)
        configure_header_footer(section)

    add_cover(doc)
    add_page_break(doc)

    doc.add_paragraph("HOW TO USE THIS VOLUME", style="Kicker")
    doc.add_paragraph("Follow the lifecycle, not the technology", style="Heading 1")
    doc.add_paragraph(
        "This manual starts with the business calendar and follows one responsibility from assignment to evidence. Administrators can use it to design a cycle; planners can use it as a day-to-day guide; reviewers can use it to understand approvals; support teams can use it to diagnose waiting, blocked, or failed work.",
        style="Lead",
    )
    add_table(
        doc,
        ["If you are...", "Start with", "Primary goal"],
        [
            ["A planner or task owner", "Chapters 5-8", "Find ready work, complete it, validate data, and submit it."],
            ["A reviewer or FP&A manager", "Chapters 5 and 9", "Manage by exception and approve or return submissions."],
            ["A Service Administrator", "Chapters 2-4 and 11-14", "Open a correct cycle, assign owners, schedule safe automation, and monitor evidence."],
            ["A support consultant", "Chapters 6-7 and 12-14", "Trace dependencies, execution attempts, failures, and recovery."],
        ],
        [2300, 1800, 5260],
    )
    add_heading(doc, "Contents")
    add_table(
        doc,
        ["Part", "Chapters", "What it explains"],
        [
            ["Foundation", "1-2", "The business/technical boundary, lifecycle objects, and role responsibilities."],
            ["Design", "3-4", "How an administrator opens a cycle and defines safe work."],
            ["Daily work", "5-8", "Home, My Work, readiness, Data Review, operations, and evidence."],
            ["Governance", "9-10", "Approval decisions and current in-application notifications."],
            ["Automation", "11", "Current unattended schedule types and safeguards."],
            ["Complete example", "12", "A practical D-2 to close monthly forecast."],
            ["Support and control", "13-15", "Monitoring, troubleshooting, design checklist, limitations, and glossary."],
        ],
        [1900, 1550, 5910],
    )
    add_callout(
        doc,
        "One sentence to remember",
        "A Planning Cycle tells people what must happen and when; Oracle Pipelines and operations perform the technical work; retained evidence connects the two.",
        fill=TEAL_LIGHT,
        accent=TEAL,
    )

    add_chapter(doc, 1, "Understand the operating model", "The platform deliberately separates business coordination from technical Oracle configuration so users receive a simple work queue without duplicating application design.")
    add_figure(doc, assets["layers"], "Figure 6.1 - Business coordination and Oracle execution.", "Two-layer diagram showing the BISP Planning Cycle launching Oracle-owned artifacts and receiving retained execution evidence.", width=6.35)
    add_heading(doc, "Planning Cycle versus Oracle Pipeline")
    add_table(
        doc,
        ["Question", "Planning Cycle", "Oracle Pipeline"],
        [
            ["Primary purpose", "Coordinate business work, owners, dates, dependencies, reviews, and progress.", "Orchestrate technical stages, jobs, mappings, calculations, and Oracle runtime variables."],
            ["Configured in", "Planning Cycles in the BISP platform.", "Oracle Data Integration / Planning."],
            ["Typical owner", "Planning administrator or FP&A process owner.", "Oracle EPM consultant or technical administrator."],
            ["What runs when created", "Nothing in Oracle. Assigned responsibilities are published immediately.", "The configured stages run only when the Pipeline is explicitly or automatically executed."],
            ["What should not be duplicated", "Do not recreate Pipeline stages as extra cycle steps merely because standalone runners exist.", "Do not use the Pipeline as the business assignment calendar."],
        ],
        [2050, 3655, 3655],
    )
    add_callout(doc, "Practical design decision", "If metadata import, data load, calculations, Data Maps, refreshes, or Oracle notifications already belong to one Pipeline, create one cycle responsibility called Run Oracle Pipeline. Do not create a second responsibility for every internal Pipeline stage.", fill=GOLD_LIGHT, accent=GOLD)
    add_heading(doc, "Current lifecycle objects")
    add_bullets(doc, [
        "Cycle: one dated forecast, budget, strategic plan, or period-close business event.",
        "Stage: an ordered business checkpoint such as Reconciliation, Planning Input, Review, or Publish.",
        "Responsibility (task): one accountable outcome assigned to exactly one named user or one platform role.",
        "Dependency: a prerequisite responsibility that must be completed before later work becomes ready.",
        "Execution attempt: the retained Oracle run linked to an executable responsibility.",
        "Validation evidence: the latest Data Review result linked to a responsibility and, when submitted, to its approval.",
        "Approval: a governed request connecting a completed submission responsibility to a dependent reviewer responsibility.",
        "Notification: a user-owned message for the current approval request or decision.",
        "Schedule occurrence: one durable, uniquely claimed unattended attempt for a configured schedule.",
    ])

    add_chapter(doc, 2, "Know who does what", "The same lifecycle looks different to a planner, reviewer, administrator, and executive. The interface shows only authorized navigation and the user's assigned work.")
    add_table(
        doc,
        ["Role", "Lifecycle responsibility", "Current platform capabilities"],
        [
            ["Service Administrator", "Own environment readiness, lifecycle design, controlled administration, and exceptions.", "All permissions, including Planning Cycle design, schedules, access control, operations, history, reports, and approvals assigned to the role."],
            ["Power User", "Operate approved processes, review data, supervise Planning work, and decide assigned approvals.", "Run processes and standalone operations, review data, generate reports, inspect history, use the assistant, and manage personal user-variable values."],
            ["User / Planner", "Complete assigned input, validation, submission, and other business responsibilities.", "Run approved task-linked processes, review data, generate reports, use the assistant, and manage personal user-variable values."],
            ["Viewer", "Consume authorized outputs and guidance.", "Generate authorized reports and use read-only assistant guidance. Do not assign operational cycle responsibilities to this role as a normal design practice."],
        ],
        [1900, 3250, 4210],
    )
    add_heading(doc, "Assignment visibility")
    add_bullets(doc, [
        "A responsibility is visible when it is assigned directly to the signed-in user or to one of that user's platform roles.",
        "My Work is not an unrestricted list of every responsibility in every cycle.",
        "A Service Administrator can see complete cycle progress in Planning Cycles, but My Work still focuses on the administrator's own assignments.",
        "The platform checks its own permission before opening a governed screen. Oracle then applies its own application security when data or jobs are requested.",
    ])
    add_callout(doc, "Exactly one owner", "Each responsibility must have either one named user or one role - never both and never neither. Role assignment is useful for shared operational queues; named assignment is better when individual accountability is required.", fill=BLUE_LIGHT, accent=BLUE)

    add_chapter(doc, 3, "Open a Planning Cycle", "Only a Service Administrator can design a lifecycle. The current wizard creates and opens the entire cycle, stages, responsibilities, and dependencies in one database transaction.")
    add_figure(doc, assets["hierarchy"], "Figure 6.2 - Cycle, stages, and responsibilities.", "Hierarchy diagram for a March Forecast FY27 cycle containing four ordered business stages and assigned responsibilities.", width=6.35)
    add_heading(doc, "Before opening the wizard")
    add_bullets(doc, [
        "Confirm the Oracle environment and Planning application are correct.",
        "Confirm users and roles already exist in Access Control.",
        "Decide which technical work is already contained in Oracle Pipelines.",
        "Define the business checkpoints, accountable outcomes, owners, dates, and approval path.",
        "Know which Data Review task, if any, should open with a prefilled cube intersection.",
    ])
    add_page_break(doc)
    add_heading(doc, "Step 1 - Set the business calendar")
    add_table(
        doc,
        ["Field", "Required?", "Current meaning"],
        [
            ["Cycle name", "Yes", "Business-friendly name, for example March Forecast FY27."],
            ["Planning year", "Yes", "Cycle context shown throughout the workspace. It does not update Oracle substitution variables by itself."],
            ["Cycle type", "Yes", "Forecast, Budget, Strategic plan, or Period close."],
            ["Start date / due date", "Yes", "Overall business calendar. Due date cannot be before start date."],
            ["Actual through", "No", "Last period treated as actual for user context."],
            ["Forecast starts", "No", "First forecast period shown as cycle and task context."],
            ["Scenario", "No", "Leave blank when Oracle category mappings already determine scenario."],
            ["Cycle code", "Yes", "Stable uppercase internal key generated from the name; letters, numbers, underscores, and hyphens only."],
        ],
        [2450, 1250, 5660],
    )
    add_callout(doc, "Year is context, not a hidden update", "The current cycle stores the Planning year so users understand the work. It does not silently change CurYr or another Oracle variable. If a variable must change, add the explicit Update substitution variables responsibility or configure that change inside the Oracle Pipeline.", fill=GOLD_LIGHT, accent=GOLD)
    add_heading(doc, "Step 2 - Define stages and accountable owners")
    add_numbered(doc, [
        "Keep each stage business-oriented: Actuals & Reconciliation, Planning Input, Review & Approval, and Publish & Reporting are the current starter pattern.",
        "Give every stage a unique code and at least one responsibility.",
        "Describe the expected outcome in plain language, not REST calls or script commands.",
        "Assign exactly one named user or one role, set an optional due date/time, and choose a priority.",
        "Choose where the user should work, then select one prerequisite in Wait for when sequencing is required.",
        "For a Data Review responsibility, optionally define the source cube, POV, rows, columns, validation method, and controls once for the planner.",
    ])
    add_heading(doc, "Step 3 - Review and open")
    add_bullets(doc, [
        "Review the cycle name, year, calendar, number of stages, number of responsibilities, owners, priorities, and dependencies.",
        "Opening is immediate: the cycle is saved in Open status and assignments appear in the relevant My Work queues.",
        "No Oracle job, Pipeline, rule, import, Data Map, or refresh starts from cycle creation.",
        "The current React wizard does not provide post-open edit or delete controls. Correct the design before selecting Open cycle.",
    ])

    add_chapter(doc, 4, "Design responsibilities that cannot confuse users", "A responsibility should state one outcome, send the user to one safe destination, and have only the prerequisites that are genuinely necessary.")
    add_figure(doc, assets["routing"], "Figure 6.3 - Current task action routes.", "Six action groups showing manual work, Data Review, approvals, governed Oracle operations, reports, and the boundary between cycle design and run-time selection.", width=6.35)
    add_heading(doc, "Current Where should the user work choices")
    add_table(
        doc,
        ["Choice", "What the user experiences", "Completion rule"],
        [
            ["Complete a business checklist task", "Start, resume, or mark complete in My Work.", "User confirmation after dependencies are complete."],
            ["Review Planning data", "Open Data Review; an administrator may prefill its grid and validation method.", "Latest validation must not fail; warnings must be acknowledged."],
            ["Generate or review a report", "Open Reports and generate or inspect an authorized output.", "Manual completion from My Work after the business outcome is confirmed."],
            ["Run an Oracle operation", "Open the modern Operations runner for Pipeline, Business Rule, Data Integration, Data Map, Planning Data Import, Metadata Import, Cube Refresh, or substitution variables.", "Automatic only after the linked retained execution succeeds."],
            ["Submit completed work for approval", "Create a governed pending request for every dependent reviewer responsibility.", "Submission action completes the submit task."],
            ["Review and decide a submission", "Open Approvals with validation evidence and business context.", "Approve completes review; return reopens the submission path."],
        ],
        [2450, 4260, 2650],
    )
    add_callout(doc, "Current operation-selection boundary", "The cycle currently chooses an operation type, not a specific Pipeline, rule, Data Map, integration, import job, or refresh job. The user selects the current Oracle artifact and inputs in the governed Operations screen. This avoids copying technical configuration into every monthly cycle.", fill=BLUE_LIGHT, accent=BLUE)
    add_heading(doc, "Dependency rules")
    add_bullets(doc, [
        "A responsibility becomes Ready only when every selected prerequisite is Completed.",
        "A waiting responsibility cannot be started, completed, submitted, or executed.",
        "The system rejects a responsibility that depends on itself, an unknown task key, or a circular chain.",
        "Removing a task or stage in the draft wizard also removes affected draft dependency links.",
        "Do not use dependencies to reproduce every technical Pipeline stage; use them for real business gates.",
    ])

    add_chapter(doc, 5, "Work from Home and My Work", "The day-to-day experience is role-aware. Home highlights priority work; My Work provides the complete assignment queue with cycle and dependency context.")
    add_heading(doc, "Home: start with attention, not navigation")
    add_table(
        doc,
        ["Persona", "What Home emphasizes"],
        [
            ["Service Administrator", "Items needing attention, failed recent runs, due-today work, active cycles, and environment health."],
            ["Power User", "Assigned reviews, due-today work, tasks in scope, cycle progress, and entity progress where available."],
            ["User / Planner", "Tasks remaining, due today, waiting on others, completed work, and the first ready responsibilities."],
            ["Viewer", "Active cycle context, progress, published reporting activity, and approved reporting navigation."],
        ],
        [2300, 7060],
    )
    add_heading(doc, "My Work controls")
    add_bullets(doc, [
        "Search by task, entity, period, or stage.",
        "Filter by Planning cycle and by All, Ready, In progress, Waiting, Blocked, or Completed.",
        "See priority, description, stage, cycle, entity, period, due timing, execution attempt, and latest Data Review evidence on the task card.",
        "Use the right action for the task: Start task, Mark complete, Resume, Open, Run operation, Retry operation, Submit for review, or View run.",
        "Use the side panel to see cycle context, stage progress, and the exact prerequisite responsible for waiting work.",
    ])
    add_callout(doc, "My Work is the primary planner experience", "A planner should not need to remember which technical page to open. Start from the assigned responsibility; the task carries its context into Data Review or Operations.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 6, "Understand task status and readiness", "Status is stored as the user's or system's progress. Readiness is calculated from that status and the current state of every prerequisite.")
    add_figure(doc, assets["readiness"], "Figure 6.4 - Task readiness and outcome states.", "Flow from Waiting to Ready and In progress, then to Completed or Blocked, with a dependency safety note.", width=6.35)
    add_table(
        doc,
        ["Displayed state", "Why it appears", "User response"],
        [
            ["Waiting", "One or more prerequisites are unfinished.", "Open the dependency explanation; do not bypass the sequence."],
            ["Ready", "No unfinished prerequisite remains.", "Start, open, submit, or run the responsibility."],
            ["In progress", "Manual work started or a linked Oracle execution is active.", "Finish the work or monitor the linked execution."],
            ["Blocked", "The user recorded a blocker or the latest Oracle execution failed.", "Correct the business or technical cause, then Resume or Retry operation."],
            ["Completed", "The task met its completion rule.", "No additional action; dependent work can become ready."],
            ["Cancelled", "A cycle designer cancelled the task.", "No further action. Note that cancellation does not satisfy a dependent task."],
        ],
        [1700, 4150, 3510],
    )
    add_heading(doc, "How progress rolls up")
    add_bullets(doc, [
        "A stage is Completed when all its responsibilities are Completed or Cancelled.",
        "A stage is In progress when any responsibility is In progress or Completed and the stage is not complete.",
        "A stage is Blocked only when all its responsibilities are Blocked.",
        "A cycle is Completed when every stage is Completed or Skipped.",
        "The displayed cycle percentage is completed stages divided by total stages. It is not a weighted task percentage.",
    ])

    add_chapter(doc, 7, "Use Data Review as a business gate", "A Data Review responsibility can carry a prefilled, read-only cube intersection so planners validate the intended data instead of rebuilding the grid from memory.")
    add_heading(doc, "Optional administrator prefill")
    add_table(
        doc,
        ["Configuration", "Purpose"],
        [
            ["Source cube", "The cube that owns the reviewed data."],
            ["POV", "One member for every fixed dimension, for example Scenario=Forecast and Year=FY27."],
            ["Rows and columns", "One or more dimensions with member lists separated by |."],
            ["Quality validation", "Check missing data, optionally warn on zero, and retain issue counts."],
            ["Comparison validation", "Compare the same intersection in a target cube using an allowed numeric difference."],
        ],
        [2700, 6660],
    )
    add_callout(doc, "Prefill is optional", "If no Data Review layout is stored in the task, the user can build the live grid in Data Review. If a layout is stored, the platform verifies that the submitted cube intersection and validation method match the assignment.", fill=BLUE_LIGHT, accent=BLUE)
    add_heading(doc, "Completion gate")
    add_numbered(doc, [
        "Open the responsibility from My Work. The task ID authorizes access and loads any saved selection.",
        "Load the live Oracle data and run the configured quality or comparison validation.",
        "If the result is Fail, correct the data or configuration and validate again.",
        "If the result is Warning, inspect the exceptions and explicitly acknowledge the warning when it is acceptable.",
        "Return to My Work and mark the responsibility complete. The latest accepted evidence is retained.",
    ])
    add_heading(doc, "Evidence passed into approval")
    add_para(doc, "When a later Submit for approval responsibility depends on validated work, the most recent validation evidence from its prerequisite chain is attached to the approval request. The reviewer can see the validation type, source and target cubes, status, checked cells, exception count, and timestamp, then open the evidence.")

    add_hard_page_break(doc)
    add_chapter(doc, 8, "Run Oracle operations from a responsibility", "Executable responsibilities do not trust a manual Mark complete. Their state is controlled by the linked, retained Oracle execution.")
    add_figure(doc, assets["evidence"], "Figure 6.5 - Responsibility-to-execution evidence chain.", "Flow linking a Planning task to its execution attempt, Oracle evidence, and automatic business progress, with success and failure outcomes.", width=6.35)
    add_heading(doc, "Current executable responsibility types")
    add_table(
        doc,
        ["Business need", "Task action", "Governed destination"],
        [
            ["Run multi-stage automation", "Run an Oracle Pipeline", "Pipelines runner"],
            ["Calculate", "Run a business rule", "Business Rules runner"],
            ["Load through Data Integration", "Run a Data Integration", "Data Integrations runner"],
            ["Push data", "Push data with a Data Map", "Data Maps runner"],
            ["Load native Planning data", "Run a Planning data import", "Planning Data Import runner"],
            ["Load metadata", "Run a metadata import", "Metadata Import runner"],
            ["Synchronize cube metadata", "Refresh the Planning cube", "Cube Refresh runner"],
            ["Change shared context", "Update substitution variables", "Substitution Variables runner; assigned to Service Administrator by the wizard"],
        ],
        [2750, 2900, 3710],
    )
    add_heading(doc, "What happens after Run operation")
    add_numbered(doc, [
        "The Operations screen receives the Planning task ID and verifies that the user owns the task or its assigned role.",
        "The user selects the live or verified Oracle artifact, resolves its current inputs, reviews the impact, and approves the operation.",
        "The request is queued durably and linked to the responsibility as attempt 1, 2, or a later retry.",
        "The task becomes In progress while the worker connects, submits, and monitors Oracle.",
        "Success completes the task automatically. Failure changes it to Blocked and exposes Retry operation while retaining every prior attempt.",
    ])
    add_callout(doc, "Never mark an executable task complete manually", "The backend rejects manual completion for executable action types. This prevents a cycle from showing green when Oracle did not actually finish the assigned operation.", fill=RED_LIGHT, accent=RED)

    add_chapter(doc, 9, "Submit, approve, return, and resubmit", "Approval is tied to explicit task dependencies. A submit responsibility needs at least one dependent reviewer responsibility configured with Review approval.")
    add_figure(doc, assets["approval"], "Figure 6.6 - Governed approval loop.", "Five-step approval flow from validation through submission and review to approval or a controlled return and resubmission path.", width=6.35)
    add_heading(doc, "Planner submission")
    add_numbered(doc, [
        "Complete every prerequisite shown under the Submit for approval responsibility.",
        "Select Submit for review. Do not use Mark complete; the platform blocks that shortcut.",
        "The submit responsibility completes and one pending approval is created for each dependent Review approval responsibility.",
        "The assigned reviewer receives an in-application approval request notification.",
    ])
    add_heading(doc, "Reviewer decision")
    add_bullets(doc, [
        "Approvals shows only requests assigned directly to the reviewer or to one of the reviewer's roles.",
        "The card shows submitter, cycle, entity, scenario/period context, and linked validation evidence when available.",
        "Approve may include an optional comment and completes the reviewer responsibility.",
        "Return for changes requires a clear comment. The submitted responsibility reopens In progress, the review responsibility returns to Not started, and later dependencies wait again.",
        "After correction, the planner can submit again. The previous decision remains in approval history.",
    ])
    add_callout(doc, "Management by exception", "A reviewer should not approve only because the process reached the queue. Review the business context and attached evidence, then open Data Review when evidence is missing or additional inspection is required.", fill=GOLD_LIGHT, accent=GOLD)

    add_chapter(doc, 10, "Use in-application notifications", "Notifications are a personal Planning inbox, separate from technical logs and separate from optional SMTP operation emails.")
    add_heading(doc, "Current notification events")
    add_table(
        doc,
        ["Event", "Recipient", "Message behavior"],
        [
            ["Approval requested", "Named reviewer or every active user in the assigned reviewer role", "Warning message opens Approvals."],
            ["Submission approved", "User who submitted the work", "Success message opens My Work."],
            ["Submission returned", "User who submitted the work", "Warning includes the review comment and opens My Work."],
        ],
        [2500, 3000, 3860],
    )
    add_heading(doc, "Inbox controls")
    add_bullets(doc, [
        "Filter to unread messages only.",
        "Open the action attached to a message; opening can also mark it read.",
        "Mark one message read or select Mark all read.",
        "Unread count is calculated from durable user-owned notification records.",
    ])
    add_callout(doc, "Current boundary", "Assignments and due dates are visible in Home and My Work. The current in-application notification generator creates approval-request and approval-decision messages; it does not yet create a separate notification for every new assignment or approaching deadline.", fill=BLUE_LIGHT, accent=BLUE)

    add_hard_page_break(doc)
    add_chapter(doc, 11, "Schedule the right unattended work", "Schedules remove repetitive button clicks, but only when all required inputs can be resolved without pausing for a user or local file picker.")
    add_figure(doc, assets["schedule"], "Figure 6.7 - Current unattended scheduling flow.", "Five-step schedule flow from configuration and live review through atomic claim, repeat validation, and execution, with misfire and duplicate controls.", width=6.35)
    add_heading(doc, "Current schedule targets")
    add_table(
        doc,
        ["Target", "Use it for", "Important behavior"],
        [
            ["Oracle Pipeline", "Recurring technical lifecycles already configured in Oracle.", "Select a registered, currently verified Pipeline. Use Oracle defaults or fixed unattended values."],
            ["Business Rule RTP registry sync", "Refreshing the assistant/platform knowledge of Calculation Manager runtime prompts.", "Oracle generates a temporary Calculation Manager snapshot; the platform downloads, parses, publishes, and removes it."],
        ],
        [2500, 3300, 3560],
    )
    add_heading(doc, "Configure and review")
    add_numbered(doc, [
        "Choose the automation type and its live Oracle target.",
        "Give the schedule a business-friendly name.",
        "For a Pipeline, choose Oracle defaults or fixed unattended values. Fixed file inputs must be existing Oracle Inbox references, never a local computer path.",
        "Choose One time, Daily, Weekly, or Monthly; set the first run and an IANA timezone such as Asia/Kolkata. Daylight-saving changes are handled through the selected timezone.",
        "Choose Run once or Skip for an occurrence missed while the service was unavailable.",
        "Select Validate & review. Save only after live Oracle validation and recurrence preview pass.",
    ])
    add_heading(doc, "Operational controls")
    add_bullets(doc, [
        "Pause and resume without deleting the definition.",
        "Edit only after the target and unattended inputs revalidate.",
        "Archive the schedule while retaining its occurrence history.",
        "Search schedules and inspect next run, input strategy, last outcome, and last error.",
        "Filter occurrence history by schedule, outcome, and date range; open the linked execution when one exists.",
        "If the same target already has active work, the occurrence is skipped instead of duplicated.",
    ])
    add_callout(doc, "What is not scheduled today", "The current Schedules page does not schedule arbitrary standalone rules, Data Maps, imports, Data Integrations, Cube Refresh jobs, reports, or Planning Cycle creation. Put multi-stage recurring work in an Oracle Pipeline when appropriate. Dynamic per-occurrence Pipeline inputs are also not enabled yet.", fill=GOLD_LIGHT, accent=GOLD)

    add_hard_page_break(doc)
    add_chapter(doc, 12, "Run a complete monthly forecast", "This example shows how to use current capabilities from two business days before the cycle begins through final publication and close.")
    add_figure(doc, assets["timeline"], "Figure 6.8 - D-2 to close monthly forecast timeline.", "Timeline covering preparation, reconciliation, planning input, review, publication, and final lifecycle completion.", width=6.35)
    add_heading(doc, "D-2 - Prepare the cycle")
    add_bullets(doc, [
        "Service Administrator checks environment health and confirms the current Planning application.",
        "Oracle EPM consultant confirms the monthly Pipeline, integrations, rules, Data Maps, jobs, forms, mappings, security, and required Inbox files in Oracle.",
        "FP&A owner confirms Planning year, actual-through period, forecast-start period, dates, owners, reviewers, and business dependencies.",
        "Service Administrator opens the cycle using the four-stage starter pattern, then adjusts responsibilities rather than copying Oracle Pipeline stages.",
    ])
    add_heading(doc, "D-1 - Reconcile actuals")
    add_bullets(doc, [
        "Assigned operator opens the first ready responsibility.",
        "If actual load and transformation are inside a Pipeline, run one Pipeline responsibility. Otherwise run the smallest assigned operation.",
        "Open Data Review and run source quality checks or a source-to-target comparison.",
        "Resolve failures; acknowledge only understood warnings; complete the validation responsibility.",
    ])
    add_heading(doc, "D0 to D2 - Complete Planning input")
    add_bullets(doc, [
        "Planners work from My Work, not from remembered technical menus.",
        "Complete assumptions in the authorized Oracle form or interface, then run assigned calculations when required.",
        "Inspect result data in Data Review and complete the responsibility only after the accepted validation result exists.",
        "The Submit for approval responsibility becomes ready after its prerequisites complete.",
    ])
    add_heading(doc, "D2 to D3 - Review and approve")
    add_bullets(doc, [
        "Planner submits completed work.",
        "Reviewer opens Approvals, checks context and validation evidence, and approves or returns with a precise correction request.",
        "Returned work reopens for the planner; resubmission creates a fresh pending decision while retaining prior history.",
    ])
    add_heading(doc, "D3 to D4 - Publish and close")
    add_bullets(doc, [
        "Run the Oracle Pipeline or Data Map responsible for reporting publication, but not both when the Pipeline already contains the Data Map.",
        "Generate the authorized business report and inspect Jobs & Activity for retained execution evidence.",
        "Complete the final business responsibility. When all responsibilities finish, every stage completes and the cycle reaches 100 percent and Completed status.",
    ])

    add_chapter(doc, 13, "Monitor and recover by evidence", "When work is late, waiting, blocked, or failed, diagnose the first incomplete responsibility and its proof instead of restarting the entire lifecycle.")
    add_heading(doc, "A simple exception sequence")
    add_numbered(doc, [
        "Open the cycle context and identify the first incomplete stage.",
        "Open My Work and filter to Waiting or Blocked.",
        "For Waiting, read the named prerequisite and ask its owner to complete it.",
        "For Blocked manual work, resolve the business issue and select Resume.",
        "For a failed Oracle operation, select View run, inspect the failed step and Oracle message, correct the cause, then use Retry operation.",
        "For a failed Data Review, correct data or the assigned selection and validate again. Do not acknowledge a Fail.",
        "For a returned approval, follow the review comment, rerun evidence when necessary, and resubmit.",
    ])
    add_heading(doc, "Evidence retained")
    add_table(
        doc,
        ["Evidence", "What it proves"],
        [
            ["Task and dependency record", "Who owned the business outcome, when it was due, and what had to finish first."],
            ["Execution attempts", "Every linked attempt number, execution ID, initiator, state, completion time, and failure message."],
            ["Jobs & Activity", "Workflow steps, trigger source, Oracle status, logs, output files, statistics, and timestamps."],
            ["Validation evidence", "Cube selection, criteria, checked cells, exceptions, warnings, acknowledgement, and performer."],
            ["Approval history", "Submitter, reviewer, decision, comment, timing, and linked validation."],
            ["Schedule occurrence history", "Scheduled time, atomic claim, resolved payload, outcome, execution ID, and error."],
        ],
        [2800, 6560],
    )
    add_callout(doc, "Browser behavior", "Closing or refreshing the browser does not cancel a queued or running execution. Reopen My Work or Jobs & Activity and inspect the retained request before attempting another run.", fill=TEAL_LIGHT, accent=TEAL)

    add_chapter(doc, 14, "Use an administrator design checklist", "A good lifecycle removes manual navigation without hiding accountability, evidence, or Oracle ownership.")
    add_heading(doc, "Cycle design checklist")
    add_bullets(doc, [
        "Use a business-recognizable cycle name and a stable unique code.",
        "Keep Planning year required; leave Scenario blank when Oracle mapping owns it.",
        "Use four to six understandable business stages rather than dozens of technical steps.",
        "Assign every responsibility to one current user or role and give high-risk work an explicit owner.",
        "Make instructions outcome-based and concise.",
        "Use dependencies only for real gates and verify there is no circular or impossible path.",
        "Use one Pipeline responsibility when Oracle already owns the complete technical sequence.",
        "Prefill Data Review when users should validate one repeatable cube intersection.",
        "Pair every Submit for approval responsibility with at least one dependent Review approval responsibility.",
        "Use an explicit substitution-variable task only when the change is not already governed inside the Oracle Pipeline.",
        "Review all fields before opening because current post-open edit/delete controls are not available.",
    ])
    add_heading(doc, "Scheduling checklist")
    add_bullets(doc, [
        "Schedule only an Oracle Pipeline or RTP registry sync supported by the current page.",
        "Confirm every required value is available from Oracle defaults or a fixed unattended value.",
        "Use only current Oracle Inbox file references; never save local paths or credentials.",
        "Select the correct timezone and missed-occurrence behavior.",
        "Keep Skip if active as duplicate protection and review the first occurrence in history.",
        "Pause schedules before planned Oracle maintenance or material Pipeline redesign, then revalidate before resuming.",
    ])
    add_heading(doc, "Security checklist")
    add_bullets(doc, [
        "Use platform roles for navigation and lifecycle authorization, and Oracle security for application data and artifacts.",
        "Never place passwords, API keys, session values, or encrypted-password file content in task instructions or action configuration; the backend rejects credential-shaped secret fields.",
        "Use a Service Administrator for cycle design and shared substitution-variable responsibilities.",
        "Use named assignments when individual accountability matters and role assignments only for a genuinely shared queue.",
    ])

    add_chapter(doc, 15, "Know the current boundaries and terms", "The current lifecycle is production-oriented but intentionally does not claim features that are not yet implemented in the modern application.")
    add_heading(doc, "Implemented now")
    add_bullets(doc, [
        "Role-aware Home, My Work, Planning Cycles, Approvals, Notifications, Operations, Schedules, Data Review, Reports, and Jobs & Activity navigation.",
        "Three-step Planning Cycle wizard with current users/roles, dates, priorities, task actions, one selected prerequisite, and optional Data Review prefill.",
        "Dependency-aware readiness, controlled manual transitions, automatic stage/cycle progress, operation-linked completion, retries, and retained validation/approval evidence.",
        "Unattended Oracle Pipeline and Business Rule RTP registry schedules with live validation, recurrence preview, pause/resume, archive, misfire handling, duplicate protection, and occurrence history.",
    ])
    add_heading(doc, "Not implemented in the current Planning Workspace")
    add_bullets(doc, [
        "Direct synchronization of Oracle Task Manager or Oracle Planning Task Lists into My Work.",
        "Automatic creation or redesign of Oracle Pipelines, forms, Data Maps, rules, integrations, import jobs, or refresh jobs from the cycle wizard.",
        "Post-open cycle editing, cloning, deletion, or archive controls in the current React Planning Cycles screen.",
        "Automatic in-application alerts for every assignment, due-soon item, or overdue item; current durable notifications cover approval requests and decisions.",
        "Scheduling of every standalone operation or arbitrary dynamic per-occurrence Pipeline values.",
        "A specific Oracle artifact selection inside each cycle responsibility; artifact and input selection occurs at governed execution time.",
    ])
    add_callout(doc, "Why these boundaries matter", "The platform should automate repetitive work without inventing Planning configuration or silently changing Oracle application context. Where Oracle or a business owner must make a design decision, the current interface keeps that decision explicit.", fill=GOLD_LIGHT, accent=GOLD)
    add_heading(doc, "Glossary")
    add_table(
        doc,
        ["Term", "Plain-language meaning"],
        [
            ["Cycle", "One dated business event such as March Forecast FY27."],
            ["Stage", "An ordered business checkpoint inside a cycle."],
            ["Responsibility / task", "One accountable outcome assigned to a person or role."],
            ["Dependency", "A task that must complete before another task becomes ready."],
            ["Readiness", "A calculated label showing whether the task can proceed."],
            ["Action type", "The governed workspace or operation opened from a responsibility."],
            ["Execution attempt", "One queued Oracle run linked to a task, including retries."],
            ["Validation evidence", "Retained proof that a selected cube intersection passed, warned, or failed configured checks."],
            ["Approval", "A submitted business decision assigned to a dependent reviewer task."],
            ["Schedule occurrence", "One unique unattended run expected at a particular time."],
            ["Misfire", "An occurrence missed while the scheduler was unavailable."],
            ["Oracle defaults", "Current values and files already provided by the Oracle Pipeline definition."],
        ],
        [2600, 6760],
    )
    add_callout(doc, "End of Document 06", "Use this volume to operate the current business lifecycle from cycle creation through close. Document 07 continues with Data Review and Reporting: live cube grids, validation, comparison, exports, registered report definitions, and generated Excel outputs.", fill=TEAL_LIGHT, accent=TEAL)

    _finish_document(doc)
    doc.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    print(build_document())
